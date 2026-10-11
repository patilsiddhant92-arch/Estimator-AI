"""The Evidence Package of the rebuilt app: the old workbook (the bid check, how the job was read, the job ledger,
the spec overview and the reading tabs, from `estimator.evidence.workbook.write_evidence`, imported and not
edited) with the new Job right after the bid check:

- "Scope lines": per system, every device with its printed line, slot, verb, who furnishes / installs / wires it,
  its quantity and the other counts (and the sum of the sheets of one floor and area when ruling L1 left it above
  the printed count: `Count.added_up`), every citation (source id, where: sheet / spec section / page, and the quote),
  its assumption (what and why) and its flag; the family counts and the block's clarification too.
- "Findings": every finding of the validators and the QC left open (rule, system, device, severity, what, fix).
- "Demolition Scope" (T-G1): one row per demolition line (family, action, count, note, refs as source id, quote, where).
- "Set aside" (T-007): one row per count the takeoff set aside (family, number, reason, note with both numbers, its ref
  as source id, quote, where), every reason, the demolition-sheet ones too.

Plain data only: it writes what the Job holds, nothing internal is invented.
"""
from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from estimator.evidence.workbook import _table
from estimator.evidence.workbook import write_evidence as write_old_evidence

from bms.model import Count, Device, Finding, Job, Ref, SetAside
from bms.render import COUNT_ADDED_UP, left_out_text

SCOPE_TAB, FINDINGS_TAB, DEMOLITION_TAB = "Scope lines", "Findings", "Demolition Scope"
SET_ASIDE_TAB = "Set aside"
SCOPE_HEADINGS = ["System", "Device id", "Kind", "Printed line", "Slot", "Device", "Verb", "Furnished by",
                  "Installed by", "Wired by", "Quantity", "Other counts", "Citations (id, where, quote)",
                  "Assumption", "Flag"]
FINDING_HEADINGS = ["Rule", "System", "Device", "Severity", "What", "Fix"]
DEMOLITION_HEADINGS = ["Family", "Action", "Count", "Note", "Refs"]
SET_ASIDE_HEADINGS = ["Family", "Set aside", "Reason", "Note (both numbers)", "Ref (id, quote, where)"]


def _ref(ref: Ref) -> str:
    return f'{ref.source_id or "-"} | {ref.where or "-"} | "{ref.quote}"'


def _refs(refs: Sequence[Ref]) -> str:
    return "\n".join(_ref(r) for r in refs)


def _quantity(count: Count) -> str:
    """The number the proposal prints (`Count.highest()`, estimator ruling D9)."""
    n = count.highest()
    return "Qty (##), not known" if n is None else str(n)


def _others(count: Count) -> str | None:
    """Every other count with its whole citation: source id, where and quote. Last, the sum ruling L1 left out of the
    count when it is above the number printed (`Count.added_up`), worded as the proposal's note words it: it is no
    count that prints, and it is named beside the printed one."""
    lines = [f"{o.value} per {_ref(o.ref)}" for o in count.others]
    up, top = count.added_up, count.highest()
    if up is not None and top is not None and up.value > top:
        lines.append(f"{COUNT_ADDED_UP.format(n=up.value)}, per {_ref(up.ref)}")
    return "\n".join(lines) or None


def _device_row(system: str, device: Device, printed: dict[str, str]) -> list:
    a = device.assumption
    return [system, device.id, "device", printed.get(device.id), device.slot, device.name, device.verb,
            device.furnished_by, device.installed_by, device.wired_by, _quantity(device.quantity),
            _others(device.quantity), _refs([*device.refs, *device.quantity.sources]),
            f"{a.what} — why: {a.why}" if a is not None else None, device.flag]


def scope_rows(job: Job, printed: dict[str, str] | None = None) -> list[list]:
    """The Scope lines rows: per system its family counts, its devices in model order (a system with none: a "left
    out" row with the writer's reason), its clarification; then the front end's devices and the exceptions."""
    printed = printed or {}
    rows: list[list] = []
    for system in job.systems:
        label = f"{system.name} ({system.key})"
        if system.alternate != "base":
            label += f", {system.alternate} alternate"
        for family in system.families:
            rows.append([label, None, "count", None, None, f"{family.name} {', '.join(family.tags)}".strip(), None,
                         None, None, None, _quantity(family.count), _others(family.count),
                         _refs(family.count.sources), None, None])
        rows += [_device_row(label, d, printed) for d in system.devices]
        if not system.devices:              # it prints nothing in the proposal: the package says why
            rows.append([label, None, "left out", left_out_text(system)] + [None] * 11)
        for clarification in system.clarifications:
            rows.append([label, None, f"clarification ({clarification.kind})", clarification.text] + [None] * 11)
        if system.sequence_refs:
            rows.append([label, None, "sequence", None, None, None, None, None, None, None, None, None,
                         _refs(system.sequence_refs), None, None])
    if job.front_end is not None:
        rows += [_device_row("Front end", d, printed) for d in job.front_end.devices]
        for clarification in job.front_end.clarifications:
            rows.append(["Front end", None, f"clarification ({clarification.kind})", clarification.text] + [None] * 11)
    for exception in job.exceptions:
        rows.append([exception.system_key or "Whole job", None, "exception", exception.text] + [None] * 8
                    + [_refs(exception.refs), None, None])
    return rows


def finding_rows(findings: Sequence[Finding]) -> list[list]:
    return [[f.rule, f.system_key or "Whole job", f.device_id, f.severity, f.what, f.fix] for f in findings]


def _demolition_refs(refs: Sequence[Ref]) -> str:
    """Each citation as source id, quote, where. One citation per line."""
    return "\n".join(f'{ref.source_id or "-"} | "{ref.quote}" | {ref.where or "-"}' for ref in refs)


def demolition_rows(job: Job) -> list[list]:
    """One row per Demolition Scope line: family, action, count, note, and its citations."""
    return [[line.family, line.action, line.count, line.note, _demolition_refs(line.refs)] for line in job.demolition]


def set_aside_rows(set_aside: Sequence[SetAside]) -> list[list]:
    """One row per set-aside, in the takeoff's order: family, the number set aside, the reason, the note, its ref."""
    return [[a.family, a.value, a.reason, a.note, _demolition_refs((a.ref,))] for a in set_aside]


def write_evidence(path: Path, catalogue, reading, ledger, job: Job, findings: Sequence[Finding] = (),
                   printed: dict[str, str] | None = None) -> Path:
    """The Evidence Package at `path`: the old workbook from the catalogue, reading and ledger, with the Job's
    Scope lines and the Findings as the second and third tabs, the Demolition Scope as the fourth and the takeoff's
    set-asides ("Set aside") as the fifth. `printed`
    maps a device id to its printed line (from `bms.render`)."""
    path = Path(path)
    write_old_evidence(path, catalogue, reading, ledger=ledger)
    book = load_workbook(str(path))
    tabs = ((SCOPE_TAB, SCOPE_HEADINGS, scope_rows(job, printed)),
            (FINDINGS_TAB, FINDING_HEADINGS, finding_rows(findings)),
            (DEMOLITION_TAB, DEMOLITION_HEADINGS, demolition_rows(job)),
            (SET_ASIDE_TAB, SET_ASIDE_HEADINGS, set_aside_rows(job.set_aside)))
    for at, (title, headings, rows) in enumerate(tabs, start=1):
        sheet = book.create_sheet(title, at)
        _table(sheet, headings, rows or [["Nothing here."] + [None] * (len(headings) - 1)])
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and "\n" in cell.value:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.save(str(path))
    return path
