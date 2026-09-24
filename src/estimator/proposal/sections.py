"""A proposal's parts as plain text — for checking one proposal against another."""
from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import _Cell

from estimator.proposal import layout
from estimator.text import squash


def _texts(path: Path) -> list[str]:
    return [squash(p.text) for p in docx.Document(str(path)).paragraphs if p.text.strip()]


def _index(texts: list[str], prefix: str) -> int:
    return next(i for i, t in enumerate(texts) if t.startswith(prefix))


def scope_texts(path: Path) -> list[str]:
    """Every line from SCOPE OF WORK down to, not including, the clarifications."""
    texts = _texts(path)
    return texts[_index(texts, layout.SCOPE_HEADING) : _index(texts, layout.CLARIFICATIONS)]


def closing_texts(path: Path) -> list[str]:
    """Every line from the clarifications to the end — the part the template owns."""
    texts = _texts(path)
    return texts[_index(texts, layout.CLARIFICATIONS) :]


def opening_texts(path: Path) -> list[str]:
    """Every line before the opening sentence — the heading area the template owns."""
    texts = _texts(path)
    return texts[: _index(texts, layout.OPENING)]


def cover_lines(path: Path) -> list[str]:
    """The cover page table, line by line."""
    document = docx.Document(str(path))
    table = document.tables[0]
    return [squash(p.text) for tc in table._tbl.iter(qn("w:tc")) for p in _Cell(tc, table).paragraphs if p.text.strip()]
