"""Read a proposal written on the AME template back into a Job.

Used to turn the proposal an estimator actually sent into test data and, from
build step 2, to score the tool against past jobs. It reads what Word draws —
bullets, numbering, indent — rather than what the styles are called, because
copy-paste between proposals leaves the same visual level built from several
different lists and styles.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from estimator.files import open_docx
from estimator.proposal import layout
from estimator.proposal.word import WordLists, deleted_text, field_result_text, merged_runs, visible_text
from estimator.scope import Alternate, Block, FrontEnd, Header, Item, Job, Run
from estimator.text import squash


class ProposalFormatError(ValueError):
    """The proposal doesn't follow the template closely enough to be read safely."""


# --- the cover page ---------------------------------------------------------------

def read_cover(document: DocxDocument) -> Header:
    """The table on the cover page: to, date, quote number, reference, engineer, drawings."""
    table = document.tables[0]
    values: dict[str, object] = {}
    for tc in table._tbl.iter(qn("w:tc")):
        lines = [p for p in _Cell(tc, table).paragraphs if p.text.strip()]
        if not lines:
            continue
        first = squash(lines[0].text)
        if first.startswith("To:"):
            values["to"] = first[len("To:"):].strip()
        elif first.startswith("Date:"):
            values["date"] = _date(field_result_text(lines[0], "DATE") or first[len("Date:"):])
        elif first.startswith("Quote"):
            values["quote_number"] = re.sub(r"^Quote\s*#?\s*:\s*", "", first)
        elif first.startswith("Reference"):
            joined = " ".join(squash(p.text) for p in lines)
            values["reference"] = re.sub(r"^Reference\s*:\s*", "", joined)
        elif first.startswith("MEP Engineer:"):
            values["mep_engineer"] = first[len("MEP Engineer:"):].strip()
        elif first.startswith("Based on Plans and Specs:"):
            rest = first[len("Based on Plans and Specs:"):].strip()
            values["based_on"] = tuple(([rest] if rest else []) + [squash(p.text) for p in lines[1:]])
    missing = {"to", "date", "quote_number", "reference", "mep_engineer", "based_on"} - values.keys()
    if missing:
        raise ProposalFormatError(f"the cover page is missing {sorted(missing)}")
    return Header(**values)


def _date(text: str) -> dt.date:
    try:
        return dt.datetime.strptime(squash(text), "%m/%d/%Y").date()
    except ValueError as error:
        raise ProposalFormatError(f"can't read the date {text!r}") from error


# --- the exclusions ---------------------------------------------------------------

def _exclusions(document: DocxDocument) -> list[str]:
    # visible_text, not p.text: a paragraph whose only text sits inside a tracked insertion (or
    # any other container p.text can't see) must still be found here, so it can be refused below
    # rather than silently missing from the list entirely.
    paragraphs = [p for p in document.paragraphs if visible_text(p).strip()]
    texts = [squash(p.text) for p in paragraphs]
    start = next((i for i, t in enumerate(texts) if t.startswith(layout.EXCLUSIONS)), None)
    end = next((i for i, t in enumerate(texts) if t.startswith(layout.BASE_PRICE)), None)
    if start is None or end is None or end <= start:
        raise ProposalFormatError("can't find the exclusions list")
    _refuse_foreign_children(document, paragraphs[start], paragraphs[end], region="among the exclusions")
    _refuse_deleted_paragraphs(document.paragraphs, paragraphs[start], paragraphs[end])
    for paragraph in paragraphs[start + 1 : end]:
        _refuse_hidden_text(paragraph)
    return texts[start + 1 : end]


def _clarifications(document: DocxDocument) -> list[str]:
    paragraphs = [p for p in document.paragraphs if visible_text(p).strip()]
    texts = [squash(p.text) for p in paragraphs]
    start = next((i for i, t in enumerate(texts) if t.startswith(layout.CLARIFICATIONS)), None)
    end = next((i for i, t in enumerate(texts) if t.startswith(layout.EXCLUSIONS)), None)
    if start is None or end is None or end <= start:
        raise ProposalFormatError("can't find the clarifications list")
    _refuse_foreign_children(document, paragraphs[start], paragraphs[end], region="among the clarifications")
    _refuse_deleted_paragraphs(document.paragraphs, paragraphs[start], paragraphs[end])
    for paragraph in paragraphs[start + 1 : end]:
        _refuse_hidden_text(paragraph)
    return texts[start + 1 : end]


def read_standard_lines(template_path: Path) -> tuple[tuple[str, ...], tuple[str, ...]]:
    """The template's standard clarifications and its standard exclusions, word for word, in
    the order the proposal prints them."""
    document = open_docx(template_path)
    return tuple(_clarifications(document)), tuple(_exclusions(document))


def read_named_exceptions(document: DocxDocument, template_path: Path) -> tuple[str, ...]:
    """Exclusions this job added after the template's standard list.

    T-009 (ruling 22): a proposal prints one of the template's two demolition lines, so it may carry both or lack
    exactly one of them; it may not lack both. Any other standard line missing, edited or reordered is refused."""
    standard = _exclusions(open_docx(template_path))
    sent = _exclusions(document)
    demolition = {squash(layout.DEMOLITION_GENERAL), squash(layout.DEMOLITION_SPECIFIC)}
    at, missing = 0, 0
    for line in standard:
        if at < len(sent) and sent[at] == line:
            at += 1
        elif line in demolition:
            missing += 1
        else:
            at = -1
            break
    if at < 0 or missing > 1:
        raise ProposalFormatError(
            "the standard exclusions were edited, reordered or cut in this proposal; the tool can "
            "only add exceptions after the standard list, so it can't reproduce this one")
    return tuple(sent[at:])


# --- the scope of work ------------------------------------------------------------

def read_proposal(path: Path, template_path: Path) -> Job:
    """A sent proposal as a Job. A proposal carries no hours or quantities, so those stay empty."""
    document = open_docx(path)
    suffix, front_end, blocks, add, deduct = read_scope(document)
    return Job(
        header=read_cover(document),
        scope_heading_suffix=suffix,
        front_end=front_end,
        blocks=blocks,
        add_alternates=add,
        deduct_alternates=deduct,
        named_exceptions=read_named_exceptions(document, template_path),
    )


def read_scope(document: DocxDocument):
    """(scope heading suffix, front end, blocks, add alternates, deduct alternates)."""
    lists = WordLists(document)
    full = document.paragraphs
    joined_after = _joined_after_map(full)
    paragraphs = [p for p in full if visible_text(p).strip()]
    start = _find(paragraphs, layout.SCOPE_HEADING)
    end = _find(paragraphs, layout.CLARIFICATIONS)
    _refuse_foreign_children(document, paragraphs[start], paragraphs[end])
    _refuse_deleted_paragraphs(full, paragraphs[start], paragraphs[end])
    _refuse_hidden_text(paragraphs[start])
    _refuse_hidden_text(paragraphs[start + 1])
    suffix = squash(paragraphs[start].text)[len(layout.SCOPE_HEADING):].strip()
    mode, existing = _front_end_mode(squash(paragraphs[start + 1].text))

    front_end_items: list[Item] = []
    blocks: list[dict] = []
    alternates: dict[str, list[dict]] = {"add": [], "deduct": []}
    section = "front end"
    items: list[Item] | None = None
    previous_depth = 0   # the first line of a block, or of the front end, is at most depth 1
    for paragraph in paragraphs[start + 2 : end]:
        _refuse_hidden_text(paragraph)
        text = squash(paragraph.text)
        if text.upper().startswith(layout.ADD_ALTERNATES):
            section, items = "add", None
            continue
        if text.upper().startswith(layout.DEDUCT_ALTERNATES):
            section, items = "deduct", None
            continue
        if section in alternates and re.match(r"ALT\b", text):
            alternates[section].append({"label": text, "blocks": []})
            items = None
            continue
        if _is_block_header(paragraph):
            block = {"header": text, "items": [], "joined": joined_after[id(paragraph)],
                     "numbered": not _is_group_heading(paragraph)}
            if section == "front end":
                section = "blocks"
            if section == "blocks":
                blocks.append(block)
            elif alternates[section]:
                alternates[section][-1]["blocks"].append(block)
            else:
                raise ProposalFormatError(f"{text!r} sits under the {section} alternates before any ALT line")
            items = block["items"]
            previous_depth = 0
            continue
        item = _read_item(paragraph, lists, in_front_end=(section == "front end"), previous_depth=previous_depth)
        if item.kind == "line":
            previous_depth = item.depth
        if section == "front end":
            front_end_items.append(item)
        elif items is not None:
            items.append(item)
        else:
            raise ProposalFormatError(f"{text!r} sits outside any block")

    front_end = FrontEnd(mode=mode, existing_system_name=existing, items=tuple(front_end_items))
    return (
        suffix,
        front_end,
        tuple(_block(b) for b in blocks),
        _alternates(alternates["add"]),
        _alternates(alternates["deduct"]),
    )


def _alternates(raw: list[dict]) -> tuple[Alternate, ...]:
    """The template's own empty "ALT #:" line is a placeholder, not an alternate — leave it out."""
    return tuple(_alternate(a) for a in raw if a["blocks"] or a["label"] != layout.ALTERNATE_LABEL)


def _find(paragraphs: list[Paragraph], heading: str) -> int:
    for index, paragraph in enumerate(paragraphs):
        if squash(paragraph.text).startswith(heading):
            return index
    raise ProposalFormatError(f"can't find {heading!r}")


# Body-level elements the reader is safe to skip over — a bookmark or a comment range marker
# carries no text of its own, so it can sit between two paragraphs without hiding anything.
_HARMLESS_BODY_CHILDREN = (qn("w:p"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"),
                          qn("w:commentRangeStart"), qn("w:commentRangeEnd"))
_FOREIGN_BODY_CHILD_NAMES = {
    qn("w:tbl"): "a table", qn("w:sdt"): "a content control", qn("w:customXml"): "custom XML",
    qn("w:altChunk"): "embedded content",
}


def _refuse_foreign_children(document: DocxDocument, first: Paragraph, last: Paragraph,
                             region: str = "inside the scope of work") -> None:
    """Anything between the two headings that isn't an ordinary paragraph — a table, a content
    control, custom XML, embedded (altChunk) content — can carry text of its own that the reader,
    which only ever walks document.paragraphs, would never see at all."""
    inside = False
    for element in document.element.body.iterchildren():
        if element is first._p:
            inside = True
        elif element is last._p:
            return
        elif inside and element.tag not in _HARMLESS_BODY_CHILDREN:
            name = _FOREIGN_BODY_CHILD_NAMES.get(element.tag, f"a {element.tag.split('}')[-1]}")
            raise ProposalFormatError(f"{name} sits {region}; the reader doesn't handle it")


def _refuse_deleted_paragraphs(full: list[Paragraph], first: Paragraph, last: Paragraph) -> None:
    """A paragraph that is entirely a tracked deletion has no ordinary run text at all — its
    words live only in <w:delText> — so it has no visible text and never reaches the paragraph
    lists read_scope and _exclusions filter down to. Left unchecked, it would simply vanish from
    the Job with no complaint, the same failure _refuse_hidden_text exists to catch on a paragraph
    that keeps some visible text."""
    inside = False
    for paragraph in full:
        if paragraph._p is first._p:
            inside = True
            continue
        if paragraph._p is last._p:
            return
        if inside and not visible_text(paragraph).strip():
            removed = deleted_text(paragraph)
            if removed.strip():
                raise ProposalFormatError(
                    f"{squash(removed)!r} is a tracked deletion — accept or reject tracked changes in Word first")


def _refuse_hidden_text(paragraph: Paragraph) -> None:
    """A tracked insertion, a content control or a text box can hide text from every reader that
    only looks at ordinary runs. Compare what the ordinary runs say against every character Word
    actually draws, so a paragraph like that is refused rather than silently read short. Hidden
    (vanish) text and unaccepted tracked deletions or moves need their own check: both leave the
    ordinary runs and the visible text agreeing with each other, just agreeing on the wrong thing."""
    stretches = merged_runs(paragraph)
    if re.sub(r"\s+", "", "".join(t for t, _ in stretches)) != re.sub(r"\s+", "", visible_text(paragraph)):
        raise ProposalFormatError(
            f"part of {squash(paragraph.text)!r} sits inside something the reader doesn't handle yet "
            "(a tracked change, a content control or a text box)")
    if _has(paragraph, "w:vanish") or _has(paragraph, "w:webHidden"):
        raise ProposalFormatError(f"{squash(paragraph.text)!r} has hidden text, which the reader can't tell apart "
                                  "from ordinary text")
    if _has(paragraph, "w:del") or _has(paragraph, "w:moveFrom"):
        raise ProposalFormatError(f"{squash(paragraph.text)!r} has an unaccepted tracked deletion or move; "
                                  "accept or reject the tracked changes before this proposal is read")


def _has(paragraph: Paragraph, tag: str) -> bool:
    return next(paragraph._p.iter(qn(tag)), None) is not None


def _front_end_mode(heading: str) -> tuple[str, str | None]:
    if heading.startswith(layout.NEW_SYSTEM_HEADING):
        return "new-system", None
    if heading.startswith(layout.EXISTING_SYSTEM_HEADING):
        name = heading[len(layout.EXISTING_SYSTEM_HEADING):].strip()
        return "extend-existing", re.sub(r"\s*BMS$", "", name).strip()
    raise ProposalFormatError(f"expected the front-end heading after SCOPE OF WORK, found {heading!r}")


def _joined_after_map(paragraphs: list[Paragraph]) -> dict[int, bool]:
    """For every paragraph (keyed by id(paragraph)): is the paragraph right before it, blanks
    included, itself non-blank? A header for which this is true runs straight on from whatever
    came before it — no blank line — so `Block.joined` should be True."""
    joined: dict[int, bool] = {}
    previous_blank = True
    for paragraph in paragraphs:
        joined[id(paragraph)] = not previous_blank
        previous_blank = not visible_text(paragraph).strip()
    return joined


def _mark_size_pt(paragraph: Paragraph) -> float | None:
    """The paragraph mark's own font size, in points — what Word draws a list number or bullet
    in. None if the mark doesn't override the style's size."""
    pPr = paragraph._p.pPr
    rPr = pPr.find(qn("w:rPr")) if pPr is not None else None
    sz = rPr.find(qn("w:sz")) if rPr is not None else None
    return int(sz.get(qn("w:val"))) / 2 if sz is not None else None


def _mark_bold(paragraph: Paragraph) -> bool:
    """Whether the paragraph mark's own run properties are bold — what Word draws a list number
    or bullet in. A bare <w:b/> (no w:val) means bold; a w:val of 0/false/off turns it back off."""
    pPr = paragraph._p.pPr
    rPr = pPr.find(qn("w:rPr")) if pPr is not None else None
    b = rPr.find(qn("w:b")) if rPr is not None else None
    return b is not None and b.get(qn("w:val")) not in ("0", "false", "off")


def _is_group_heading(paragraph: Paragraph) -> bool:
    """An unnumbered group heading such as "Kitchen Exhaust System": PR2, numbering explicitly
    switched off (list 0), with the paragraph mark bold and sized like a block header's — unlike
    a note, which also sits at list 0 but never carries that mark formatting."""
    return (WordLists.written_numbering(paragraph) == (0, 0)
            and _mark_size_pt(paragraph) == layout.GROUP_HEADING.size
            and _mark_bold(paragraph))


def _is_block_header(paragraph: Paragraph) -> bool:
    """A PR2 paragraph Word numbers by its own outline list, or an unnumbered group heading."""
    if paragraph.style.name != layout.OUTLINE_STYLE:
        return False
    return WordLists.written_numbering(paragraph) is None or _is_group_heading(paragraph)


def _read_item(paragraph: Paragraph, lists: WordLists, *, in_front_end: bool, previous_depth: int) -> Item:
    stretches = merged_runs(paragraph)
    runs = tuple(Run(text=re.sub(r"\s+", " ", text), bold=bold) for text, bold in stretches)
    if lists.numbering(paragraph) is None:
        return Item(kind="note", runs=runs)
    base, step = ((layout.FRONT_END_LINE_INDENT, layout.FRONT_END_LINE_STEP) if in_front_end
                  else (layout.BLOCK_LINE_INDENT, layout.BLOCK_LINE_STEP))
    depth = max(1, 1 + round((lists.left_indent(paragraph) - base) / step))
    # A line is never more than one level deeper than the line before it in the same block (the
    # first line of a block or front end is at most depth 1, since previous_depth starts at 0) —
    # copy-paste between proposals sometimes leaves an indent that jumps two or more levels at
    # once, which is never really a new depth of nesting, just a paragraph indented wrong.
    depth = min(depth, previous_depth + 1)
    limit = 2 if in_front_end else 5
    if depth > limit:
        raise ProposalFormatError(f"{squash(paragraph.text)!r} is indented to depth {depth}; the most is {limit}")
    return Item(kind="line", depth=depth, runs=runs)


def _block(raw: dict) -> Block:
    return Block(header=raw["header"], items=tuple(raw["items"]), joined=raw["joined"], numbered=raw["numbered"])


def _alternate(raw: dict) -> Alternate:
    return Alternate(label=raw["label"], blocks=tuple(_block(b) for b in raw["blocks"]))


# --- the older "(xN)" style ---------------------------------------------------------------

_XN_HEADER = re.compile(r"^\*?\s*(?:\(x\s*\d+\)|Qty\s*\(\s*[\d#]+\s*\))", re.I)
_XN_END = ("ADD ALTERNATE", "DEDUCT ALTERNATE", "CLARIFICATIONS", "EXCEPTIONS AND EXCLUSIONS", "BASE PRICE")
_XN_STEP = 360                          # twips between one depth and the next


def read_blocks_fallback(document: DocxDocument) -> list[dict]:
    """A proposal written in the older "(xN)" style, which `read_scope` cannot read: a header line starting "(xN)"
    or "Qty (N)", then the lines under it until the next header. A deeper "(xN)" line right under a header is a
    sub-header, joined to it with " / ". Each block: {"header", "lines": [(text, depth)], "notes"}, the texts as
    written (a line's own "- (xN)" is kept)."""
    lists = WordLists(document)
    started, header_indent = False, None
    blocks: list[dict] = []
    for paragraph in document.paragraphs:
        text = squash(paragraph.text)
        if not text:
            continue
        if not started:
            started = text.upper().startswith("SCOPE OF WORK")
            continue
        if text.upper().startswith(_XN_END):
            break
        indent = lists.left_indent(paragraph)
        if _XN_HEADER.match(text) and (header_indent is None or indent <= header_indent):
            header_indent = indent if header_indent is None else header_indent
            blocks.append({"header": text, "lines": [], "notes": [], "base": None, "previous": 0})
            continue
        if not blocks:
            continue                                    # the front end
        block = blocks[-1]
        if _XN_HEADER.match(text):
            if not block["lines"]:
                block["header"] += " / " + text
            continue
        if text.startswith("*"):
            block["notes"].append(text)
            continue
        if block["base"] is None:
            block["base"] = indent
        depth = 1 + max(0, round((indent - block["base"]) / _XN_STEP))
        depth = max(1, min(depth, block["previous"] + 1, 5))
        block["previous"] = depth
        block["lines"].append((text, depth))
    return [{"header": b["header"], "lines": b["lines"], "notes": b["notes"]} for b in blocks]
