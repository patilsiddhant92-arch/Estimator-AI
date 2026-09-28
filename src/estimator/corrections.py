"""The estimator's corrections, read back from the Evidence Package and applied on the next run.

The Evidence Package's Scope lines and Review list tabs carry two columns for the estimator: "Estimator decision"
(keep, delete or change) and "Change to / why" (for a change, the new text, then "why:" and the reason; for keep or
delete, the reason). `estimator corrections EVIDENCE.xlsx` reads them into out/corrections/<job>.json (never into
the repository: the jobs are under NDA) and appends each reason given, made general, to
assets/gem/ame_practice_proposed.txt for the user to approve. The next run of the same job applies the file: a
deleted line is taken out of the proposal (with a review row saying so), a changed line takes the new text, a kept
line stays as it is; a review row the estimator deleted is taken off the review list.
"""
from __future__ import annotations

import dataclasses
import json
import re
from pathlib import Path

from openpyxl import load_workbook

KEEP, DELETE, CHANGE = "keep", "delete", "change"
LINES, REVIEW = "lines", "review"
DECISION, CHANGE_TO = "Estimator decision", "Change to / why"
SCOPE_TAB, REVIEW_TAB, BID_TAB = "Scope lines", "Review list", "Bid check"
DELETED = "Deleted by the estimator"
CHANGED = "Changed by the estimator"
_WHY = re.compile(r"\bwhy\s*:\s*", re.I)
_BAD_FILENAME = re.compile(r'[<>:"/\\|?*]')


def _squash(text) -> str:
    return " ".join(str(text or "").split())


def _key(text) -> str:
    return _squash(text).casefold()


def decision(value) -> str | None:
    """"Keep" / "delete" / "Change" (or their first letter) as the decision word; None when the cell says none."""
    word = _key(value)
    if not word:
        return None
    for name in (KEEP, DELETE, CHANGE):
        if name.startswith(word) or word.startswith(name) or (word == "remove" and name == DELETE):
            return name
    return None


def split_change(decided: str, value) -> tuple[str, str]:
    """(the new text, the reason) from the "Change to / why" cell. A change is its new text, then "why:" and the
    reason (or the two split by the last " / "); for keep or delete the whole cell is the reason."""
    text = _squash(value)
    if decided != CHANGE:
        return "", _WHY.sub("", text, count=1).strip()
    parts = _WHY.split(text, maxsplit=1)
    if len(parts) == 2:
        return parts[0].strip(" /-–—;"), parts[1].strip()
    if " / " in text:
        new, why = text.rsplit(" / ", 1)
        return new.strip(), why.strip()
    return text, ""


def _columns(ws) -> dict[str, int]:
    return {_squash(c.value): n for n, c in enumerate(ws[1], start=1) if c.value}


def _job_name(book, path: Path) -> str:
    if BID_TAB in book.sheetnames:
        title = _squash(book[BID_TAB]["A1"].value)
        m = re.match(r"^Bid check\s*[—–-]\s*(.+)$", title)
        if m:
            return m.group(1).strip()
    return re.sub(r"\s*Evidence Package$", "", path.stem).strip()


def _read_tab(ws, kind: str, text_heading: str) -> tuple[list[dict], list[str]]:
    columns = _columns(ws)
    if DECISION not in columns or text_heading not in columns:
        return [], []
    out, problems = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        def cell(heading):
            at = columns.get(heading)
            return row[at - 1] if at and at - 1 < len(row) else None
        raw = cell(DECISION)
        if raw is None or not _squash(raw):
            continue
        text = _squash(cell(text_heading))
        decided = decision(raw)
        if decided is None:
            problems.append(f"{ws.title}: “{_squash(raw)}” for “{text}” is not keep, delete or change; left out")
            continue
        if not text:
            continue
        change_to, why = split_change(decided, cell(CHANGE_TO))
        if decided == CHANGE and not change_to:
            problems.append(f"{ws.title}: “{text}” is marked change but gives no new text; left out")
            continue
        out.append({"kind": kind, "system": _squash(cell("System")), "text": text, "decision": decided,
                    "change_to": change_to, "why": why})
    return out, problems


def read_corrections(path: Path) -> tuple[dict, list[str]]:
    """The decisions the estimator filled in: ({"job", "lines", "review"}, problems in plain words)."""
    book = load_workbook(str(path), read_only=True, data_only=True)
    try:
        lines, problems = ([], []) if SCOPE_TAB not in book.sheetnames else _read_tab(book[SCOPE_TAB], LINES, "Line")
        review, more = (([], []) if REVIEW_TAB not in book.sheetnames
                        else _read_tab(book[REVIEW_TAB], REVIEW, "What to decide"))
        return {"job": _job_name(book, Path(path)), LINES: lines, REVIEW: review}, problems + more
    finally:
        book.close()


def corrections_file(folder: Path, job: str) -> Path:
    return Path(folder) / f"{_BAD_FILENAME.sub('-', job)}.json"


def _entry_key(entry: dict) -> tuple:
    return entry.get("kind"), _key(entry.get("system")), _key(entry.get("text"))


def save_corrections(corrections: dict, folder: Path) -> Path:
    """Write the job's corrections file, merged with the one already there: a new decision on the same line replaces
    the old one; the old ones on lines no longer in the package (a line deleted before) are kept."""
    path = corrections_file(folder, corrections["job"])
    have = load_corrections(path)
    merged: dict[tuple, dict] = {}
    for entry in (*have.get(LINES, ()), *have.get(REVIEW, ()), *corrections.get(LINES, ()),
                  *corrections.get(REVIEW, ())):
        merged[_entry_key(entry)] = entry
    data = {"job": corrections["job"], LINES: [e for e in merged.values() if e.get("kind") == LINES],
            REVIEW: [e for e in merged.values() if e.get("kind") == REVIEW]}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def load_corrections(path: Path) -> dict:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


# --- the reasons, as rules proposed for the user to approve ------------------------------------------------------

_QUOTED = re.compile(r"“[^”]*”|\"[^\"]*\"|‘[^’]*’|(?<!\w)'[^']*'(?!\w)")
_SOURCE_ID = re.compile(r"\[\s*[A-Z]{1,3}\d+\s*\]")
_JOB_NUMBER = re.compile(r"\b\d{2}-\d{4}\b")
_TAG = re.compile(r"(?<![\w-])[A-Z]{1,6}-[A-Z0-9]+(?:[-.][A-Z0-9]+)*(?:\s*(?:thru|through|to|,|&|and)\s*"
                  r"(?:[A-Z]{1,6}-)?[A-Z0-9]+(?:[-.][A-Z0-9]+)*)*(?![\w-])")
_SHEET = re.compile(r"(?<![\w-])[A-Z]{1,2}-?\d{3}(?:\.\d+)?[A-Z]?(?![\w-])")


def general_rule(why: str, job: str = "") -> str | None:
    """The estimator's reason as a plain rule for every job: no job name or number, no equipment tags or sheet
    numbers, no quotes from the job, no source ids. None when nothing general is left."""
    text = _squash(why)
    for name in {job, job.split(",")[0]} - {""}:
        text = re.sub(re.escape(name), " ", text, flags=re.I)
    for pattern in (_QUOTED, _SOURCE_ID, _JOB_NUMBER, _TAG, _SHEET):
        text = pattern.sub(" ", text)
    text = re.sub(r"[(\[]\s*(?:(?:job|sheets?|notes?|see|per|on|and)\b\s*[,;.&]?\s*|[,;.&]\s*)*[)\]]", " ", text,
                  flags=re.I)                      # a bracket left with nothing in it: "(Job , )"
    text = _squash(re.sub(r"\s+([,.;:])", r"\1", text)).strip(" ,;:-–—.")
    while True:                                    # a word left hanging at the end: "… a separate point on"
        cut = re.sub(r"\s+(?:on|in|at|for|of|per|to|see|sheet|note|and|with|from|by|the|a|an)$", "", text, flags=re.I)
        if cut == text:
            break
        text = cut.strip(" ,;:-–—")
    if len(re.findall(r"[A-Za-z]{2,}", text)) < 3:
        return None
    text = text[0].upper() + text[1:]
    return text if text.endswith((".", "!", "?")) else text + "."


def propose_rules(corrections: dict, path: Path) -> list[str]:
    """Append each reason the estimator gave, made general, to the proposed rules file (once each). Returns the
    lines added."""
    path = Path(path)
    try:
        have = path.read_text(encoding="utf-8-sig")
    except OSError:
        have = ("AME PRACTICE — PROPOSED RULES (from the estimators' corrections)\n"
                "Each line below is a reason an estimator gave for a correction, made general. Approve a rule by\n"
                "moving it into ame_practice.txt; delete the ones that do not hold for every job.\n")
    known = {_key(line.lstrip("- ")) for line in have.splitlines()}
    added = []
    for entry in (*corrections.get(LINES, ()), *corrections.get(REVIEW, ())):
        rule = general_rule(entry.get("why", ""), corrections.get("job", ""))
        if rule and _key(rule) not in known:
            known.add(_key(rule))
            added.append(f"- {rule}")
    if added:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(have.rstrip("\n") + "\n" + "\n".join(added) + "\n", encoding="utf-8")
    return added


# --- applying them on the next run --------------------------------------------------------------------------------

def _match(entries: list[dict], system_names: set[str], text: str) -> dict | None:
    """The correction for this line: the one naming its text and system, else the one naming its text alone."""
    same = [e for e in entries if _key(e.get("text")) == _key(text)]
    if not same:
        return None
    return next((e for e in same if _key(e.get("system")) in system_names), None) or (same[0] if len(same) == 1
                                                                                     else None)


def _items(items, decide):
    out = []
    for item in items:
        if item.kind != "line":
            out.append(item)
            continue
        entry = decide(item.text)
        if entry is None or entry["decision"] == KEEP:
            out.append(item)
        elif entry["decision"] == CHANGE:
            out.append(item.model_copy(update={"runs": (item.runs[0].model_copy(update={"text": entry["change_to"]}),)}))
    return tuple(out)


def apply_corrections(job, systems, review, corrections: dict):
    """(job, systems, review) with the estimator's corrections applied: a deleted line out of the proposal (kept in
    the Evidence Package, not printed, with a review row), a changed line with its new text, a review row the
    estimator deleted taken off the list. Each ReviewItem and ScopeLine is changed by its field names only."""
    entries = [e for e in corrections.get(LINES, ()) if e.get("decision") in (DELETE, CHANGE)]
    dropped_rows = [e for e in corrections.get(REVIEW, ()) if e.get("decision") == DELETE]
    if not entries and not dropped_rows:
        return job, systems, review
    applied: list[tuple[str, dict]] = []
    replaced: dict[int, object] = {}          # each system block, by identity: the job holds the same blocks
    new_systems = []
    for system in systems:
        names = {_key(system.bundle_key)} | {_key(line.system) for line in system.lines}
        lines, rows = [], list(system.review)
        for line in system.lines:
            entry = _match(entries, names, line.text) if line.client else None
            if entry is None:
                lines.append(line)
                continue
            applied.append((system.bundle_key, entry))
            reason = f" — why: {entry['why']}" if entry.get("why") else ""
            if entry["decision"] == DELETE:
                lines.append(dataclasses.replace(line, client=False, source=DELETED))
                rows.append(_review_row(system.bundle_key, f"{DELETED}: “{line.text}”", entry.get("why", "")))
            else:
                lines.append(dataclasses.replace(line, text=entry["change_to"], source=CHANGED,
                                                 evidence=_squash(f"was “{line.text}”{reason}; {line.evidence}")))
        decide = lambda text, names=names: _match(entries, names, text)    # noqa: E731
        blocks = tuple(b.model_copy(update={"items": _items(b.items, decide)}) for b in system.blocks)
        replaced.update({id(old): new for old, new in zip(system.blocks, blocks)})
        new_systems.append(dataclasses.replace(system, blocks=blocks, lines=tuple(lines), review=tuple(rows)))

    def everywhere(text: str):
        return next((e for _, e in applied if _key(e.get("text")) == _key(text)), None)

    def block(b):
        return replaced.get(id(b)) or b.model_copy(update={"items": _items(b.items, everywhere)})

    job = job.model_copy(update={
        "blocks": tuple(block(b) for b in job.blocks),
        "front_end": job.front_end.model_copy(update={"items": _items(job.front_end.items, everywhere)}),
        "add_alternates": tuple(a.model_copy(update={"blocks": tuple(block(b) for b in a.blocks)})
                                for a in job.add_alternates),
        "deduct_alternates": tuple(a.model_copy(update={"blocks": tuple(block(b) for b in a.blocks)})
                                   for a in job.deduct_alternates)})
    review = [r for r in review if not any(_key(e.get("text")) == _key(r.decide)
                                            and (not e.get("system") or _key(e["system"]) == _key(r.system))
                                            for e in dropped_rows)]
    return job, new_systems, review


def _review_row(system: str, decide: str, why: str):
    from estimator.writing.write import ReviewItem
    return ReviewItem(system, "", decide, why or "no reason given", "Nothing to do: the estimator's correction, "
                      "applied from the corrections file", tier=5)
