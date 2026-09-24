"""The price sheet as rows: what the tool writes, what a sent workup holds, and how to line them up.

The rows are the proposal's own lines, in the proposal's order — the same list
the proposal is written from (the one-list, two-documents guarantee).
"""
from __future__ import annotations

import difflib
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from estimator.proposal import layout
from estimator.scope import Block, Item, Job
from estimator.text import squash, strip_list_marker

FIRST_ROW, LAST_ROW = 6, 396                               # the rows that hold Material Total = Unit Price × Quantity
HOURS, QUANTITY, SCOPE, UNIT_PRICE, FLAG = 2, 3, 4, 5, 8   # columns B, C, D, E, H


@dataclass(frozen=True)
class Row:
    text: str = ""                   # column D; empty for a blank spacer row
    hours: float | None = None       # column B
    quantity: int | None = None      # column C
    item: Item | None = None         # set for a line or note
    ref: tuple = ()                  # where in Job.model_dump() this row came from


@dataclass(frozen=True)
class Sheet:
    title: str
    rows: tuple[Row, ...]


def sheet_plan(job: Job) -> list[Sheet]:
    """Every price sheet the job needs, in workbook order: the base bid, then each alternate."""
    numbers = iter(range(1, 10_000))
    base = [
        Row(layout.scope_heading(job), hours=job.front_end.tech_hours, ref=("front_end",)),
        Row(f"{layout.FRONT_END_HEADING_MARKER}  {layout.front_end_heading(job.front_end)}"),
    ]
    markers = layout.FRONT_END_LINE_MARKER[job.front_end.mode]
    for j, item in enumerate(job.front_end.items):
        base.append(_item_row(item, markers, ("front_end", "items", j)))
    for i, block in enumerate(job.blocks):
        base += _block_rows(block, next(numbers), ("blocks", i))
    sheets = [Sheet("Price Sheet 1", tuple(base))]
    for field, heading, name in (("add_alternates", layout.ADD_ALTERNATES, "ADD ALT"),
                                 ("deduct_alternates", layout.DEDUCT_ALTERNATES, "DEDUCT ALT")):
        alternates = getattr(job, field)
        for k, alternate in enumerate(alternates):
            rows = [Row(heading), Row(), Row(f"{k + 1}.  {alternate.label}")]
            for i, block in enumerate(alternate.blocks):
                rows += _block_rows(block, next(numbers), (field, k, "blocks", i))
            title = name if len(alternates) == 1 else f"{name} {k + 1}"
            sheets.append(Sheet(title, tuple(rows)))
    return sheets


def _block_rows(block: Block, number: int, ref: tuple) -> list[Row]:
    rows = [Row(), Row(f"{number}.  {squash(block.header)}", hours=block.tech_hours, ref=ref)]
    for j, item in enumerate(block.items):
        rows.append(_item_row(item, layout.BLOCK_LINE_MARKER, ref + ("items", j)))
    return rows


def _item_row(item: Item, markers: dict[int, str], ref: tuple) -> Row:
    text = squash(item.text) if item.kind == "note" else f"{markers[item.depth]}  {squash(item.text)}"
    return Row(text, hours=item.tech_hours, quantity=item.quantity, item=item, ref=ref)


@dataclass(frozen=True)
class SentRow:
    sheet: str
    row: int
    text: str
    hours: float | None
    quantity: int | None
    unit_price: float | None
    hours_raw: object = None       # the cell's own value, for reporting what couldn't be used
    quantity_raw: object = None


def _number(value) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _whole(value) -> int | None:
    number = _number(value)
    return int(number) if number is not None and number.is_integer() else None


def read_sent_rows(path: Path) -> dict[str, list[SentRow]]:
    """Every price sheet's scope rows, as written — markers, odd spacing and all."""
    workbook = openpyxl.load_workbook(str(path), data_only=True)
    sheets: dict[str, list[SentRow]] = {}
    for ws in workbook.worksheets:
        if ws.title == "Summary":
            continue
        rows = []
        for r in range(FIRST_ROW, LAST_ROW + 1):
            text = ws.cell(r, SCOPE).value
            if isinstance(text, str) and text.strip():
                hours_cell, quantity_cell = ws.cell(r, HOURS).value, ws.cell(r, QUANTITY).value
                rows.append(SentRow(ws.title, r, text, _number(hours_cell), _whole(quantity_cell),
                                    _number(ws.cell(r, UNIT_PRICE).value), hours_cell, quantity_cell))
        sheets[ws.title] = rows
    return sheets


def _key(text: str) -> str:
    return strip_list_marker(text).casefold()


def align(ours: list[str], theirs: list[str]) -> list[tuple[int, int]]:
    """Pairs (our index, their index) of rows that say the same thing, in order."""
    matcher = difflib.SequenceMatcher(None, [_key(t) for t in ours], [_key(t) for t in theirs], autojunk=False)
    return [(block.a + k, block.b + k) for block in matcher.get_matching_blocks() for k in range(block.size)]
