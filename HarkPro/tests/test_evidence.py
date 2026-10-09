import json
from pathlib import Path

import openpyxl
import pytest

from harkpro import coverage, evidence
from harkpro.doctrine import load_lines

FIX = Path(__file__).parent / "fixtures"


class StubPrices:
    def unit(self, text):
        return 100.0 if "Leak Sensor" in text else None


@pytest.fixture
def built(tmp_path):
    ledger = json.loads((FIX / "ledger.json").read_text())
    job = json.loads((FIX / "job.json").read_text())
    cov = coverage.check(ledger, job, load_lines(FIX / "lines.json"), printed={})
    path = evidence.write(tmp_path / "e.xlsx", job, ledger, cov, printed={}, prices=StubPrices())
    return openpyxl.load_workbook(path), evidence.build(job, ledger, cov, {}, StubPrices())


def test_four_tabs_with_frozen_headers_and_filters(built):
    wb, _ = built
    assert wb.sheetnames == ["Summary", "Scope", "Questions & Flags", "Sources"]
    for name in wb.sheetnames[1:]:
        ws = wb[name]
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
        assert ws.column_dimensions["B"].width > 8


def test_scope_rows_prices_and_marks(built):
    wb, data = built
    ws = wb["Scope"]
    rows = {}
    for r in range(2, ws.max_row + 1):
        text = str(ws.cell(r, 2).value or "")
        if "Leak Sensor" in text: rows["RTU.d2"] = r
        if "Motorized Damper Actuator" in text: rows["RTU.d3"] = r
    leak = rows["RTU.d2"]
    assert ws.cell(leak, 10).value == 100.0 and ws.cell(leak, 11).value == 200.0
    damper = rows["RTU.d3"]
    assert ws.cell(damper, 10).value == evidence.PENDING
    assert ws.cell(damper, 12).value == "*Confirm" and ws.cell(damper, 12).font.color.rgb.endswith("ED7D31")
    notes = [r for r in range(2, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("*Confirm")]
    assert notes and ws.cell(notes[0], 2).font.color.rgb.endswith("ED7D31")
    header = [r for r in range(2, ws.max_row + 1) if str(ws.cell(r, 2).value).startswith("Rooftop Unit")][0]
    assert ws.cell(header, 8).value == 2 and ws.cell(header, 9).value == 4      # block qty, tech hours


def test_mark_colours():
    assert evidence.mark_of("** client note") == ("**", evidence.RED)
    assert evidence.mark_of("*Confirm x") == ("*Confirm", evidence.ORANGE)
    assert evidence.mark_of("*Assumed: x") == ("*Assumed", evidence.BLUE)


def test_questions_carry_coverage_count_gaps_and_contradictions(built):
    wb, data = built
    kinds = [q["kind"] for q in data["questions"]]
    assert any(k.startswith("Count gap") for k in kinds) and "Sources disagree" in kinds and "Line flag: quantity" in kinds
    assert data["questions"][0]["severity"] == "Must fix"
    ws = wb["Questions & Flags"]
    assert ws.cell(2, 2).value == "Must fix" and ws.cell(2, 2).font.color.rgb.endswith("EE0000")


def test_evidence_printed_inline_without_ids(built):
    """The estimator reads the proof on the line itself; no source ids or line ids to chase (user, 2026-10-09)."""
    wb, data = built
    quote = "MOTORIZED DAMPER ACTUATORS SHALL BE PROVIDED BY ATC CONTRACTOR."
    ws = wb["Scope"]
    row = next(r for r in range(2, ws.max_row + 1) if "Motorized Damper Actuator" in str(ws.cell(r, 2).value))
    assert quote in str(ws.cell(row, 13).value) and "M-101" in str(ws.cell(row, 13).value)
    heads = [c.value for c in ws[1]] + [c.value for c in wb["Questions & Flags"][1]] + [c.value for c in wb["Sources"][1]]
    assert not any(h and ("id" in str(h).lower().split()) for h in heads)
    assert not any(h in ("Source ids", "Line id", "Source id") for h in heads)
    src = wb["Sources"]
    keys = [(src.cell(r, 1).value, src.cell(r, 3).value) for r in range(2, src.max_row + 1)]
    assert any(q == quote for _, q in keys)
    used = [str(src.cell(r, 4).value) for r in range(2, src.max_row + 1) if src.cell(r, 3).value == quote]
    assert used and used[0].startswith("#")                # scope line number and text, not "RTU.d3"


def test_summary_totals(built):
    _, data = built
    t = data["summary"]["totals"]["base"]
    assert t["lines"] == 6 and t["hours"] == 13.0 and t["priced"] == 200.0 and t["pending"] == 5
    assert data["summary"]["coverage"]["passed"] is False


def test_debug_workbook(tmp_path):
    ledger = json.loads((FIX / "ledger.json").read_text())
    job = json.loads((FIX / "job.json").read_text())
    path = evidence.write_debug(tmp_path / "d.xlsx", job, ledger, coverage.check(ledger, job, printed={}))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Ledger facts", "Ledger sources", "Coverage"]
    assert wb["Ledger facts"].max_row == len(ledger["facts"]) + 1


def test_printed_notes_replace_clarifications_and_carry_the_why(tmp_path):
    ledger = json.loads((FIX / "ledger.json").read_text())
    job = json.loads((FIX / "job.json").read_text())
    notes = {"RTU": ["*Confirm RTU integration", "*Assumed: one damper actuator per unit, because no tags"],
             "FE": ["*Confirm the existing BMS vendor"]}
    data = evidence.build(job, ledger, None, {}, StubPrices(), notes)
    texts = [(r["text"], r["mark"]) for r in data["scope"] if r["type"] == "note"]
    assert ("*Assumed: one damper actuator per unit, because no tags", "*Assumed") in texts
    assert ("*Confirm the existing BMS vendor", "*Confirm") in texts           # the front end block shows
    flag = next(q for q in data["questions"] if "Motorized Damper Actuator" in q["about"])
    assert "because" not in flag["what"]                                        # the why is not repeated
    wb = openpyxl.load_workbook(evidence.write(tmp_path / "e.xlsx", job, ledger, None, {}, StubPrices(), notes))
    ws = wb["Scope"]
    row = next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, 2).value).startswith("*Assumed"))
    assert ws.cell(row, 2).font.color.rgb.endswith("2F5597")
