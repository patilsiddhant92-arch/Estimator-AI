"""score.py is the run-to-run regression diff (no golden answer keys)."""
import json
from pathlib import Path

import docx
import pytest

from harkpro import score

FIX = Path(__file__).parent / "fixtures"


def proposal(path, blocks, front=("Seamlessly Tie in to the Existing [existing BMS name] BMS",)):
    d = docx.Document()
    d.add_paragraph("Scope and Quotation")
    d.add_paragraph("SCOPE OF WORK – ")
    for line in front:
        d.add_paragraph(line)
    for item in blocks:
        if item == "ADD":
            d.add_paragraph("ADD ALTERNATES:")
            d.add_paragraph("ALT #1: Sample")
            continue
        head, lines = item
        d.add_paragraph(head)
        for line in lines:
            d.add_paragraph(line)
    d.add_paragraph("Clarifications and Inclusions considered:")
    d.add_paragraph("All labor will be performed during normal business working hours.")
    d.save(str(path))
    return path


V1 = [
    ("Qty (2) Rooftop Unit RTU-1, 2", ["Furnish BMS Communication Wiring to RTU Factory Controls",
                                       "Wire Leak Sensor (Furnished by Mechanical Contractor)",
                                       "Furnish and Install Duct Smoke Detector",
                                       "*Confirm: RTU integration"]),
    ("Qty (3) VAV Box VAV-1 thru 3", ["Furnish DDC Controller to the Factory for Installation",
                                      "Furnish and Install Space Temperature Sensor (Qty. 2 Total)"]),
]
V2 = [
    ("Qty (2) Rooftop Unit RTU-1, 2", ["Furnish BMS Communication Wiring to RTU Factory Controls",
                                       "Furnish and Install Leak Sensor", "Furnish and Install Local RAM Module",
                                       "*Confirm: RTU integration priced with AME as the BMS vendor"]),
    ("Qty (3) VAV Box VAV-1 thru 3", ["Furnish DDC Controller to the Factory for Installation",
                                      "Furnish and Install Space Temperature Sensor (Qty. 3 Total)"]),
    ("Qty (1) Exhaust Fan EF-1", ["Furnish and Install Space Thermostat"]),
]


def test_same_file_is_unchanged(tmp_path):
    a = proposal(tmp_path / "a.docx", V2)
    r = score.diff_files(a, a)
    assert r.unchanged and all(b.change == "same" for b in r.blocks)
    assert "No change" in score.markdown(r)


def test_added_removed_changed(tmp_path):
    r = score.diff_files(proposal(tmp_path / "v1.docx", V1), proposal(tmp_path / "v2.docx", V2))
    lines = {(l.old or l.new): l for b in r.blocks for l in b.lines}
    assert lines["Wire Leak Sensor (Furnished by Mechanical Contractor)"].change == "changed"
    assert "verb Wire -> Furnish and Install" in lines["Wire Leak Sensor (Furnished by Mechanical Contractor)"].what
    assert "qty 2 -> 3" in lines["Furnish and Install Space Temperature Sensor (Qty. 2 Total)"].what
    assert lines["Furnish and Install Duct Smoke Detector"].change == "removed"
    assert lines["Furnish and Install Local RAM Module"].change == "added"
    ef = next(b for b in r.blocks if "Exhaust" in b.new)
    assert ef.change == "added"
    rtu = next(b for b in r.blocks if "Rooftop" in b.new)
    assert rtu.notes_added and rtu.notes_removed
    t = r.totals
    assert (t["blocks_added"], t["lines_added"], t["lines_removed"], t["lines_changed"]) == (1, 2, 1, 2)
    md = score.markdown(r)
    assert "ADDED: Qty (1) Exhaust Fan EF-1" in md and "**−** Furnish and Install Duct Smoke Detector" in md


def test_blocks_match_by_tags_and_alternates_and_heading_qty(tmp_path):
    old = proposal(tmp_path / "o.docx", V2[:1] + ["ADD"] + V2[1:2])
    new = proposal(tmp_path / "n.docx", [("Qty (4) Variable Air Volume Box VAV-1, 2, 3", V2[1][1]), V2[0]])
    assert [b.alternate for b in score.read_docx(old)] == ["base", "base", "add"]
    r = score.diff_files(old, new)
    vav = next(b for b in r.blocks if "VAV" in b.old)
    assert vav.change == "changed" and "heading qty 3 -> 4" in vav.what and "alternate add -> base" in vav.what
    assert not [l for l in vav.lines]                          # same lines, just a new heading


def test_job_json_side(tmp_path):
    job = json.loads((FIX / "job.json").read_text())
    blocks = score.read_job(job, printed={})
    rtu = next(b for b in blocks if "RTU1" in b.tags)
    assert rtu.quantity == 2 and len(rtu.lines) == 3
    assert blocks[0].lines[0].text.startswith("Seamlessly Tie in")
    assert not any("EF1" in b.tags for b in blocks)            # a left-out block prints nothing
    r = score.diff_files(FIX / "job.json", FIX / "job.json")
    assert r.unchanged


def test_run_writes_json_and_md(tmp_path):
    a = proposal(tmp_path / "a.docx", V1)
    b = proposal(tmp_path / "b.docx", V2)
    score.run(a, b, tmp_path / "out", "v1-v2")
    data = json.loads((tmp_path / "out" / "v1-v2 diff.json").read_text())
    assert data["unchanged"] is False and data["totals"]["blocks_added"] == 1
    assert (tmp_path / "out" / "v1-v2 diff.md").exists()


def test_no_scope_heading_is_an_error(tmp_path):
    d = docx.Document()
    d.add_paragraph("nothing here")
    d.save(str(tmp_path / "x.docx"))
    with pytest.raises(ValueError):
        score.read_docx(tmp_path / "x.docx")
