"""T-002: the price book read from the revised AME pricing workbook's own layout.

Every workbook, device name and price here is invented. The real workbook's content is checked only for what the
app relies on (its sheets load, every alias points at an item, the app's copy is the estimator's file).
"""
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from estimator.price_sheet.price_book import BookItem, PriceBook

ROOT = Path(__file__).resolve().parents[1]


def _book(path: Path, parts=(), panels=(), valves=(), side=(), skip=()) -> Path:
    """A workbook in the revised layout: name in column B, price in column C (a note in column A); valves may carry a
    second table in columns H (name) and I (price). `skip` leaves a sheet out."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Bid Overview")["B2"] = "Bid Summary / Overview:"
    for title, rows in (("Std Parts Pricing", parts), ("BMS-DDC-Panels", panels), ("Control Valves & Meters", valves)):
        if title in skip:
            continue
        ws = wb.create_sheet(title)
        ws["B2"] = "Estimation Price"
        for r, row in enumerate(rows, start=3):
            for col, value in zip("ABC", row):
                if value is not None:
                    ws[f"{col}{r}"] = value
        if title == "Control Valves & Meters":
            for r, (name, price) in enumerate(side, start=3):
                ws[f"H{r}"] = name
                if price is not None:
                    ws[f"I{r}"] = price
    wb.save(path)
    return path


def _aliases(path: Path, aliases: dict) -> Path:
    path.write_text(json.dumps({"about": "test", "aliases": aliases}), encoding="utf-8")
    return path


# --- a price is found by name and by alias -------------------------------------------------------------------------

def test_a_device_is_priced_by_its_book_name(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Widget Sensor", 77)]),
                          _aliases(tmp_path / "a.json", {}))
    assert book.lookup("o  Furnish and Install Widget Sensor (Qty. 4 Total)").price == 77


def test_a_device_is_priced_by_a_confirmed_alias(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Gizmo Thermostat for Boxes", 120)]),
                          _aliases(tmp_path / "a.json", {"Box Room Sensor": "Gizmo Thermostat for Boxes"}))
    match = book.lookup("Furnish and Install Box Room Sensor (Qty. 9 Total)")
    assert (match.name, match.price) == ("Gizmo Thermostat for Boxes", 120)


# --- a device with no price is reported, never priced at 0 ---------------------------------------------------------

def test_a_device_the_book_does_not_have_is_reported(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Widget Sensor", 77)]),
                          _aliases(tmp_path / "a.json", {}))
    lines = ["Furnish and Install Widget Sensor", "Furnish and Install Sprocket Relay (Qty. 2 Total)"]
    assert book.lookup(lines[1]) is None
    assert book.unpriced(lines) == ["Furnish and Install Sprocket Relay (Qty. 2 Total)"]


def test_a_book_row_with_no_price_or_a_zero_price_is_not_an_item(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Empty Gadget", None), (None, "Free Gadget", 0),
                                                             (None, "Odd Gadget", "TBD")]),
                          _aliases(tmp_path / "a.json", {}))
    assert book.items == []
    assert book.unpriced(["Furnish Empty Gadget", "Furnish Free Gadget", "Furnish Odd Gadget"]) == [
        "Furnish Empty Gadget", "Furnish Free Gadget", "Furnish Odd Gadget"]


def test_an_alias_to_a_row_with_no_price_is_refused(tmp_path):
    with pytest.raises(ValueError, match="isn't in the price book"):
        PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Free Gadget", 0)]),
                       _aliases(tmp_path / "a.json", {"Gadget": "Free Gadget"}))


def test_the_price_sheet_leaves_an_unpriced_line_empty_and_flags_it(small_job, workup_template, tmp_path):
    from estimator.price_sheet.write import PRICE_PENDING, write_price_sheet
    from estimator.text import strip_list_marker
    book = PriceBook([BookItem("BMS Server", 1111, "invented", 1, False)], {})
    out = write_price_sheet(small_job, workup_template, book, tmp_path / "sheet.xlsx")
    ws = openpyxl.load_workbook(str(out))["Price Sheet 1"]
    priced = {strip_list_marker(ws.cell(r, 4).value): (ws.cell(r, 5).value, ws.cell(r, 8).value)
              for r in range(6, 60) if ws.cell(r, 3).value is not None}
    assert priced["Furnish BMS Server"] == (1111, None)
    others = [v for k, v in priced.items() if k != "Furnish BMS Server"]
    assert others and all(value == (None, PRICE_PENDING) for value in others)   # never 0, always flagged


# --- the revised workbook's layout ---------------------------------------------------------------------------------

def test_all_three_price_sheets_are_read_and_the_bid_sheets_are_not(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Widget Sensor", 77)],
                                panels=[(None, "Widget Controller", 900)],
                                valves=[(None, "Widget Meter - 2\"", 3000)]), _aliases(tmp_path / "a.json", {}))
    assert {(i.name, i.sheet) for i in book.items} == {("Widget Sensor", "Std Parts Pricing"),
                                                       ("Widget Controller", "BMS-DDC-Panels"),
                                                       ("Widget Meter - 2\"", "Control Valves & Meters")}


def test_a_second_price_on_a_row_with_no_name_is_not_an_item(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[(None, "Widget Actuator", 300), (None, None, 250)]),
                          _aliases(tmp_path / "a.json", {}))
    assert [(i.name, i.price) for i in book.items] == [("Widget Actuator", 300)]


def test_a_valve_size_takes_its_heading(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", valves=[(None, "Gizmo Valves", None), (None, '1/2"', 210),
                                                             (None, '1-1/4 & 1-1/2"', 420), (None, "Sprocket Meters", None),
                                                             (None, 'Sprocket Meter - 4"', 4000)]),
                          _aliases(tmp_path / "a.json", {}))
    assert [(i.name, i.price) for i in book.items] == [('Gizmo Valves 1/2"', 210), ('Gizmo Valves 1-1/4 & 1-1/2"', 420),
                                                       ('Sprocket Meter - 4"', 4000)]


def test_a_variant_row_takes_the_item_above(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", valves=[(None, '4" Gizmo Valve', 1300), (None, "High Performance", 2600)]),
                          _aliases(tmp_path / "a.json", {}))
    assert [i.name for i in book.items] == ['4" Gizmo Valve', '4" Gizmo Valve - High Performance']


def test_the_valve_sheets_second_table_is_read(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", valves=[(None, '4" Gizmo Valve', 1300)],
                                side=[("Other Valves", None), ('6" Sprocket Valve', 1900), (None, 0)]),
                          _aliases(tmp_path / "a.json", {}))
    assert {i.name: i.price for i in book.items} == {'4" Gizmo Valve': 1300, '6" Sprocket Valve': 1900}


def test_rows_under_the_books_not_completed_note_need_review(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", panels=[(None, "Widget Controller", 900),
                                                             (None, "Below is not completed - Need more information", None),
                                                             (None, "Plant Widget with 32 UIO", 2900)]),
                          _aliases(tmp_path / "a.json", {}))
    assert {i.name: i.needs_review for i in book.items} == {"Widget Controller": False, "Plant Widget with 32 UIO": True}


def test_a_need_note_in_column_a_still_marks_the_price(tmp_path):
    book = PriceBook.load(_book(tmp_path / "b.xlsx", parts=[("Need review from Engg Team", "Widget Switch", 50)]),
                          _aliases(tmp_path / "a.json", {}))
    assert book.lookup("Furnish Widget Switch").needs_review is True


def test_a_workbook_without_a_price_sheet_is_refused_in_plain_words(tmp_path):
    with pytest.raises(ValueError, match="Std Parts Pricing"):
        PriceBook.load(_book(tmp_path / "b.xlsx", skip=("Std Parts Pricing",)), _aliases(tmp_path / "a.json", {}))


def test_an_alias_that_would_hide_a_different_book_item_is_refused():
    items = [BookItem("Widget Thermostat", 355, "a", 1, False), BookItem("Widget Thermostat (for Heaters)", 255, "a", 2, False)]
    with pytest.raises(ValueError, match="is itself a price book item"):
        PriceBook(items, {"Widget Thermostat": "Widget Thermostat (for Heaters)"})


# --- the real workbook and alias file ------------------------------------------------------------------------------

def test_the_apps_copy_is_the_estimators_revised_workbook():
    """One source of truth: assets/price_book holds Brain's workbook byte for byte (tools/update_price_book.py)."""
    brain = ROOT / "Brain" / "AME Std Pricing Spreadsheet.xlsx.xlsx"
    app = ROOT / "assets" / "price_book" / "AME Std Pricing Spreadsheet.xlsx"
    assert app.read_bytes() == brain.read_bytes(), "run: python tools/update_price_book.py"


def test_every_alias_points_at_a_priced_item(price_book):
    aliases = json.loads((ROOT / "assets" / "price_book" / "aliases.json").read_text(encoding="utf-8"))["aliases"]
    names = {i.name for i in price_book.items}
    assert aliases and all(target in names for target in aliases.values())


# --- the tool ------------------------------------------------------------------------------------------------------

def _tool(*args):
    return subprocess.run([sys.executable, str(ROOT / "tools" / "update_price_book.py"), *map(str, args)],
                          capture_output=True, text=True, cwd=ROOT)


def test_the_tool_copies_the_workbook_and_check_says_when_it_is_stale(tmp_path):
    source = _book(tmp_path / "source.xlsx", parts=[(None, "Widget Sensor", 77)])
    target = tmp_path / "app" / "book.xlsx"
    aliases = _aliases(tmp_path / "a.json", {"Widget Probe": "Widget Sensor"})
    common = ("--source", source, "--target", target, "--aliases", aliases)
    stale = _tool("--check", *common)
    assert stale.returncode == 1 and "out of date" in stale.stdout + stale.stderr
    done = _tool(*common)
    assert done.returncode == 0, done.stdout + done.stderr
    assert target.read_bytes() == source.read_bytes()
    assert "1 priced item" in done.stdout
    assert _tool("--check", *common).returncode == 0


def test_the_tool_refuses_a_workbook_that_breaks_an_alias_and_leaves_the_app_copy(tmp_path):
    source = _book(tmp_path / "source.xlsx", parts=[(None, "Widget Sensor", 77)])
    target = tmp_path / "book.xlsx"
    target.write_bytes(b"old")
    aliases = _aliases(tmp_path / "a.json", {"Widget Probe": "Renamed Widget Sensor"})
    done = _tool("--source", source, "--target", target, "--aliases", aliases)
    assert done.returncode == 1 and "Renamed Widget Sensor" in done.stdout + done.stderr
    assert target.read_bytes() == b"old"
