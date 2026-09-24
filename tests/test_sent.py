import openpyxl

from conftest import PROPOSAL_TEMPLATE, SENT_PROPOSAL, SENT_WORKUP, TEST_JOB_FILE
from estimator.price_sheet.rows import FIRST_ROW, HOURS, LAST_ROW, SentRow, sheet_plan
from estimator.price_sheet.write import write_price_sheet
from estimator.scope import Job
from estimator.sent import import_matches, job_from_sent


def _imported():
    return job_from_sent(SENT_PROPOSAL, SENT_WORKUP, PROPOSAL_TEMPLATE)


def test_hours_come_from_the_workup():
    job, _ = _imported()
    assert job.front_end.tech_hours == 140.0
    assert (job.blocks[0].header, job.blocks[0].tech_hours) == ("*Qty (1) BACnet Data Cabinet", 10.0)
    assert job.blocks[1].tech_hours == 25.0


def test_quantities_come_from_the_workup():
    job, _ = _imported()
    assert [i.quantity for i in job.front_end.items] == [1, 1, 1, 1, 1, 1, None, None, None, None, None, None]
    assert [i.quantity for i in job.blocks[0].items] == [1, 1, 1, None, None, None, None, None]
    router_cabinet = job.blocks[1].items
    assert (router_cabinet[0].quantity, router_cabinet[1].quantity) == (5, 7)
    assert router_cabinet[1].text.startswith("Furnish and Install BACnet IP Router")


def test_alternates_get_their_hours_and_quantities():
    job, _ = _imported()
    add, deduct = job.add_alternates[0].blocks[0], job.deduct_alternates[0].blocks[0]
    assert (add.tech_hours, [i.quantity for i in add.items]) == (4.0, [1, 1])
    assert (deduct.tech_hours, [i.quantity for i in deduct.items]) == (4.0, [1, 1, None])


def test_every_line_matches_and_nothing_is_dropped():
    _, report = _imported()
    assert report.matched == 273
    assert report.unmatched_ours == ()
    assert report.unmatched_theirs == ()
    assert report.ignored == ()


def test_the_committed_test_job_is_what_the_importer_makes():
    assert Job.model_validate_json(TEST_JOB_FILE.read_text(encoding="utf-8")) == _imported()[0]


def _column_b_total(path) -> dict[str, float]:
    workbook = openpyxl.load_workbook(str(path), data_only=False)
    totals = {}
    for sheet_name in ("Price Sheet 1", "ADD ALT", "DEDUCT ALT"):
        ws = workbook[sheet_name]
        totals[sheet_name] = sum(
            ws.cell(r, HOURS).value for r in range(FIRST_ROW, LAST_ROW + 1)
            if isinstance(ws.cell(r, HOURS).value, (int, float)) and not isinstance(ws.cell(r, HOURS).value, bool))
    return totals


def test_the_generated_price_sheet_carries_every_sent_hour(tmp_path, workup_template, price_book):
    """Every tech hour in the sent workup — on the SCOPE OF WORK row, block headers,
    and lines or notes alike — must reach the generated price sheet. Column B summed
    over rows 6-396 is the template's own B401 = SUM(B4:B400)."""
    job, _ = _imported()
    sheet = write_price_sheet(job, workup_template, price_book, tmp_path / "sheet.xlsx")
    assert _column_b_total(sheet) == {"Price Sheet 1": 942.5, "ADD ALT": 4, "DEDUCT ALT": 4}
    assert _column_b_total(sheet) == _column_b_total(SENT_WORKUP)


# --- nothing that lands nowhere is ever dropped silently -------------------------------------

def _number(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _whole(value):
    number = _number(value)
    return int(number) if number is not None and number.is_integer() else None


def _fake_sent(small_job, **raw) -> dict[str, list[SentRow]]:
    """A workup, word for word the same as small_job's own Price Sheet 1, with one row's
    hours/quantity cell replaced by whatever raw value the test wants to check is reported —
    coerced the same way read_sent_rows coerces a real cell."""
    ours = [r.text for r in sheet_plan(small_job)[0].rows if r.text]
    rows = [SentRow("Price Sheet 1", 6 + i, text, None, None, None, None, None) for i, text in enumerate(ours)]
    for text, cell in raw.items():
        index = ours.index(text)
        kind, value = cell
        hours_raw = value if kind == "hours" else None
        quantity_raw = value if kind == "quantity" else None
        rows[index] = SentRow("Price Sheet 1", 6 + index, text, _number(hours_raw), _whole(quantity_raw),
                              None, hours_raw, quantity_raw)
    return {"Price Sheet 1": rows}


def test_a_fractional_quantity_is_reported_not_dropped(small_job, without_prices):
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"o  Furnish and Install DDC Controller": ("quantity", 2.5)})
    job, report = import_matches(blank, sent)
    assert job.blocks[0].items[0].quantity is None
    assert report.ignored == ("Price Sheet 1 row 12: quantity 2.5 is not a whole number",)


def test_a_quantity_typed_as_text_is_reported_not_dropped(small_job, without_prices):
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"o  Furnish and Install DDC Controller": ("quantity", "two")})
    job, report = import_matches(blank, sent)
    assert job.blocks[0].items[0].quantity is None
    assert report.ignored == ("Price Sheet 1 row 12: quantity 'two' is not a number",)


def test_a_quantity_on_a_note_is_reported_not_dropped(small_job, without_prices):
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"*DAMPERS BY OTHERS": ("quantity", 1)})
    job, report = import_matches(blank, sent)
    assert report.ignored == ("Price Sheet 1 row 14: quantity 1 is on a note, which has no quantity",)


def test_a_quantity_on_a_block_header_is_reported_not_dropped(small_job, without_prices):
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"1.  Qty (2) Exhaust Fan EF-1, 2": ("quantity", 3)})
    job, report = import_matches(blank, sent)
    assert job.blocks[0].tech_hours is None
    assert report.ignored == ("Price Sheet 1 row 11: quantity 3 is on a header, which has no quantity",)


def test_hours_typed_as_text_are_reported_not_dropped(small_job, without_prices):
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"o  Furnish and Install DDC Controller": ("hours", "n/a")})
    job, report = import_matches(blank, sent)
    assert job.blocks[0].items[0].tech_hours is None
    assert report.ignored == ("Price Sheet 1 row 12: hours 'n/a' is not a number",)


def test_hours_on_a_note_land_on_the_note(small_job, without_prices):
    """C1: hours are allowed on notes, not only on headers and lines."""
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"*DAMPERS BY OTHERS": ("hours", 3.5)})
    job, report = import_matches(blank, sent)
    assert job.blocks[0].items[2].tech_hours == 3.5
    assert report.ignored == ()


def test_a_cell_on_a_row_the_job_has_no_place_for_is_reported_not_dropped(small_job, without_prices):
    """The "* Front End..." heading row (Row.ref == ()) is matched but carries no data of its own."""
    blank = without_prices(small_job)
    sent = _fake_sent(blank, **{"•  Front End Web Enabled BMS System": ("hours", 99)})
    job, report = import_matches(blank, sent)
    assert report.ignored == ("Price Sheet 1 row 7: hours 99 on a line the job has no place for",)
