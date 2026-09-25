"""The scope review sections of the Spec overview: potential scope AME needs to carry. Every case
here is a synthetic sentence; the words come from assets/reading_lists/scope_review.json."""
from estimator.catalogue.overview import (FOUND, NOT_FOUND, RULED_OUT, STANDARD_CLARIFICATION, STANDARD_EXCLUSION,
                                          build_overview, entries_without_a_template_line,
                                          template_lines_without_an_entry)
from estimator.catalogue.register import PageRecord
from estimator.catalogue.sheet_types import Sheet
from estimator.catalogue.spec_sections import SpecSection
from estimator.proposal.read import read_standard_lines
from conftest import PROPOSAL_TEMPLATE

DOC = "2. Takeoffs/1. Base Takeoff/Specs.pdf"
DRAWINGS = "2. Takeoffs/1. Base Takeoff/Drawings.pdf"


def _pages(*texts: str) -> list[PageRecord]:
    return [PageRecord(DOC, n, "text", t, 0) for n, t in enumerate(texts, start=1)]


def _controls_spec(*texts: str):
    """Each text on its own page of a 25 05 00 controls section."""
    return build_overview(_pages(*texts), [], (SpecSection("25 05 00", "INTEGRATED AUTOMATION", DOC, 1, len(texts)),))


def _line(overview, starts: str):
    return next((f for f in overview.standard_lines if f.item.startswith(starts)), None)


# --- the template's lines -----------------------------------------------------------------

def test_the_template_lines_are_read_from_the_proposal_template():
    clarifications, exclusions = read_standard_lines(PROPOSAL_TEMPLATE)
    assert clarifications[0].startswith("All labor will be performed during normal business working hours")
    assert "All water flow meters." in exclusions
    assert "All electric power meters." in exclusions
    assert not any(line.startswith("Exceptions and Exclusions") for line in clarifications + exclusions)


def test_every_template_line_has_an_entry_in_the_reading_list():
    assert template_lines_without_an_entry() == ()


def test_every_reading_list_entry_matches_a_template_line():
    assert entries_without_a_template_line() == ()


# --- the status of a row --------------------------------------------------------------------

def test_bacnet_ip_not_allowed_reads_ruled_out():
    overview = _controls_spec("Controllers shall be integrated via Modbus TCP. BACnet IP is not allowed anywhere "
                              "in the building.")
    bip = next(t for t in overview.topics if t.item == "BACnet/IP")
    assert bip.status == RULED_OUT
    assert bip.quote == "BACnet IP is not allowed anywhere in the building."
    assert bip.places and bip.places[0].page == 1
    assert not any(line.startswith("BACnet/IP") for line in overview.not_found)


def test_a_sentence_that_asks_for_it_reads_found():
    overview = _controls_spec("Provide BACnet/IP to each network controller.")
    assert next(t for t in overview.topics if t.item == "BACnet/IP").status == FOUND


def test_the_strongest_place_decides_ruled_out_over_a_weaker_mention():
    specs = PageRecord(DOC, 1, "text", "BACnet IP is not allowed anywhere in the building.", 0)
    drawing = PageRecord(DRAWINGS, 1, "text", "GATEWAY SHALL BE PROVIDED FOR BACNET IP INTEGRATION", 0)
    sheets = [Sheet(DRAWINGS, 1, "M-601", "SCHEDULES", "Mechanical", ())]
    overview = build_overview([specs, drawing], sheets, (SpecSection("25 05 00", "AUTOMATION", DOC, 1, 1),))
    bip = next(t for t in overview.topics if t.item == "BACnet/IP")
    assert bip.status == RULED_OUT and len(bip.places) == 2


def test_found_somewhere_on_the_strongest_level_is_found():
    overview = _controls_spec("BACnet IP is not allowed on the field bus.", "Provide a BACnet/IP router to the owner.")
    assert next(t for t in overview.topics if t.item == "BACnet/IP").status == FOUND


def test_not_found_has_no_places_and_no_quote():
    overview = _controls_spec("Nothing of interest here.")
    lon = next(t for t in overview.topics if t.item == "LonWorks")
    assert (lon.status, lon.places, lon.quote) == (NOT_FOUND, (), "")


# --- A. AME normally excludes this — the job asks for it --------------------------------------

def test_the_controls_contractor_furnishing_flow_meters_is_an_exclusion_hit():
    overview = _controls_spec("The BMS contractor shall furnish and install all chilled water flow meters.")
    found = _line(overview, "All water flow meters")
    assert found is not None and found.status == FOUND and found.section == STANDARD_EXCLUSION
    assert found.quote == "The BMS contractor shall furnish and install all chilled water flow meters."
    assert _line(overview, "The installation of dampers, valves, wells, water flow meters") is not None


def test_flow_meters_by_others_is_not_an_exclusion_hit():
    overview = _controls_spec("Flow meters by others.", "Flow meters shall be furnished and installed by others.",
                              "The Owner will provide the flow meters.",
                              "Flow meters shall be provided by the mechanical contractor.")
    assert _line(overview, "All water flow meters") is None


def test_valves_furnished_by_bms_but_installed_by_mechanical_is_not_an_installation_hit():
    overview = _controls_spec("Control valves shall be furnished by the BMS contractor and installed by the "
                              "mechanical contractor.")
    assert _line(overview, "The installation of control valves") is None
    assert _line(overview, "The installation of dampers, valves") is None


def test_an_exclusion_the_job_forbids_is_not_listed():
    overview = _controls_spec("The BMS contractor shall not provide power meters.")
    assert _line(overview, "All electric power meters") is None


def test_power_quality_meters_in_a_list_after_a_colon_are_a_hit():
    overview = _controls_spec("B. Power Quality Meters (PQM):\n1. Shall be installed at all transformer inlets.")
    found = _line(overview, "All electric power meters")
    assert found is not None
    assert "Shall be installed at all transformer inlets" in found.quote


def test_blue_bms_conduit_is_asked_for():
    overview = _controls_spec("All BMS conduit shall be of blue color.")
    assert _line(overview, "EMT or RGS conduit (if specified)") is not None


def test_gateways_and_hydrogen_detection():
    overview = _controls_spec("Provide all necessary protocol converters, gateways, or other interface devices.",
                              "Hydrogen detection panel shall be provided with a hardwired connection to the BMS.")
    assert _line(overview, "All equipment manufacturer integration/communication devices") is not None
    assert _line(overview, "Any gas detection system or components") is not None


def test_a_mechanical_spec_hit_needs_the_controls_contractor_named():
    mech = "2. Takeoffs/1. Base Takeoff/Mech Specs.pdf"
    pages = [PageRecord(mech, 1, "text", "Provide chilled water flow meters at each chiller.", 0),
             PageRecord(mech, 2, "text", "The BMS contractor shall provide flow meters at each chiller.", 0)]
    sections = (SpecSection("23 21 13", "HYDRONIC PIPING", mech, 1, 2),)
    found = _line(build_overview(pages, [], sections), "All water flow meters")
    assert found is not None and [p.page for p in found.places] == [2]
    assert found.places[0].strength == "Mechanical spec"


def test_a_clarification_is_listed_when_the_job_speaks_to_the_same_subject():
    overview = _controls_spec("All work in the data hall shall be performed after hours and on weekends.")
    found = _line(overview, "All labor will be performed during normal business working hours")
    assert found is not None and found.section == STANDARD_CLARIFICATION and found.status == FOUND


def test_a_clarification_the_job_forbids_reads_ruled_out():
    overview = _controls_spec("Plenum rated cable is not permitted in this building.")
    found = _line(overview, "All unexposed low voltage wiring outside of the mechanical equipment rooms")
    assert found is not None and found.status == RULED_OUT


def test_a_clarification_near_miss():
    overview = _controls_spec("The contractor shall coordinate with the architect.")
    assert overview.standard_lines == ()


def test_a_page_outside_the_places_searched_finds_nothing():
    pages = _pages("The BMS contractor shall furnish all water flow meters.")
    assert build_overview(pages, [], ()).standard_lines == ()


def test_a_title_block_issued_for_permit_is_not_a_permit_fee():
    overview = _controls_spec("ISSUE FOR PERMIT 1 03.13.2026 DESIGN DEVELOPMENT 1 Furnish and Install Controls",
                              "The BMS contractor shall obtain all required permits and pay all permit fees.")
    found = _line(overview, "Any Taxes and/or Permit Fees")
    assert found is not None and [p.page for p in found.places] == [2]


def test_patch_leads_are_not_patching():
    overview = _controls_spec("Provide project-specific installation kit (CAT6 patch leads, mounting hardware).")
    assert _line(overview, "All cutting, patching, and/or painting") is None


# --- B. Who does what ---------------------------------------------------------------------------

from estimator.catalogue.overview import DIVISION_1, MATRIX_ROW, SPLIT_ROW  # noqa: E402


def _who(overview, says: str):
    return [f for f in overview.who_does_what if says.lower() in f.quote.lower()]


def test_a_sentence_giving_work_to_the_owner_is_a_who_does_what_row():
    overview = _controls_spec("Server hardware by Owner.")
    [row] = _who(overview, "Server hardware by Owner")
    assert row.status == FOUND and row.item.lower() == "by owner"


def test_the_owner_provided_list_is_quoted_with_its_items():
    overview = _controls_spec("2.1 OWNER PROVIDED PRODUCTS:\nA. The Owner will provide, install, and configure the "
                              "following products:\n1. Servers\n2. Switches\n3. Licensing for Owner Infrastructure")
    rows = _who(overview, "The Owner will provide")
    assert rows and "Servers" in rows[0].quote and "Switches" in rows[0].quote


def test_work_given_to_the_controls_contractor_is_not_a_who_does_what_row():
    overview = _controls_spec("All sensors shall be provided by the controls contractor.")
    assert overview.who_does_what == ()


def test_a_mechanical_spec_by_others_needs_a_controls_word():
    mech = "2. Takeoffs/1. Base Takeoff/Mech Specs.pdf"
    pages = [PageRecord(mech, 1, "text", "Access doors furnished by others.", 0),
             PageRecord(mech, 2, "text", "BMS panels furnished by others.", 0)]
    overview = build_overview(pages, [], (SpecSection("23 05 00", "COMMON WORK RESULTS", mech, 1, 2),))
    assert [f.quote for f in overview.who_does_what] == ["BMS panels furnished by others."]


def test_one_row_per_unique_sentence_with_every_place():
    overview = _controls_spec("IDF cabinet by others.", "IDF cabinet by others.")
    [row] = _who(overview, "IDF cabinet")
    assert [p.page for p in row.places] == [1, 2]


def test_more_than_sixty_sentences_are_counted_not_listed():
    overview = _controls_spec(*[f"Panel {n} by others." for n in range(65)])
    assert len([f for f in overview.who_does_what if f.item not in (SPLIT_ROW, MATRIX_ROW)]) == 60
    assert overview.who_does_what_more == 5


def test_a_responsibility_matrix_names_the_sheet():
    page = PageRecord(DRAWINGS, 1, "text", "SCOPE DESCRIPTION QTS\nBMS CONTRACTOR RESPONSIBILITY MATRIX\nNOTES", 0)
    sheets = [Sheet(DRAWINGS, 1, "IC00-00", "CONTROLS LEGEND", "Controls", ())]
    overview = build_overview([page], sheets, ())
    [row] = [f for f in overview.who_does_what if f.item == MATRIX_ROW]
    assert row.places[0].sheet == "IC00-00" and "RESPONSIBILITY MATRIX" in row.quote


def test_several_controls_contractors_are_one_row():
    overview = _controls_spec("These Division 25 sections apply to the BMS as provided by the Design Systems "
                              "Integrator (DSI), SCADA System Integrator (SSI), Panel System Integrator (PSI).",
                              "The CSI shall install all devices.")
    [row] = [f for f in overview.who_does_what if f.item == SPLIT_ROW]
    assert SPLIT_ROW == "The job splits controls between several contractors — say which part AME is pricing"
    assert row.quote.startswith("These Division 25 sections apply to the BMS as provided by the Design Systems")


def test_pressure_in_psi_and_the_csi_format_are_not_several_contractors():
    overview = _controls_spec("Pressure sensors shall be rated for 150 PSI.", "Submittals shall follow CSI format.")
    assert not [f for f in overview.who_does_what if f.item == SPLIT_ROW]


# --- C. Other systems to connect to -----------------------------------------------------------------

def _system(overview, item: str):
    return next((f for f in overview.integrations if f.item.startswith(item)), None)


def test_a_system_with_an_integration_word_and_its_protocol():
    overview = _controls_spec("The UPS shall be monitored via Modbus TCP.",
                              "Generator run status shall be hardwired to the BMS.",
                              "Chillers shall be integrated via BACnet.",
                              "Leak detection panels shall report alarm points to the BMS.",
                              "Monitor the switchgear breakers.")
    assert _system(overview, "UPS").protocol == "Modbus TCP"
    assert _system(overview, "Generators").protocol == "hardwired"
    assert _system(overview, "Chillers").protocol == "BACnet (IP or MS/TP not stated)"
    assert _system(overview, "Leak detection").status == FOUND
    assert _system(overview, "Switchgear").protocol == "not stated"


def test_a_system_without_an_integration_word_is_not_listed():
    overview = _controls_spec("Provide a generator for emergency power.")
    assert _system(overview, "Generators") is None
    assert not any(line.startswith("Generators") for line in overview.not_found)


def test_a_start_up_is_not_a_ups():
    overview = _controls_spec("Monitor all equipment start-ups.")
    assert _system(overview, "UPS") is None


# --- D. Services that cost hours -------------------------------------------------------------------

def _topic(rows, item: str):
    return next(f for f in rows if f.item == item)


def test_training_hours_and_warranty_years_are_captured():
    overview = _controls_spec("Provide eight (8) hours of on-site training for the owner.",
                              "The warranty period shall be two (2) years from acceptance.",
                              "All Software and Licenses required by the CSI shall be provided by the CSI.")
    training = _topic(overview.services, "Training")
    assert training.status == FOUND and training.amount == "eight (8) hours"
    assert _topic(overview.services, "Warranty").amount == "two (2) years"
    assert _topic(overview.services, "Software licences").status == FOUND


def test_a_service_not_asked_is_listed_as_not_found():
    overview = _controls_spec("Provide 40 hours of training.")
    assert _topic(overview.services, "Training").amount == "40 hours"
    assert _topic(overview.services, "Seasonal testing").status == NOT_FOUND
    assert any(line.startswith("Seasonal testing — not found") for line in overview.not_found)


def test_mechanical_training_without_the_controls_contractor_is_not_found():
    mech = "2. Takeoffs/1. Base Takeoff/Mech Specs.pdf"
    pages = [PageRecord(mech, 1, "text", "Provide training on the chiller for the owner.", 0)]
    overview = build_overview(pages, [], (SpecSection("23 64 00", "CHILLERS", mech, 1, 1),))
    assert _topic(overview.services, "Training").status == NOT_FOUND


# --- E. Special systems ------------------------------------------------------------------------------

def test_special_systems_are_found():
    overview = _controls_spec("Control panels shall be built to UL 508A.", "1.10 CYBERSECURITY REQUIREMENTS",
                              "Graphics shall be hosted on the SCADA servers.",
                              "Provide redundant controllers on a device level ring.")
    for item in ("Panels built to UL 508A", "Cybersecurity", "SCADA (supervisory control and data acquisition)",
                 "Redundant controllers or a ring network"):
        assert _topic(overview.special_systems, item).status == FOUND, item


def test_a_special_system_near_miss():
    overview = _controls_spec("Panels shall be UL listed.")
    assert _topic(overview.special_systems, "Panels built to UL 508A").status == NOT_FOUND


# --- F. Project conditions that change the price ---------------------------------------------------------

def _division_1(*texts: str):
    return build_overview(_pages(*texts), [], (SpecSection("01 11 00", "SUMMARY", DOC, 1, len(texts)),))


def test_division_1_is_searched_for_project_conditions():
    overview = _division_1("The Contractor shall perform all work in occupied areas during off-hours.",
                           "Prevailing wage rates apply to all subcontractors.")
    hours = _topic(overview.conditions, "Working hours: off-hours, night or weekend work")
    assert hours.status == FOUND and hours.places[0].strength == DIVISION_1
    assert _topic(overview.conditions, "Prevailing wage").status == FOUND


def test_a_division_1_condition_needs_a_contractor_word():
    overview = _division_1("Normal working hours are 7 AM to 3 PM.")
    assert _topic(overview.conditions, "Working hours: off-hours, night or weekend work").status == NOT_FOUND


def test_division_1_is_searched_only_for_project_conditions():
    overview = _division_1("The BMS contractor shall furnish all water flow meters and provide training. "
                           "Provide BACnet MS/TP.")
    assert overview.standard_lines == ()
    assert _topic(overview.services, "Training").status == NOT_FOUND
    assert next(t for t in overview.topics if t.item == "BACnet MS/TP").status == NOT_FOUND
    assert any("general requirements (Division 1)" in line for line in overview.not_found)


# --- the tab and the summary ---------------------------------------------------------------------------

import re  # noqa: E402
from dataclasses import replace  # noqa: E402

import openpyxl  # noqa: E402

from estimator.catalogue.job import summary_lines  # noqa: E402
from estimator.evidence.workbook import write_evidence  # noqa: E402
from test_evidence import _catalogue  # noqa: E402

HEADINGS = ["1. Controls platform", "2. AME normally excludes this — the job asks for it", "3. Who does what",
            "4. Other systems to connect to", "5. Services that cost hours", "6. Special systems",
            "7. Project conditions that change the price", "8. Protocols", "9. Network and front end",
            "10. Electrical and conduit", "11. Manufacturers named on this job", "12. Searched for but not found"]


def _scope_catalogue():
    overview = _controls_spec("The BMS contractor shall furnish and install all water flow meters.",
                              "The UPS shall be monitored via Modbus TCP.", "IDF cabinet by others.",
                              "Provide 40 hours of training.", "Panels shall be built to UL 508A.")
    return replace(_catalogue(), overview=overview)


def test_the_tab_has_the_twelve_sections_in_order(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _scope_catalogue()))["Spec overview"]
    headings = [r[0].value for r in ws.iter_rows() if isinstance(r[0].value, str) and re.match(r"\d+\. ", r[0].value)]
    assert headings == HEADINGS


def test_the_tab_rows_for_each_new_section(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _scope_catalogue()))["Spec overview"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert ["System", "Status", "Protocol named in the same sentence", "Where", "How strong",
            "What the job says (word for word)"] in [r[:6] for r in rows]
    ups = next(r for r in rows if str(r[0]).startswith("UPS"))
    assert ups[1:3] == ["Found", "Modbus TCP"] and ups[5] == "The UPS shall be monitored via Modbus TCP."
    training = next(r for r in rows if r[0] == "Training")
    assert training[1:3] == ["Found", "40 hours"]
    idf = next(r for r in rows if r[0] == "by others")
    assert idf[1] == "Found" and idf[4] == "IDF cabinet by others."
    ul = next(r for r in rows if r[0] == "Panels built to UL 508A")
    assert ul[1] == "Found"
    seasonal = next(r for r in rows if r[0] == "Seasonal testing")
    assert seasonal[1] == "Not found"


def test_the_summary_has_a_scope_review_line():
    lines = summary_lines(_scope_catalogue())
    index = next(i for i, line in enumerate(lines) if line.startswith("Spec overview:"))
    # "furnish and install all water flow meters" is both "All water flow meters" and "The installation of … water
    # flow meters".
    assert lines[index + 1] == ("Scope review: 2 standard exclusions the job asks for; 1 other systems to connect to; "
                                "1 who-does-what lines.")
