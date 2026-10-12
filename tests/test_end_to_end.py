import openpyxl
import pytest

from conftest import TEST_JOB_FILE
from estimator.__main__ import _build as build_job
from estimator.__main__ import main
from estimator.price_sheet.rows import read_sent_rows
from estimator.price_sheet.write import PriceSheetError
from estimator.proposal.read import read_proposal
from estimator.proposal.sections import closing_texts, cover_lines, opening_texts, scope_blank_pattern, scope_texts
from estimator.scope import Alternate
from estimator.text import strip_list_marker


# T-009 (ruling 22): the proposal prints one demolition exclusion. These jobs have no Demolition Scope block, so the
# general line stays and the template's specific line is gone; everything else is as sent.
SPECIFIC_DEMOLITION = "Any demolition work not specifically mentioned"


def _one_demolition_line(texts):
    return [t for t in texts if t != SPECIFIC_DEMOLITION]


def _build(tmp_path):
    main(["build", str(TEST_JOB_FILE), "--out", str(tmp_path)])
    return tmp_path / "26-0246 Proposal.docx", tmp_path / "26-0246 Price Sheet.xlsx"


def test_the_proposal_matches_what_was_sent(tmp_path, sent_proposal, proposal_template, test_job, without_prices):
    proposal, _ = _build(tmp_path)
    assert scope_texts(proposal) == scope_texts(sent_proposal)
    assert closing_texts(proposal) == _one_demolition_line(closing_texts(sent_proposal))
    assert opening_texts(proposal) == opening_texts(sent_proposal)
    assert read_proposal(proposal, proposal_template) == without_prices(test_job)


def test_the_blank_line_pattern_matches_what_was_sent(tmp_path, sent_proposal):
    """Words alone don't show grouping: a header that runs straight on from the one before it
    (no blank line) reads the same as one behind a blank line, unless the blanks are compared too."""
    proposal, _ = _build(tmp_path)
    assert scope_blank_pattern(proposal) == scope_blank_pattern(sent_proposal)


def test_the_cover_page_differs_only_where_the_sent_copy_left_the_template(tmp_path, sent_proposal):
    proposal, _ = _build(tmp_path)
    ours, sent = cover_lines(proposal), cover_lines(sent_proposal)
    assert len(ours) == len(sent)
    assert [i for i, (a, b) in enumerate(zip(ours, sent)) if a != b] == [2]
    assert (ours[2], sent[2]) == ("Quote #: 26-0246", "Quote: 26-0246")   # the template's label; the sent copy was edited


def test_the_price_sheet_is_the_proposal_row_for_row(tmp_path):
    proposal, sheet = _build(tmp_path)
    rows = read_sent_rows(sheet)
    in_sheet = [strip_list_marker(r.text) for title in ("Price Sheet 1", "ADD ALT", "DEDUCT ALT") for r in rows[title]]
    assert in_sheet == scope_texts(proposal)


def test_hours_and_quantities_reach_the_price_sheet(tmp_path):
    _, sheet = _build(tmp_path)
    ws = openpyxl.load_workbook(str(sheet))["Price Sheet 1"]
    assert (ws["D6"].value, ws["B6"].value) == ("SCOPE OF WORK – PREWAILING WAGE", 140)
    header = next(r for r in range(6, 397) if ws.cell(r, 4).value == "1.  *Qty (1) BACnet Data Cabinet")
    assert ws.cell(header, 2).value == 10
    assert (ws.cell(header + 1, 4).value, ws.cell(header + 1, 3).value) == ("o  Furnish and Install UPS", 1)


def test_the_alternates_are_numbered_as_the_proposal_numbers_them(tmp_path):
    _, sheet = _build(tmp_path)
    wb = openpyxl.load_workbook(str(sheet))
    assert "45.  *Qty (1) Can Wash Kitchen Exhaust Fan EF-9" in [c.value for c in wb["ADD ALT"]["D"]]
    assert "46.  Qty (1) Shower Room Exhaust Fan EF-2" in [c.value for c in wb["DEDUCT ALT"]["D"]]


def test_a_job_that_does_not_fit_leaves_no_proposal_behind(tmp_path, small_job, proposal_template,
                                                            workup_template, price_book):
    """I2: the fit is checked before anything is written, and both files are written under
    temporary names and only renamed into place together — so a job that needs more price
    sheets than the template has never leaves a proposal on disk with no price sheet beside it."""
    too_many = small_job.model_copy(update={
        "add_alternates": tuple(Alternate(label=f"ALT {i}: x", blocks=small_job.add_alternates[0].blocks)
                                for i in range(1, 11))})
    with pytest.raises(PriceSheetError):
        build_job(too_many, proposal_template, workup_template, price_book, tmp_path)
    assert list(tmp_path.iterdir()) == []
