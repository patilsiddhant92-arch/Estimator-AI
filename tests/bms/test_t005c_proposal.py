"""T-005c (estimator rulings of 2026-10-11, 3, 4, 5 and the printing of 7): the proposal's header from the job folder,
a red *** for an existing BMS that is not one of AME's three, "No front end information found", a long clarification
as short points, and the red FUTURE / notice lines. Every job and folder here is invented."""
from __future__ import annotations

import datetime as dt
import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn

from bms.model import (MARKS, Clarification, Count, Device, FrontEnd, Job, ModelError, Ref, json_schema,
                       schema_errors)
from bms.playbook import load_playbook
from bms.render import AME_ONLY_FLAG, NO_FRONT_END, render

FIXTURES = Path(__file__).resolve().parent / "fixtures"
REF = Ref(source_id="S1", quote="q", where="M-1, Drawings.pdf page 1")
TODAY = dt.date(2026, 10, 11)
RED = "FF0000"


@pytest.fixture(scope="module")
def playbook():
    return load_playbook(FIXTURES / "sample_playbook")


@pytest.fixture
def job() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def _device(id="FE.d1", slot="link", name="Link", **kw) -> Device:
    given = {"verb": "Furnish and Install", "furnished_by": "AME", "installed_by": "AME", "wired_by": "AME", **kw}
    return Device(id=id, slot=slot, name=name, quantity=Count(value=1), refs=(REF,), **given)


# ---- ruling 3: the header from the job folder -----------------------------------------------------------------------

def _catalogue(root: Path):
    return SimpleNamespace(job=SimpleNamespace(name=root.name, root=root), pages=(), sheets=(), sections=(),
                           overview=SimpleNamespace(topics=(), who_does_what=(), services=()))


def _header(root: Path):
    from estimator.writing.front_end import job_header
    return job_header(_catalogue(root), SimpleNamespace(versions=()), TODAY)


JOB = "26-7001, Invented County Courthouse - Rplc of Air Handling & Fan Coil Units, NJ"


@pytest.mark.parametrize("root", [Path("Jobs") / JOB / "2. Takeoffs" / "1. Base Takeoff",
                                  Path("Jobs") / JOB / "2. Takeoffs",
                                  Path("Jobs") / JOB / "2. Takeoffs" / "1. Base Takeoff" / "Plans"])
def test_the_header_reads_the_job_folder_that_holds_the_takeoffs(root):
    header, flags = _header(root)
    assert header.quote_number == "26-7001"
    assert header.reference == "Invented County Courthouse - Rplc of Air Handling & Fan Coil Units, NJ"
    assert not any("project name" in f or "job number" in f for f in flags)


def test_the_header_never_prints_the_takeoff_folders_name_or_the_bracket_fill_in():
    header, flags = _header(Path("Jobs") / "Loose PDFs" / "1. Base Takeoff")
    assert header.quote_number == "***" and header.reference == "***"
    assert "Base Takeoff" not in header.quote_number + header.reference
    assert "[project name]" not in header.reference
    assert any("job number" in f and "project name" in f and "2. Takeoffs" in f for f in flags)


def test_a_job_folder_with_only_its_number_gives_a_red_placeholder_reference_and_a_question():
    header, flags = _header(Path("Jobs") / "26-7002" / "2. Takeoffs" / "1. Base Takeoff")
    assert header.quote_number == "26-7002"
    assert header.reference == "***"
    assert any("project name" in f for f in flags)


def test_a_job_folder_given_directly_still_reads_its_own_name():
    # the old app is given the job folder itself ("<job number>, <project name>"): that is a job folder, not a
    # takeoff folder
    header, _ = _header(Path("Jobs") / "26-7003, Invented School")
    assert (header.quote_number, header.reference) == ("26-7003", "Invented School")


# ---- the model: mark and mark_note ---------------------------------------------------------------------------------

def test_a_device_carries_an_optional_mark_and_mark_note():
    assert MARKS == ("", "future", "notice")
    plain = _device()
    assert plain.mark == "" and plain.mark_note == ""
    future = _device(mark="future", mark_note="FUTURE AS PER control schematic on sheet M-710")
    assert Device.from_dict(future.to_dict()) == future
    with pytest.raises(ModelError):
        Device.from_dict({**future.to_dict(), "mark": "FUTURE"})
    schema = json_schema(Device)
    assert schema["properties"]["mark"]["enum"] == ["", "future", "notice"]
    assert schema["properties"]["mark_note"]["type"] == "string"
    # additive: an AI answer written before the mark (no mark, no mark_note) still fits the schema
    assert "mark" not in schema["required"] and "mark_note" not in schema["required"]
    older = {k: v for k, v in plain.to_dict().items() if k not in ("mark", "mark_note")}
    assert schema_errors(older, schema) == []
    assert schema_errors({**older, "mark": "later"}, schema)


def test_a_device_saved_before_the_mark_still_loads():
    data = _device().to_dict()
    del data["mark"], data["mark_note"]
    assert Device.from_dict(data).mark == ""


# ---- ruling 7 printed: FUTURE and notice lines in red --------------------------------------------------------------

def _front_items(job: Job, playbook, front: FrontEnd):
    return render(replace(job, front_end=front), playbook).scope.front_end.items


def test_a_future_line_prints_its_note_in_brackets_in_red(job, playbook):
    note = "FUTURE AS PER control schematic on sheet M-710"
    front = FrontEnd(mode="new-system", devices=(
        _device(name="Duct-Mounted Supply Static Pressure Sensor", mark="future", mark_note=note),
        _device(id="FE.d2", name="Link")))
    items = _front_items(job, playbook, front)
    future = next(i for i in items if "Static Pressure" in i.text)
    assert future.text.endswith(f"({note})")
    assert all(r.red for r in future.runs)
    plain = next(i for i in items if i.text.endswith("Link"))
    assert not any(r.red for r in plain.runs)


def test_a_future_line_with_no_note_still_says_future(job, playbook):
    front = FrontEnd(mode="new-system", devices=(_device(name="Airflow Station", mark="future"),))
    # (T-005e: the new front end's Master Template default lines print too, ruling 19)
    (item,) = [i for i in _front_items(job, playbook, front) if i.kind == "line" and "Airflow" in i.text]
    assert item.text.endswith("(FUTURE)") and item.runs[0].red


def test_a_notice_line_prints_in_red_with_no_bracket(job, playbook):
    front = FrontEnd(mode="new-system", devices=(_device(name="Integral Drain Pan Level Sensor", mark="notice",
                                                         furnished_by="Unit Mfr."),))
    (item,) = [i for i in _front_items(job, playbook, front) if i.kind == "line" and "Drain Pan" in i.text]
    assert item.runs[0].red and "FUTURE" not in item.text and item.text.endswith("Unit Mfr.)")    # its owner only


def test_a_marked_line_in_a_system_block_prints_red_too(job, playbook):
    system = job.systems[1]
    device = replace(system.devices[0], mark="future", mark_note="FUTURE AS PER sheet M-710")
    marked = replace(job, demolition=(), systems=tuple(
        replace(s, devices=(device, *s.devices[1:])) if s.key == system.key else s for s in job.systems))
    draft = render(marked, playbook)
    lines = [i for b in draft.scope.blocks for i in b.items if i.kind == "line"]
    red = [i for i in lines if i.runs[0].red]
    assert any(i.text.endswith("(FUTURE AS PER sheet M-710)") for i in red)
    # G-004: an unknown device count is red too. Every other red line is that "Qty (##)" line.
    assert all(i.text.endswith("(FUTURE AS PER sheet M-710)") or i.text.startswith("Qty (##) ") for i in red)
    assert any(line.text == red[0].text for line in draft.lines)        # the traced line carries the same words


# ---- ruling 4: the existing BMS ------------------------------------------------------------------------------------

def test_an_existing_bms_that_is_not_one_of_ames_three_gets_one_red_flag_line(job, playbook):
    front = FrontEnd(mode="extend-existing", existing_system_name="Schneider Electric", devices=(_device(),))
    items = _front_items(job, playbook, front)
    flags = [i for i in items if i.text == AME_ONLY_FLAG]
    assert len(flags) == 1 and flags[0].kind == "note" and flags[0].runs[0].red
    assert "Honeywell WEBS" in AME_ONLY_FLAG and "Johnson FX" in AME_ONLY_FLAG and "Niagara Tridium" in AME_ONLY_FLAG


@pytest.mark.parametrize("name", ["Honeywell WEBs", "Niagara N4", "Tridium Niagara", "Johnson Controls FX",
                                  "Facility Explorer", "Honeywell / Niagara N4"])
def test_one_of_ames_three_gets_no_flag(job, playbook, name):
    front = FrontEnd(mode="extend-existing", existing_system_name=name, devices=(_device(),))
    assert not any(i.text == AME_ONLY_FLAG for i in _front_items(job, playbook, front))


def test_an_unnamed_existing_bms_gets_no_flag_line_the_planner_asks_for_the_name(job, playbook):
    front = FrontEnd(mode="extend-existing", existing_system_name=None, devices=(_device(),))
    assert not any(i.text == AME_ONLY_FLAG for i in _front_items(job, playbook, front))


def test_ame_bms_check():
    from estimator.proposal.write import ame_bms
    assert ame_bms("Honeywell WEBs") and ame_bms("Niagara N4") and ame_bms("Johnson Controls FX")
    assert not ame_bms("Schneider Electric") and not ame_bms("Honeywell") and not ame_bms("Metasys")
    assert not ame_bms("[existing BMS name]") and not ame_bms("")


# ---- ruling 4: no front end ----------------------------------------------------------------------------------------

def test_no_front_end_in_the_documents_prints_no_front_end_information_found(job, playbook):
    for front in (None, FrontEnd(mode="new-system"), FrontEnd(mode="new-system", no_scope_reason="none called for")):
        draft = render(replace(job, front_end=front), playbook)
        assert [i.text for i in draft.scope.front_end.items] == [NO_FRONT_END] == ["No front end information found"]
        assert any(line.text == NO_FRONT_END and line.system_key == "" for line in draft.lines)


def test_a_front_end_the_documents_call_for_but_no_line_was_written_is_not_called_missing(job, playbook):
    front = FrontEnd(mode="new-system", lines=("Provide graphics",), no_scope_reason="waiting")
    items = render(replace(job, front_end=front), playbook).scope.front_end.items
    assert NO_FRONT_END not in [i.text for i in items]


# ---- ruling 5: a long clarification as short points ----------------------------------------------------------------

LONG = ("The BMS integration of the new units is by the County's integrator under a separate contract. AME provides "
        "the controllers with BACnet MS/TP communication for that integration. Graphics, trending and alarms at the "
        "existing front end are by others; AME assists with the point mapping for one day")


def test_a_long_clarification_prints_as_short_points(job, playbook):
    front = FrontEnd(mode="new-system", devices=(_device(),), clarifications=(Clarification("confirm", LONG),))
    notes = [i for i in _front_items(job, playbook, front) if i.kind == "note"]
    (note,) = notes
    points = note.text.split("\n")
    assert len(points) >= 3
    assert points[0].startswith("*")              # the kind's label, as the Playbook words it
    assert all(p.startswith("- ") for p in points[1:])
    assert all(len(p) <= 120 for p in points)
    words = lambda t: " ".join(t.replace("\n- ", " ").replace(";", " ").replace(".", " ").split())
    assert words(LONG)[-40:] in words(note.text) and words(LONG)[:40] in words(note.text)     # nothing cut


def test_a_short_clarification_stays_one_line(job, playbook):
    front = FrontEnd(mode="new-system", devices=(_device(),),
                     clarifications=(Clarification("confirm", "integration by others."),))
    (note,) = [i for i in _front_items(job, playbook, front) if i.kind == "note"]
    assert "\n" not in note.text


def test_an_abbreviation_does_not_split_a_point():
    from bms.render import short_points
    text = "*Confirm Map Available Points Via Mfr. Provided BMS Interface per Spec. Section 23 09 00 " + "x" * 200
    assert all(not p.endswith("Mfr.") and not p.endswith("Spec.") for p in short_points(text).split("\n"))


# ---- write.py: the red runs and the red *** -------------------------------------------------------------------------

def _write(job: Job, playbook, tmp_path) -> Document:
    from estimator.__main__ import PROPOSAL_TEMPLATE
    from estimator.proposal.write import write_proposal
    out = write_proposal(render(job, playbook).scope, PROPOSAL_TEMPLATE, tmp_path / "p.docx")
    return Document(str(out))


def _colour(run) -> str:
    colour = run._r.find(f"{qn('w:rPr')}/{qn('w:color')}")
    return (colour.get(qn("w:val")) or "").upper() if colour is not None else ""


def _paragraph(document, start: str):
    return next(p for p in document.paragraphs if " ".join(p.text.split()).startswith(start))


def _red_text(paragraph) -> str:
    return "".join(r.text for r in paragraph.runs if _colour(r) == RED)


def test_a_non_ame_existing_bms_prints_red_with_stars_in_the_heading_and_the_opening(job, playbook, tmp_path):
    front = FrontEnd(mode="extend-existing", existing_system_name="Schneider Electric", devices=(_device(),))
    document = _write(replace(job, front_end=front), playbook, tmp_path)
    heading = _paragraph(document, "Seamlessly Tie in to the Existing")
    assert " ".join(heading.text.split()) == "Seamlessly Tie in to the Existing ***Schneider Electric*** BMS"
    assert _red_text(heading).strip() == "***Schneider Electric***"
    opening = next(p for p in document.paragraphs if "an extension of the existing" in p.text)
    assert "***Schneider Electric***" in _red_text(opening)
    flag = _paragraph(document, AME_ONLY_FLAG)
    assert all(_colour(r) == RED for r in flag.runs if r.text.strip())


def test_the_unnamed_placeholder_prints_red_with_stars(job, playbook, tmp_path):
    front = FrontEnd(mode="extend-existing", existing_system_name=None, devices=(_device(),))
    heading = _paragraph(_write(replace(job, front_end=front), playbook, tmp_path), "Seamlessly Tie in")
    assert _red_text(heading).strip() == "***[existing BMS name]***"


def test_an_ame_existing_bms_prints_plain(job, playbook, tmp_path):
    front = FrontEnd(mode="extend-existing", existing_system_name="Niagara N4", devices=(_device(),))
    heading = _paragraph(_write(replace(job, front_end=front), playbook, tmp_path), "Seamlessly Tie in")
    assert " ".join(heading.text.split()) == "Seamlessly Tie in to the Existing Niagara N4 BMS"
    assert _red_text(heading) == ""


def test_a_placeholder_header_value_prints_red(job, playbook, tmp_path):
    header = replace(job.header, quote_number="***", reference="***")
    document = _write(replace(job, header=header), playbook, tmp_path)
    cells = [c for t in document.tables[:1] for row in t.rows for c in row.cells]
    for label in ("Quote #:", "Reference:"):
        paragraph = next(c.paragraphs[0] for c in cells if c.paragraphs and c.paragraphs[0].text.startswith(label[:5]))
        assert _red_text(paragraph).strip() == "***", label


def test_a_named_header_value_keeps_its_own_colour(job, playbook, tmp_path):
    document = _write(job, playbook, tmp_path)
    cells = [c for t in document.tables[:1] for row in t.rows for c in row.cells]
    paragraph = next(c.paragraphs[0] for c in cells if c.paragraphs and c.paragraphs[0].text.startswith("Refer"))
    assert job.header.reference in paragraph.text and job.header.reference not in _red_text(paragraph)


def test_a_red_line_and_short_points_print_in_the_document(job, playbook, tmp_path):
    note = "FUTURE AS PER control schematic on sheet M-710"
    front = FrontEnd(mode="new-system", devices=(_device(name="Airflow Station", mark="future", mark_note=note),),
                     clarifications=(Clarification("confirm", LONG),))
    document = _write(replace(job, front_end=front), playbook, tmp_path)
    line = next(p for p in document.paragraphs if note in p.text)
    assert all(_colour(r) == RED for r in line.runs if r.text.strip())
    points = next(p for p in document.paragraphs if "County's integrator" in p.text)
    assert points._p.findall(f".//{qn('w:br')}")                    # one paragraph, the points on their own lines
