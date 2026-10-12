"""G-005 ruling 26: off-hours work and prevailing wage are red sub-bullets.

Invented fixtures only (the sample job, the sample Playbook, the repo template).
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from bms.model import Job
from bms.playbook import load_playbook
from bms.render import render
from estimator.proposal import layout
from estimator.scope import LaborCondition
from estimator.text import squash

FIXTURES = Path(__file__).resolve().parent / "fixtures"
OFF_PARENT = "All off-hours or overtime labor."
WAGE_PARENT = "All labor will be performed during normal business working hours"
HOURS_LINE = "The documents require work between 6 PM and 6 AM (Spec Section 01 14 00)."
NO_HOURS_LINE = "The documents require off-hours work (Spec Section 01 14 00)."
WAGE_LINE = ("All labor is priced at prevailing wage rates, as the documents require "
             "(As per Spec Section 00 73 46).")
RULING_22 = "Any demolition work not"

OFF = LaborCondition(
    kind="off-hours",
    text="Off-Hours Labor (6 PM to 6 AM) Included (As per Spec Section 01 14 00)",
    source="Spec Section 01 14 00",
    quote="Work between 6 PM and 6 AM is included.",
)
WAGE = LaborCondition(
    kind="prevailing-wage",
    text="Prevailing Wage Labor Rates Included (As per Spec Section 00 73 46)",
    source="Spec Section 00 73 46",
    quote="Pay the invented prevailing wage.",
)
OFF_NO_HOURS = LaborCondition(
    kind="off-hours",
    text="Off-Hours Labor Included (As per Spec Section 01 14 00)",
    source="Spec Section 01 14 00",
    quote="Off-hours labor is included.",
)


@pytest.fixture(scope="module")
def scope():
    job = Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))
    return render(job, load_playbook(FIXTURES / "sample_playbook")).scope


@pytest.fixture(scope="module")
def template_doc():
    from estimator.__main__ import PROPOSAL_TEMPLATE
    return Document(str(PROPOSAL_TEMPLATE))


def _red(paragraph) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(runs) and all(
        r.font.color is not None and str(r.font.color.rgb) == "FF0000" for r in runs)


def _proposal(scope, tmp_path, conditions, template=None, named=()):
    from estimator.__main__ import PROPOSAL_TEMPLATE
    from estimator.proposal.write import write_proposal
    # The sample job's named exception is not a ruling 26 sub-bullet. Clear it unless a test sets one.
    job = scope.model_copy(update={"labor_conditions": tuple(conditions), "named_exceptions": tuple(named)})
    path = write_proposal(job, template or PROPOSAL_TEMPLATE, tmp_path / "p.docx")
    return Document(str(path))


def _region(document):
    """Paragraphs from the clarifications heading up to, and not including, Base Price."""
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    start = next(i for i, p in enumerate(paragraphs) if squash(p.text).startswith(layout.CLARIFICATIONS))
    end = next(i for i, p in enumerate(paragraphs) if squash(p.text).startswith(layout.BASE_PRICE))
    return paragraphs[start:end]


def _body(paragraphs):
    """Section points, without the two headings and without the ruling 22 demolition lines."""
    kept = []
    for paragraph in paragraphs:
        text = squash(paragraph.text)
        if text.startswith(layout.CLARIFICATIONS) or text.startswith(layout.EXCLUSIONS):
            continue
        if text.startswith(RULING_22):
            continue
        kept.append(paragraph)
    return kept


def _standard_texts(template_doc):
    return [squash(p.text) for p in _body(_region(template_doc))]


def _added(document, template_doc):
    known = set(_standard_texts(template_doc))
    return [p for p in _body(_region(document)) if squash(p.text) not in known]


def _same_points(document, template_doc):
    known = set(_standard_texts(template_doc))
    kept = [squash(p.text) for p in _body(_region(document)) if squash(p.text) in known]
    assert kept == _standard_texts(template_doc)


def _next_after(document, prefix: str):
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    index = next(i for i, p in enumerate(paragraphs) if squash(p.text).startswith(prefix))
    return paragraphs[index + 1]


def test_both_sections_keep_every_template_point_and_add_two_red_sub_bullets(scope, template_doc, tmp_path):
    document = _proposal(scope, tmp_path, (OFF, WAGE))
    _same_points(document, template_doc)
    added = _added(document, template_doc)
    assert len(added) == 2
    assert [p.style.name for p in added] == ["PR3", "PR3"]
    assert all(_red(p) for p in added)
    assert all(not squash(p.text).startswith(("a.", "b.")) for p in added)


def test_each_sub_bullet_follows_its_parent(scope, tmp_path):
    document = _proposal(scope, tmp_path, (OFF, WAGE))
    off = _next_after(document, OFF_PARENT)
    wage = _next_after(document, WAGE_PARENT)
    assert squash(off.text) == HOURS_LINE
    assert off.style.name == "PR3" and _red(off)
    assert squash(wage.text) == WAGE_LINE
    assert wage.style.name == "PR3" and _red(wage)


def test_the_five_template_sub_bullets_stay_under_their_parent(scope, template_doc, tmp_path):
    parent = "The Division 23 BMS electrical work shall be performed"
    document = _proposal(scope, tmp_path, (OFF, WAGE))

    def five(document):
        paragraphs = [p for p in document.paragraphs if p.text.strip()]
        index = next(i for i, p in enumerate(paragraphs) if squash(p.text).startswith(parent))
        return paragraphs[index + 1:index + 6]

    template_five = five(template_doc)
    proposal_five = five(document)
    assert [p.style.name for p in template_five] == ["PR3"] * 5
    assert [squash(p.text) for p in proposal_five] == [squash(p.text) for p in template_five]
    assert [p.style.name for p in proposal_five] == ["PR3"] * 5


def test_the_normal_hours_point_and_the_off_hours_exclusion_stay(scope, tmp_path):
    document = _proposal(scope, tmp_path, (OFF, WAGE))
    texts = [squash(p.text) for p in document.paragraphs]
    assert any(t.startswith(WAGE_PARENT) for t in texts)
    assert any(t.startswith(OFF_PARENT) for t in texts)
    assert not any(t.startswith("All labor for the off-hours work") for t in texts)


def test_off_hours_with_no_range_names_the_work_and_the_source(scope, tmp_path):
    document = _proposal(scope, tmp_path, (OFF_NO_HOURS,))
    found = _next_after(document, OFF_PARENT)
    assert squash(found.text) == NO_HOURS_LINE
    assert found.style.name == "PR3" and _red(found)


def test_one_condition_adds_one_sub_bullet(scope, template_doc, tmp_path):
    off_only = _proposal(scope, tmp_path, (OFF,))
    added = _added(off_only, template_doc)
    assert [squash(p.text) for p in added] == [HOURS_LINE]
    assert _next_after(off_only, OFF_PARENT).style.name == "PR3"
    assert _next_after(off_only, WAGE_PARENT).text.startswith("All Motor starters require")

    wage_only = _proposal(scope, tmp_path, (WAGE,))
    added = _added(wage_only, template_doc)
    assert [squash(p.text) for p in added] == [WAGE_LINE]
    assert _next_after(wage_only, WAGE_PARENT).style.name == "PR3"
    assert _next_after(wage_only, OFF_PARENT).text.startswith("All temporary heating or cooling")


def test_no_labor_condition_leaves_both_sections_as_the_template(scope, template_doc, tmp_path):
    document = _proposal(scope, tmp_path, ())
    assert _added(document, template_doc) == []
    _same_points(document, template_doc)


def test_a_sub_bullet_is_not_a_named_exception(scope, tmp_path):
    from estimator.__main__ import PROPOSAL_TEMPLATE
    from estimator.proposal.read import read_named_exceptions
    from estimator.proposal.write import write_proposal

    plain = scope.model_copy(update={"labor_conditions": (OFF, WAGE), "named_exceptions": ()})
    written = write_proposal(plain, PROPOSAL_TEMPLATE, tmp_path / "plain.docx")
    assert read_named_exceptions(Document(str(written)), PROPOSAL_TEMPLATE) == ()

    named = "EXCEPTION TO ANY WORK WITH the invented fan coil the documents give this work to others"
    extra = scope.model_copy(update={"labor_conditions": (OFF, WAGE), "named_exceptions": (named,)})
    written = write_proposal(extra, PROPOSAL_TEMPLATE, tmp_path / "extra.docx")
    assert read_named_exceptions(Document(str(written)), PROPOSAL_TEMPLATE) == (named,)


def test_a_missing_parent_raises(scope, tmp_path):
    from estimator.__main__ import PROPOSAL_TEMPLATE
    from estimator.files import open_docx
    from estimator.proposal.write import TemplateError

    document = open_docx(PROPOSAL_TEMPLATE)
    parent = next(p for p in document.paragraphs if squash(p.text).startswith(OFF_PARENT))
    parent._p.getparent().remove(parent._p)
    template = tmp_path / "template.docx"
    document.save(str(template))
    with pytest.raises(TemplateError, match="off-hours"):
        _proposal(scope, tmp_path, (OFF,), template=template)


@pytest.mark.parametrize("text, hours", [
    ("Off-Hours Labor (6 PM to 6 AM) Included (As per Spec Section 01 14 00)", "6 PM to 6 AM"),
    ("Off-Hours Labor (6:00 PM to 6:00 AM) Included (As per Sheet M-001)", "6:00 PM to 6:00 AM"),
    ("Off-Hours Labor (7 PM to 5 AM) Included (As per the documents)", "7 PM to 5 AM"),
    ("Off-Hours Labor Included (As per Spec Section 01 14 00)", ""),
    ("Off-Hours Labor Included (As per Sheet M-001)", ""),
    ("Off-Hours Labor Included (As per the documents)", ""),
])
def test_the_hours_reader_takes_a_time_range_in_the_first_brackets_only(text, hours):
    from estimator.proposal.write import hours_between
    assert hours_between(text) == hours
