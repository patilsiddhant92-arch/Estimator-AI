"""Workstream 4: the renderer, golden-text style, on the invented sample job and the sample Playbook only."""
from __future__ import annotations

import json
import re
from dataclasses import replace
from pathlib import Path

import pytest

from bms.model import (Assumption, Clarification, Count, Device, Family, FrontEnd, Job, OtherCount, Ref,
                       ScopeException, System)
from bms.playbook import load_playbook
from bms.render import ProposalDraft, RenderedLine, RenderError, render

FIXTURES = Path(__file__).resolve().parent / "fixtures"
ROOT = Path(__file__).resolve().parents[2]
REF = Ref(source_id="S1", quote="q", where="M-1, Drawings.pdf page 1")


@pytest.fixture(scope="module")
def playbook():
    return load_playbook(FIXTURES / "sample_playbook")


@pytest.fixture
def job() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


@pytest.fixture
def draft(job, playbook) -> ProposalDraft:
    return render(job, playbook)


def texts(draft: ProposalDraft, key: str) -> list[str]:
    return [l.text for l in draft.lines if l.system_key == key]


def device(slot="controller", name="DDC Controller", qty: int | None = 2, **kw) -> Device:
    base = dict(id="X.d1", slot=slot, name=name, verb="Furnish and Install", furnished_by="AME", installed_by="AME",
                wired_by="AME", quantity=Count(value=qty), refs=(REF,))
    base.update(kw)
    return Device(**base)


def system(devices=(), families=None, **kw) -> System:
    base = dict(key="EF", name="Exhaust Fans", shape="Exhaust Fan", control="field",
                families=families or (Family("Exhaust Fan", ("EF-1", "EF-2"), Count(value=2)),), devices=tuple(devices))
    base.update(kw)
    return System(**base)


def one(job: Job, sys: System) -> Job:
    return replace(job, systems=(sys,))


# ---- golden text ---------------------------------------------------------------------------------------------

def test_rooftop_block_golden(draft):
    assert texts(draft, "RTU") == [
        "Qty (2) Rooftop Units RTU-1, 2",
        "Furnish BMS Communication Wiring to RTU Factory Controls",
        "Map Available Points Via Mfr. Provided BMS Interface",
        "AME requires mfr. Rep to assist in integration, cost to be incurred by mech contractor",
        "Install and Wire Duct Smoke Detector (Furnished by Unit Mfr.)",
        "*one detector per unit, because the schedule note names a detector for each unit without a count",
        "*Unit count carried at 2; the roof plan shows 1",
        "*Qty. (2) per M-001; 1 per M-101.",
    ]


def test_exhaust_fan_block_golden(draft):
    assert texts(draft, "EF") == [
        "Qty (3) Exhaust Fans EF-1 thru 3",
        "Furnish and Install DDC Controller",
        "Furnish Fan Start / Stop and Status",
        "Furnish and Install Motorized Damper Actuator (Qty. 2)",
    ]


def test_vrf_stacks_one_header_per_family_over_one_set_of_lines(draft):
    assert texts(draft, "VRF")[:3] == ["Qty (4) Indoor Unit FCU-1 thru 4", "Qty (1) Condensing Unit CU-1",
                                       "Qty (2) Branch Selector BS-1, 2"]
    blocks = [b for b in draft.scope.blocks if b.header.startswith(("Qty (4) Indoor", "Qty (1) Cond", "Qty (2) Branch"))]
    assert [b.joined for b in blocks] == [False, True, True]
    assert blocks[0].items == () and blocks[1].items == () and blocks[2].items
    assert blocks[0].tech_hours == 6.0 and blocks[2].tech_hours is None
    assert "Furnish BMS Communication Wiring to VRF Factory Controls" in texts(draft, "VRF")
    assert "Wire BACnet Gateway (Furnished by Unit Mfr.)" in texts(draft, "VRF")


def test_unknown_count_prints_hashes_and_quantity_unknown(draft):
    cuh = texts(draft, "CUH")
    assert cuh[0] == "Qty (##) Cabinet Unit Heaters"
    assert cuh[1] == "Furnish and Install Space Temperature Sensor (Qty. ##)"
    assert cuh[2].startswith("*one sensor per heater, because ")
    assert cuh[3] == "*Cabinet unit heater quantity as per drawings"


def test_confirm_note_and_count_note_for_vrf(draft):
    notes = [l.text for l in draft.lines if l.system_key == "VRF" and l.kind == "note"]
    assert notes == ["*Branch selector count: 2 on the plan picture, 1 on the schedule - to be confirmed",
                     "*Qty. (2) per M-102; 1 per M-002."]


def test_front_end_exceptions_and_alternates(draft):
    scope = draft.scope
    assert [i.text for i in scope.front_end.items] == ["Furnish and Install BMS Server Software and Graphics"]
    assert scope.front_end.mode == "new-system" and scope.front_end.tech_hours == 8.0
    assert scope.named_exceptions == ("Fire alarm system by others",)
    assert [b.header for b in scope.blocks if "Duct Heaters" in b.header] == []
    assert len(scope.add_alternates) == 1 and scope.deduct_alternates == ()
    alt = scope.add_alternates[0].blocks[0]
    assert alt.header == "Qty (2) Duct Heaters DH-1, 2" and alt.tech_hours == 4.0
    assert [l.text for l in draft.lines if l.system_key == "DH"][0] == alt.header
    assert [l.system_key for l in draft.lines][-1] == "DH"          # alternates print after the base scope


def test_header_comes_from_the_job(draft):
    h = draft.scope.header
    assert (h.quote_number, h.reference, h.mep_engineer, h.to) == ("00-0000", "Sample Clinic (invented test job)",
                                                                  "Sample Engineering", "All Bidders")
    assert h.date.isoformat() == "2026-10-04" and len(h.based_on) == 2


# ---- the price sheet rows ------------------------------------------------------------------------------------

def test_tech_hours_go_on_the_block_header(draft):
    by_header = {b.header: b.tech_hours for b in draft.scope.blocks}
    assert by_header["Qty (2) Rooftop Units RTU-1, 2"] == 4.0
    assert by_header["Qty (3) Exhaust Fans EF-1 thru 3"] == 12.0
    assert by_header["Qty (##) Cabinet Unit Heaters"] is None


def test_item_quantities_price_only_what_ame_acts_on_and_never_under_a_hash_count(draft):
    blocks = {b.header: b for b in draft.scope.blocks}
    rtu = blocks["Qty (2) Rooftop Units RTU-1, 2"]
    assert [i.quantity for i in rtu.items][:4] == [2, 2, None, 2]      # the rep line is the mechanical contractor's
    assert [i.depth for i in rtu.items][:4] == [1, 1, 2, 1]            # the rep line sits one level in, under its map line
    assert all(i.quantity is None for i in blocks["Qty (##) Cabinet Unit Heaters"].items)
    assert all(i.quantity is None for i in rtu.items if i.kind == "note")


def test_rendered_lines_trace_to_devices(draft):
    by_id = {l.device_id: l for l in draft.lines if l.device_id}
    assert by_id["RTU.d4"].verb == "Install and Wire" and by_id["RTU.d4"].quantity == 2
    assert by_id["EF.d3"].quantity == 2 and by_id["CUH.d1"].quantity is None
    assert by_id["RTU.d1"].verb == "Furnish"                           # the trio line leads with its verb
    assert by_id["RTU.d2"].verb is None                                # "Map ..." has none
    assert {l.kind for l in draft.lines} == {"header", "line", "note"}
    assert all(l.device_id is None for l in draft.lines if l.kind != "line")


# ---- rules ---------------------------------------------------------------------------------------------------

def test_pure_same_job_same_text(job, playbook):
    assert render(job, playbook) == render(Job.from_dict(job.to_dict()), playbook)


def test_lines_follow_the_shape_order_not_the_device_order(job, playbook):
    ef = job.system("EF")
    shuffled = replace(ef, devices=tuple(reversed(ef.devices)))
    assert texts(render(one(job, shuffled), playbook), "EF")[1:] == [
        "Furnish and Install DDC Controller", "Furnish Fan Start / Stop and Status",
        "Furnish and Install Motorized Damper Actuator (Qty. 2)"]


def test_the_higher_count_prints_even_when_value_is_lower(job, playbook):
    fam = Family("Exhaust Fan", ("EF-1", "EF-2"), Count(value=2, sources=(Ref("S0", "q", "M-2, Drawings.pdf page 2"),),
                                                       others=(OtherCount(5, REF),)))
    out = render(one(job, system([device()], families=(fam,))), playbook)
    assert texts(out, "EF")[0] == "Qty (5) Exhaust Fans EF-1, 2"
    assert texts(out, "EF")[-1] == "*Qty. (5) per M-1; 2 per M-2."


def test_count_note_names_the_place_of_every_count(job, playbook):
    fam = Family("Exhaust Fan", ("EF-1",), Count(value=4, sources=(Ref("S1", "q", "M-201, Drawings.pdf page 3"),),
                                                 others=(OtherCount(3, Ref("S2", "q", "M-301, Specs.pdf page 8")),
                                                         OtherCount(5, Ref("S3", "q", "M-401, Drawings.pdf page 9")))))
    note = texts(render(one(job, system([device()], families=(fam,))), playbook), "EF")[-1]
    assert note == "*Qty. (5) per M-401; 4 per M-201; 3 per M-301."


def test_count_note_uses_sheet_numbers_never_file_names_and_folds_one_place(job, playbook):
    pb = load_playbook(ROOT / "playbook")
    fam = Family("Exhaust Fan", ("EF-1",), Count(
        value=7, sources=(Ref("S1", "q", "equipment counts for EF"),),
        others=(OtherCount(3, Ref("S2", "q", "Q-210.10, Invented Set - Combined.pdf page 15")),
                OtherCount(3, Ref("S3", "q", "Q-210.10 whole sheet, Invented Set - Combined.pdf page 15 (picture)")),
                OtherCount(2, Ref("S4", "q", "INVENTED SCHEDULE, Q-702.00, Invented Set - Combined.pdf page 51")),
                OtherCount(1, Ref("S5", "q", "Q-305.00, Invented Set.pdf page 4; Q-306.00, Invented Set.pdf page 5")),
                OtherCount(0, Ref("S6", "q", "Q-305.00 whole sheet, Invented Set.pdf page 4 (picture)")),
                OtherCount(4, Ref("S7", "q", "99 99 99 Invented Section, invented-spec_Rev_2.pdf page 3")))))
    out = texts(render(one(job, system([device(qty=7)], families=(fam,), shape="")), pb), "EF")
    qty = [t for t in out if t.startswith("*Qty.")]
    assert qty == ["*Qty. (7); 4 per 99 99 99 Invented Section; 3 per Q-210.10; 2 per Q-702.00; "
                   "1 and 0 per Q-305.00."]
    assert not any("the documents" in t for t in out)
    assert not any(".pdf" in t for t in out)


def test_count_note_is_one_line_per_block_beside_the_confirm(job, playbook):
    pb = load_playbook(ROOT / "playbook")
    fam = Family("Exhaust Fan", ("EF-1",), Count(value=3, others=(OtherCount(2, Ref("S2", "q", "Q-1, A.pdf page 1")),)))
    d = device(qty=3, quantity=Count(3, others=(OtherCount(2, Ref("S2", "q", "Q-1, A.pdf page 1")),)))
    clar = (Clarification("confirm", "Confirm quantity: 3 fans carried per the schedule"),)
    out = texts(render(one(job, system([d], families=(fam,), shape="", clarifications=clar)), pb), "EF")
    assert out[-2:] == ["*Confirm quantity: 3 fans carried per the schedule", "*Qty. (3); 2 per Q-1."]


# ---- the clarification labels (the text's own label is not printed twice) -------------------------------------

@pytest.mark.parametrize("kind, text, printed", [
    ("confirm", "Confirm ownership: invented relay by others", "*Confirm ownership: invented relay by others"),
    ("confirm", "Confirm quantity: 4 invented units", "*Confirm quantity: 4 invented units"),
    ("confirm", "*Confirm quantities: invented sensors", "*Confirm quantities: invented sensors"),
    ("confirm", "quantity: 4 invented units", "*Confirm quantity: 4 invented units"),
    ("confirm", "Assumed: invented panel by the maker", "*Confirm invented panel by the maker"),
    ("assumed", "*Assumed: one invented switch, as the sequence asks", "*Assumed: one invented switch, as the sequence asks"),
    ("assumed", "Assumed: one invented switch", "*Assumed: one invented switch"),
    ("per_drawings", "Quantity per plumbing drawings", "*Quantity per plumbing drawings"),
    ("per_drawings", "*Quantity per plumbing drawings", "*Quantity per plumbing drawings"),
    ("per_drawings", "plumbing", "*Quantity per plumbing drawings"),
    ("per_drawings", "quantity per the electrical drawings", "*Quantity per electrical drawings"),
    ("per_drawings", "Mechanical drawings", "*Quantity per mechanical drawings"),
])
def test_a_label_the_text_already_carries_prints_once(job, kind, text, printed):
    pb = load_playbook(ROOT / "playbook")
    out = texts(render(one(job, system([device()], clarifications=(Clarification(kind, text),))), pb), "EF")
    assert out[-1] == printed


def test_an_assumption_that_carries_its_own_label_prints_it_once(job):
    pb = load_playbook(ROOT / "playbook")
    a = device(assumption=Assumption("Assumed: one invented sensor", "the invented sequence asks for it"))
    out = texts(render(one(job, system([a])), pb), "EF")
    assert out[-1] == "*Assumed: one invented sensor, because the invented sequence asks for it"


def test_a_long_clarification_prints_whole(job):
    pb = load_playbook(ROOT / "playbook")
    long_text = "Confirm quantity: " + "invented words " * 30 + "end"
    out = texts(render(one(job, system([device()], clarifications=(Clarification("confirm", long_text),))), pb), "EF")
    assert out[-1] == "*" + long_text


# ---- source ids: the renderer rewrites no sentence; an issued id left in a stale text goes, token only ----------
# (the intake refuses such a text and the validator reports it as must_fix before render: tests/bms/test_validate.py)

def test_an_issued_id_in_an_assumption_is_taken_out_token_only(job):
    pb = load_playbook(ROOT / "playbook")
    refs = (Ref("S41", "q", "Q-701.00, Invented Set.pdf page 50"),
            Ref("S7", "q", "sequence section B. Invented fans, 99 99 99 INVENTED SEQUENCE, invented.pdf page 3"))
    a = device(refs=refs, assumption=Assumption(
        "one invented controller", "S41 note 1 asks for it and S7 paragraph 2 names no party"))
    out = texts(render(one(job, system([a])), pb), "EF")
    assert out[-1] == "*Assumed: one invented controller, because note 1 asks for it and paragraph 2 names no party"


def test_issued_ids_with_no_place_are_taken_out_and_no_word_around_them(job):
    pb = load_playbook(ROOT / "playbook")
    a = device(assumption=Assumption("one invented sensor per S900", "the schedule (S901, S902) shows one per unit"))
    b = device(id="X.d2", slot="damper_actuator", name="Damper",
               refs=(REF,) + tuple(Ref(i, "x", "ledger fact") for i in ("S900", "S901", "S902", "S903")),
               assumption=Assumption("one damper", "the plan shows it in S903"))
    out = texts(render(one(job, system([a, b])), pb), "EF")
    assert out[-1] == ("*Assumed: one invented sensor per, because the schedule ( , ) shows one per unit; "
                       "one damper, because the plan shows it in")
    assert "S9" not in out[-1] and "  " not in out[-1]


def test_keynote_drawing_fact_and_count_ids_do_not_print(job):
    pb = load_playbook(ROOT / "playbook")
    clar = (Clarification("confirm", "Confirm the keynote Q3, drawing I12, fact F8 and count C2a"),)
    d = device(refs=(REF,) + tuple(Ref(i, "x", "ledger fact") for i in ("Q3", "I12", "F8", "C2a")))
    out = texts(render(one(job, system([d], clarifications=clar)), pb), "EF")
    assert "Q3" not in out[-1] and "I12" not in out[-1] and "F8" not in out[-1] and "C2a" not in out[-1]


def test_issued_ids_in_a_clarification_are_taken_out_token_only(job):
    pb = load_playbook(ROOT / "playbook")
    d = device(refs=(Ref("S12", "q", "Q-100.00, Invented.pdf page 3"), Ref("S99", "q", "ledger fact")))
    clar = (Clarification("confirm", "Confirm quantity: 4 invented units per S12, 2 per S99"),)
    out = texts(render(one(job, system([d], clarifications=clar)), pb), "EF")
    assert out[-1] == "*Confirm quantity: 4 invented units per , 2 per"


def test_a_clarification_prints_its_kind_label_once_and_keeps_its_content_words(job):
    """I3: the renderer strips only the label of the clarification's own kind."""
    pb = load_playbook(ROOT / "playbook")
    for kind, text, printed in [
        ("confirm", "Confirm ownership of the duct smoke detectors with the electrical contractor.",
         "*Confirm ownership of the duct smoke detectors with the electrical contractor"),
        ("confirm", "Confirm quantity: 4 invented dampers", "*Confirm quantity: 4 invented dampers"),
        ("note", "Confirm with the owner the invented fan count", "**Confirm with the owner the invented fan count"),
        ("note", "Note: the invented fans are by others", "**the invented fans are by others"),
        ("assumed", "Assumed: one invented damper per fan", "*Assumed: one invented damper per fan"),
    ]:
        out = texts(render(one(job, system([device()], clarifications=(Clarification(kind, text),))), pb), "EF")
        assert out[-1] == printed, text


def test_a_where_that_holds_an_id_never_prints(job):
    """I4: a ledger fact, a picture file or a where naming the id itself gives no place; the id never prints."""
    pb = load_playbook(ROOT / "playbook")
    refs = (REF, Ref("F17", "x", "Ledger fact F17"), Ref("I3", "x", "I3 picture, M-201.png"), Ref("S8", "x", "S8"))
    a = device(refs=refs, assumption=Assumption("one invented actuator", "the actuator is shown on I3 and in F17"))
    clar = (Clarification("confirm", "the invented damper count per S8"),)
    fams = (Family("Exhaust Fan", ("EF-1",), Count(3, (Ref("F17", "x", "Ledger fact F17"),),
                                                  (OtherCount(2, Ref("S2", "q", "M-101, Invented.pdf page 2")),))),)
    out = texts(render(one(job, system([a], families=fams, clarifications=clar)), pb), "EF")
    assert "*Confirm the invented damper count per" in out
    assert "*Assumed: one invented actuator, because the actuator is shown on and in" in out
    for line in out:
        assert not re.search(r"\b(?:F17|I3|S8)\b", line), line


@pytest.mark.parametrize("issued", [(), (Ref("F1", "x", "ledger fact"),)])
def test_tags_and_sheet_numbers_shaped_like_ids_print_unchanged(job, issued):
    """I5: F1, F2 are not issued ids (or F1 is a family tag) and C101 (unissued) reads as a sheet number."""
    pb = load_playbook(ROOT / "playbook")
    fams = (Family("Exhaust Fan", ("F1", "F2"), Count(value=2)),)
    clar = (Clarification("note", "Fans F1 and F2 serve the kitchen; see civil sheet C101."),)
    d = device(refs=(REF,) + issued)
    out = texts(render(one(job, system([d], families=fams, clarifications=clar)), pb), "EF")
    assert out[-1] == "**Fans F1 and F2 serve the kitchen; see civil sheet C101"


def test_a_family_tag_shaped_like_an_id_prints_even_when_it_is_an_issued_id(job):
    pb = load_playbook(ROOT / "playbook")
    tags = ("F1", "C2", "I3", "Q4", "S5")
    d = device(refs=tuple(Ref(t, "x", "ledger fact") for t in tags))
    clar = (Clarification("confirm", "the controls for units F1, C2, I3, Q4 and S5"),)
    out = texts(render(one(job, system([d], families=(Family("Unit", tags, Count(value=5)),),
                                       clarifications=clar)), pb), "EF")
    assert out[-1] == "*Confirm the controls for units F1, C2, I3, Q4 and S5"


# ---- the safety net: tags T kept, an issued id I taken out (token only), an unissued token printed as written ----

S141 = Ref("S141", "q", "M-101, Invented.pdf page 2")


def _note_of(job, kind, text, tags=("EF-1", "EF-2"), refs=(S141,), assumption=None):
    pb = load_playbook(ROOT / "playbook")
    d = device(refs=refs, assumption=assumption)
    clar = (Clarification(kind, text),) if text else ()
    return texts(render(one(job, system([d], families=(Family("Exhaust Fan", tags, Count(value=2)),),
                                        clarifications=clar)), pb), "EF")[-1]


@pytest.mark.parametrize("kind, text, printed", [
    ("confirm", "Confirm valves per S141 and S999.", "*Confirm valves per and S999"),
    ("note", "Valves open per S999.", "**Valves open per S999"),
    ("note", "Valves open [S999].", "**Valves open [S999]"),
    ("note", "Valves open (S901, S902).", "**Valves open (S901, S902)"),
])
def test_render_takes_out_only_issued_tokens_and_rewrites_nothing(job, kind, text, printed):
    """No sentence rewriting at render: an issued id goes (token only); an unissued token prints as written (the
    intake refused it in a citation; the validator reports it)."""
    assert _note_of(job, kind, text) == printed


@pytest.mark.parametrize("text, printed", [("Valves per s141.", "**Valves per"),
                                           ("Valves per S141A.", "**Valves per S141A")])
def test_ids_match_whatever_their_case_and_suffix(job, text, printed):
    """I3: "s141" is the issued S141; "S141A" is another token (unissued, printed; refused at intake)."""
    assert _note_of(job, "note", text) == printed


@pytest.mark.parametrize("text, printed", [("Fans F1-F4 run per S141.", "**Fans F1-F4 run per"),
                                           ("Fans F1 thru F4 per S141.", "**Fans F1 thru F4 per"),
                                           ("Fan F3 runs per S141.", "**Fan F3 runs per")])
def test_a_tag_range_protects_every_member(job, text, printed):
    """I1: the family tag "F1-F4" protects F1, F2, F3 and F4, even when F1 is also an issued id."""
    assert _note_of(job, "note", text, tags=("F1-F4",), refs=(S141, Ref("F1", "x", "ledger fact"))) == printed


@pytest.mark.parametrize("kind, text, printed", [
    ("confirm", "**Confirm**: ownership of valves", "*Confirm ownership of valves"),
    ("confirm", "Confirm - ownership of valves", "*Confirm ownership of valves"),
    ("assumed", "Assumed - two valves", "*Assumed: two valves"),
    ("assumed", "Assumption-based count", "*Assumed: Assumption-based count"),
    ("note", "Note: - the fans are by others", "**the fans are by others"),
])
def test_no_punctuation_is_left_after_a_label(job, kind, text, printed):
    """M1: a label's colon, dash or asterisks go with it; a label word joined by a dash is a word, not a label."""
    assert _note_of(job, kind, text) == printed


@pytest.mark.parametrize("text", [
    "per S141 " + "and the " * 10_000 + ".",              # 20,000 words after an id
    "Valves per S141. " * 5_000,                          # 5,000 ids
    "Valves (S141) open [S999]. " * 2_500,                # 5,000 ids in 5,000 bracket groups
    "S141" + " " * 40_000 + "x",                          # review 3 M1: a long whitespace run
    "([" * 8_000 + "I3" + "])" * 8_000,                   # review 3 M1: deep brackets
], ids=["20000-words", "5000-ids", "5000-bracketed-ids", "40000-spaces", "8000-deep-brackets"])
@pytest.mark.slow
def test_the_detector_and_the_safety_net_are_linear_on_a_long_text(text):
    """M3: a runaway answer never stalls a run."""
    from timing import grows_linearly

    from bms.citations import SourceIds
    ids = SourceIds.of(["S141", "I3"], ())
    out = ids.without_issued(text)
    ids.refusal(text)
    assert "S141" not in out and "I3" not in out

    def work(n):                            # growth, not a wall-clock limit (it flaked under load): n characters
        sample = (text * (2 + n // len(text)))[:n]
        ids.without_issued(sample)
        ids.refusal(sample)
    assert grows_linearly(work, len(text) // 4)


@pytest.mark.parametrize("text", ["the fans run (see C2 in the lab).", "the fans run (C2 in the lab).",
                                  "the fans run, see C2 in the lab."])
@pytest.mark.parametrize("tags", [("EF-1",), ("C2",)], ids=["C2-no-tag", "C2-a-tag"])
def test_an_unissued_id_in_a_bracket_or_after_a_cue_word_prints_as_written(job, text, tags):
    """The intake refuses "(see C2 in the lab)" when C2 is no tag; the renderer rewrites nothing either way."""
    assert _note_of(job, "note", text, tags=tags) == "**" + text.rstrip(".")


def test_device_names_and_named_exceptions_lose_an_issued_id_token_only(job):
    """I4: a device name and a named exception go through the same safety net."""
    pb = load_playbook(ROOT / "playbook")
    d = device(name="Control Valve per S141", refs=(S141,))
    j = replace(one(job, system([d])), exceptions=(ScopeException("EF", "Fire alarm by others per S141", (S141,)),))
    out = render(j, pb)
    assert texts(out, "EF")[1] == "Furnish and Install Control Valve per"
    assert out.scope.named_exceptions == ("Fire alarm by others per",)
    j2 = replace(one(job, system([device(name="Control Valve per S999", refs=(S141,))])),
                 exceptions=(ScopeException("EF", "Fire alarm by others (S141)", ()),))
    out2 = render(j2, pb)
    assert texts(out2, "EF")[1] == "Furnish and Install Control Valve per S999"         # unissued: as written
    assert out2.scope.named_exceptions == ("Fire alarm by others ( )",)                 # the token only


def test_an_id_in_a_refs_where_is_never_a_printed_place(job):
    """M2: "schedule S9" names no place, issued or not."""
    pb = load_playbook(ROOT / "playbook")
    for issued in ((), (Ref("S9", "x", "ledger fact"),)):
        fam = Family("Exhaust Fan", ("EF-1",), Count(3, (S141,), (OtherCount(5, Ref("S5", "q", "schedule S9")),)))
        out = texts(render(one(job, system([device(refs=(S141,) + issued)], families=(fam,))), pb), "EF")
        assert not any("S9" in line for line in out), out


# ---- a system with no lines is left out ------------------------------------------------------------------------

def test_a_system_with_no_lines_prints_no_header_and_is_kept_for_the_evidence_package(job, playbook):
    empty = system([], key="EMPTY", name="Invented Units",
                   clarifications=(Clarification("note", "the invented units have no BMS work"),))
    out = render(replace(job, systems=(job.system("EF"), empty)), playbook)
    assert [b.header for b in out.scope.blocks] == ["Qty (3) Exhaust Fans EF-1 thru 3"]
    left = [l for l in out.lines if l.system_key == "EMPTY"]
    assert [(l.kind, l.device_id) for l in left] == [("note", None)]
    assert "no lines written" in left[0].text and "the invented units have no BMS work" in left[0].text


def test_an_alternate_with_no_lines_adds_no_alternate(job, playbook):
    empty = system([], key="EMPTY", alternate="add")
    out = render(replace(job, systems=(job.system("EF"), empty)), playbook)
    assert out.scope.add_alternates == ()


# ---- dampers under their fan ----------------------------------------------------------------------------------

def test_a_damper_family_under_a_fan_prints_no_header_of_its_own(job, playbook):
    ef = job.system("EF")
    fan_with_damper = replace(ef, families=ef.families + (Family("Motorized Damper", (), Count(value=2)),))
    damper_only = system([device(id="MD.d1", slot="damper_actuator", name="Motorized Damper Actuator", qty=4)],
                         key="MD", name="Motorized Dampers", shape="",
                         families=(Family("Motorized Damper", (), Count(value=4)),))
    out = render(replace(job, systems=(fan_with_damper, damper_only)), playbook)
    ef_lines = texts(out, "EF")
    assert [l.text for l in out.lines if l.system_key == "EF" and l.kind == "header"] == [
        "Qty (3) Exhaust Fans EF-1 thru 3"]
    assert ef_lines[1:] == ["Furnish and Install DDC Controller", "Furnish Fan Start / Stop and Status",
                            "Furnish and Install Motorized Damper Actuator (Qty. 2)"]
    assert texts(out, "MD")[0] == "Qty (4) Motorized Dampers"
    assert [b.header for b in out.scope.blocks] == ["Qty (3) Exhaust Fans EF-1 thru 3", "Qty (4) Motorized Dampers"]


# ---- the header never undershoots ------------------------------------------------------------------------------

def test_each_family_header_prints_its_highest_count(job, playbook):
    fams = (Family("Indoor Unit", ("IU-1",), Count(value=4, others=(OtherCount(6, REF), OtherCount(1, REF)))),
            Family("Condensing Unit", ("CU-1",), Count(value=None, others=(OtherCount(2, REF),))))
    out = render(one(job, system([device()], families=fams, shape="")), playbook)
    assert [l.text for l in out.lines if l.kind == "header"] == ["Qty (6) Indoor Unit IU-1", "Qty (2) Condensing Unit CU-1"]


def test_a_quantity_prints_only_when_it_differs_from_the_blocks_count(job, playbook):
    out = render(one(job, system([device(qty=2), device(qty=1, id="X.d2", slot="damper_actuator", name="Damper"),
                                  device(qty=2, id="X.d3", slot="damper_actuator", name="Damper B")])), playbook)
    assert texts(out, "EF")[1:] == ["Furnish and Install DDC Controller",
                                    "Furnish and Install Damper (Qty. 1)", "Furnish and Install Damper B"]


def test_every_assumption_says_why_and_a_reason_may_lead_with_its_own_word(job, playbook):
    a = device(assumption=Assumption("one per fan", "since the sequence names each fan"), flag="assumed")
    b = device(id="X.d2", slot="damper_actuator", name="Damper", assumption=Assumption("a damper per fan.", "the schedule"))
    note = texts(render(one(job, system([a, b])), playbook), "EF")[-1]
    assert note == "*one per fan, since the sequence names each fan; a damper per fan, because the schedule"


def test_an_assumption_without_its_why_is_refused(job, playbook):
    bad = device(assumption=Assumption("one per fan", "  "))
    with pytest.raises(RenderError, match="what and why"):
        render(one(job, system([bad])), playbook)


def test_one_star_line_per_kind_in_a_fixed_order(job, playbook):
    clar = (Clarification("note", "n1"), Clarification("per_drawings", "p1"), Clarification("confirm", "c1"),
            Clarification("confirm", "c2"), Clarification("assumed", "a1"), Clarification("note", "n2"))
    out = texts(render(one(job, system([device()], clarifications=clar)), playbook), "EF")
    assert out[2:] == ["*c1; c2 - to be confirmed", "*a1", "*p1 as per drawings", "*n1; n2"]


def test_owner_brackets_come_from_the_wording(job, playbook):
    d = device(furnished_by="Mechanical Contractor", installed_by="AME", verb="Install and Wire")
    assert texts(render(one(job, system([d])), playbook), "EF")[1] == (
        "Install and Wire DDC Controller (Furnished by Mechanical Contractor)")


def test_a_shapeless_system_prints_in_the_generic_form_in_device_order(job, playbook):
    d1 = device(name="Gas Detector", qty=3, furnished_by="Others", slot="x")
    d2 = device(id="X.d2", name="Panel", qty=None, slot="y")
    out = texts(render(one(job, system([d1, d2], shape="", families=(Family("Gas Detector", ("GD-1", "GD-2"),
                                                                              Count(value=3)),))), playbook), "EF")
    assert out == ["Qty (3) Exhaust Fans GD-1, 2", "Furnish and Install Gas Detector (Furnished by Others)",
                   "Furnish and Install Panel (Qty. ##)"]


def test_a_slot_the_shape_lacks_still_prints_after_the_known_slots(job, playbook):
    odd = device(id="X.d2", slot="not_a_slot", name="Extra Thing")
    out = texts(render(one(job, system([odd, device()])), playbook), "EF")
    assert out[1:] == ["Furnish and Install DDC Controller", "Furnish and Install Extra Thing"]


def test_a_system_with_no_families_prints_an_unknown_count(job, playbook):
    s = system([device()])
    out = render(one(job, replace(s, families=())), playbook)
    assert texts(out, "EF")[0] == "Qty (##) Exhaust Fans"
    assert out.scope.blocks[0].items[0].quantity is None


def test_unknown_shape_and_missing_header_are_refused(job, playbook):
    with pytest.raises(RenderError, match="shape"):
        render(one(job, system([device()], shape="No Such Shape")), playbook)
    with pytest.raises(RenderError, match="header"):
        render(replace(job, header=None), playbook)


def test_extend_existing_front_end_and_deduct_alternate(job, playbook):
    fe = FrontEnd(mode="extend-existing", existing_system_name="Old BMS", devices=(device(slot="x", name="Link", qty=1),),
                  clarifications=(Clarification("note", "tie to the old head end"),), tech_hours=2.0)
    out = render(replace(job, front_end=fe, systems=(system([device()], alternate="deduct"),)), playbook)
    assert out.scope.front_end.existing_system_name == "Old BMS"
    assert [i.text for i in out.scope.front_end.items] == ["Furnish and Install Link", "*tie to the old head end"]
    assert len(out.scope.deduct_alternates) == 1 and out.scope.blocks == ()


def test_nothing_internal_reaches_the_client_text(draft, job):
    printed = "\n".join(l.text for l in draft.lines)
    for system_ in job.systems:
        for d in system_.devices:
            for ref in d.refs:
                assert ref.source_id + " " not in printed and ref.quote not in printed
    assert "Is the existing lobby fan to be kept?" not in printed


# ---- the bridge to the old writers ---------------------------------------------------------------------------

def test_the_old_word_and_price_sheet_writers_take_the_draft(draft, tmp_path):
    from estimator.__main__ import ALIASES, PRICE_BOOK, PROPOSAL_TEMPLATE, WORKUP_TEMPLATE
    from estimator.price_sheet.price_book import PriceBook
    from estimator.price_sheet.write import check_fit, write_price_sheet
    from estimator.proposal.write import write_proposal
    check_fit(draft.scope)
    proposal = write_proposal(draft.scope, PROPOSAL_TEMPLATE, tmp_path / "p.docx")
    sheet = write_price_sheet(draft.scope, WORKUP_TEMPLATE, PriceBook.load(PRICE_BOOK, ALIASES), tmp_path / "s.xlsx")
    from docx import Document
    words = "\n".join(p.text for p in Document(str(proposal)).paragraphs)
    assert "Qty (3) Exhaust Fans EF-1 thru 3" in words and "Fire alarm system by others" in words
    assert "Qty (2) Duct Heaters DH-1, 2" in words
    assert sheet.exists()


# ---- sub-lines (the shape's "under") -------------------------------------------------------------------------

def test_sub_line_prints_one_level_in_directly_after_its_parent_with_no_verb():
    from bms.playbook import load_playbook as repo_playbook
    from bms.model import Header
    pb = repo_playbook()
    devices = (
        device(id="GD.d1", qty=1, slot="sensor", name="CO Sensor", verb="Furnish and Install"),
        device(id="GD.d2", qty=1, slot="panel", name="Gas Detection Panel", verb="Furnish and Install"),
        device(id="GD.d3", qty=1, slot="damper_actuator", name="Damper Actuator", verb="Furnish and Install"),
        device(id="GD.d4", qty=1, slot="horn_strobe", name="Horn / Strobe", verb="Furnish and Install"),
    )
    gd = System(key="GD", name="Gas Detection", shape="Gas Detection", control="field",
                families=(Family("Gas Detection", ("GD-1",), Count(value=1)),), devices=devices)
    j = Job(name="Invented", systems=(gd,), header=Header("00-0000", "Invented", "Invented", ("A",), "2026-10-04"))
    draft = render(j, pb)
    lines = [l for l in draft.lines if l.system_key == "GD" and l.kind == "line"]
    assert [l.device_id for l in lines] == ["GD.d2", "GD.d1", "GD.d4", "GD.d3"]
    assert [l.text for l in lines][:3] == ["Furnish and Install Gas Detection Panel", "CO Sensor", "Horn / Strobe"]
    assert [l.verb for l in lines][:3] == ["Furnish and Install", None, None]
    items = draft.scope.blocks[0].items
    assert [i.depth for i in items if i.kind == "line"] == [1, 2, 2, 1]


# ---- second-run fixes: block order, fixed text, casing, quantity-note labels, singular kinds -------------------

def _repo_job(*systems_) -> Job:
    from bms.model import Header
    return Job(name="Invented", systems=tuple(systems_),
               header=Header("00-0000", "Invented", "Invented", ("A",), "2026-10-04"))


def test_the_playbook_gives_each_shape_a_block_order():
    pb = load_playbook(ROOT / "playbook")
    assert pb.shape("Global Points").order < pb.shape("Exhaust Fan").order
    assert all(isinstance(s.order, int) for s in pb.shapes.values())


def test_global_points_print_first_and_qty_hash_systems_after_the_equipment():
    pb = load_playbook(ROOT / "playbook")
    seq = System(key="SEQ", name="Invented Boiler Plant", shape="Integrated Equipment", control="integration",
                 families=(Family("Invented Boiler Plant", (), Count(value=None)),),
                 sequence_refs=(REF,), devices=(device(id="SEQ.d1", slot="comm_wiring", qty=None,
                                                       name="BMS Communication Wiring to Boiler Controls",
                                                       verb="Furnish"),))
    ef1 = system([device(id="EF.d1")])
    ef2 = system([device(id="EF2.d1")], key="EF2", name="Exhaust Fans",
                 families=(Family("Exhaust Fan", ("EF-8", "EF-9"), Count(value=2)),))
    gp = System(key="GP", name="Global Points", shape="Global Points", control="integration",
                families=(Family("Global Points", (), Count(value=1)),),
                devices=(device(id="GP.d1", slot="oa_temperature", name="Outdoor Air Temperature Sensor", qty=1),))
    out = render(_repo_job(seq, ef1, gp, ef2), pb)
    headers = [b.header for b in out.scope.blocks]
    assert headers[0] == "Qty (1) Global Points"
    assert headers[1].endswith("EF-1, 2") and headers[2].endswith("EF-8, 9")       # the plan's order kept
    assert headers[3].startswith("Qty (##)")
    keys = list(dict.fromkeys(l.system_key for l in out.lines))
    assert keys == ["GP", "EF", "EF2", "SEQ"]


def test_fixed_house_lines_print_with_no_verb_whatever_verb_the_device_carries():
    pb = load_playbook(ROOT / "playbook")
    devices = (
        device(id="X.d1", slot="integration_wiring", name="BMS Communication Wiring to Heater Controller",
               verb="Furnish", qty=1),
        device(id="X.d2", slot="integration_map", name="Map Available Points Via Mfr. Provided BMS Interface",
               verb="Furnish", qty=1),
        device(id="X.d3", slot="rep_assist", name="AME requires mfr. Rep to assist",
               verb="Furnish", qty=1),
        device(id="X.d4", slot="vfd_integration_rep",
               name="AME requires mfr. rep to assist in integration, cost to be incurred by mech contractor.",
               verb="Furnish", qty=1),
    )
    shapeless = system(devices, shape="", families=(Family("Invented Heater", ("H-1",), Count(value=1)),))
    out = render(_repo_job(shapeless), pb)
    lines = [l for l in out.lines if l.kind == "line"]
    rep = pb.wording.integration_trio[2]
    assert [l.text for l in lines] == ["Furnish BMS Communication Wiring to Heater Controller",
                                       pb.wording.integration_trio[1], rep, rep]
    assert [l.verb for l in lines] == ["Furnish", None, None, None]
    # a shaped block's fixed slots print their house text too
    shaped = system([device(id="EF.d5", slot="map_points", verb="Furnish", qty=2, name="map points")])
    assert texts(render(_repo_job(shaped), pb), "EF")[1] == pb.wording.integration_trio[1]


@pytest.mark.parametrize("name, printed", [
    ("space thermostat", "Space Thermostat"),
    ("bacnet interface for co and no sensors w/ vfd", "BACnet Interface for CO and No Sensors w/ VFD"),
    ("Erv discharge air temperature sensor", "ERV Discharge Air Temperature Sensor"),
    ("Interlock Wiring between Invented Fans and the Dampers", "Interlock Wiring between Invented Fans and the Dampers"),
    ("fan start/stop and status", "Fan Start/Stop and Status"),
    ("it room hpac scr control", "IT Room HPAC SCR Control"),
])
def test_device_names_print_in_title_case_with_acronyms_kept(job, name, printed):
    pb = load_playbook(ROOT / "playbook")
    out = texts(render(one(job, system([device(name=name, slot="x", qty=2)], shape="")), pb), "EF")
    assert out[1] == f"Furnish and Install {printed}"


def test_a_header_name_prints_its_acronyms():
    pb = load_playbook(ROOT / "playbook")
    s = system([device()], key="ITXF", name="It Room Transfer Fan", shape="",
               families=(Family("It Room Transfer Fan", ("ITXF-1",), Count(value=1)),))
    assert texts(render(_repo_job(s), pb), "ITXF")[0] == "Qty (1) IT Room Transfer Fan ITXF-1"


def _note(count: Count) -> list[str]:
    pb = load_playbook(ROOT / "playbook")
    fam = Family("Exhaust Fan", ("EF-1",), count)
    out = texts(render(_repo_job(system([device(qty=count.highest())], families=(fam,), shape="")), pb), "EF")
    return [t for t in out if t.startswith("*Qty.")]


def test_count_note_places_are_human_labels_never_the_documents():
    counts_line = "equipment counts for EF"
    count = Count(value=16, sources=(Ref("S1", "located on the sheet text", "Q-100.00, Invented.pdf page 3"),),
                  others=(OtherCount(15, Ref("S2", "distinct tags 15", counts_line)),
                          OtherCount(6, Ref("S3", "riser 6", counts_line)),
                          OtherCount(12, Ref("S4", "plan placements 12", counts_line)),
                          OtherCount(9, Ref("S5", "schedule rows 9", counts_line)),
                          OtherCount(2, Ref("S6", "counted on the sheet picture", "pictures of the invented set"))))
    assert _note(count) == ["*Qty. (16) per Q-100.00; 12 per plan placements; 9 per schedule rows; "
                            "6 per the riser; 2 per the sheet pictures."]


def test_count_note_reads_the_counts_line_part_that_gives_the_value():
    line = "XF (Invented Fan): schedule rows 3; QUANTITY column total 7; distinct tags 3; plan placements 5; riser 4"
    count = Count(value=8, sources=(Ref("S1", "q", "Q-101.00, Invented.pdf page 2"),),
                  others=(OtherCount(7, Ref("S2", line, "equipment counts for XF")),
                          OtherCount(5, Ref("S3", line, "equipment counts for XF")),
                          OtherCount(4, Ref("S4", line, "equipment counts for XF")),
                          OtherCount(3, Ref("S5", line, "equipment counts for XF"))))
    assert _note(count) == ["*Qty. (8) per Q-101.00; 7 per the schedule; 5 per plan placements; 4 per the riser; "
                            "3 per schedule rows."]


def test_count_note_names_every_count_even_one_with_no_label_and_has_no_cap():
    # estimator ruling D6, 2026-10-07: every other count prints; one with no place gets the plain fallback words
    counts_line = "equipment counts for EF"
    count = Count(value=20, sources=(Ref("S1", "distinct tags 20", counts_line),),
                  others=(OtherCount(15, Ref("S2", "distinct tags 15", counts_line)),))
    assert _note(count) == ["*Qty. (20); 15 elsewhere in the documents."]
    many = Count(value=20, sources=(Ref("S1", "q", "Q-1.00, A.pdf page 1"),),
                 others=tuple(OtherCount(n, Ref(f"S{n}", "q", f"Q-{n}.00, A.pdf page {n}")) for n in (3, 9, 5, 7, 11, 2)))
    assert _note(many) == ["*Qty. (20) per Q-1.00; 11 per Q-11.00; 9 per Q-9.00; 7 per Q-7.00; 5 per Q-5.00; "
                           "3 per Q-3.00; 2 per Q-2.00."]


def test_a_two_family_count_note_names_each_family():
    pb = load_playbook(ROOT / "playbook")
    fans = Family("Exhaust Fan", ("EF-1",), Count(value=20, sources=(Ref("S1", "q", "Q-1.00, A.pdf page 1"),),
                                                 others=(OtherCount(15, Ref("S2", "q", "Q-2.00, A.pdf page 2")),)))
    heaters = Family("Unit Heater", ("UH-1",), Count(value=4, sources=(Ref("S3", "q", "Q-3.00, A.pdf page 3"),),
                                                    others=(OtherCount(3, Ref("S4", "q", "Q-4.00, A.pdf page 4")),)))
    out = texts(render(_repo_job(system([device(qty=20)], families=(fans, heaters), shape="")), pb), "EF")
    assert [t for t in out if t.startswith("*Qty.")] == [
        "*Qty. Exhaust Fan (20) per Q-1.00; 15 per Q-2.00; Unit Heater (4) per Q-3.00; 3 per Q-4.00."]
    # a device that carries its family's count does not print that count a second time
    fan_device = replace(device(qty=20), quantity=fans.count)
    out = texts(render(_repo_job(system([fan_device], families=(fans, heaters), shape="")), pb), "EF")
    assert [t for t in out if t.startswith("*Qty.")] == [
        "*Qty. Exhaust Fan (20) per Q-1.00; 15 per Q-2.00; Unit Heater (4) per Q-3.00; 3 per Q-4.00."]


def test_count_note_folds_the_same_value_from_the_same_label():
    count = Count(value=4, sources=(Ref("S1", "q", "Q-2.00, A.pdf page 1"),),
                  others=(OtherCount(2, Ref("S2", "located on the sheet text", "Q-3.00, A.pdf page 2")),
                          OtherCount(2, Ref("S3", "counted on the sheet picture", "Q-3.00 whole sheet, A.pdf page 2")),
                          OtherCount(2, Ref("S4", "riser 2", "equipment counts for EF"))))
    assert _note(count) == ["*Qty. (4) per Q-2.00; 2 per Q-3.00; 2 per the riser."]


@pytest.mark.parametrize("shape, name, tags, n, header", [
    ("Exhaust Fan", "Refuse Room Exhaust Fan", ("RXF-1",), 1, "Qty (1) Exhaust Fan RXF-1"),
    ("Exhaust Fan", "Exhaust Fans", ("EF-1", "EF-2"), 2, "Qty (2) Exhaust Fans EF-1, 2"),
    ("", "Toilet Exhaust Fans", ("TX-1",), 1, "Qty (1) Toilet Exhaust Fan TX-1"),
    ("", "Indoor Units", ("IU-1",), 3, "Qty (3) Indoor Unit IU-1"),
    ("", "Unit Heaters", ("UH-1", "UH-2"), 1, "Qty (1) Unit Heater UH-1, 2"),
    ("Global Points", "Global Points", (), 1, "Qty (1) Global Points"),
])
def test_a_single_unit_header_reads_singular(shape, name, tags, n, header):
    pb = load_playbook(ROOT / "playbook")
    slot = "oa_temperature" if shape == "Global Points" else "controller"
    s = system([device(slot=slot, qty=n)], name=name, shape=shape, families=(Family(name, tags, Count(value=n)),))
    assert texts(render(_repo_job(s), pb), "EF")[0] == header


def test_a_line_whose_quantity_differs_from_the_block_count_is_priced_from_the_book(tmp_path):
    from openpyxl import load_workbook
    from estimator.__main__ import ALIASES, PRICE_BOOK, WORKUP_TEMPLATE
    from estimator.price_sheet.price_book import PriceBook
    from estimator.price_sheet.write import write_price_sheet
    pb = load_playbook(ROOT / "playbook")
    sensor = device(id="EF.d2", slot="space_temp", name="Space Temperature Sensor", qty=6)
    out = render(_repo_job(system([device(id="EF.d1", qty=2), sensor])), pb)
    line = next(l.text for l in out.lines if l.device_id == "EF.d2")
    assert "(Qty. 6 Total)" in line and "–" not in line
    path = write_price_sheet(out.scope, WORKUP_TEMPLATE, PriceBook.load(PRICE_BOOK, ALIASES), tmp_path / "s.xlsx")
    ws = load_workbook(path)["Price Sheet 1"]
    row = next(r for r in range(6, 397) if "Space Temperature Sensor" in str(ws.cell(r, 4).value or ""))
    assert ws.cell(row, 3).value == 6
    assert ws.cell(row, 5).value == 142.0 and ws.cell(row, 8).value is None


def test_a_line_whose_quantity_differs_from_the_block_count_is_priced_from_the_book(tmp_path):
    from openpyxl import load_workbook
    from estimator.__main__ import ALIASES, PRICE_BOOK, WORKUP_TEMPLATE
    from estimator.price_sheet.price_book import PriceBook
    from estimator.price_sheet.write import write_price_sheet
    pb = load_playbook(ROOT / "playbook")
    sensor = device(id="EF.d2", slot="space_temp", name="Space Temperature Sensor", qty=6)
    out = render(_repo_job(system([device(id="EF.d1", qty=2), sensor])), pb)
    line = next(l.text for l in out.lines if l.device_id == "EF.d2")
    assert "(Qty. 6 Total)" in line and "–" not in line
    path = write_price_sheet(out.scope, WORKUP_TEMPLATE, PriceBook.load(PRICE_BOOK, ALIASES), tmp_path / "s.xlsx")
    ws = load_workbook(path)["Price Sheet 1"]
    row = next(r for r in range(6, 397) if "Space Temperature Sensor" in str(ws.cell(r, 4).value or ""))
    assert ws.cell(row, 3).value == 6
    assert ws.cell(row, 5).value == 142.0 and ws.cell(row, 8).value is None


# ---- RND-3: a draft that cannot be printed is a RenderError, never a crash further on ---------------------------

def _gas_detection(*devices) -> Job:
    from bms.model import Header
    gd = System(key="GD", name="Gas Detection", shape="Gas Detection", control="field",
                families=(Family("Gas Detection", ("GD-1",), Count(value=1)),), devices=tuple(devices))
    return Job(name="Invented", systems=(gd,), header=Header("00-0000", "Invented", "Invented", ("A",), "2026-10-04"))


def test_a_blank_named_sub_line_is_a_render_error():
    pb = load_playbook(ROOT / "playbook")
    job = _gas_detection(device(id="GD.d1", qty=1, slot="panel", name="Gas Detection Panel"),
                         device(id="GD.d2", qty=1, slot="sensor", name=""))
    with pytest.raises(RenderError, match="GD"):
        render(job, pb)


def test_negative_tech_hours_are_a_render_error():
    pb = load_playbook(ROOT / "playbook")
    with pytest.raises(RenderError, match="EF"):
        render(_repo_job(system([device(id="EF.d1", qty=2)], tech_hours=-1.0)), pb)


def test_more_alternates_than_the_price_sheet_holds_are_a_render_error():
    pb = load_playbook(ROOT / "playbook")
    adds = [system([device(id=f"A{n}.d1", qty=2)], key=f"A{n}", alternate="add",
                   families=(Family("Exhaust Fan", (f"EF-{n}",), Count(value=1)),)) for n in range(10)]
    with pytest.raises(RenderError, match="price sheets"):
        render(_repo_job(system([device(id="EF.d1", qty=2)]), *adds), pb)


def test_a_control_character_in_a_name_is_printed_without_it(tmp_path):
    from estimator.__main__ import PROPOSAL_TEMPLATE
    from estimator.proposal.write import write_proposal
    pb = load_playbook(ROOT / "playbook")
    out = render(_repo_job(system([device(id="EF.d1", qty=2, name="Control\x01 Valve")],
                                  name="Exhaust\x02 Fans")), pb)
    texts_ = [l.text for l in out.lines]
    assert any("Control Valve" in t for t in texts_)
    assert not any(ch in t for t in texts_ for ch in "\x01\x02")
    assert write_proposal(out.scope, PROPOSAL_TEMPLATE, tmp_path / "p.docx").is_file()


def test_a_system_with_no_lines_carries_its_no_scope_reason(job, playbook):
    """RND-8: the writer's no_scope_reason is the reason the left-out note gives."""
    empty = system([], key="EMPTY", name="Invented Units", no_scope_reason="the invented units are owner furnished")
    out = render(replace(job, systems=(job.system("EF"), empty)), playbook)
    left = [l for l in out.lines if l.system_key == "EMPTY"]
    assert len(left) == 1 and "the invented units are owner furnished" in left[0].text
