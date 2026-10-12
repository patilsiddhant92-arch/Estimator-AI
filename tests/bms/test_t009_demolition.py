"""T-009 task 2 (estimator ruling 22): the proposal prints one demolition exclusion, never both template lines.
A job with a Demolition Scope block prints "Any demolition work not specifically mentioned"; any other job prints
"Any demolition work not mentioned." Every job here is invented."""
from __future__ import annotations

import copy
import json
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from bms.model import DemolitionLine, Job, Ref
from bms.playbook import load_playbook
from bms.render import render
from estimator.proposal import layout
from estimator.proposal.read import ProposalFormatError, read_named_exceptions, read_standard_lines
from estimator.proposal.write import write_proposal
from estimator.scope import Block, Item, Run
from estimator.text import squash

FIXTURES = Path(__file__).resolve().parent / "fixtures"
GENERAL = "Any demolition work not mentioned."
SPECIFIC = "Any demolition work not specifically mentioned"


@pytest.fixture(scope="module")
def playbook():
    return load_playbook(FIXTURES / "sample_playbook")


@pytest.fixture
def job() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def _template() -> Path:
    from estimator.__main__ import PROPOSAL_TEMPLATE
    return Path(PROPOSAL_TEMPLATE)


def _write(scope, tmp_path, template=None) -> Document:
    out = write_proposal(scope, template or _template(), tmp_path / "p.docx")
    return Document(str(out))


def _exclusion_texts(document) -> list[str]:
    texts = [squash(p.text) for p in document.paragraphs if squash(p.text)]
    start = next(i for i, t in enumerate(texts) if t.startswith(layout.EXCLUSIONS))
    end = next(i for i, t in enumerate(texts) if t.startswith(layout.BASE_PRICE))
    return texts[start + 1:end]


def _demolition_lines(document) -> list[str]:
    return [t for t in _exclusion_texts(document) if t.startswith("Any demolition work not")]


def _with_demolition(job: Job) -> Job:
    ref = Ref(source_id="S1", quote="q", where="MD-1, Drawings.pdf page 1")
    return replace(job, demolition=(DemolitionLine(family="EF", action="demo", count=2, refs=(ref,)),))


def _without_demolition(job: Job) -> Job:
    return replace(job, demolition=())


def test_a_job_with_a_demolition_scope_prints_the_specific_line_once(job, playbook, tmp_path):
    scope = render(_with_demolition(job), playbook).scope
    assert any(b.key == "demolition" for b in scope.blocks)
    assert _demolition_lines(_write(scope, tmp_path)) == [SPECIFIC]


def test_a_job_with_no_demolition_prints_the_general_line_once(job, playbook, tmp_path):
    scope = render(_without_demolition(job), playbook).scope
    assert not any(b.key == "demolition" for b in scope.blocks)
    assert _demolition_lines(_write(scope, tmp_path)) == [GENERAL]


def test_the_block_is_found_by_its_heading_when_it_has_no_key(job, playbook, tmp_path):
    scope = render(_without_demolition(job), playbook).scope
    block = Block(header="Demolition Scope", items=(Item(kind="line", depth=1, runs=(Run(text="x"),)),),
                  numbered=False)
    scope = scope.model_copy(update={"blocks": scope.blocks + (block,)})
    assert _demolition_lines(_write(scope, tmp_path)) == [SPECIFIC]


@pytest.mark.parametrize("with_block", [True, False])
def test_every_other_standard_exclusion_is_there_in_order_and_a_named_exception_follows(job, playbook, tmp_path,
                                                                                         with_block):
    base = _with_demolition(job) if with_block else _without_demolition(job)
    scope = render(base, playbook).scope
    scope = scope.model_copy(update={"named_exceptions": ("Any invented named exception.",)})
    texts = _exclusion_texts(_write(scope, tmp_path))
    _, standard = read_standard_lines(_template())
    wanted = SPECIFIC if with_block else GENERAL
    # the template's order, with the one demolition line that stays where the template had it
    expected = [t for t in standard if not t.startswith("Any demolition work not") or t == wanted]
    assert texts[:-1] == expected
    assert texts[-1] == "Any invented named exception."


def _template_without(tmp_path, *, drop: str) -> Path:
    document = Document(str(_template()))
    for p in list(document.paragraphs):
        if squash(p.text).startswith(drop):
            p._p.getparent().remove(p._p)
    out = tmp_path / "one_line_template.docx"
    document.save(str(out))
    return out


@pytest.mark.parametrize("drop, with_block, left", [
    (SPECIFIC, False, GENERAL), (SPECIFIC, True, GENERAL),
    (GENERAL, True, SPECIFIC), (GENERAL, False, SPECIFIC)])
def test_a_template_with_only_one_of_the_lines_writes_and_prints_that_line(job, playbook, tmp_path, drop,
                                                                           with_block, left):
    template = _template_without(tmp_path, drop=drop)
    base = _with_demolition(job) if with_block else _without_demolition(job)
    document = _write(render(base, playbook).scope, tmp_path, template)
    assert _demolition_lines(document) == [left]


def test_a_template_with_neither_line_writes_with_no_error(job, playbook, tmp_path):
    document = Document(str(_template()))
    for p in list(document.paragraphs):
        if squash(p.text).startswith("Any demolition work not"):
            p._p.getparent().remove(p._p)
    template = tmp_path / "none.docx"
    document.save(str(template))
    assert _demolition_lines(_write(render(job, playbook).scope, tmp_path, template)) == []


def _sent(tmp_path, mutate) -> Document:
    document = Document(str(_template()))
    mutate(document)
    return document


def test_read_named_exceptions_accepts_both_lines_or_either_missing(tmp_path):
    template = _template()
    assert read_named_exceptions(Document(str(template)), template) == ()
    for drop in (SPECIFIC, GENERAL):
        def cut(document, drop=drop):
            for p in list(document.paragraphs):
                if squash(p.text).startswith(drop):
                    p._p.getparent().remove(p._p)
        assert read_named_exceptions(_sent(tmp_path, cut), template) == ()


def test_read_named_exceptions_still_returns_the_named_exceptions_after_a_written_proposal(job, playbook, tmp_path):
    scope = render(_without_demolition(job), playbook).scope
    scope = scope.model_copy(update={"named_exceptions": ("Any invented named exception.",)})
    document = _write(scope, tmp_path)
    assert read_named_exceptions(document, _template()) == ("Any invented named exception.",)


def test_read_named_exceptions_refuses_a_proposal_that_lacks_another_standard_line(tmp_path):
    def cut(document):
        for p in list(document.paragraphs):
            if squash(p.text).startswith("Any gas detection system"):
                p._p.getparent().remove(p._p)
    with pytest.raises(ProposalFormatError):
        read_named_exceptions(_sent(tmp_path, cut), _template())


def test_the_two_template_texts_are_in_layout():
    assert layout.DEMOLITION_GENERAL == GENERAL.rstrip(".") + "." and layout.DEMOLITION_SPECIFIC == SPECIFIC


# --- review of record: only the exclusions section is searched, never the scope lines ---------------------------

def _with_scope_line(scope, text: str):
    block = Block(header="Invented Other Work", items=(Item(kind="line", depth=1, runs=(Run(text=text),)),),
                  numbered=False)
    return scope.model_copy(update={"blocks": scope.blocks + (block,)})


def _paragraphs_with(document, text: str) -> list[str]:
    return [p.text for p in document.paragraphs if squash(p.text) == squash(text)]


def test_a_scope_line_with_the_dropped_text_stays_when_the_job_has_no_demolition_block(job, playbook, tmp_path):
    scope = _with_scope_line(render(_without_demolition(job), playbook).scope, SPECIFIC)
    document = _write(scope, tmp_path)
    assert len(_paragraphs_with(document, SPECIFIC)) == 1          # the scope line; the specific exclusion is dropped
    assert _demolition_lines(document) == [GENERAL]               # the exclusions list still has exactly one


def test_a_scope_line_with_the_dropped_text_stays_when_the_job_has_a_demolition_block(job, playbook, tmp_path):
    scope = _with_scope_line(render(_with_demolition(job), playbook).scope, GENERAL)
    document = _write(scope, tmp_path)
    assert len(_paragraphs_with(document, GENERAL)) == 1           # the scope line; the general exclusion is dropped
    assert _demolition_lines(document) == [SPECIFIC]


def test_the_scope_line_is_not_taken_for_the_line_to_keep(job, playbook, tmp_path):
    # a scope line that has the text of the line to keep must not stop the exclusion from being dropped or kept right:
    # the template's own exclusion line is in the exclusions list, and the job's scope line is not counted for it
    scope = _with_scope_line(render(_without_demolition(job), playbook).scope, GENERAL)
    document = _write(scope, tmp_path)
    assert len(_paragraphs_with(document, GENERAL)) == 2           # scope line plus the kept general exclusion
    assert _demolition_lines(document) == [GENERAL]
