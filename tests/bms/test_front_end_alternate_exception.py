"""Four fixes from a real run (no real job text here; every ledger, dossier and job is invented):

1. the front end: the spec pass's front-end lines and client notes reach the plan, one writer request writes the
   front end, and a must_fix finding about the front end reaches a revise;
2. a bid alternate the documents name in a schedule's title, a schedule note or a plan note ("ADD ALTERNATE",
   "ADD ALT", "ALTERNATE #1", "DEDUCT ALTERNATE") makes the block an add / deduct alternate, with its citation;
3. a block whose equipment the documents list but which ends with no AME line prints the Playbook's named exception
   ("EXCEPTION TO ANY WORK WITH Qty (n) <tags> - <reason>");
4. the proposal's extend-existing opening names the existing system (or the placeholder), never the template's
   sample words.
"""
from __future__ import annotations

import copy
import dataclasses
import json
import sys
from pathlib import Path

import docx
import pytest

from estimator.ledger.facts import SETTLED, Ledger
from estimator.proposal import layout
from estimator.proposal.write import write_proposal
from estimator.scope import FrontEnd as ScopeFrontEnd
from estimator.text import squash
from estimator.writing.dossier import Dossier

HERE = Path(__file__).resolve().parent
for folder in (HERE, HERE.parent):
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

from bms import ai, run
from bms.model import (Clarification, Count, Device, Family, Finding, FrontEnd, Job, Header, Ref, ScopeException,
                       System)
from bms.plan import plan_job
from bms.playbook import load_playbook
from bms.render import exception_reason, render, scope_exceptions

from test_run import answer_all, job_folder  # noqa: E402,F401  (test_run.py's invented job folder)

PLAYBOOK = load_playbook(HERE / "fixtures" / "sample_playbook")

SOURCES = {
    "S1": "M-401, Drawings.pdf page 7 (schedule)",
    "S2": "M-601, Drawings.pdf page 9 (sheet text)",
    "S3": "M-602, Drawings.pdf page 10 (sheet text)",
    "S4": "equipment counts for AP (schedules and plans)",
    "S5": "M-101, Drawings.pdf page 2 (sheet text)",
}


def _ledger() -> Ledger:
    ledger = Ledger(dict(SOURCES))
    ledger.job_name = "Invented Job"
    ledger.add("front_end", "mode", "new-system", "S2", "PROVIDE A NEW DDC SYSTEM", "spec", SETTLED)
    return ledger


def _family(ledger, prefix, name, tags, count, source="S1", quote=None):
    ledger.add("family", prefix, name, source, quote or f"{prefix} {name}", "schedules", SETTLED)
    ledger.add("unit", prefix, tags, source, f"Tags: {tags}", "schedules", SETTLED)
    ledger.add("count", prefix, count, source, f"{prefix} rows {count}", "schedules", SETTLED)


def _dossier(text: str = "") -> Dossier:
    return Dossier(job_name="Invented Job", text=text, images=(), sources=dict(SOURCES), families=(),
                   device_totals={}, counts={}, stats={})


def _plan(ledger, text: str = "") -> Job:
    return plan_job(ledger, None, _dossier(text), PLAYBOOK)


# ---- 1. the front end -----------------------------------------------------------------------------------------------

def _front_ledger() -> Ledger:
    ledger = _ledger()
    ledger.add("front_end", "line", "Prewired local DDC control panel in the equipment room", "S2",
               "PROVIDE A PREWIRED LOCAL CONTROL PANEL", "spec", SETTLED)
    ledger.add("front_end", "line", "Operator workstation graphics", "S3", "GRAPHICS PROGRAMMING", "spec", SETTLED)
    ledger.add("front_end", "note", "** Lighting integration protocol to be confirmed", "S3",
               "INTEGRATE THE LIGHTING CONTROL SYSTEM", "spec", SETTLED)
    return ledger


def test_the_plan_carries_the_front_end_lines_notes_and_their_citations():
    job = _plan(_front_ledger())
    fe = job.front_end
    assert fe.mode == "new-system"
    assert fe.lines == ("Prewired local DDC control panel in the equipment room", "Operator workstation graphics")
    assert fe.notes == ("** Lighting integration protocol to be confirmed",)
    assert {(r.source_id, r.quote) for r in fe.refs} >= {("S2", "PROVIDE A PREWIRED LOCAL CONTROL PANEL"),
                                                         ("S3", "GRAPHICS PROGRAMMING"),
                                                         ("S3", "INTEGRATE THE LIGHTING CONTROL SYSTEM")}
    assert fe.devices == () and fe.clarifications == ()          # the writer writes them
    assert Job.from_dict(job.to_dict()) == job
    assert FrontEnd.from_dict({"mode": "new-system"}).lines == ()   # an older Job.json still loads


class _Queue:
    def __init__(self, *answers):
        self.answers, self.asked, self.waiting = list(answers), [], 0

    def ask(self, kind, name, instructions, system_prompt, prompt, schema, extra=None, images=()):
        self.asked.append(dict(kind=kind, name=name, prompt=prompt, extra=extra))
        if not self.answers:
            raise ai.NoAnswer("no answer queued", waiting=True)
        return copy.deepcopy(self.answers.pop(0))


_TEXT = ("== 8. CONTROLS SPECIFICATION ==\n--- M-601, Drawings.pdf page 9 ---\n"
         "[S2] PROVIDE A PREWIRED LOCAL CONTROL PANEL WITH DIRECT DIGITAL CONTROL.\n"
         "--- M-602, Drawings.pdf page 10 ---\n"
         "[S3] FURNISH GRAPHICS PROGRAMMING. INTEGRATE THE LIGHTING CONTROL SYSTEM.\n")


def _fe_answer(devices=None, note="Lighting integration protocol to be confirmed"):
    device = {"id": "x", "slot": "network_panel", "name": "Prewired Local DDC Control Panel",
              "verb": "Furnish and Install", "furnished_by": "AME", "installed_by": "AME", "wired_by": "AME",
              "quantity": {"value": 1, "sources": [], "others": [], "basis": "highest"},
              "refs": [{"source_id": "S2", "quote": "PROVIDE A PREWIRED LOCAL CONTROL PANEL", "where": "M-601"}],
              "assumption": None, "flag": "none"}
    return {"key": ai.FRONT_END_KEY, "name": ai.FRONT_END_NAME, "shape": "", "control": "field", "families": [],
            "sequence_refs": [], "devices": devices if devices is not None else [device],
            "clarifications": [{"kind": "note", "text": note}], "tech_hours": None, "alternate": "base",
            "no_scope_reason": None, "serves": []}


def test_one_writer_request_writes_the_front_end_from_its_facts():
    front = _plan(_front_ledger(), _TEXT).front_end
    evidence = ai.front_end_evidence(front, _dossier(_TEXT))
    assert "Operator workstation graphics" in evidence and "[S2] PROVIDE A PREWIRED" in evidence
    queue = _Queue(_fe_answer())
    written = ai.write_front_end(front, evidence, PLAYBOOK, queue)
    (asked,) = queue.asked
    assert asked["kind"] == ai.WRITE_SYSTEM and asked["extra"]["plan_key"] == ai.FRONT_END_KEY
    assert asked["extra"]["step"] == "write"
    assert [d.name for d in written.devices] == ["Prewired Local DDC Control Panel"]
    assert written.devices[0].id == f"{ai.FRONT_END_KEY}.d1"
    assert [c.text for c in written.clarifications] == ["Lighting integration protocol to be confirmed"]
    assert (written.mode, written.lines, written.refs) == (front.mode, front.lines, front.refs)
    # written once: a front end that has its lines is not asked again; one with no facts is never asked
    assert ai.write_front_end(written, evidence, PLAYBOOK, queue) is written
    bare = FrontEnd(mode="new-system")
    assert ai.write_front_end(bare, "", PLAYBOOK, queue) is bare and len(queue.asked) == 1


def test_front_end_findings_are_told_from_whole_job_findings():
    front = FrontEnd(mode="new-system", devices=(Device.from_dict(dict(
        _fe_answer()["devices"][0], id="front-end.d1")),))
    about = [Finding("Front end lines cite the specification", "", None, "must_fix", "no workstation line", "add"),
             Finding("Every line cites this job", "", "front-end.d1", "must_fix", "bad quote", "fix"),
             Finding("One clarification", "", None, "must_fix", "The front end has two notes", "join them")]
    other = [Finding("Named exceptions", "", None, "must_fix", "the named exception prints S4", "name it"),
             Finding("Front end lines cite the specification", "EF", None, "must_fix", "front end", "x")]
    assert all(ai.is_front_end_finding(f, front) for f in about)
    assert not any(ai.is_front_end_finding(f, front) for f in other)


def test_a_must_fix_on_the_front_end_reaches_its_revise():
    front = _plan(_front_ledger(), _TEXT).front_end
    evidence = ai.front_end_evidence(front, _dossier(_TEXT))
    written = ai.write_front_end(front, evidence, PLAYBOOK, _Queue(_fe_answer()))
    finding = Finding("Front end lines cite the specification", "", None, "must_fix",
                      "The front end has no line for the operator workstation graphics the spec calls for",
                      "Add the graphics line citing the spec.")
    graphics = dict(_fe_answer()["devices"][0], id="front-end.d2", slot="graphics", name="Graphics Programming",
                    verb="Furnish", refs=[{"source_id": "S3", "quote": "FURNISH GRAPHICS PROGRAMMING",
                                           "where": "M-602"}])
    answer = _fe_answer([dict(_fe_answer()["devices"][0], id="front-end.d1"), graphics])
    queue = _Queue(answer)
    revised = ai.revise_front_end(written, [finding], evidence, PLAYBOOK, queue)
    (asked,) = queue.asked
    assert asked["extra"]["step"] == "revise" and asked["extra"]["plan_key"] == ai.FRONT_END_KEY
    assert "operator workstation graphics" in asked["prompt"]
    assert [d.name for d in revised.devices] == ["Prewired Local DDC Control Panel", "Graphics Programming"]
    # no front-end must_fix: not asked
    whole = Finding("Named exceptions", "", None, "must_fix", "the named exception prints S4", "name it")
    assert ai.revise_front_end(written, [whole], evidence, PLAYBOOK, queue) is written and len(queue.asked) == 1
    # the run closes a front-end QC finding the revise answered; a whole-job one stays open
    before, after = Job("J", front_end=written), Job("J", front_end=revised)
    assert not run._untouched(finding, before, after)
    assert run._untouched(whole, before, after)
    assert run._untouched(finding, before, before)


def test_the_run_writes_and_revises_the_front_end(job_folder, tmp_path, monkeypatch):
    """End to end on test_run.py's invented job: the planner's front-end lines become one write request, its
    devices are in Job.json and the proposal, and a must_fix about the front end makes a revise request."""
    folder = job_folder
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    real_plan = run.plan_job

    def with_front_end(ledger, catalogue, dossier, playbook):
        job = real_plan(ledger, catalogue, dossier, playbook)
        ref = next(r for s in job.systems for f in s.families for r in f.count.sources)
        return dataclasses.replace(job, front_end=dataclasses.replace(
            job.front_end, lines=("Local DDC control panel",), refs=(ref,)))
    monkeypatch.setattr(run, "plan_job", with_front_end)
    real_validate = run.validate
    monkeypatch.setattr(run, "validate", lambda job, playbook: [*real_validate(job, playbook), Finding(
        "Front end lines cite the specification", "", None, "must_fix", "the front end misses a line",
        "write it")])
    out, cache = tmp_path / "out", tmp_path / "cache"
    kinds: list[tuple[str, str]] = []
    result = None
    for _ in range(25):
        result = run.estimate(folder, out, cache, "chat")
        if result.status == "done":
            break
        for path in sorted((cache / "judge-chat" / "requests").glob("*.json")):
            request = json.loads(path.read_text(encoding="utf-8"))
            kinds.append((request.get("plan_key", ""), request.get("step", "")))
        answer_all(cache)
    assert result is not None and result.status == "done", result
    assert (ai.FRONT_END_KEY, "write") in kinds and (ai.FRONT_END_KEY, "revise") in kinds
    saved = json.loads(result.job_file.read_text(encoding="utf-8"))["front_end"]
    assert saved["devices"] and saved["devices"][0]["id"].startswith(f"{ai.FRONT_END_KEY}.")
    assert saved["lines"] == ["Local DDC control panel"]
    text = [squash(p.text) for p in docx.Document(result.proposal).paragraphs]
    assert any("Space Temperature Sensor" in t for t in text)


# ---- 2. an alternate the schedule's title, a schedule note or a plan note names ---------------------------------------

_SCHEDULE = ("== 4. SCHEDULES ==\n"
             "[S1] AIR PURIFIER (ADD ALTERNATE) — M-401, Drawings.pdf page 7 (2 rows)\n"
             "    Columns: TAG | SERVES\n"
             "[S6] AP-1 | AHU-1\n"
             "[S7] AP-2 | AHU-2\n")


def test_a_schedule_title_that_names_an_add_alternate_makes_an_add_alternate_with_its_citation():
    ledger = _ledger()
    _family(ledger, "AP", "Air Purifier", "AP-1, 2", 2, quote="AIR PURIFIER")
    _family(ledger, "EF", "Exhaust Fan", "EF-1", 1, source="S5")
    job = _plan(ledger, _SCHEDULE)
    ap, ef = job.system("AP"), job.system("EF")
    assert ap.alternate == "add" and ef.alternate == "base"
    assert Ref("S1", "AIR PURIFIER (ADD ALTERNATE)", "M-401, Drawings.pdf page 7") in ap.sequence_refs
    assert job.questions == ()


def test_the_title_in_the_family_facts_place_is_read_without_the_dossier_text():
    ledger = _ledger()
    ledger.sources["S8"] = "M-401 AIR PURIFIER (ADD ALT), Drawings.pdf page 7"
    _family(ledger, "AP", "Air Purifier", "AP-1, 2", 2, source="S8", quote="AIR PURIFIER")
    ap = _plan(ledger).system("AP")
    assert ap.alternate == "add" and ap.sequence_refs[-1].source_id == "S8"


@pytest.mark.parametrize("note, kind, asked", [
    ("PROVIDE AP-1 AND AP-2 AS DEDUCT ALTERNATE NO. 2.", "deduct", False),
    ("AIR PURIFIERS AP-1, AP-2 ARE ADD ALT #1.", "add", False),
    ("AP-1 AND AP-2: ALTERNATE #1.", "add", True),
    ("AP-1 AND AP-2 SHALL ALTERNATE LEAD AND LAG.", "base", False),
    ("AP-1 AND AP-2 ARE NOT AN ALTERNATE.", "base", False),
])
def test_a_plan_or_schedule_note_naming_the_alternate(note, kind, asked):
    ledger = _ledger()
    _family(ledger, "AP", "Air Purifier", "AP-1, 2", 2, quote="AIR PURIFIER")
    ledger.add("note", "general", note, "S5", note, "plans", SETTLED)
    job = _plan(ledger)
    ap = job.system("AP")
    assert ap.alternate == kind
    if kind != "base":
        assert ap.sequence_refs[-1] == Ref("S5", note, "M-101, Drawings.pdf page 2")
    else:
        assert ap.sequence_refs == ()
    assert bool([q for q in job.questions if "add or a deduct" in q]) == asked


def test_an_alternate_note_still_wins_and_adds_no_second_ref():
    ledger = _ledger()
    _family(ledger, "AP", "Air Purifier", "AP-1, 2", 2, quote="AIR PURIFIER")
    ledger.add("note", "alternate", "deduct AP", "S5", "AP AS A DEDUCT", "sequence", SETTLED)
    ap = _plan(ledger, _SCHEDULE).system("AP")
    assert ap.alternate == "deduct" and ap.sequence_refs == ()


# ---- 3. a listed block with no AME line prints a named exception ------------------------------------------------------

_COUNT_REF = Ref("S4", "EF rows 2", "M-401, Drawings.pdf page 7")


def _empty(key="EF", reason="EF-1 and EF-2 run on a wall mounted thermostat the exhaust fan schedule on M-401 gives "
           "to the mechanical contractor, and no document gives the BMS any work on them.", tags=("EF-1", "EF-2")):
    return System(key=key, name="Exhaust Fan", shape="", control="field",
                  families=(Family("Exhaust Fan", tags, Count(2, (_COUNT_REF,), basis="schedule")),),
                  no_scope_reason=reason)


def test_a_listed_block_with_no_lines_becomes_the_playbooks_named_exception():
    job = Job("J", systems=(_empty(),), front_end=FrontEnd(mode="new-system"))
    (x,) = scope_exceptions(job)
    assert x.system_key == "EF" and x.refs == (_COUNT_REF,)
    assert x.text == "EXCEPTION TO ANY WORK WITH Qty (2) EF-1, 2 - WALL MOUNTED THERMOSTAT BY MECHANICAL CONTRACTOR"
    # idempotent: the run puts it in the job, the renderer adds no second one
    assert scope_exceptions(dataclasses.replace(job, exceptions=(x,))) == (x,)


def test_the_exception_prints_in_the_proposal_scope():
    header = Header("26-0001", "Invented", "Invented MEP", ("M-401",), "2026-10-09")
    job = Job("J", systems=(_empty(),), front_end=FrontEnd(mode="new-system"), header=header)
    draft = render(job, PLAYBOOK)
    assert draft.scope.named_exceptions == (
        "EXCEPTION TO ANY WORK WITH Qty (2) EF-1, 2 - WALL MOUNTED THERMOSTAT BY MECHANICAL CONTRACTOR",)
    assert any("left out" in line.text for line in draft.lines)            # still noted for the Evidence Package


def test_blocks_with_lines_with_their_own_exception_or_with_no_listed_equipment_get_none():
    device = Device.from_dict(dict(_fe_answer()["devices"][0], id="EF.d1"))
    written = dataclasses.replace(_empty(), devices=(device,), no_scope_reason=None)
    own = ScopeException("EF", "EXCEPTION TO ANY WORK WITH Qty (2) EF-1, 2 - INTEGRAL THERMOSTAT")
    sequence_only = System(key="HT", name="Heat Trace", shape="", control="field",
                           families=(Family("Heat Trace", (), Count(None)),), no_scope_reason="nothing for the BMS")
    assert scope_exceptions(Job("J", systems=(written,))) == ()
    assert scope_exceptions(Job("J", systems=(_empty(),), exceptions=(own,))) == (own,)
    assert scope_exceptions(Job("J", systems=(sequence_only,))) == ()


@pytest.mark.parametrize("reason, words", [
    ("The units have an integral thermostat furnished with the unit.", "INTEGRAL THERMOSTAT"),
    ("Each heater is on a line voltage thermostat by the electrical contractor.", "LINE-VOLTAGE THERMOSTAT"),
    ("The fans run from the manufacturer's on/off switch.", "MANUFACTURER ON/OFF SWITCH"),
    ("The schedule marks these fans not in contract.", "NOT IN CONTRACT"),
    ("Controls are by the unit manufacturer and the documents ask nothing of the BMS.", "CONTROLS BY UNIT MFR."),
    ("No document gives the BMS any work on them.", "NO BMS WORK IN THE CONTRACT DOCUMENTS"),
])
def test_the_exception_reason_is_read_in_the_estimators_words(reason, words):
    assert exception_reason(_empty(reason=reason)) == words


# ---- 4. the extend-existing opening names the existing system -------------------------------------------------------

def _opening(job, template, tmp_path) -> str:
    texts = [squash(p.text) for p in docx.Document(str(write_proposal(job, template, tmp_path / "p.docx"))).paragraphs]
    (opening,) = [t for t in texts if t.startswith(layout.OPENING)]
    return opening


@pytest.mark.parametrize("name, said", [
    ("Johnson Controls Metasys", "an extension of the existing Johnson Controls Metasys BMS system."),
    ("Honeywell WEBs BMS", "an extension of the existing Honeywell WEBs BMS system."),
    ("[existing BMS name]", "an extension of the existing [existing BMS name] BMS system."),
])
def test_the_existing_opening_names_the_existing_system(bare_job, proposal_template, tmp_path, name, said):
    job = bare_job.model_copy(update={"front_end": ScopeFrontEnd(mode="extend-existing", existing_system_name=name)})
    opening = _opening(job, proposal_template, tmp_path)
    assert opening.endswith(said) and "District Wide" not in opening
    texts = [squash(p.text) for p in docx.Document(str(tmp_path / "p.docx")).paragraphs]
    assert any(t.startswith(f"Seamlessly Tie in to the Existing {name.removesuffix(' BMS')} BMS") for t in texts)


def test_a_new_system_keeps_the_templates_new_opening(bare_job, proposal_template, tmp_path):
    opening = _opening(bare_job, proposal_template, tmp_path)
    assert "extension" not in opening and "District Wide" not in opening
    assert "Automated Temperature Controls system" in opening
