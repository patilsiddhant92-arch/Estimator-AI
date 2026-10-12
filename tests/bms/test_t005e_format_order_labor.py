"""T-005e: the Master Template's scope-of-work format, the GEM order of points, off-hours work and prevailing wage.

Invented fixtures only (the sample job and sample Playbook in tests/bms/fixtures, the repo's own templates).
"""
from __future__ import annotations

import json
from dataclasses import dataclass, replace
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from bms import order
from bms.labor import (HEADING, LABOR_KEY, calls_for_off_hours, hours_of, labor_block, labor_conditions, state_of,
                       with_labor)
from bms.model import Count, Device, Family, FrontEnd, Job, Ref, System
from bms.playbook import load_playbook
from bms.render import render

FIXTURES = Path(__file__).resolve().parent / "fixtures"
REF = Ref(source_id="S1", quote="q", where="M-1, Drawings.pdf page 1")


@pytest.fixture(scope="module")
def playbook():
    return load_playbook(FIXTURES / "sample_playbook")


@pytest.fixture(scope="module")
def house():
    return load_playbook()          # the real Playbook (shapes from the Master Template)


@pytest.fixture
def job() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def dev(n: int, name: str, verb: str = "Furnish and Install", slot: str = "other_device", qty: int | None = 2,
        key: str = "X", **kw) -> Device:
    base = dict(id=f"{key}.d{n}", slot=slot, name=name, verb=verb, furnished_by="AME", installed_by="AME",
                wired_by="AME", quantity=Count(value=qty), refs=(REF,))
    base.update(kw)
    return Device(**base)


def texts(draft, key: str) -> list[str]:
    return [l.text for l in draft.lines if l.system_key == key]


# ---- 2. the GEM order of points ----------------------------------------------------------------------------------

@pytest.mark.parametrize("text, verb, expected", [
    ("Furnish and Install DDC Controller", "Furnish and Install", 1),
    ("Furnish DDC Controller to the Factory for Installation", "Furnish", 1),
    ("Furnish and Install Discharge Air Temperature Sensor", "Furnish and Install", 2),
    ("Furnish and Install High Static Pressure Switch", "Furnish and Install", 3),
    ("Furnish and Install Freezestat", "Furnish and Install", 3),
    ("Furnish and Install Leak Sensor", "Furnish and Install", 4),
    ("Furnish and Install Current Switch", "Furnish and Install", 5),
    ("Furnish and Install Interposing Relay", "Furnish and Install", 6),
    ("Furnish and Install Outside Air Damper Actuator", "Furnish and Install", 7),
    ("Furnish Supply Fan Start/Stop and Status", "Furnish", 8),
    ("Furnish Supply Fan VFD Control", "Furnish", 8),
    ("Furnish Air-Cooled Chiller Enable (Qty. 2 Total)", "Furnish", 9),
    ("Furnish Open / Close Status to BMS via FSD End Switch", "Furnish", 10),
    ("Furnish Electric Heat Enable (DO)", "Furnish", 11),
    ("Furnish and Wire Filter DP Switch", "Furnish and Wire", 14),
    ("Furnish and Wire Outside Air Flow Measuring Station", "Furnish and Wire", 15),
    ("Furnish and Wire Hot Water Control Valve and Actuator", "Furnish and Wire", 16),
    ("Install and Wire Space Thermostat (Furnished by Unit Mfr.)", "Install and Wire", 17),
    ("Wire Supply Fan Flow Transmitter (Furnished by Others)", "Wire", 18),
    ("Wire Duct Smoke Detector Relay (Relay Furnished by FA)", "Wire", 13),
    ("Furnish BMS Communication Wiring to Factory Controls", "Furnish", 20),     # lead ruling: Playbook, last
    ("Map Available Points Via Mfr. Provided BMS Interface", None, 20),
])
def test_each_line_takes_its_gem_position(text, verb, expected):
    assert order.position(text, verb) == expected


def test_slot_kinds_place_interlocks_wiring_and_integration():
    assert order.position("Interlock with Associated AHU Operation", None, "other_interlock", "interlock") == 18
    # lead ruling 2026-10-11 (the Playbook wins, rules.md:486, :604): plain work statements (19) before the
    # integration group (20)
    assert order.position("Coordination of Testing with the TAB Contractor", None, "other_wiring", "wiring") == 19
    assert order.position("Wire Existing Panel to New Controller", "Wire", "other_wiring", "wiring") == 18
    assert order.position("Run 24V Power Loop to VAVs", None, "power_loop", "wiring") == 12
    assert order.position("anything", "Furnish", "comm_wiring", "integration") == 20


def test_an_air_handler_prints_in_the_master_template_order_not_the_writers(job, house):
    # the writer's order, scrambled: integration first, valve, start/stop, sensor, damper, controller, filter, Install
    # and Wire, safety switch, leak
    devices = [
        dev(1, "BMS Communication Wiring to Factory Controls", "Furnish", "comm_wiring", key="AHU"),
        dev(2, "Hot Water Control Valve and Actuator", "Furnish and Wire", "hot_water_valve", key="AHU"),
        dev(3, "Supply Fan Start/Stop and Status", "Furnish", "supply_fan_start_stop", key="AHU"),
        dev(4, "Mixed Air Temperature Sensor", "Furnish and Install", "mixed_air_temperature", key="AHU"),
        dev(5, "Outside Air Damper Actuator", "Furnish and Install", "oa_damper", key="AHU"),
        dev(6, "DDC Controller", "Furnish and Install", "controller", key="AHU"),
        dev(7, "Filter DP Switch", "Furnish and Wire", "filter_dp", key="AHU"),
        dev(8, "Duct Smoke Detector", "Install and Wire", "other_install_and_wire", key="AHU",
            furnished_by="Unit Mfr."),
        dev(9, "Freezestat", "Furnish and Install", "freezestat", key="AHU"),
        dev(10, "Leak Sensor", "Furnish and Install", "leak_sensor", key="AHU"),
        dev(11, "Supply Air Temperature Sensor", "Furnish and Install", "discharge_temperature", key="AHU"),
    ]
    ahu = System(key="AHU", name="Air Handling Units", shape="Air Handling Unit", control="field",
                 families=(Family("Air Handling Unit", ("AHU-1", "AHU-2"), Count(value=2)),), devices=tuple(devices))
    out = texts(render(replace(job, systems=(ahu,)), house), "AHU")
    # T-005e addendum (GEM v4 ~396-402): the Master Template entry (the shape) sets the order of its own slots; the
    # catch-all Install and Wire line is placed by the standard order (17: after the valves); integration last
    assert out[1:] == [
        "Furnish and Install DDC Controller",
        "Furnish and Install Supply Air Temperature Sensor",
        "Furnish and Install Mixed Air Temperature Sensor",
        "Furnish Supply Fan Start/Stop and Status",
        "Furnish and Install Freezestat",
        "Furnish and Wire Filter DP Switch",
        "Furnish and Install Outside Air Damper Actuator",
        "Furnish and Wire Hot Water Control Valve and Actuator",
        "Furnish and Install Leak Sensor",
        "Install and Wire Duct Smoke Detector (Furnished by Unit Mfr.)",
        "Furnish BMS Communication Wiring to Factory Controls",
    ]


def test_a_shapeless_block_follows_the_gem_order_too(job, playbook):
    devices = [dev(1, "Control Valve", "Furnish and Wire", "x"), dev(2, "Fan Status", "Furnish", "y"),
               dev(3, "Supply Temperature Sensor", "Furnish and Install", "z"), dev(4, "DDC Controller", slot="w")]
    sys_ = System(key="Q", name="Odd Units", shape="", control="field",
                  families=(Family("Odd Unit", ("OU-1", "OU-2"), Count(value=2)),), devices=tuple(devices))
    assert texts(render(replace(job, systems=(sys_,)), playbook), "Q")[1:] == [
        "Furnish and Install DDC Controller", "Furnish and Install Supply Temperature Sensor",
        "Furnish Fan Status", "Furnish and Wire Control Valve"]


def test_the_rep_line_stays_under_its_map_line_and_the_trio_prints_last(job, playbook):
    lines = [l for l in render(job, playbook).lines if l.system_key == "RTU" and l.kind == "line"]
    assert [l.text for l in lines][-3:] == [
        "Furnish BMS Communication Wiring to RTU Factory Controls",
        "Map Available Points Via Mfr. Provided BMS Interface",
        "AME requires mfr. Rep to assist in integration, cost to be incurred by mech contractor"]


# ---- 1. the Master Template's format: sub-headings, the existing BMS front end -------------------------------------

def test_a_grouped_system_gets_its_entrys_group_line(job, house):  # T-005e evening correction: as the entry prints
    pumps = System(key="CHW", name="Chilled Water System", shape="Chilled Water System", control="field",
                   families=(Family("Chiller", ("CH-1", "CH-2"), Count(value=2)),
                             Family("Chilled Water Pump", ("CHWP-1", "CHWP-2"), Count(value=2))),
                   devices=(dev(1, "DDC Controller", slot="controller", key="CHW", qty=1),))
    draft = render(replace(job, systems=(pumps,), demolition=()), house)
    headers = [b.header for b in draft.scope.blocks]
    assert headers[0] == "Chilled Water System" and not draft.scope.blocks[0].numbered
    assert headers[1].startswith("Qty (2) Chiller") and headers[2].startswith("Qty (2) Chilled Water Pump")
    assert all(b.numbered and b.joined for b in draft.scope.blocks[1:])


def test_a_single_family_block_gets_no_sub_heading(job, house):
    one = System(key="CHW", name="Chilled Water System", shape="Chilled Water System", control="field",
                 families=(Family("Chiller", ("CH-1",), Count(value=1)),),
                 devices=(dev(1, "DDC Controller", slot="controller", key="CHW", qty=1),))
    assert all(b.numbered for b in render(replace(job, systems=(one,), demolition=()), house).scope.blocks)


def _fe(n: int, name: str, verb: str = "Furnish") -> Device:
    return dev(n, name, verb, slot=name.lower().replace(" ", "_"), qty=1, key="FE")


def test_an_existing_bms_front_end_prints_the_master_template_lines_in_its_order(job, playbook):
    # T-005e (merge with T-005c): one of AME's three BMS names, so T-005c's red *** flag stays out of the lines
    front = FrontEnd(mode="extend-existing", existing_system_name="Niagara N4",
                     devices=(_fe(1, "Necessary Training on New Controls"), _fe(2, "PC Workstation")))
    items = render(replace(job, front_end=front), playbook).scope.front_end.items
    assert [i.text for i in items] == [
        "Furnish Network Level Control Panel",                  # the template's, the job does not speak to it
        "Furnish PC Workstation",                               # the writer's, citing the job
        "Furnish 3-D Graphics to match existing",               # the template's
        "Furnish Necessary Training on New Controls",           # the writer's, citing the job: not printed twice
    ]
    assert [i.quantity for i in items] == [1, 1, 1, 1]


def test_an_existing_front_end_the_documents_give_to_others_gets_no_template_line(job, playbook):
    # T-005e (merge with T-005c): one of AME's three BMS names, so T-005c's red *** flag stays out of the lines
    front = FrontEnd(mode="extend-existing", existing_system_name="Niagara N4", lines=("front end by the owner",),
                     no_scope_reason="The controls specification gives the whole front end to the owner.")
    assert render(replace(job, front_end=front), playbook).scope.front_end.items == ()


def test_an_existing_front_end_with_nothing_in_the_documents_gets_no_template_line(job, playbook):
    # T-005e (merge with T-005c): one of AME's three BMS names, so T-005c's red *** flag stays out of the lines
    front = FrontEnd(mode="extend-existing", existing_system_name="Niagara N4")     # (T-005c prints its own flag)
    assert not any(i.text.startswith("Furnish") for i in render(replace(job, front_end=front),
                                                                   playbook).scope.front_end.items)


def test_an_existing_front_end_the_reading_found_but_no_line_written_gets_the_template_lines(job, playbook):
    # T-005e (merge with T-005c): one of AME's three BMS names, so T-005c's red *** flag stays out of the lines
    front = FrontEnd(mode="extend-existing", existing_system_name="Niagara N4", lines=("tie in to the existing BMS",))
    assert [i.text for i in render(replace(job, front_end=front), playbook).scope.front_end.items] == [
        "Furnish Network Level Control Panel", "Furnish 3-D Graphics to match existing",
        "Furnish Necessary Training on New Controls"]


def test_a_new_system_front_end_is_ordered_and_gets_the_template_default_lines(job, playbook):
    front = FrontEnd(mode="new-system", devices=(_fe(1, "Necessary Training on New Controls"), _fe(2, "BMS Server"),
                                                 _fe(3, "Network Level Control Panel")))
    # T-005e: estimator ruling 19: the New front end's "Default" lines (3-D Graphics, training) print unless written
    assert [i.text for i in render(replace(job, front_end=front), playbook).scope.front_end.items] == [
        "Furnish Network Level Control Panel", "Furnish BMS Server", "Furnish 3-D Graphics",
        "Furnish Necessary Training on New Controls"]


# ---- 3. off-hours work and prevailing wage ----------------------------------------------------------------------

@dataclass(frozen=True)
class _Place:
    document: str = "Specs.pdf"
    page: int = 12
    sheet: str | None = None
    section: str | None = "01 10 00"
    strength: str = "Division 1"


@dataclass(frozen=True)
class _Found:
    section: str
    item: str
    quote: str
    status: str = "Found"
    places: tuple = (_Place(),)


@dataclass(frozen=True)
class _Overview:
    conditions: tuple = ()
    standard_lines: tuple = ()


OFF = _Found("Project conditions that change the price", "Working hours: off-hours, night or weekend work",
             "All work in occupied areas shall be performed between 6:00 PM and 6:00 AM, Monday through Friday.")
WAGE = _Found("Project conditions that change the price", "Prevailing wage",
              "The Contractor shall pay the New Jersey prevailing wage rates to all workers.",
              places=(_Place(section="00 73 46", page=40),))
NORMAL = _Found("Project conditions that change the price", "Working hours: off-hours, night or weekend work",
                "Work shall be performed during normal working hours, 7:00 AM to 3:30 PM.")


def test_the_quote_decides_off_hours():
    assert calls_for_off_hours(OFF.quote) and hours_of(OFF.quote) == "6:00 PM to 6:00 AM"
    assert calls_for_off_hours("Overtime and weekend work shall be included as required.")
    assert not calls_for_off_hours(NORMAL.quote)
    assert state_of(WAGE.quote) == "NJ"


def test_labor_conditions_from_the_spec_overview():
    found = labor_conditions(_Overview(conditions=(NORMAL, OFF, WAGE)))
    assert [c.kind for c in found] == ["off-hours", "prevailing-wage"]
    assert found[0].text == "Off-Hours Labor (6:00 PM to 6:00 AM) Included (As per Spec Section 01 10 00)"
    assert found[1].text == "NJ Prevailing Wage Labor Rates Included (As per Spec Section 00 73 46)"
    assert found[0].quote == OFF.quote
    assert labor_conditions(_Overview(conditions=(NORMAL,))) == () and labor_conditions(None) == ()


def test_a_not_found_row_is_not_a_condition():
    assert labor_conditions(_Overview(conditions=(replace(OFF, status="Not found"),))) == ()


def test_the_labor_block_is_red_and_unpriced():
    block = labor_block(labor_conditions(_Overview(conditions=(OFF, WAGE))))
    assert block.header == HEADING and block.numbered is False
    assert all(run.red for item in block.items for run in item.runs)
    assert all(item.quantity is None for item in block.items)
    assert labor_block(()) is None


def test_with_labor_adds_the_block_last_in_the_base_bid(job, playbook):
    draft = render(job, playbook)
    conditions = labor_conditions(_Overview(conditions=(OFF, WAGE)))
    out = with_labor(draft, conditions)
    assert out.scope.blocks[:-1] == draft.scope.blocks and out.scope.blocks[-1].header == HEADING
    assert out.scope.labor_conditions == conditions
    assert out.scope.add_alternates == draft.scope.add_alternates
    assert [l.text for l in out.lines if l.system_key == LABOR_KEY][1:] == [c.text for c in conditions]
    assert with_labor(draft, ()) is draft


def _written(job, playbook, tmp_path, overview):
    from estimator.__main__ import ALIASES, PRICE_BOOK, PROPOSAL_TEMPLATE, WORKUP_TEMPLATE
    from estimator.price_sheet.price_book import PriceBook
    from estimator.price_sheet.write import write_price_sheet
    from estimator.proposal.write import write_proposal
    draft = with_labor(render(job, playbook), labor_conditions(overview))
    docx = write_proposal(draft.scope, PROPOSAL_TEMPLATE, tmp_path / "p.docx")
    xlsx = write_price_sheet(draft.scope, WORKUP_TEMPLATE, PriceBook.load(PRICE_BOOK, ALIASES), tmp_path / "s.xlsx")
    return Document(str(docx)), load_workbook(str(xlsx))


def _red(paragraph) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(runs) and all(r.font.color is not None and str(r.font.color.rgb) == "FF0000" for r in runs)


def test_the_proposal_keeps_both_labor_points_and_adds_red_sub_bullets(job, playbook, tmp_path):
    document, _ = _written(job, playbook, tmp_path, _Overview(conditions=(OFF, WAGE)))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    text = [p.text.strip() for p in paragraphs]
    line = "Off-Hours Labor (6:00 PM to 6:00 AM) Included (As per Spec Section 01 10 00)"
    assert line in text and _red(paragraphs[text.index(line)])
    assert "Labor Conditions" in text
    # Ruling 26 keeps the template points. The red lines are sub-bullets under them, not replacements.
    assert any(t.startswith("All off-hours or overtime labor") for t in text)
    assert any(t.startswith("All labor will be performed during normal business working hours") for t in text)
    off = next(p for p in paragraphs if p.text.startswith("The documents require work between"))
    wage = next(p for p in paragraphs if p.text.startswith("All labor is priced at prevailing wage rates"))
    assert _red(off) and "6:00 PM and 6:00 AM" in off.text and "Spec Section 01 10 00" in off.text
    assert _red(wage) and "Spec Section 00 73 46" in wage.text
    exclusion = next(i for i, t in enumerate(text) if t.startswith("All off-hours or overtime labor"))
    normal = next(i for i, t in enumerate(text) if t.startswith("All labor will be performed during normal"))
    assert text[exclusion + 1] == off.text.strip()
    assert text[normal + 1] == wage.text.strip()


def test_without_conditions_the_template_boilerplate_stays(job, playbook, tmp_path):
    document, _ = _written(job, playbook, tmp_path, _Overview(conditions=(NORMAL,)))
    text = [p.text.strip() for p in document.paragraphs]
    assert any(t.startswith("All off-hours or overtime labor") for t in text)
    assert any(t.startswith("All labor will be performed during normal business working hours") for t in text)
    assert "Labor Conditions" not in text


def test_the_price_sheet_carries_the_lines_in_red_with_a_flag(job, playbook, tmp_path):
    _, workbook = _written(job, playbook, tmp_path, _Overview(conditions=(OFF, WAGE)))
    ws = workbook["Price Sheet 1"]
    rows = {ws.cell(r, 4).value: r for r in range(6, 397) if ws.cell(r, 4).value}
    found = [t for t in rows if isinstance(t, str) and "Off-Hours Labor" in t]
    assert found
    r = rows[found[0]]
    assert str(ws.cell(r, 4).font.color.rgb).endswith("FF0000")
    assert ws.cell(r, 8).value == "LABOR PREMIUM - ESTIMATOR TO PRICE"
    assert ws.cell(r, 3).value is None                       # not priced as material: the estimator adds the premium
    assert "Prevailing wage is called for by the documents" in workbook["Summary"]["B34"].value
