"""G-002 review of record (lead, 2026-10-12): a block of 5 units, one device line counted 33 and one counted 5.

Invented counts and names only. The sample Playbook, and the real Playbook's VAV shape for the house forms.
"""
from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from bms.model import Count, Device, Family, Job, OtherCount, Ref
from bms.playbook import load_playbook
from bms.render import render
from estimator.price_sheet.price_book import BookItem, PriceBook
from estimator.price_sheet.write import QUANTITY_PENDING, write_price_sheet

FIXTURES = Path(__file__).resolve().parent / "fixtures"
PLAYBOOK = Path(__file__).resolve().parents[2] / "playbook"
REF = Ref(source_id="S1", quote="invented count", where="M-1, Drawings.pdf page 1")
TAGS = ("EF-1", "EF-2", "EF-3", "EF-4", "EF-5")


def _device(id, slot, name, qty) -> Device:
    return Device(id=id, slot=slot, name=name, verb="Furnish and Install", furnished_by="AME", installed_by="AME",
                  wired_by="AME", quantity=Count(value=qty), refs=(REF,))


def _job(key: str, shape: str, devices: tuple[Device, ...]) -> Job:
    job = Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))
    system = replace(job.system("EF"), key=key, shape=shape, name="Exhaust Fans",
                     families=(Family("Exhaust Fan", TAGS, Count(value=5)),), devices=devices, clarifications=(),
                     sequence_refs=())
    return replace(job, systems=(system,))


def _lines(draft, key):
    return [line.text for line in draft.lines if line.system_key == key]


def _ef_devices(damper: int):
    return (_device("EF.d1", "controller", "DDC Controller", 5),
            _device("EF.d2", "damper_actuator", "Motorized Damper Actuator", damper))


def test_a_33_line_prints_the_end_form_and_a_5_line_prints_none():
    playbook = load_playbook(FIXTURES / "sample_playbook")
    lines = _lines(render(_job("EF", "Exhaust Fan", _ef_devices(33)), playbook), "EF")
    assert lines[0].startswith("Qty (5) Exhaust Fans")
    assert "Furnish and Install DDC Controller" in lines
    assert "Furnish and Install Motorized Damper Actuator (Qty. 33 Total)" in lines
    assert not any(line.startswith("Qty (") and "Furnish" in line for line in lines)


def test_the_end_form_does_not_move_a_line():
    """The GEM order reads the printed text: the end form must not change where a line prints."""
    playbook = load_playbook(FIXTURES / "sample_playbook")
    plain = _lines(render(_job("EF", "Exhaust Fan", _ef_devices(5)), playbook), "EF")
    marked = _lines(render(_job("EF", "Exhaust Fan", _ef_devices(33)), playbook), "EF")
    assert [t.removesuffix(" (Qty. 33 Total)") for t in marked] == plain


def test_the_real_vav_shape_prints_the_end_form_only_on_a_different_count():
    playbook = load_playbook(PLAYBOOK)
    devices = (_device("V.d1", "controller", "DDC Controller", 5),
               _device("V.d2", "space_sensor", "Space Temperature Sensor", 33),
               _device("V.d3", "discharge_temperature", "Discharge Air Temperature Sensor", 5))
    lines = _lines(render(_job("VAV", "VAV / Fan VAV", devices), playbook), "VAV")
    assert lines[0].startswith("Qty (5) ")
    assert "Furnish and Install Space Temperature Sensor (Qty. 33 Total)" in lines
    assert "Furnish and Install DDC Controller" in lines
    assert "Furnish and Install Discharge Air Temperature Sensor" in lines
    plain = _lines(render(_job("VAV", "VAV / Fan VAV", devices[:1] + (replace(devices[1], quantity=Count(value=5)),)
                               + devices[2:]), playbook), "VAV")
    assert [t.removesuffix(" (Qty. 33 Total)") for t in lines] == plain


def test_a_line_with_a_count_note_prints_the_total_before_the_note():
    """Ruling 24: a note (other counts) prints "(Qty. N Total)" even when N is the block's count. The note follows."""
    playbook = load_playbook(FIXTURES / "sample_playbook")
    plan = Ref(source_id="S2", quote="invented count 4", where="M-601, Drawings.pdf page 2")

    def damper_line(qty: int) -> str:
        count = Count(value=qty, sources=(REF,), others=(OtherCount(4, plan),))
        devices = (_device("EF.d1", "controller", "DDC Controller", 5),
                   replace(_device("EF.d2", "damper_actuator", "Motorized Damper Actuator", qty), quantity=count))
        lines = _lines(render(_job("EF", "Exhaust Fan", devices), playbook), "EF")
        assert lines[0].startswith("Qty (5) Exhaust Fans")
        return next(line for line in lines if "Motorized Damper Actuator" in line)

    assert damper_line(5) == ("Furnish and Install Motorized Damper Actuator (Qty. 5 Total) "
                              "(5 per M-1; 4 per M-601)")
    assert damper_line(33) == ("Furnish and Install Motorized Damper Actuator (Qty. 33 Total) "
                               "(33 per M-1; 4 per M-601)")


def test_the_price_sheet_reads_both_lines(workup_template, tmp_path):
    playbook = load_playbook(FIXTURES / "sample_playbook")
    book = PriceBook([BookItem("DDC Controller", 900.0, "Std Parts Pricing", 2, False),
                      BookItem("Motorized Damper Actuator", 210.0, "Std Parts Pricing", 3, False)], {})
    draft = render(_job("EF", "Exhaust Fan", _ef_devices(33)), playbook)
    path = write_price_sheet(draft.scope, workup_template, book, tmp_path / "s.xlsx")
    ws = load_workbook(path)["Price Sheet 1"]
    rows = {str(ws.cell(r, 4).value): r for r in range(6, 80) if ws.cell(r, 4).value}
    damper = next(r for t, r in rows.items()
                  if t.endswith("Furnish and Install Motorized Damper Actuator (Qty. 33 Total)"))
    controller = next(r for t, r in rows.items() if t.endswith("Furnish and Install DDC Controller"))
    assert (ws.cell(damper, 3).value, ws.cell(damper, 5).value) == (33, 210.0)
    assert (ws.cell(controller, 3).value, ws.cell(controller, 5).value) == (5, 900.0)
    assert ws.cell(damper, 8).value != QUANTITY_PENDING and ws.cell(controller, 8).value != QUANTITY_PENDING
