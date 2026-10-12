"""G-005 Task 1: a device line ends with "(Qty. N Total)". Invented data only.

The block header, the Demolition Scope and a named exception keep "Qty (n)".
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from bms.model import (Count, DemolitionLine, Device, Family, FrontEnd, Job, OtherCount, Ref, ScopeException, System)
from bms.playbook import load_playbook
from bms.render import render
from estimator.price_sheet.price_book import BookItem, PriceBook, price_key
from estimator.price_sheet.write import PRICE_PENDING, QUANTITY_PENDING, write_price_sheet
from estimator.reading.inventory import COUNT_CAUSES_LEAD
from estimator.scope import Block, Item, Run

FIXTURES = Path(__file__).resolve().parent / "fixtures"
ROOT = Path(__file__).resolve().parents[2]
REF = Ref("S1", "invented count", "M-1, Invented.pdf page 1")
SHEET_601 = Ref("S2", "counted on the sheet", "M-601, Invented.pdf page 2")
SHEET_401 = Ref("S3", "the other count", "M-401, Invented.pdf page 3")
HEADER = (
    "Qty (55) Fan Coil Unit FCU-1 (55 per M901.00 (likely a typical unit at each place); "
    "1 per schedule rows; 1 per M902.00)"
)
LINE_55 = "Furnish and Install Wall-Mounted Thermostat (Qty. 55 Total)"
LINE_NOTE = "Furnish and Install Wall-Mounted Thermostat (Qty. 4 Total) (4 per M-601; 2 per M-401)"
LINE_PLAIN = "Furnish and Install Wall-Mounted Thermostat"


def _device(name: str, qty: int | None, id: str, slot: str = "space_temperature", mark: str = "",
            count: Count | None = None) -> Device:
    quantity = count if count is not None else Count(qty, (REF,))
    return Device(id=id, slot=slot, name=name, verb="Furnish and Install", furnished_by="AME", installed_by="AME",
                  wired_by="AME", quantity=quantity, refs=(REF,), mark=mark)


def _base() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def _job(*devices: Device, count: int = 6, shape: str = "Unit Heater", name: str = "Unit Heaters",
         key: str = "UH", family: str = "Cabinet Unit Heater", tags: tuple[str, ...] = ("UH-1",),
         family_count: Count | None = None, front_end: FrontEnd | None = None,
         demolition: tuple[DemolitionLine, ...] = (), exceptions: tuple[ScopeException, ...] = ()) -> Job:
    system = System(key=key, name=name, shape=shape, control="field",
                    families=(Family(family, tags, family_count or Count(count, (REF,))),),
                    devices=devices)
    return replace(_base(), systems=(system,), demolition=demolition, exceptions=exceptions,
                   front_end=front_end or FrontEnd(mode="new-system"))


def _playbook():
    return load_playbook(FIXTURES / "sample_playbook")


def _lines(job: Job, playbook=None) -> list:
    block = render(job, playbook or _playbook()).scope.blocks[0]
    return [item for item in block.items if item.kind == "line"]


def _by_name(items, name: str):
    return next(item for item in items if name in item.text)


def _fan_count() -> Count:
    cause = f"plan placements 55; {COUNT_CAUSES_LEAD} a typical unit at each place"
    return Count(
        55,
        (Ref("S1", "counted on the sheet", "M901.00, Invented.pdf page 2"),),
        (OtherCount(55, Ref("S2", cause, "equipment counts for FCU")),
         OtherCount(1, Ref("S3", "schedule rows 1", "equipment counts for FCU")),
         OtherCount(1, Ref("S4", "one unit", "M902.00, Invented.pdf page 4"))),
    )


def test_a_different_count_prints_at_the_end():
    items = _lines(_job(_device("Wall-Mounted Thermostat", 55, "UH.d1")))
    line = _by_name(items, "Wall-Mounted Thermostat")
    assert line.text == LINE_55
    assert not line.text.startswith("Qty (")


def test_the_block_count_with_no_note_prints_no_quantity_words():
    items = _lines(_job(_device("Wall-Mounted Thermostat", 6, "UH.d1"), count=6))
    line = _by_name(items, "Wall-Mounted Thermostat")
    assert line.text == LINE_PLAIN
    assert "Qty" not in line.text


def test_the_block_count_with_a_note_prints_the_total_then_the_other_counts():
    count = Count(4, (SHEET_601,), (OtherCount(2, SHEET_401),))
    device = _device("Wall-Mounted Thermostat", 4, "UH.d1", count=count)
    items = _lines(_job(device, count=4))
    line = _by_name(items, "Wall-Mounted Thermostat")
    assert line.text == LINE_NOTE
    assert line.text.index("(Qty. 4 Total)") < line.text.index("(4 per M-601")


def test_an_unknown_count_ends_red_and_a_future_line_stays_red():
    known = _device("Space Temperature Sensor", 4, "UH.d1")
    unknown = _device("Current Switch", None, "UH.d2", slot="control_valve")
    future = _device("Discharge Air Sensor", 2, "UH.d3", mark="future")
    items = _lines(_job(known, unknown, future, count=6))
    known_line = _by_name(items, "Space Temperature Sensor")
    unknown_line = _by_name(items, "Current Switch")
    future_line = _by_name(items, "Discharge Air Sensor")
    assert known_line.text == "Furnish and Install Space Temperature Sensor (Qty. 4 Total)"
    assert not known_line.runs[0].red and not known_line.text.startswith("Qty (")
    assert unknown_line.text.endswith(" (Qty. ##)") and not unknown_line.text.startswith("Qty (")
    assert unknown_line.runs[0].red and unknown_line.quantity is None
    assert future_line.runs[0].red and future_line.text.endswith("(FUTURE)")
    assert not future_line.text.startswith("Qty (")


def test_the_fan_coil_header_keeps_the_source_breakdown():
    device = _device("Wall-Mounted Thermostat", 55, "FCU.d1")
    job = _job(device, shape="Fan Coil Unit", name="Fan Coil Units", key="FCU", family="Fan Coil Unit",
               tags=("FCU-1",), family_count=_fan_count())
    block = render(job, load_playbook(ROOT / "playbook")).scope.blocks[0]
    assert block.header == HEADER


def test_a_front_end_line_uses_the_same_end_form():
    other = _device("Wall-Mounted Thermostat", 55, "FE.d1", slot="other_device")
    same = _device("Space Temperature Sensor", 1, "FE.d2", slot="other_device")
    unknown = _device("Current Switch", None, "FE.d3", slot="other_device")
    job = _job(front_end=FrontEnd(mode="new-system", devices=(other, same, unknown)))
    items = [item for item in render(job, _playbook()).scope.front_end.items if item.kind == "line"]
    assert _by_name(items, "Wall-Mounted Thermostat").text == LINE_55
    same_line = _by_name(items, "Space Temperature Sensor")
    assert same_line.text == "Furnish and Install Space Temperature Sensor" and "Qty" not in same_line.text
    unknown_line = _by_name(items, "Current Switch")
    assert unknown_line.text.endswith(" (Qty. ##)") and not unknown_line.text.startswith("Qty (")
    assert unknown_line.runs[0].red and unknown_line.quantity is None


def test_demolition_and_named_exceptions_keep_qty_n():
    demo = DemolitionLine("AHU", "demo", 2, (REF,))
    text = "EXCEPTION TO ANY WORK WITH Qty (2) FCU-1 the documents give this work to others"
    job = _job(_device("Wall-Mounted Thermostat", 6, "UH.d1"), demolition=(demo,),
               exceptions=(ScopeException("FCU", text, (REF,)),))
    out = render(job, _playbook())
    block = out.scope.blocks[-1]
    assert block.header == "Demolition Scope"
    assert block.items[0].text == "Qty (2) AHU demo"
    assert "(Qty." not in block.items[0].text
    assert out.scope.named_exceptions == (text,)


def test_the_price_sheet_reads_the_end_form_and_the_old_prefix(workup_template, tmp_path):
    thermostat = _device("Wall-Mounted Thermostat", 55, "UH.d1")
    unknown = _device("Current Switch", None, "UH.d2", slot="control_valve")
    book = PriceBook([
        BookItem("Wall-Mounted Thermostat", 80.0, "Std Parts Pricing", 1, False),
        BookItem("Current Switch", 12.0, "Std Parts Pricing", 2, False),
    ], {})
    draft = render(_job(thermostat, unknown, count=6), _playbook())
    old = Block(header="Qty (1) Spare Parts", items=(
        Item(kind="line", quantity=33, runs=(Run(text="Qty (33) Furnish and Install Wall-Mounted Thermostat"),)),
    ))
    scope = draft.scope.model_copy(update={"blocks": draft.scope.blocks + (old,)})
    path = write_price_sheet(scope, workup_template, book, tmp_path / "sheet.xlsx")
    ws = load_workbook(path)["Price Sheet 1"]
    rows = {str(ws.cell(r, 4).value): r for r in range(6, 80) if ws.cell(r, 4).value}

    def rgb(row: int) -> str:
        color = ws.cell(row, 4).font.color
        if color is None or color.type == "theme" or color.rgb is None:
            return ""
        return str(color.rgb)

    end = next(r for text, r in rows.items() if "(Qty. 55 Total)" in text and "Wall-Mounted Thermostat" in text)
    pending = next(r for text, r in rows.items() if "(Qty. ##)" in text and "Current Switch" in text)
    prefix = next(r for text, r in rows.items() if "Qty (33) Furnish and Install Wall-Mounted Thermostat" in text)
    assert ws.cell(end, 3).value == 55 and ws.cell(end, 5).value == 80.0
    assert ws.cell(end, 8).value != PRICE_PENDING
    assert ws.cell(pending, 8).value == QUANTITY_PENDING and ws.cell(pending, 3).value is None
    assert rgb(pending).endswith("FF0000")
    assert ws.cell(prefix, 3).value == 33 and ws.cell(prefix, 5).value == 80.0
    end_key = price_key("Furnish and Install Wall-Mounted Thermostat (Qty. 55 Total)")
    prefix_key = price_key("Qty (33) Furnish and Install Wall-Mounted Thermostat")
    unknown_key = price_key("Furnish and Install Wall-Mounted Thermostat (Qty. ##)")
    assert end_key == prefix_key == unknown_key == "wall-mounted thermostat"


def test_the_real_playbook_loads_and_lints_clean():
    playbook = load_playbook(ROOT / "playbook")
    assert playbook.wording.quantity["line"] == " (Qty. {n} Total)"
    assert playbook.wording.quantity["unknown"] == " (Qty. ##)"
    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src")
    proc = subprocess.run([sys.executable, "-m", "bms", "playbook-lint"], cwd=ROOT, env=env,
                          capture_output=True, text=True, check=False)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "is clean." in proc.stdout
    assert "problem" not in proc.stdout.lower()
