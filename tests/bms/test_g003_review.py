"""Reviewer's test for G-003 (price-sheet polish). Every name here is invented."""
from __future__ import annotations

from copy import copy
from dataclasses import replace

from docx import Document
from openpyxl import load_workbook

from bms.labor import labor_conditions, with_labor
from bms.model import DemolitionLine, Ref
from bms.render import render
from estimator.__main__ import ALIASES, PRICE_BOOK, PROPOSAL_TEMPLATE, WORKUP_TEMPLATE
from estimator.price_sheet.price_book import PriceBook
from estimator.price_sheet.rows import FLAG, HOURS, QUANTITY, SCOPE
from estimator.price_sheet.write import write_price_sheet
from estimator.proposal.write import write_proposal
from test_demolition import PLAYBOOK, _sample
from test_t005e_format_order_labor import OFF, WAGE, _Overview


def test_the_key_survives_the_run_path_and_both_red_row_kinds_are_painted_alike(tmp_path):
    """The run renders, then adds the Labor Conditions (run.py), then writes. The Demolition Scope key survives that
    path. Its labor row and the T-005e off-hours and prevailing-wage rows are painted alike: the description red, the
    flag bold dark red, and the hours and quantity cells in the template's own font. The key prints nowhere in the
    proposal."""
    shown = Ref("S1", "shown for removal", "MD-101, Drawings.pdf page 6")
    job = replace(_sample(), demolition=(DemolitionLine("AHU", "demo", 2, (shown,)),))
    draft = with_labor(render(job, PLAYBOOK), labor_conditions(_Overview(conditions=(OFF, WAGE))))
    assert [b.key for b in draft.scope.blocks if b.header == "Demolition Scope"] == ["demolition"]
    assert [b.key for b in draft.scope.blocks if b.header != "Demolition Scope"] == [""] * (len(draft.scope.blocks) - 1)

    book = PriceBook.load(PRICE_BOOK, ALIASES)
    ws = load_workbook(write_price_sheet(draft.scope, WORKUP_TEMPLATE, book, tmp_path / "s.xlsx"))["Price Sheet 1"]
    rows = {ws.cell(r, SCOPE).value: r for r in range(6, 397) if ws.cell(r, SCOPE).value}
    demo = rows["AHU demo"]
    premium = [r for text, r in rows.items() if "Off-Hours Labor" in text or "Prevailing Wage" in text]
    blank = next(r for r in range(6, 397) if ws.cell(r, SCOPE).value is None)
    assert len(premium) == 2
    assert ws.cell(demo, FLAG).value == "HOURS BY ESTIMATOR" and ws.cell(demo, QUANTITY).value == 2
    assert [ws.cell(r, FLAG).value for r in premium] == ["LABOR PREMIUM - ESTIMATOR TO PRICE"] * 2
    for row in (demo, *premium):
        assert str(ws.cell(row, SCOPE).font.color.rgb).endswith("FF0000")
        assert str(ws.cell(row, FLAG).font.color.rgb).endswith("C00000") and ws.cell(row, FLAG).font.b
        for column in (HOURS, QUANTITY):
            # "keep their own font": the same font as a row nobody wrote, not merely a colour that is not red
            assert copy(ws.cell(row, column).font) == copy(ws.cell(blank, column).font)

    proposal = write_proposal(draft.scope, PROPOSAL_TEMPLATE, tmp_path / "p.docx")
    words = [p.text.strip() for p in Document(str(proposal)).paragraphs]
    assert "demolition" not in words and words.count("Demolition Scope") == 1
