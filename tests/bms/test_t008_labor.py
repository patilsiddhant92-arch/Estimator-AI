"""T-008 task 2: a conditional sentence is no ban; the ledger facts and plan key notes can call for off-hours work.

Invented fixtures only.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from bms.labor import calls_for_off_hours, calls_for_prevailing_wage, hours_of, labor_conditions, with_labor
from bms.model import Job
from bms.playbook import load_playbook
from bms.render import render
from estimator.ledger.facts import Ledger
from estimator.reading.keynotes import Note

FIXTURES = Path(__file__).resolve().parent / "fixtures"
SENTENCE = "If no hours are named in the permit then the work hours shall be Monday to Friday 7:00 pm - 5:00 am."


def test_a_conditional_sentence_calls_for_off_hours():
    assert calls_for_off_hours(SENTENCE)
    assert hours_of(SENTENCE) == "7:00 PM to 5:00 AM"


@pytest.mark.parametrize("ban", ["No overtime is permitted.", "Overtime work shall not be done without written approval."])
def test_a_ban_stays_a_ban(ban):
    assert not calls_for_off_hours(ban)


def _ledger(*quotes, kind="note"):
    ledger = Ledger({"S1": "M-101, Drawings.pdf page 3 (a sheet)"})
    for quote in quotes:
        ledger.add(kind, "hours", "x", "S1", quote, "spec")
    return ledger


def _note(text, controls=False):
    return Note(sheet="M-201", document="Drawings.pdf", page=2, block="KEY NOTES", number="1", text=text,
                controls=controls, who="")


def test_a_ledger_fact_gives_the_off_hours_condition():
    found = labor_conditions(None, ledger=_ledger(SENTENCE))
    assert len(found) == 1 and found[0].kind == "off-hours"
    assert "(7:00 PM to 5:00 AM)" in found[0].text and "M-101" in found[0].source


def test_a_key_note_gives_the_off_hours_condition_even_when_it_is_not_a_controls_note():
    found = labor_conditions(None, notes=(_note("ALL WORK DURING OFF HOURS.", controls=False),))
    assert len(found) == 1 and found[0].kind == "off-hours" and found[0].source == "Sheet M-201"


def test_a_note_without_a_sheet_names_the_document_and_page():
    note = Note(sheet=None, document="C:/x/Drawings.pdf", page=4, block="KEY NOTES", number="1",
                text="Work during off hours.", controls=True, who="")
    assert labor_conditions(None, notes=(note,))[0].source == "Drawings.pdf page 4"


@pytest.mark.parametrize("text", ["OFF HOURS", "work 6:00 pm - 6:00 am", "nights, weekends"])
def test_the_calls_a_fact_can_make(text):
    assert labor_conditions(None, ledger=_ledger(text))[0].kind == "off-hours"


def test_a_prevailing_wage_fact():
    found = labor_conditions(None, ledger=_ledger("Prevailing wage applies."))
    assert [c.kind for c in found] == ["prevailing-wage"]


def test_a_single_time_and_a_ban_give_nothing():
    assert labor_conditions(None, ledger=_ledger("7:00 AM to 3:30 PM", "No overtime is permitted.")) == ()
    assert labor_conditions(None, notes=(_note("No work during off hours."),)) == ()
    assert not calls_for_prevailing_wage("normal hours")


class _Row:
    status = "Found"
    quote = "Work hours shall be 6:00 pm - 6:00 am."
    places = ()


class _Overview:
    conditions = (_Row(),)
    standard_lines = ()


def test_the_overview_rows_still_win_and_the_old_call_works():
    both = labor_conditions(_Overview(), ledger=_ledger(SENTENCE), notes=(_note("off hours"),))
    assert len(both) == 1 and "6:00 PM to 6:00 AM" in both[0].text
    assert labor_conditions(_Overview())[0].kind == "off-hours"


def test_at_most_one_condition_of_each_kind():
    found = labor_conditions(None, ledger=_ledger(SENTENCE, "off hours", "Prevailing wage."), notes=(_note("off hours"),))
    assert [c.kind for c in found] == ["off-hours", "prevailing-wage"]


@pytest.fixture(scope="module")
def playbook():
    return load_playbook(FIXTURES / "sample_playbook")


@pytest.fixture
def job():
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def _written(job, playbook, tmp_path, conditions):
    from estimator.__main__ import ALIASES, PRICE_BOOK, PROPOSAL_TEMPLATE, WORKUP_TEMPLATE
    from estimator.price_sheet.price_book import PriceBook
    from estimator.price_sheet.write import write_price_sheet
    from estimator.proposal.write import write_proposal
    draft = with_labor(render(job, playbook), conditions)
    docx = write_proposal(draft.scope, PROPOSAL_TEMPLATE, tmp_path / "p.docx")
    xlsx = write_price_sheet(draft.scope, WORKUP_TEMPLATE, PriceBook.load(PRICE_BOOK, ALIASES), tmp_path / "s.xlsx")
    return Document(str(docx)), load_workbook(str(xlsx))


def test_end_to_end_the_sentence_gives_a_red_line_and_a_red_row(job, playbook, tmp_path):
    document, workbook = _written(job, playbook, tmp_path, labor_conditions(None, ledger=_ledger(SENTENCE)))
    paras = [p for p in document.paragraphs if p.text.startswith("Off-Hours Labor")]
    assert paras and "Included" in paras[0].text
    assert all(str(r.font.color.rgb) == "FF0000" for r in paras[0].runs if r.text.strip())
    ws = workbook["Price Sheet 1"]
    rows = [r for r in range(6, 397) if isinstance(ws.cell(r, 4).value, str) and "Off-Hours Labor" in ws.cell(r, 4).value]
    assert rows and ws.cell(rows[0], 8).value == "LABOR PREMIUM - ESTIMATOR TO PRICE"
    assert str(ws.cell(rows[0], 4).font.color.rgb).endswith("FF0000")


def test_end_to_end_a_ban_gives_neither(job, playbook, tmp_path):
    document, workbook = _written(job, playbook, tmp_path, labor_conditions(None, ledger=_ledger("No overtime is permitted.")))
    assert not any(p.text.startswith("Off-Hours Labor") for p in document.paragraphs)
    ws = workbook["Price Sheet 1"]
    assert not any(isinstance(ws.cell(r, 4).value, str) and "Off-Hours Labor" in ws.cell(r, 4).value
                   for r in range(6, 397))
