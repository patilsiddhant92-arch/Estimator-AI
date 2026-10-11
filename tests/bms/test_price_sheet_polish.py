"""Price-sheet polish (G-003). Every name here is invented."""
from __future__ import annotations

import re
from dataclasses import replace

from openpyxl import load_workbook

from bms.model import DemolitionLine, Ref
from bms.render import render
from estimator.__main__ import ALIASES, PRICE_BOOK, WORKUP_TEMPLATE
from estimator.price_sheet.check import price_check
from estimator.price_sheet.price_book import PriceBook
from estimator.price_sheet.rows import FLAG, HOURS, QUANTITY, SCOPE, UNIT_PRICE
from estimator.price_sheet.write import write_price_sheet
from estimator.scope import Block, Item, Run
from test_demolition import PLAYBOOK, _sample
from test_t005e_format_order_labor import OFF, WAGE, _Overview, _written

_REF = Ref("S1", "shown for removal", "MD-101, Drawings.pdf page 6")
_BOOK = PriceBook.load(PRICE_BOOK, ALIASES)
_PENDING_COUNT = re.compile(r"(\d+) lines? with a quantity (?:has|have) no price yet")
_CHECK_COUNT = re.compile(r"\| No price book match yet \(PRICE PENDING\) \| (\d+) \|")


def _rgb(cell) -> str:
    color = cell.font.color
    if color is None or color.type == "theme":
        return ""
    return "" if color.rgb is None else str(color.rgb)


def _red(cell) -> bool:
    return _rgb(cell).endswith("FF0000") or _rgb(cell).endswith("C00000")


def _demo_job():
    line = DemolitionLine("AHU", "demo", 2, (_REF,))
    return render(replace(_sample(), demolition=(line,)), PLAYBOOK)


def _with_widget(scope):
    widget = Block(
        header="Qty (1) Imaginary Units",
        items=(Item(kind="line", depth=1, runs=(Run(text="Furnish and Install Imaginary Widget"),), quantity=1),),
    )
    return scope.model_copy(update={"blocks": (*scope.blocks, widget)})


def _write(scope, tmp_path, name="sheet.xlsx"):
    path = write_price_sheet(scope, WORKUP_TEMPLATE, _BOOK, tmp_path / name)
    return path, load_workbook(path)


def _row(ws, text: str) -> int:
    found = [r for r in range(6, 397) if ws.cell(r, SCOPE).value == text]
    assert found, text
    return found[0]


def test_a_known_labor_count_is_hours_by_estimator_and_not_priced(tmp_path):
    """A Demolition Scope labor row with a count is hours for the estimator, not a device price."""
    _, wb = _write(_demo_job().scope, tmp_path)
    ws = wb["Price Sheet 1"]
    row = _row(ws, "AHU demo")
    assert ws.cell(row, FLAG).value == "HOURS BY ESTIMATOR"
    assert ws.cell(row, HOURS).value is None
    assert ws.cell(row, UNIT_PRICE).value is None
    assert ws.cell(row, QUANTITY).value == 2
    assert ws.cell(row, FLAG).value != "PRICE PENDING"


def test_an_unknown_labor_count_is_quantity_pending(tmp_path):
    """A labor row whose count is unknown shows QUANTITY PENDING and no price."""
    draft = _demo_job()
    item = Item(kind="line", depth=1, runs=(Run(text="Qty (##) AHU demo", red=True),), quantity=None)
    blocks = []
    for block in draft.scope.blocks:
        if block.header == "Demolition Scope":
            blocks.append(block.model_copy(update={"items": (item,)}))
        else:
            blocks.append(block)
    _, wb = _write(draft.scope.model_copy(update={"blocks": tuple(blocks)}), tmp_path)
    ws = wb["Price Sheet 1"]
    row = next(r for r in range(6, 397) if ws.cell(r, SCOPE).value and "AHU" in str(ws.cell(r, SCOPE).value)
               and ws.cell(r, FLAG).value == "QUANTITY PENDING")
    assert ws.cell(row, UNIT_PRICE).value is None
    assert ws.cell(row, FLAG).value != "PRICE PENDING"
    assert ws.cell(row, FLAG).value != "HOURS BY ESTIMATOR"


def test_an_unmatched_device_stays_price_pending_and_a_labor_row_does_not(tmp_path):
    """The Summary and the price check still count a device with no book match. They do not count a labor row."""
    scope = _with_widget(_demo_job().scope)
    path, wb = _write(scope, tmp_path)
    ws = wb["Price Sheet 1"]
    flags = [ws.cell(r, FLAG).value for r in range(6, 397)]
    priced_pending = sum(
        1
        for sheet in wb.worksheets if sheet.title != "Summary"
        for r in range(6, 397) if sheet.cell(r, FLAG).value == "PRICE PENDING"
    )
    assert priced_pending >= 1
    assert "HOURS BY ESTIMATOR" in flags
    summary = wb["Summary"]["B34"].value or ""
    said = _PENDING_COUNT.search(summary)
    assert said is not None, summary
    assert int(said.group(1)) == priced_pending
    report = price_check(path, path, _BOOK, "Invented Hall")
    counted = _CHECK_COUNT.search(report)
    assert counted is not None, report
    assert int(counted.group(1)) == priced_pending
    unmatched = report.split("## No price book match yet", 1)[1]
    assert "Imaginary Widget" in unmatched
    assert "AHU demo" not in unmatched


def test_an_off_hours_row_is_red_only_on_the_description_and_the_flag(tmp_path):
    """The description and the flag are red. The hours cell and the quantity cell keep their own font."""
    _, workbook = _written(_sample(), PLAYBOOK, tmp_path, _Overview(conditions=(OFF, WAGE)))
    ws = workbook["Price Sheet 1"]
    row = _row(ws, next(t for t in (ws.cell(r, SCOPE).value for r in range(6, 397))
                        if isinstance(t, str) and "Off-Hours Labor" in t))
    assert _rgb(ws.cell(row, SCOPE)).endswith("FF0000")
    assert _rgb(ws.cell(row, FLAG)).endswith("C00000")
    assert ws.cell(row, FLAG).font.bold is True
    assert not _red(ws.cell(row, HOURS))
    assert not _red(ws.cell(row, QUANTITY))


def test_a_demolition_labor_row_is_red_only_on_the_description_and_the_flag(tmp_path):
    """A Demolition Scope labor row uses the same red cells as an off-hours row."""
    _, wb = _write(_demo_job().scope, tmp_path)
    ws = wb["Price Sheet 1"]
    row = _row(ws, "AHU demo")
    assert _rgb(ws.cell(row, SCOPE)).endswith("FF0000")
    assert _rgb(ws.cell(row, FLAG)).endswith("C00000")
    assert ws.cell(row, FLAG).font.bold is True
    assert not _red(ws.cell(row, HOURS))
    assert not _red(ws.cell(row, QUANTITY))


def test_render_gives_the_demolition_block_a_key():
    """The Demolition Scope block carries a stable key. The proposal text does not print it."""
    draft = _demo_job()
    found = [block for block in draft.scope.blocks if block.header == "Demolition Scope"]
    assert len(found) == 1
    assert found[0].key == "demolition"
    assert all("demolition" != line.text for line in draft.lines)


def test_the_key_selects_labor_rows_when_the_heading_text_differs(tmp_path):
    """A different heading still prices as labor when the block key is demolition."""
    draft = _demo_job()
    blocks = []
    for block in draft.scope.blocks:
        if block.key == "demolition":
            blocks.append(block.model_copy(update={"header": "Removal work"}))
        else:
            blocks.append(block)
    _, wb = _write(draft.scope.model_copy(update={"blocks": tuple(blocks)}), tmp_path)
    ws = wb["Price Sheet 1"]
    row = _row(ws, "AHU demo")
    assert ws.cell(row, FLAG).value == "HOURS BY ESTIMATOR"
    assert ws.cell(row, QUANTITY).value == 2


def test_a_job_saved_without_the_key_loads_with_an_empty_key():
    """A Job saved before the key field loads, and the key is empty."""
    draft = render(replace(_sample(), demolition=()), PLAYBOOK)
    raw = draft.scope.model_dump(mode="json")
    for block in raw["blocks"]:
        block.pop("key", None)
    loaded = type(draft.scope).model_validate(raw)
    assert loaded.blocks
    assert all(block.key == "" for block in loaded.blocks)
