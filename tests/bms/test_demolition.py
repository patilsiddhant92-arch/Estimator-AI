"""The Demolition Scope (T-G1). Every name here is invented."""
from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import takeoff_fixture as fx

from bms.model import DemolitionLine, Job, Ref, Unit
from bms.playbook import load_playbook
from bms.render import render
from bms.takeoff import demolition_lines

FIXTURES = Path(__file__).resolve().parent / "fixtures"
PLAYBOOK = load_playbook(FIXTURES / "sample_playbook")


def _sample() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def _unit(family: str, status: str, count: int = 1, tag: str = "", reason: str = "",
          refs: tuple[Ref, ...] = ()) -> Unit:
    tag = tag or f"{family}-1"
    held = refs or (Ref("S1", tag, "MD-101, Mech.pdf page 6"),)
    return Unit(family, tag, (tag,), status, count, (), held, reason=reason)


def test_the_fixture_answer_is_one_ahu_demo_and_one_ef_demo():
    """answer.json: AHU-6 and EF-6 are demolish, count 1. Every other unit is new work."""
    lines = demolition_lines(tuple(Unit.from_dict(row) for row in fx.answer()["units"]))
    assert [(line.family, line.action, line.count) for line in lines] == [("AHU", "demo", 1), ("EF", "demo", 1)]
    assert [line.note for line in lines] == ["", ""]
    assert all(line.refs and line.refs[0].source_id == "F13" for line in lines)


def test_lines_group_by_family_and_action():
    shown = Ref("S1", "AHU-6 shown for demolition", "MD-101, Mech.pdf page 6")
    work = Ref("S2", "remove the existing controls", "MD-101, Mech.pdf page 6")
    remain = Ref("S3", "FCU-2 to remain", "M-201, Mech.pdf page 7")
    move = Ref("S4", "VAV-3 relocate", "M-202, Mech.pdf page 8")
    units = (
        _unit("AHU", "demolish", 1, "AHU-6", refs=(shown,)),
        _unit("AHU", "demolish_with_work", 2, "AHU-7", "existing controls removed", (work,)),
        _unit("AHU", "new", 1, "AHU-1", refs=(shown,)),
        _unit("FCU", "existing_remain", 4, "FCU-2", refs=(remain,)),
        _unit("VAV", "relocate", 1, "VAV-3", refs=(move,)),
        _unit("EF", "future", 1, "EF-9", refs=(shown,)),
    )
    lines = demolition_lines(units)
    assert [(line.family, line.action, line.count, line.note) for line in lines] == [
        ("AHU", "demo", 3, "existing controls removed"),
        ("FCU", "existing to remain", 4, ""),
        ("VAV", "relocation", 1, ""),
    ]
    assert lines[0].refs == (shown, work)


def test_named_work_reasons_join_in_unit_order():
    ref = Ref("S1", "demo", "MD-1, Mech.pdf page 1")
    units = (
        _unit("AHU", "demolish_with_work", reason="existing controls removed", refs=(ref,)),
        _unit("AHU", "demolish_with_work", tag="AHU-2", reason="pneumatic tubing removed", refs=(ref,)),
        _unit("AHU", "demolish_with_work", tag="AHU-3", reason="  ", refs=(ref,)),
    )
    (line,) = demolition_lines(units)
    assert line.count == 3 and line.refs == (ref,)
    assert line.note == "existing controls removed; pneumatic tubing removed"


def test_no_units_gives_no_lines():
    assert demolition_lines(()) == ()


def test_two_demolition_lines_render_as_the_last_block():
    ref = Ref("S1", "shown for removal", "MD-101, Drawings.pdf page 6")
    lines = (
        DemolitionLine("AHU", "demo", 2, (ref,), "existing controls removed"),
        DemolitionLine("FCU", "existing to remain", 3, (ref,)),
    )
    out = render(replace(_sample(), demolition=lines), PLAYBOOK)
    block = out.scope.blocks[-1]
    assert block.header == "Demolition Scope"
    assert [item.text for item in block.items] == [
        "Qty (2) AHU demo, existing controls removed",
        "Qty (3) FCU Existing to Remain",
    ]
    assert all(item.runs[0].red and item.quantity == line.count for item, line in zip(block.items, lines))
    moved = render(replace(_sample(), demolition=(DemolitionLine("VAV", "relocation", 1, (ref,)),)), PLAYBOOK)
    assert [item.text for item in moved.scope.blocks[-1].items] == ["Qty (1) VAV Relocation"]


def test_a_job_with_no_demolition_lines_renders_no_block():
    out = render(replace(_sample(), demolition=()), PLAYBOOK)
    assert [block.header for block in out.scope.blocks if block.header == "Demolition Scope"] == []
    assert all(line.text != "Demolition Scope" for line in out.lines)


def test_the_evidence_package_has_two_demolition_rows(tmp_path):
    import sys

    import openpyxl

    from bms.evidence import DEMOLITION_HEADINGS, DEMOLITION_TAB, write_evidence

    tests = Path(__file__).resolve().parents[1]
    if str(tests) not in sys.path:
        sys.path.insert(0, str(tests))
    from test_evidence import _catalogue

    shown = Ref("S1", "shown for removal", "MD-101, Drawings.pdf page 6")
    remain = Ref("S2", "FCU-2 remains", "M-201, Mech.pdf page 7")
    lines = (
        DemolitionLine("AHU", "demo", 2, (shown,), "existing controls removed"),
        DemolitionLine("FCU", "existing to remain", 1, (remain,)),
    )
    path = write_evidence(tmp_path / "e.xlsx", _catalogue(), None, None, Job("Invented", demolition=lines))
    book = openpyxl.load_workbook(path)
    assert book.sheetnames[1:3] == ["Scope lines", "Findings"]
    assert DEMOLITION_TAB in book.sheetnames
    rows = [[cell.value for cell in row] for row in book[DEMOLITION_TAB].iter_rows()]
    assert rows[0] == DEMOLITION_HEADINGS
    assert rows[1][:4] == ["AHU", "demo", 2, "existing controls removed"]
    assert rows[1][4] == 'S1 | "shown for removal" | MD-101, Drawings.pdf page 6'
    assert rows[2][:3] == ["FCU", "existing to remain", 1]
    assert rows[2][3] in ("", None)
    assert 'S2 | "FCU-2 remains" | M-201, Mech.pdf page 7' in rows[2][4]
    assert len(rows) == 3
