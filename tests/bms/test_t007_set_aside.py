"""T-007 Task 2.6 (reduced): the set-asides of the takeoff in the proposal and the Evidence Package.

A duplicate of one unit ("same unit labelled again", "part plan of one area") prints as one short '*' line in its
block, with both sheet numbers (ruling 14: one short line per item). A "demolition sheet" set-aside prints no '*'
line: the Demolition Scope block carries those units. The Evidence Package has a "Set aside" tab with every
set-aside. Invented data only: the sample fixtures and invented refs."""
from __future__ import annotations

import json
import sys
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from bms.model import Job, Ref, SetAside
from bms.playbook import load_playbook
from bms.render import _count_note, render
from bms.citations import SourceIds

FIXTURES = Path(__file__).resolve().parent / "fixtures"
PART_PLANS = Ref("C1", "EF-1 drawn on both part plans", "M-403 and M-404, Invented Set.pdf page 12")
LABELLED = Ref("C2", "EF-2 (EF-2A)", "M-201, Invented Set.pdf page 5")
DEMO = Ref("C3", "EF-3 shown for removal", "MD-101, Invented Set.pdf page 4")


@pytest.fixture(scope="module")
def playbook():
    return load_playbook(FIXTURES / "sample_playbook")


def _job(*set_aside: SetAside) -> Job:
    job = Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))
    return replace(job, set_aside=tuple(set_aside))


def _ef_notes(job: Job, playbook) -> list[str]:
    draft = render(job, playbook)
    return [l.text for l in draft.lines if l.kind == "note" and l.system_key == "EF"]


def _star(notes: list[str]) -> list[str]:
    return [n for n in notes if "counted once" in n]


def test_a_part_plan_set_aside_prints_one_short_star_line_with_both_sheets(playbook):
    job = _job(SetAside("EF", 2, PART_PLANS, "part plan of one area"))
    assert _star(_ef_notes(job, playbook)) == ["*EF counted once on part plans M-403 and M-404: 2 set aside, 3 counted"]


def test_a_unit_labelled_again_prints_one_short_star_line(playbook):
    job = _job(SetAside("EF", 4, LABELLED, "same unit labelled again", note="EF-2 labelled again as EF-2A on M-202"))
    assert _star(_ef_notes(job, playbook)) == ["*EF counted once, labelled again on M-201 and M-202: 4 set aside, 3 counted"]


def test_a_demolition_sheet_set_aside_prints_no_star_line(playbook):
    job = _job(SetAside("EF", 1, DEMO, "demolition sheet", note="1 unit on MD-101 set aside; 3 on the new-work sheets"))
    plain = _ef_notes(_job(), playbook)
    assert _ef_notes(job, playbook) == plain
    assert not any("MD-101" in n for n in _ef_notes(job, playbook))


def test_each_duplicate_is_its_own_line_and_only_in_its_family_block(playbook):
    job = _job(SetAside("EF", 2, PART_PLANS, "part plan of one area"),
               SetAside("EF", 4, LABELLED, "same unit labelled again"),
               SetAside("RTU", 2, PART_PLANS, "part plan of one area"))
    assert len(_star(_ef_notes(job, playbook))) == 2
    draft = render(job, playbook)
    rtu = [l.text for l in draft.lines if l.kind == "note" and l.system_key == "RTU"]
    assert _star(rtu) == ["*RTU counted once on part plans M-403 and M-404: 2 set aside, 2 counted"]
    vrf = [l.text for l in draft.lines if l.kind == "note" and l.system_key == "VRF"]
    assert _star(vrf) == []


def test_the_star_line_is_a_note_item_in_the_block_and_carries_no_tech_hours(playbook):
    job = _job(SetAside("EF", 2, PART_PLANS, "part plan of one area"))
    draft = render(job, playbook)
    block = next(b for b in draft.scope.blocks if "EF-1" in b.header)
    texts = [i.runs[0].text for i in block.items if i.kind == "note"]
    assert "*EF counted once on part plans M-403 and M-404: 2 set aside, 3 counted" in texts
    assert not any("hour" in t.lower() for t in texts if "counted once" in t)


def test_count_note_skips_a_future_item_and_a_set_aside_of_another_family():
    job = _job()
    ef = next(s for s in job.systems if s.key == "EF")
    ids = SourceIds.of_job(job)
    aside = (SetAside("EF", 1, DEMO, "future item"), SetAside("FCU", 2, PART_PLANS, "part plan of one area"),
             SetAside("EF", 1, DEMO, "not a unit of this family"))
    assert _count_note(ef, aside, ids) == []


def test_a_set_aside_with_no_sheet_still_names_both_numbers():
    job = _job()
    ef = next(s for s in job.systems if s.key == "EF")
    ref = Ref("C4", "EF plan placements 6", "equipment counts for EF")
    out = _count_note(ef, (SetAside("EF", 6, ref, "part plan of one area"),), SourceIds.of_job(job))
    assert out == ["*EF counted once on part plans: 6 set aside, 3 counted"]


def test_the_evidence_package_has_a_set_aside_tab_with_every_set_aside(tmp_path):
    from bms.evidence import SET_ASIDE_HEADINGS, SET_ASIDE_TAB, write_evidence

    tests = Path(__file__).resolve().parents[1]
    if str(tests) not in sys.path:
        sys.path.insert(0, str(tests))
    from test_evidence import _catalogue

    aside = (SetAside("EF", 2, PART_PLANS, "part plan of one area", "2 placements set aside; 1 unit"),
             SetAside("AHU", 1, DEMO, "demolition sheet"),
             SetAside("FCU", 9, LABELLED, "future item", "9 future thermostats; 30 counted"))
    path = write_evidence(tmp_path / "e.xlsx", _catalogue(), None, None, Job("Invented", set_aside=aside))
    book = openpyxl.load_workbook(path)
    assert book.sheetnames[1:3] == ["Scope lines", "Findings"]
    assert SET_ASIDE_TAB == "Set aside" and SET_ASIDE_TAB in book.sheetnames
    rows = [[cell.value for cell in row] for row in book[SET_ASIDE_TAB].iter_rows()]
    assert rows[0] == SET_ASIDE_HEADINGS
    assert [[c if c != "" else None for c in r[:4]] for r in rows[1:]] == [
        ["EF", 2, "part plan of one area", "2 placements set aside; 1 unit"],
                                         ["AHU", 1, "demolition sheet", None],
                                         ["FCU", 9, "future item", "9 future thermostats; 30 counted"]]
    assert rows[1][4] == 'C1 | "EF-1 drawn on both part plans" | M-403 and M-404, Invented Set.pdf page 12'


def test_no_set_asides_leaves_the_tab_saying_nothing_here(tmp_path):
    from bms.evidence import SET_ASIDE_TAB, write_evidence

    tests = Path(__file__).resolve().parents[1]
    if str(tests) not in sys.path:
        sys.path.insert(0, str(tests))
    from test_evidence import _catalogue

    path = write_evidence(tmp_path / "e.xlsx", _catalogue(), None, None, Job("Invented"))
    rows = [[cell.value for cell in row] for row in openpyxl.load_workbook(path)[SET_ASIDE_TAB].iter_rows()]
    assert rows[1][0] == "Nothing here."


def test_the_ahu_family_has_one_review_card_and_no_device_card():
    """Already done by T-005b (`review._questions` folds a line flagged only for the family count into the family
    card); this checks it holds with set-asides on the job."""
    from bms import review
    from bms.model import Count, Device, Family, FrontEnd, OtherCount, System

    plans = Ref("C1", "plan placements 21", "M-101, Invented Set.pdf page 2")
    count = Count(21, (plans,), (OtherCount(5, Ref("S1", "AHU rows 5", "M-501, Invented Set.pdf page 6")),))
    devices = tuple(Device(id=f"AHU.d{i}", slot="other_device", name=name, verb="Furnish and Install",
                           furnished_by="AME", installed_by="AME", wired_by="AME", quantity=count, refs=(plans,),
                           flag="quantity") for i, name in enumerate(("DDC Controller", "Freezestat"), start=1))
    system = System(key="AHU", name="Air Handling Unit", shape="Air Handling Unit", control="field",
                    families=(Family("Air Handling Unit", ("AHU-1", "AHU-2"), count),), devices=devices)
    job = Job(name="Invented", systems=(system,), front_end=FrontEnd(mode="new-system"),
              set_aside=(SetAside("AHU", 16, plans, "part plan of one area"),))
    assert [q.level for q in review._questions(job, ())] == ["family"]
