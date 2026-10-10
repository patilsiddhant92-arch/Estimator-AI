"""Integration: `bms.run.estimate` end to end on an invented job folder, in chat mode, with a fake answerer that
answers every request the run writes (reading passes, plan count sheets, write and revise, QC) from canned invented
answers. No real job data: the PDFs are built here."""
from __future__ import annotations

import dataclasses
import json
import re
import sys
from pathlib import Path

import openpyxl
import pytest

from bms.run import estimate

TESTS = Path(__file__).resolve().parents[1]
if str(TESTS) not in sys.path:
    sys.path.insert(0, str(TESTS))

from pdfs import FILLER, Page, make_pdf  # noqa: E402
from test_schedules import grid  # noqa: E402

W, H = 3024, 2160
JOB = "00-9999"
VAV_ROWS = [["VAV BOX SCHEDULE", "", "", "", ""],
            ["TAG", "SERVES", "MAX CFM", "HW GPM", "MANUFACTURER"],
            ["VAV-1", "ROOM 101", "600", "1.5", "INVENTED CO"],
            ["VAV-2", "ROOM 102", "700", "1.5", "INVENTED CO"],
            ["VAV-3", "ROOM 103", "500", "1.0", "INVENTED CO"],
            ["VAV-4", "ROOM 104", "500", "1.0", "INVENTED CO"]]


def _schedule_sheet() -> Page:
    lines, texts = grid(300, 300, [90, 120, 70, 70, 100], VAV_ROWS)
    return Page(W, H, ((200, 200, FILLER, 12), *texts,
                       (2640, 1860, "HVAC SCHEDULES", 20), (2800, 2000, "M-601", 40)), lines=lines)


PLAN = Page(W, H, ((200, 300, FILLER, 12), (900, 900, "VAV-1", 12), (900, 1000, "VAV-2", 12),
                   (900, 1100, "VAV-3", 12), (2640, 1860, "SECOND FLOOR PLAN", 20), (2800, 2000, "M-201", 40)))
SEQUENCE = Page(texts=((72, 72, "SECTION 23 09 93 - SEQUENCE OF OPERATIONS\n\n"
                                "A. VAV boxes: each VAV box shall maintain the space temperature setpoint with a "
                                "space temperature sensor furnished and installed by the controls contractor.\n\n"
                                "23 09 93 - 1", 10),))


@pytest.fixture
def job_folder(tmp_path) -> Path:
    root = tmp_path / f"{JOB}, 1 Invented Street"
    base = root / "2. Takeoffs" / "1. Base Takeoff"
    make_pdf(base / "Mech.pdf", _schedule_sheet(), PLAN)
    make_pdf(base / "230993.pdf", SEQUENCE)
    return root


# ---- the fake answerer: one canned, invented answer per request kind ------------------------------------------------

def _evidence_ref(prompt: str) -> dict:
    """A ref to the first source of the evidence that has words after its id, quoting its first words."""
    evidence = prompt.split("The evidence:", 1)[-1]
    for line in evidence.splitlines():
        found = re.match(r"\s*\[(S\d+)\]\s+(.*)", line)          # a document source, not a count record
        if found:
            words = re.findall(r"[A-Za-z0-9]+", found.group(2))[:4]
            if words:
                return {"source_id": found.group(1), "quote": " ".join(words), "where": "invented place"}
    raise AssertionError("the write request carries no evidence with an id")


def _first_slot(prompt: str) -> str:
    found = re.search(r"^- ([a-z0-9_]+) \((?:device|point)\)", prompt, re.M)
    return found.group(1) if found else "controller"


def _system_answer(request: dict) -> dict:
    prompt = request["prompt"]
    ref = _evidence_ref(prompt)
    key = request.get("plan_key") or "X"
    device = {"id": f"{key}.d1", "slot": _first_slot(prompt), "name": "Space Temperature Sensor",
              "verb": "Furnish and Install", "furnished_by": "AME", "installed_by": "AME", "wired_by": "AME",
              "quantity": {"value": None, "sources": [], "others": []}, "refs": [ref], "assumption": None,
              "flag": "quantity"}
    return {"key": key, "name": "answered", "shape": "", "control": "none", "families": [], "sequence_refs": [],
            "devices": [device], "clarifications": [{"kind": "per_drawings", "text": "Quantity as per drawings"}],
            "tech_hours": None, "alternate": "base"}


def _answer(request: dict) -> dict:
    kind = request["kind"]
    if kind == "read_pass":
        return {"facts": [], "contradictions": [], "confirmations": [], "understanding": "invented reading"}
    if kind == "count_tags":
        examined = list(range(1, len(request.get("tiles", [])) + 1))          # every tile opened
        return {"patterns": [{"pattern": p, "tiles_examined": examined, "tags": []}
                             for p in request.get("patterns", [])]}
    if kind == "write_system":
        return _system_answer(request)
    if kind == "qc_job":
        return {"findings": []}
    raise AssertionError(f"a request of an unknown kind: {kind}")


def answer_all(cache: Path) -> dict[str, int]:
    """Write an answer for every waiting request; how many of each kind."""
    kinds: dict[str, int] = {}
    for path in sorted((cache / "judge-chat" / "requests").glob("*.json")):
        request = json.loads(path.read_text(encoding="utf-8"))
        answer = Path(request["answer_file"])
        answer.parent.mkdir(parents=True, exist_ok=True)
        answer.write_text(json.dumps(_answer(request)), encoding="utf-8")
        kinds[request["kind"]] = kinds.get(request["kind"], 0) + 1
    return kinds


# ---- the run ---------------------------------------------------------------------------------------------------------

def test_a_run_waits_for_its_answers_then_writes_all_four_outputs(job_folder, tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"

    first = estimate(job_folder, out, cache, "chat")
    assert first.status == "waiting" and first.stage == "read" and first.waiting >= 1
    assert first.proposal is None

    stages, kinds, result = [first.stage], {}, first
    for _ in range(25):
        for kind, n in answer_all(cache).items():
            kinds[kind] = kinds.get(kind, 0) + n
        result = estimate(job_folder, out, cache, "chat")
        if result.status == "done":
            break
        assert result.status == "waiting" and result.waiting >= 1
        stages.append(result.stage)
    assert result.status == "done" and result.stage == "done", stages

    assert stages.index("read") < stages.index("count") < stages.index("write") < stages.index("qc")
    assert {"read_pass", "count_tags", "write_system", "qc_job"} <= set(kinds)
    assert result.proposal.name == f"{JOB} Proposal.docx" and result.proposal.exists()
    assert result.price_sheet.name == f"{JOB} Price Sheet.xlsx" and result.price_sheet.exists()
    assert result.evidence.name == f"{JOB} Evidence Package.xlsx" and result.evidence.exists()
    assert result.review == out / "Review.html" and result.review.exists()
    assert result.job_file.exists()
    assert [f for f in result.findings if f.severity == "must_fix"] == []

    assert (out / "count-proofs" / "proofs.json").exists()          # the count proofs, beside the review page
    assert 'src="count-proofs/' in result.review.read_text(encoding="utf-8")

    book = openpyxl.load_workbook(result.evidence)
    assert book.sheetnames[:3] == ["Summary", "Scope lines", "Findings"]
    rows = [[c.value for c in row] for row in book["Scope lines"].iter_rows()]
    at = {h: i for i, h in enumerate(rows[0])}
    device = rows[1]
    assert device[at["Verb"]] == "Furnish and Install"
    assert [device[at[h]] for h in ("Furnished by", "Installed by", "Wired by")] == ["AME", "AME", "AME"]
    assert "invented place" in device[at["Cited at"]] and device[at["Flag"]] == "quantity"

    # a further run with every answer cached finishes straight away, writing nothing new to ask
    again = estimate(job_folder, out, cache, "chat")
    assert again.status == "done" and again.waiting == 0
    assert list((cache / "judge-chat" / "requests").glob("*.json")) == []


# ---- what stays open after the one revise round ----------------------------------------------------------------------

def _reach(job_folder, out, cache, stage):
    """Run until the estimate is waiting at `stage`. Answers every earlier request."""
    result = None
    for _ in range(25):
        result = estimate(job_folder, out, cache, "chat")
        if result.status == "done" or result.stage == stage:
            return result
        assert result.status == "waiting" and result.waiting >= 1
        assert answer_all(cache)
    raise AssertionError(result)


def test_every_writer_is_given_the_tags_of_the_whole_job(job_folder, tmp_path, monkeypatch):
    """I2: the intake protects the same tag set as the renderer and the validator: every tag of the job."""
    from bms import ai, citations, run
    from bms.model import Job
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    seen: list[tuple] = []

    def spy(system, evidence, playbook, transport, **kw):
        seen.append((system.key, kw.get("tags")))
        raise ai.NoAnswer("spied", waiting=True)
    monkeypatch.setattr(run.ai, "write_system", spy)
    out, cache = tmp_path / "out", tmp_path / "cache"
    result = _reach(job_folder, out, cache, "write")
    assert result.stage == "write"
    job = Job.from_json(result.job_file.read_text(encoding="utf-8"))
    assert len(seen) == len(job.systems) >= 1, seen
    every = {t for s in job.systems for f in s.families for t in f.tags}
    for _, tags in seen:
        assert tags == tuple(citations.job_tags(job)) and set(tags) >= every, (tags, every)


def test_the_decision_rewrite_is_given_the_tags_of_the_whole_job(job_folder, tmp_path, monkeypatch):
    """A family-count change asks the writer again, and that ask gets every tag of the job. A device edit does not."""
    from bms import ai, citations, run
    from bms.model import Job
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    reached = _reach(job_folder, out, cache, "write")
    assert reached.stage == "write"
    saved = json.loads(reached.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item.get("families"))
    changed, family = system["key"], system["families"][0]["name"]
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "change", "family_name": family,
        "count": {"value": 99, "sources": [], "others": []}, "reason": "invented",
    }]), encoding="utf-8")
    assert answer_all(cache)
    real, seen = ai.rewrite_for_count, []

    def spy(system, evidence, playbook, transport, **kw):
        if "The estimator's count decisions" in evidence:
            seen.append(kw.get("tags"))
        return real(system, evidence, playbook, transport, **kw)
    monkeypatch.setattr(run.ai, "rewrite_for_count", spy)
    result = estimate(job_folder, out, cache, "chat")
    job = Job.from_json(result.job_file.read_text(encoding="utf-8"))
    assert seen and all(tags == tuple(citations.job_tags(job)) and tags for tags in seen), seen


def test_the_revise_is_given_the_tags_of_the_whole_job(job_folder, tmp_path, monkeypatch):
    """I2: the revise round's writer gets every tag of the job too."""
    from bms import ai, citations, run
    from bms.model import Finding, Job
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    seen: list = []

    def must_fix(job, playbook):
        return [Finding("Every line cites this job", s.key, None, "must_fix", "invented", "invented")
                for s in job.systems[:1]]

    def spy(system, findings, evidence, playbook, transport, **kw):
        seen.append(kw.get("tags"))
        return system
    monkeypatch.setattr(run, "validate", must_fix)
    monkeypatch.setattr(run.ai, "revise_system", spy)
    out, cache = tmp_path / "out", tmp_path / "cache"
    result = _reach(job_folder, out, cache, "never")
    job = Job.from_json(result.job_file.read_text(encoding="utf-8"))
    assert seen and all(tags == tuple(citations.job_tags(job)) and tags for tags in seen), seen


def _planned_ids(job) -> set[str]:
    """The source ids the planner's refs carry (every system's sequence refs and family counts)."""
    from bms.citations import plain_id
    return {plain_id(r.source_id) for s in job.systems
            for r in [*s.sequence_refs, *(r for f in s.families for r in f.count.sources)]}


def test_every_writer_is_given_the_source_ids_of_the_whole_job(job_folder, tmp_path, monkeypatch):
    """The intake refuses a client text that prints any source id of the job, not only this system's: the run
    passes every one (`citations.job_ids`) and every id of every system's evidence, one set for all writers, so a
    text the first writer accepts is never an issued id once a later system cites it (secondary review M1)."""
    from bms import ai, ai_evidence, citations, run
    from bms.model import Job
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    seen: list = []
    evidence_ids: set = set()

    def spy(system, evidence, playbook, transport, **kw):
        seen.append(kw.get("ids"))
        evidence_ids.update(ai_evidence.cited_ids(evidence))
        raise ai.NoAnswer("spied", waiting=True)
    monkeypatch.setattr(run.ai, "write_system", spy)
    out, cache = tmp_path / "out", tmp_path / "cache"
    result = _reach(job_folder, out, cache, "write")
    job = Job.from_json(result.job_file.read_text(encoding="utf-8"))
    assert len(seen) == len(job.systems) >= 1, seen
    assert evidence_ids, "the evidence names no source id"
    for ids in seen:
        assert ids == seen[0] and set(ids) >= set(citations.job_ids(job)) | _planned_ids(job) | evidence_ids, ids


def test_the_decision_rewrite_and_the_revise_are_given_the_source_ids_of_the_whole_job(job_folder, tmp_path,
                                                                                         monkeypatch):
    from bms import ai, run
    from bms.model import Finding, Job
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    reached = _reach(job_folder, out, cache, "write")
    assert reached.stage == "write"
    saved = json.loads(reached.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item.get("families"))
    changed, family = system["key"], system["families"][0]["name"]
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "change", "family_name": family,
        "count": {"value": 99, "sources": [], "others": []}, "reason": "invented",
    }]), encoding="utf-8")
    assert answer_all(cache)
    real, rewritten, revised = ai.rewrite_for_count, [], []

    def spy(system, evidence, playbook, transport, **kw):
        if "The estimator's count decisions" in evidence:
            rewritten.append(kw.get("ids"))
        return real(system, evidence, playbook, transport, **kw)

    def revise_spy(system, findings, evidence, playbook, transport, **kw):
        revised.append(kw.get("ids"))
        return system

    def must_fix(job, playbook):
        return [Finding("Every line cites this job", s.key, None, "must_fix", "invented", "invented")
                for s in job.systems[:1]]
    monkeypatch.setattr(run.ai, "rewrite_for_count", spy)
    monkeypatch.setattr(run.ai, "revise_system", revise_spy)
    monkeypatch.setattr(run, "validate", must_fix)
    result = _reach(job_folder, out, cache, "never")
    job = Job.from_json(result.job_file.read_text(encoding="utf-8"))
    for seen in (rewritten, revised):
        assert seen and all(ids and set(ids) >= _planned_ids(job) for ids in seen), seen


def test_every_writer_the_decision_rewrite_and_the_revise_are_given_the_sheets_of_the_whole_plan(job_folder, tmp_path,
                                                                                                 monkeypatch):
    """K, the sheet numbers the planner's wheres name (`citations.job_sheets`), is the whole plan's and is passed to
    every writer, the count rewrite and the revise, so a sheet only another system's where names is accepted at
    intake as the validator and the renderer accept it (review of record, cleanup round 4, L-B). It is worked out per
    call: the writers all get the same K, the count rewrite and the revise the K of the job after the decisions
    (lead/detector-fix3, R-1). Since the merge of 018c a device decision is applied without asking again, so the
    decision here is a family-count change, which asks `ai.rewrite_for_count`; the real one is called, and it takes
    `sheets` (on rebuild 7befe74 it did not, and the count rewrite stopped with a TypeError)."""
    from bms import ai, citations, run
    from bms.model import Finding, Job
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    real_write, real_rewrite, written, rewritten, revised = ai.write_system, ai.rewrite_for_count, [], [], []

    def spy(system, evidence, playbook, transport, **kw):
        written.append(kw.get("sheets"))
        return real_write(system, evidence, playbook, transport, **kw)

    def rewrite_spy(system, evidence, playbook, transport, **kw):
        rewritten.append(kw.get("sheets"))
        return real_rewrite(system, evidence, playbook, transport, **kw)

    def revise_spy(system, findings, evidence, playbook, transport, **kw):
        revised.append(kw.get("sheets"))
        return system

    def must_fix(job, playbook):
        return [Finding("Every line cites this job", s.key, None, "must_fix", "invented", "invented")
                for s in job.systems[:1]]
    monkeypatch.setattr(run.ai, "write_system", spy)
    monkeypatch.setattr(run.ai, "rewrite_for_count", rewrite_spy)
    out, cache = tmp_path / "out", tmp_path / "cache"
    first = _reach(job_folder, out, cache, "write")
    assert first.stage == "write"
    planned = Job.from_json(first.job_file.read_text(encoding="utf-8"))
    system = planned.systems[0]
    family = system.families[0]
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": system.key, "action": "change", "family_name": family.name,
        "count": dataclasses.replace(family.count, value=99).to_dict(), "reason": "invented",
    }]), encoding="utf-8")
    assert answer_all(cache)
    monkeypatch.setattr(run.ai, "revise_system", revise_spy)
    monkeypatch.setattr(run, "validate", must_fix)
    result = _reach(job_folder, out, cache, "never")
    job = Job.from_json(result.job_file.read_text(encoding="utf-8"))
    assert written and rewritten and revised, (written, rewritten, revised)
    assert all(sheets is not None and tuple(sheets) == tuple(written[0]) for sheets in written), written
    assert set(written[0]) == set(citations.job_sheets(planned)), (written[0], citations.job_sheets(planned))
    for sheets in (*rewritten, *revised):
        assert sheets is not None and set(sheets) == set(citations.job_sheets(job)), (sheets, job)


def test_a_killed_write_resumes_from_cached_answers_without_asking_again(job_folder, tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    waiting = _reach(job_folder, out, cache, "write")
    assert waiting.stage == "write"
    folder = cache / "judge-chat" / "requests"
    requests = sorted(folder.glob("*.json"))
    assert requests and {json.loads(p.read_text(encoding="utf-8"))["kind"] for p in requests} == {"write_system"}
    saved = {p.name: p.read_bytes() for p in requests}
    stamped = {p.name: p.stat().st_mtime_ns for p in requests}

    # killed after the requests were written, before any answer: the same files, not a second ask
    resumed = estimate(job_folder, out, cache, "chat")
    assert resumed.status == "waiting" and resumed.stage == "write"
    assert {p.name: p.read_bytes() for p in folder.glob("*.json")} == saved
    assert {p.name: p.stat().st_mtime_ns for p in folder.glob("*.json")} == stamped
    assert not any(p.name in saved for p in (folder / "old").glob("*.json"))

    # those answers are now cached, but a kill left the request files behind: resume must drop them
    for path in folder.glob("*.json"):
        request = json.loads(path.read_text(encoding="utf-8"))
        answer = Path(request["answer_file"])
        answer.parent.mkdir(parents=True, exist_ok=True)
        answer.write_text(json.dumps(_answer(request)), encoding="utf-8")
    consumed = estimate(job_folder, out, cache, "chat")
    assert consumed.stage != "write"
    for name, payload in saved.items():
        (folder / name).write_bytes(payload)
    estimate(job_folder, out, cache, "chat")
    assert not any(p.name in saved for p in folder.rglob("*.json"))


def test_a_malformed_decisions_file_stops_with_a_plain_message(job_folder, tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    (out / "decisions.json").write_text("{", encoding="utf-8")
    assert answer_all(cache)
    result = estimate(job_folder, out, cache, "chat")
    assert result.messages and any("decisions.json" in line and "Stopped" in line for line in result.messages)
    assert result.proposal is None


def _plant(cache: Path, names_for: dict[str, tuple[str, ...]]) -> None:
    """Answer every waiting request. A listed system gets those device names, in order, as d1, d2, ..."""
    for path in sorted((cache / "judge-chat" / "requests").glob("*.json")):
        request = json.loads(path.read_text(encoding="utf-8"))
        answer = _answer(request)
        key = request.get("plan_key")
        if request.get("kind") == "write_system" and key in names_for:
            base = dict(answer["devices"][0])
            answer["devices"] = []
            for number, name in enumerate(names_for[key], start=1):
                device = dict(base)
                device["id"] = f"{key}.d{number}"
                device["name"] = name
                answer["devices"].append(device)
        dest = Path(request["answer_file"])
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(json.dumps(answer), encoding="utf-8")


def test_a_device_decision_is_not_sent_back_to_the_writer(job_folder, tmp_path, monkeypatch):
    """A delete is applied to the written device. The writer is not asked again, so it cannot put the line back."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    changed = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "delete", "device_id": f"{changed}.d1", "reason": "invented",
    }]), encoding="utf-8")
    assert answer_all(cache)
    result = estimate(job_folder, out, cache, "chat")
    assert result.stage == "qc"
    fresh = [json.loads(p.read_text(encoding="utf-8"))
             for p in (cache / "judge-chat" / "requests").glob("*.json")]
    assert all(item.get("kind") != "write_system" for item in fresh)
    saved = json.loads(result.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item["key"] == changed)
    assert system["devices"] == []


def test_deleting_the_damper_keeps_the_start_stop_when_the_writer_would_renumber(job_folder, tmp_path, monkeypatch):
    """The written devices are Controller, Start Stop, Damper Actuator. The estimator deletes the damper.
    If the writer were asked again it would number the damper d2 and drop the start/stop. The damper is the
    device that ends up gone."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    changed = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "delete", "device_id": f"{changed}.d3", "reason": "invented",
    }]), encoding="utf-8")
    _plant(cache, {changed: ("Controller", "Start Stop", "Damper Actuator")})
    result = estimate(job_folder, out, cache, "chat")
    if result.stage == "write":
        _plant(cache, {changed: ("Controller", "Damper Actuator")})
        result = estimate(job_folder, out, cache, "chat")
    saved = json.loads(result.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item["key"] == changed)
    names = [device["name"] for device in system["devices"]]
    assert names == ["Controller", "Start Stop"]
    assert names.count("Damper Actuator") == 0 and len(names) == len(set(names))


def test_an_accepted_quantity_flag_stays_cleared_on_the_rerun(job_folder, tmp_path, monkeypatch):
    """An estimator accepts a quantity flag. On re-running, the device flag remains 'none'."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    changed = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    _plant(cache, {changed: ("Controller", "Start Stop")})
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "accept", "device_id": f"{changed}.d1", "reason": "invented",
    }]), encoding="utf-8")
    result = estimate(job_folder, out, cache, "chat")
    saved = json.loads(result.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item["key"] == changed)
    assert system["devices"][0]["flag"] == "none"


def test_a_delete_made_on_a_revised_page_removes_the_line_shown_on_the_page(job_folder, tmp_path, monkeypatch):
    """The review page showed the Damper as d2 with was snapshot.
    Deleting it removes the Damper from the first-write job even though Damper was d3 there."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    changed = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    _plant(cache, {changed: ("Controller", "Start Stop", "Damper Actuator")})
    answer_path = Path(json.loads(requests[0].read_text(encoding="utf-8"))["answer_file"])
    planted = json.loads(answer_path.read_text(encoding="utf-8"))
    damper = next(d for d in planted["devices"] if d["name"] == "Damper Actuator")
    damper_was = dict(damper)
    damper_was["id"] = f"{changed}.d2"  # as shown on the revised review page
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "delete", "device_id": f"{changed}.d2", "was": damper_was, "reason": "invented",
    }]), encoding="utf-8")
    result = estimate(job_folder, out, cache, "chat")
    saved = json.loads(result.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item["key"] == changed)
    names = [device["name"] for device in system["devices"]]
    assert names == ["Controller", "Start Stop"]


def test_a_device_delete_stays_honored_because_the_writer_is_not_asked(job_folder, tmp_path, monkeypatch):
    """The deleted line does not come back, and the review page does not call the decision unhonored."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    changed = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": changed, "action": "delete", "device_id": f"{changed}.d1", "reason": "invented",
    }]), encoding="utf-8")
    result = None
    for _ in range(25):
        assert answer_all(cache)
        result = estimate(job_folder, out, cache, "chat")
        if result.status == "done":
            break
    assert result is not None and result.status == "done", result
    assert not any(finding.what.startswith("decision not honored") for finding in result.findings)
    saved = json.loads(result.job_file.read_text(encoding="utf-8"))
    system = next(item for item in saved["systems"] if item["key"] == changed)
    assert system["devices"] == []


def _reworded_damper_run(job_folder, tmp_path, monkeypatch):
    """A run to done where QC's must_fix makes the revise re-word the damper line, as the review page then shows it.
    Returns (out, cache, system key, the damper as shown on the page)."""
    import dataclasses

    from bms import ai, run
    from bms.model import Finding

    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    key = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    _plant(cache, {key: ("Controller", "Start Stop", "Damper Actuator")})
    real_validate = run.validate
    monkeypatch.setattr(run, "validate", lambda job, playbook: [*real_validate(job, playbook), Finding(
        "Every line cites this job", key, None, "must_fix", "invented must_fix", "invented fix")])

    def reword(system, findings, evidence, playbook, transport, **kw):   # stands for the cached revise answer
        return dataclasses.replace(system, devices=tuple(
            dataclasses.replace(d, name="Damper Actuator, motorized") if d.name == "Damper Actuator" else d
            for d in system.devices))
    monkeypatch.setattr(ai, "revise_system", reword)
    first = _reach(job_folder, out, cache, "done")
    assert first.status == "done"
    shown = next(s for s in json.loads(first.job_file.read_text(encoding="utf-8"))["systems"] if s["key"] == key)
    return out, cache, key, next(d for d in shown["devices"] if d["name"] == "Damper Actuator, motorized")


def test_a_delete_of_a_line_the_revise_reworded_removes_that_line(job_folder, tmp_path, monkeypatch):
    """The estimator deletes the line as the review page showed it, after the revise re-worded it. The next run
    applies the delete to the revised job too, so the line is gone and nothing is reported."""
    out, cache, key, damper = _reworded_damper_run(job_folder, tmp_path, monkeypatch)
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": key, "action": "delete", "device_id": damper["id"], "was": damper,
        "reason": "invented"}]), encoding="utf-8")
    again = _reach(job_folder, out, cache, "done")
    assert again.status == "done"
    final = next(s for s in json.loads(again.job_file.read_text(encoding="utf-8"))["systems"] if s["key"] == key)
    assert [d["name"] for d in final["devices"]] == ["Controller", "Start Stop"]
    assert not [f for f in again.findings if f.rule == "decision not honored"]


def test_a_decision_on_a_line_no_run_has_shows_on_the_review_page(job_folder, tmp_path, monkeypatch):
    """A saved change names a line that neither the write nor the revise has (the answers changed since the page
    was made). The run cannot apply it, and the review page says so."""
    out, cache, key, damper = _reworded_damper_run(job_folder, tmp_path, monkeypatch)
    stale = dict(damper, name="Damper Actuator, as an older run wrote it")
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": key, "action": "change", "device_id": damper["id"], "was": stale,
        "device": dict(stale, name="Damper Actuator, edited"), "reason": "invented"}]), encoding="utf-8")
    again = _reach(job_folder, out, cache, "done")
    assert again.status == "done"
    assert [f.what for f in again.findings if f.rule == "decision not honored"] == [
        f"decision could not be applied: change {damper['id']}"]
    assert "decision could not be applied" in again.review.read_text(encoding="utf-8")


def test_a_qc_must_fix_the_revise_left_as_it_was_stays_open():
    import dataclasses

    from bms.model import Finding, Job
    from bms.run import _untouched

    job = Job.from_dict(json.loads((TESTS / "bms" / "fixtures" / "sample_job.json").read_text(encoding="utf-8")))
    ef = job.system("EF")
    changed = job.with_system(dataclasses.replace(ef, devices=(dataclasses.replace(ef.devices[0], name="New"),
                                                               *ef.devices[1:])))
    on = lambda device=None, key="EF": Finding("Every line cites this job", key, device, "must_fix", "w", "f")
    assert _untouched(on("EF.d2"), job, changed) and not _untouched(on("EF.d1"), job, changed)
    assert not _untouched(on(None), job, changed) and _untouched(on(None, "RTU"), job, changed)
    assert _untouched(on(None, ""), job, changed)


def test_revised_keeps_accepted_decisions_without_index_error(monkeypatch):
    import dataclasses
    from bms.model import Decision, Job
    from bms.run import _revised
    from bms.playbook import load_playbook
    import bms.ai as ai

    job = Job.from_dict(json.loads((TESTS / "bms" / "fixtures" / "sample_job.json").read_text(encoding="utf-8")))
    ef = job.system("EF")
    pb = load_playbook()
    monkeypatch.setattr(ai, "revise_system", lambda system, findings, evidence, playbook, transport, **kwargs:
                        dataclasses.replace(system, devices=tuple(
                            dataclasses.replace(d, flag="quantity") if d.id == "EF.d1" else d for d in system.devices)))
    decisions = [Decision("EF", "accept", device_id="EF.d1")]
    res = _revised(ef, [], "", pb, None, job=job, decisions=decisions)
    assert res.key == "EF"
    assert next(d for d in res.devices if d.id == "EF.d1").flag == "none"


# ---- API mode end to end (TASK-005) ----------------------------------------------------------------------------------

class _ApiBlock:
    def __init__(self, text):
        self.type, self.text = "text", text


class _ApiResponse:
    def __init__(self, text, stop="end_turn", model="claude-opus-5-5"):
        self.content, self.stop_reason, self.model = [_ApiBlock(text)], stop, model


class _ApiStream:
    def __init__(self, response):
        self.response = response

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self.response


class FakeAnthropicClient:
    def __init__(self, qc_must_fix_once: bool = True):
        self.calls: list[dict] = []
        self.messages = self
        self.qc_must_fix_once = qc_must_fix_once
        self.qc_count = 0

    def stream(self, **params):
        self.calls.append(params)
        output_config = params.get("output_config", {})
        schema = output_config.get("format", {}).get("schema", {})
        props = schema.get("properties", {})
        messages = params.get("messages", [])
        content = messages[0]["content"] if messages else ""
        if isinstance(content, list):
            prompt = content[0]["text"] if content and content[0].get("type") == "text" else ""
            n_images = sum(1 for b in content if b.get("type") == "image")
        else:
            prompt = str(content)
            n_images = 0

        # 1. read_pass (from Judge in staged reading)
        if "facts" in props and "understanding" in props:
            answer = {"facts": [], "contradictions": [], "confirmations": [], "understanding": "invented reading"}
            return _ApiStream(_ApiResponse(json.dumps(answer)))

        # 2. count_tags (from counts_vision)
        if "patterns" in props:
            patterns = [m.group(1).strip() for m in re.finditer(r"^Pattern \d+:\s*(.+)$", prompt, re.M)]
            tiles_examined = list(range(1, n_images + 1)) if n_images else [1]
            answer = {"patterns": [{"pattern": p, "tiles_examined": tiles_examined, "tags": []} for p in patterns]}
            return _ApiStream(_ApiResponse(json.dumps(answer)))

        # 3. qc_job (from ai.qc_job)
        if "findings" in props and "properties" in schema:
            self.qc_count += 1
            if self.qc_must_fix_once and self.qc_count == 1:
                m = re.search(r"- ([A-Z0-9_-]+):", prompt)
                sys_key = m.group(1) if m else "VAV"
                finding = {"rule": "Every line cites this job", "system_key": sys_key, "device_id": f"{sys_key}.d1",
                           "severity": "must_fix", "what": "citation needs verification", "fix": "check quote"}
                answer = {"findings": [finding]}
            else:
                answer = {"findings": []}
            return _ApiStream(_ApiResponse(json.dumps(answer)))

        # 4. write_system / revise_system (from ai.write_system / ai.revise_system)
        if "devices" in props and "families" in props:
            ref = _evidence_ref(prompt)
            m = re.search(r'"key":\s*"([^"]+)"', prompt)
            key = m.group(1) if m else "VAV"
            slot = _first_slot(prompt)
            device = {"id": f"{key}.d1", "slot": slot, "name": "Space Temperature Sensor",
                      "verb": "Furnish and Install", "furnished_by": "AME", "installed_by": "AME", "wired_by": "AME",
                      "quantity": {"value": None, "sources": [], "others": []}, "refs": [ref], "assumption": None,
                      "flag": "none"}
            answer = {"key": key, "name": "answered", "shape": "", "control": "none", "families": [], "sequence_refs": [],
                      "devices": [device], "clarifications": [{"kind": "per_drawings", "text": "Quantity as per drawings"}],
                      "tech_hours": None, "alternate": "base"}
            return _ApiStream(_ApiResponse(json.dumps(answer)))

        raise AssertionError(f"Unknown call: schema props={list(props)}")


def test_api_mode_end_to_end_with_fake_client(job_folder, tmp_path, monkeypatch):
    import anthropic
    from bms.playbook import load_playbook

    fake_client = FakeAnthropicClient(qc_must_fix_once=True)
    monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: fake_client)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "fake-key-for-test")

    out, cache = tmp_path / "out_api", tmp_path / "cache_api"

    result = estimate(job_folder, out, cache, mode="api")
    assert result.status == "done" and result.stage == "done"

    # All outputs are written
    assert result.proposal.name == f"{JOB} Proposal.docx" and result.proposal.exists()
    assert result.price_sheet.name == f"{JOB} Price Sheet.xlsx" and result.price_sheet.exists()
    assert result.evidence.name == f"{JOB} Evidence Package.xlsx" and result.evidence.exists()
    assert result.review == out / "Review.html" and result.review.exists()
    assert result.job_file.exists()

    # In API mode, no chat requests are created
    assert not (cache / "judge-chat").exists() or list((cache / "judge-chat" / "requests").glob("*.json")) == []

    # Prompt caching is set on the first system block for every call from ApiTransport
    playbook = load_playbook()
    from bms.counts_vision import INSTRUCTIONS as COUNT_INSTRUCTIONS
    api_calls = [c for c in fake_client.calls if isinstance(c.get("system"), list)]
    assert len(api_calls) >= 4  # count_tags, write_system, qc_job, revise_system
    for call in api_calls:
        first_block = call["system"][0]
        assert first_block.get("cache_control") == {"type": "ephemeral"}
        # Write, qc, and revise cache the Playbook rules; count_tags caches count instructions
        assert first_block["text"] in (playbook.rules_text, COUNT_INSTRUCTIONS)
    rules_calls = [c for c in api_calls if c["system"][0]["text"] == playbook.rules_text]
    assert len(rules_calls) >= 3  # at least write, qc, revise

    # Second run with all answers cached finishes immediately with no additional calls
    calls_before = len(fake_client.calls)
    cached_run = estimate(job_folder, out, cache, mode="api")
    assert cached_run.status == "done" and cached_run.stage == "done"
    assert len(fake_client.calls) == calls_before


def test_api_mode_prompt_caching_and_refusal_retry(tmp_path):
    from bms import ai
    from bms.model import json_schema, System
    from bms.playbook import load_playbook

    # A device count of 0 is a domain error caught by _check_system, causing refusal
    bad_answer = {"key": "VAV", "name": "VAV", "shape": "", "control": "none", "families": [], "sequence_refs": [],
                  "devices": [{"id": "VAV.d1", "slot": "controller", "name": "Sensor", "verb": "Furnish",
                               "furnished_by": "AME", "installed_by": "AME", "wired_by": "AME",
                               "quantity": {"value": 0, "sources": [], "others": []},
                               "refs": [{"source_id": "S1", "quote": "quote", "where": "where"}],
                               "assumption": None, "flag": "none"}],
                  "clarifications": [], "tech_hours": None, "alternate": "base"}

    # Client returns bad answer on first call, good answer on retry
    good_answer = dict(bad_answer, devices=[dict(bad_answer["devices"][0],
                                                 quantity={"value": 1, "sources": [], "others": []})])
    client = FakeAnthropicClient()
    # Mock responses for client: 1 bad then 1 good
    responses = [_ApiResponse(json.dumps(bad_answer)), _ApiResponse(json.dumps(good_answer))]
    client.stream = lambda **params: (client.calls.append(params), _ApiStream(responses.pop(0)))[1]

    transport = ai.ApiTransport(tmp_path / "retry_cache", model="claude-test", client=client)

    system = System("VAV", "VAV", "", "none", ())
    # Check that immediate retry succeeds after 1 refusal
    written = ai.write_system(system, "[S1] quote from where", load_playbook(), transport)
    assert len(client.calls) == 2
    # Check prompt caching was set on every attempt
    for call in client.calls:
        assert call["system"][0]["cache_control"] == {"type": "ephemeral"}
    # Check that retry prompt included the failure reason
    assert "Your earlier answer to this question could not be used" in client.calls[1]["messages"][0]["content"][0]["text"]

    # Now verify that failing the immediate retry raises NoAnswer
    client_fail = FakeAnthropicClient()
    client_fail.stream = lambda **params: (client_fail.calls.append(params), _ApiStream(_ApiResponse(json.dumps(bad_answer))))[1]
    transport_fail = ai.ApiTransport(tmp_path / "fail_cache", model="claude-test", client=client_fail)

    with pytest.raises(ai.NoAnswer) as exc:
        ai.write_system(system, "[S1] quote from where", load_playbook(), transport_fail)
    assert len(client_fail.calls) == 2
    assert "could not be used" in exc.value.reason


def test_a_revised_line_quoting_words_its_source_lacks_stays_open(job_folder, tmp_path, monkeypatch):
    """AI-4: QC is not asked again after the revise, so the run checks the revised job's quotes itself: a revised
    line whose quote its source does not hold is a must_fix finding on the review page, never closed."""
    from bms import ai, run
    from bms.model import Finding

    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    requests = sorted((cache / "judge-chat" / "requests").glob("*.json"))
    key = json.loads(requests[0].read_text(encoding="utf-8"))["plan_key"]
    _plant(cache, {key: ("Controller", "Start Stop", "Damper Actuator")})
    real_validate = run.validate
    monkeypatch.setattr(run, "validate", lambda job, playbook: [*real_validate(job, playbook), Finding(
        "Every line cites this job", key, None, "must_fix", "invented must_fix", "invented fix")])
    cited = []

    def misquote(system, findings, evidence, playbook, transport, **kw):   # stands for a revise answer let through
        def bad(device):
            ref = dataclasses.replace(device.refs[0], quote="invented words no source of this job says")
            cited.append(ref.source_id)
            return dataclasses.replace(device, name="Damper Actuator, motorized", refs=(ref,))
        return dataclasses.replace(system, devices=tuple(bad(d) if d.name == "Damper Actuator" else d
                                                         for d in system.devices))
    monkeypatch.setattr(ai, "revise_system", misquote)
    result = _reach(job_folder, out, cache, "done")
    assert result.status == "done" and cited
    open_quotes = [f for f in result.findings if f.severity == "must_fix" and f.system_key == key
                   and "invented words no source of this job says" in f.what]
    assert len(open_quotes) == 1 and cited[0].strip("[]") in open_quotes[0].what


def test_a_line_left_at_the_old_family_count_after_a_count_rewrite_takes_the_new_count():
    """AI-1 safety net: the estimator changes the Exhaust Fan count from 3 to 5 and the count rewrite's answer
    leaves a line counted from that family at 3 (or drops it, and it is put back at 3). The run gives it 5. A line
    at another count, or one not counted from the family, keeps its count."""
    from bms import run
    from bms.model import Job
    sample = Job.from_dict(json.loads((TESTS / "bms" / "fixtures" / "sample_job.json").read_text(encoding="utf-8")))
    before = sample.system("EF")
    family = before.families[0]
    after = dataclasses.replace(before, families=(dataclasses.replace(
        family, count=dataclasses.replace(family.count, value=5)),))
    carried = run._carry_new_counts(before, after, after)
    assert [(d.id, d.quantity.value) for d in carried.devices] == [("EF.d1", 5), ("EF.d2", 3), ("EF.d3", 2)]
    assert carried.devices[0].quantity.sources == before.devices[0].quantity.sources
    assert run._carry_new_counts(before, before, before) == before          # no family count changed


def test_a_line_the_estimator_changed_keeps_its_typed_count_through_a_count_rewrite():
    """AI-1: the estimator changes the Exhaust Fan count to 5 and, in the same save, types 3 on EF.d1. Neither the
    rewrite's answer (5) nor the safety net replaces the typed count."""
    from bms import run
    from bms.model import Job
    sample = Job.from_dict(json.loads((TESTS / "bms" / "fixtures" / "sample_job.json").read_text(encoding="utf-8")))
    before = sample.system("EF")
    family = before.families[0]
    after = dataclasses.replace(before, families=(dataclasses.replace(
        family, count=dataclasses.replace(family.count, value=5)),))
    first = before.devices[0]
    answered = dataclasses.replace(after, devices=(
        dataclasses.replace(first, quantity=dataclasses.replace(first.quantity, value=5)), *after.devices[1:]))
    carried = run._carry_new_counts(before, after, answered, frozenset({"EF.d1"}))
    assert carried.devices[0].quantity == first.quantity


def test_a_renamed_line_still_takes_a_changed_family_count():
    """AI-1 (lead review): a rename sets no count, so it does not freeze the line's count through a count rewrite.
    The estimator renames EF.d1 (its count as shown) and changes the Exhaust Fan count 3 -> 5; the answer leaves
    EF.d1 at 3: it carries 5. A change that typed a count (EF.d1 at 4) keeps it."""
    from bms import run
    from bms.model import Decision, Job
    sample = Job.from_dict(json.loads((TESTS / "bms" / "fixtures" / "sample_job.json").read_text(encoding="utf-8")))
    before = sample.system("EF")
    family = before.families[0]
    first = before.devices[0]
    renamed = dataclasses.replace(first, name="Exhaust Fan Controller")
    decisions = [Decision("EF", "change", device_id="EF.d1", device=renamed, was=first),
                 Decision("EF", "change", family_name=family.name,
                          count=dataclasses.replace(family.count, value=5))]
    assert run._typed_counts(decisions, "EF", sample) == frozenset()
    after = dataclasses.replace(before, families=(dataclasses.replace(family, count=decisions[1].count),),
                                devices=(renamed, *before.devices[1:]))
    carried = run._carry_new_counts(before, after, after, run._typed_counts(decisions, "EF", sample))
    assert (carried.devices[0].name, carried.devices[0].quantity.value) == ("Exhaust Fan Controller", 5)
    counted = dataclasses.replace(first, quantity=dataclasses.replace(first.quantity, value=4))
    typed = [Decision("EF", "change", device_id="EF.d1", device=counted, was=first)]
    assert run._typed_counts(typed, "EF", sample) == frozenset({"EF.d1"})


def test_a_run_says_it_wrote_each_file_it_touched_under_out(job_folder, tmp_path, monkeypatch, capsys):
    """Job.json, the decisions history, the count proofs' manifest and the run state each get a "Wrote" line."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert _reach(job_folder, out, cache, "write").stage == "write"
    key = json.loads(sorted((cache / "judge-chat" / "requests").glob("*.json"))[0].read_text(encoding="utf-8"))["plan_key"]
    _plant(cache, {key: ("Controller", "Start Stop")})
    (out / "decisions.json").write_text(json.dumps([{
        "system_key": key, "action": "accept", "device_id": f"{key}.d1", "reason": "invented"}]), encoding="utf-8")
    capsys.readouterr()
    result = estimate(job_folder, out, cache, "chat")
    for _ in range(25):
        if result.status == "done":
            break
        assert answer_all(cache)
        result = estimate(job_folder, out, cache, "chat")
    assert result.status == "done"
    printed = capsys.readouterr().out
    for path in (result.job_file, out / "decisions history.json", out / "count-proofs" / "proofs.json",
                 out / "run state.json"):
        assert path.exists(), path
        assert f"Wrote {path}" in printed.splitlines(), path
