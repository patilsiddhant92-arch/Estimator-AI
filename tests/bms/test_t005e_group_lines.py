"""T-005e evening correction (lead, NOTES-2026-10-11 "Corrections"): the group line follows the Master Template entry.

"The Master Template has no sub-headings" was wrong as a general rule. Where the system's Master Template entry prints a
group line over its stacked Qty headers (Chilled Water System, Hot Water System, VRF System, Kitchen Exhaust
Ventilation), the block keeps it; where the entry prints none (Heat Exchanger System, Condenser Water System, an air
handler, ...), the block prints none.

Invented fixtures only: invented tags, names and counts; the real Playbook's shapes (from the Master Template).
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from bms import order
from bms.model import Count, Device, Family, Job, Ref, System
from bms.playbook import load_playbook
from bms.render import render

FIXTURES = Path(__file__).resolve().parent / "fixtures"
REF = Ref(source_id="S1", quote="q", where="M-9, Invented Drawings.pdf page 3")


@pytest.fixture(scope="module")
def house():
    return load_playbook()          # the real Playbook (shapes from the Master Template)


@pytest.fixture
def job() -> Job:
    return Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))


def _controller(key: str) -> Device:
    return Device(id=f"{key}.d1", slot="controller", name="DDC Controller", verb="Furnish and Install",
                  furnished_by="AME", installed_by="AME", wired_by="AME", quantity=Count(value=1), refs=(REF,))


def _system(key: str, name: str, shape: str, families: list[tuple[str, tuple[str, ...]]]) -> System:
    return System(key=key, name=name, shape=shape, control="field",
                  families=tuple(Family(f, tags, Count(value=len(tags))) for f, tags in families),
                  devices=(_controller(key),), tech_hours=6.0)


def _draft(job, house, system):
    from dataclasses import replace
    return render(replace(job, systems=(system,), demolition=()), house)


# the entries whose Master Template text prints a group line over the stacked Qty headers, with invented families
GROUPED = [
    ("ZCW", "Chilled Water System", [("Chiller", ("ZCH-7", "ZCH-8")), ("Chilled Water Pump", ("ZCP-1", "ZCP-2"))],
     "Chilled Water System"),
    ("ZHW", "Hot Water System", [("Boiler", ("ZB-3",)), ("Hot Water Pump", ("ZHP-4", "ZHP-5"))], "Hot Water System"),
    ("ZVR", "VRF System", [("Indoor Unit", ("ZIU-1", "ZIU-2", "ZIU-3")), ("Condensing Unit", ("ZCU-1",))],
     "VRF System"),
    ("ZKX", "Kitchen Exhaust Ventilation", [("Kitchen Exhaust System", ("ZKEF-1",)), ("Makeup Air Unit", ("ZMAU-2",))],
     "Kitchen Exhaust Ventilation"),
]


@pytest.mark.parametrize("key, shape, families, group", GROUPED)
def test_an_entry_with_a_group_line_keeps_it_over_its_stacked_headers(job, house, key, shape, families, group):
    draft = _draft(job, house, _system(key, shape, shape, families))
    blocks = draft.scope.blocks
    assert blocks[0].header == group and blocks[0].numbered is False and blocks[0].items == ()
    assert blocks[0].tech_hours is None                       # the hours stay on the first Qty header (price sheet)
    qty = blocks[1:]
    assert [b.header.split(" ")[0] for b in qty] == ["Qty"] * len(families)
    assert all(b.numbered for b in qty) and all(b.joined for b in qty)
    assert qty[0].tech_hours == 6.0 and qty[-1].items
    # traced like a header, before the Qty headers
    headers = [l.text for l in draft.lines if l.system_key == key and l.kind == "header"]
    assert headers[0] == group and len(headers) == len(families) + 1


@pytest.mark.parametrize("key, shape, families", [
    ("ZHX", "Heat Exchanger System", [("Heat Exchanger", ("ZHX-2",)), ("Hot Water Pump", ("ZHWP-1", "ZHWP-2"))]),
    ("ZCN", "Condenser Water System", [("Cooling Tower", ("ZCT-1",)), ("Condenser Water Pump", ("ZCWP-1", "ZCWP-2"))]),
    ("ZAH", "Air Handling Unit", [("Air Handling Unit", ("ZAHU-1",)), ("Return Fan", ("ZRF-1",))]),
])
def test_an_entry_without_a_group_line_prints_none(job, house, key, shape, families):
    draft = _draft(job, house, _system(key, shape, shape, families))
    assert all(b.numbered for b in draft.scope.blocks)
    assert all(b.header.startswith("Qty (") for b in draft.scope.blocks)
    assert shape not in [l.text for l in draft.lines if l.kind == "header"]


def test_a_single_family_block_gets_no_group_line(job, house):
    draft = _draft(job, house, _system("ZCW", "Chilled Water System", "Chilled Water System",
                                       [("Chiller", ("ZCH-7",))]))
    assert all(b.numbered for b in draft.scope.blocks)
    assert len(draft.scope.blocks) == 1 and draft.scope.blocks[0].header.startswith("Qty (1) ")


def test_the_group_line_words_come_from_the_asset_and_name_real_shapes(house):
    groups = order.group_lines()
    assert groups == {"Chilled Water System": "Chilled Water System", "Hot Water System": "Hot Water System",
                      "VRF System": "VRF System", "Kitchen Exhaust Ventilation": "Kitchen Exhaust Ventilation"}
    assert set(groups) <= set(house.shapes)
    assert order.group_line("Heat Exchanger System") == "" and order.group_line(None) == ""


def test_the_price_sheet_prints_the_group_line_unnumbered_and_unpriced(job, house):
    from estimator.price_sheet.rows import sheet_plan
    draft = _draft(job, house, _system("ZHW", "Hot Water System", "Hot Water System",
                                       [("Boiler", ("ZB-3",)), ("Hot Water Pump", ("ZHP-4", "ZHP-5"))]))
    rows = [r for r in sheet_plan(draft.scope)[0].rows if r.text]
    at = next(i for i, r in enumerate(rows) if r.text == "Hot Water System")
    assert rows[at].hours is None and rows[at].quantity is None
    assert rows[at + 1].text.endswith("Qty (1) Boiler ZB-3") and rows[at + 1].hours == 6.0
    assert rows[at + 2].text.split("  ", 1)[1].startswith("Qty (2) Hot Water Pump")


# ---- the evening correction's other two points: '*' lines per item (rulings 14, 16), no tech-hour line -----------

def test_one_short_star_line_per_item_several_per_block_and_no_tech_hour_line(job, house, tmp_path):
    from dataclasses import replace
    from docx import Document
    from bms.model import Clarification
    from estimator.__main__ import PROPOSAL_TEMPLATE
    from estimator.proposal.write import write_proposal
    system = replace(_system("ZHW", "Hot Water System", "Hot Water System",
                             [("Boiler", ("ZB-3",)), ("Hot Water Pump", ("ZHP-4", "ZHP-5"))]),
                     clarifications=(Clarification(kind="confirm", text="invented boiler count on sheet M-9"),
                                     Clarification(kind="confirm", text="invented pump tags on sheet M-8"),
                                     Clarification(kind="note", text="Invented pumps are by the boiler maker")))
    draft = _draft(job, house, system)
    notes = [l.text for l in draft.lines if l.system_key == "ZHW" and l.kind == "note"]
    assert len(notes) == 3 and all(n.startswith("*") and "\n" not in n for n in notes)   # one line per item
    block = next(b for b in draft.scope.blocks if b.items)
    assert [i.kind for i in block.items].count("note") == 3
    text = "\n".join(p.text for p in Document(write_proposal(draft.scope, PROPOSAL_TEMPLATE,
                                                             tmp_path / "p.docx")).paragraphs)
    assert "Hot Water System" in text
    assert "tech hour" not in text.lower() and "6.0" not in text                 # ruling 16: hours on the sheet only
    assert "Demolition Scope" not in text                                         # the lead adds that block (T-007)
