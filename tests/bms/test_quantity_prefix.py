"""G-002: a device line prints its own count in front, "Qty (N) <device line>".

The sample job and the sample Playbook only. Counts and names here are invented.
"""
from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

from bms.model import Count, Device, Family, Job, Ref
from bms.playbook import load_playbook
from bms.render import render
from bms.validate import validate
from estimator.price_sheet.price_book import BookItem, PriceBook, price_key
from estimator.price_sheet.write import PRICE_PENDING, QUANTITY_PENDING, write_price_sheet
from estimator.scope import Block, Item, Run

FIXTURES = Path(__file__).resolve().parent / "fixtures"
REF = Ref(source_id="S1", quote="invented count", where="M-1, Drawings.pdf page 1")


def _device(slot, name, qty, id="UH.d1", verb="Furnish and Install", furnished_by="AME") -> Device:
    return Device(id=id, slot=slot, name=name, verb=verb, furnished_by=furnished_by, installed_by="AME",
                  wired_by="AME", quantity=Count(value=qty), refs=(REF,))


def _job() -> Job:
    job = Job.from_dict(json.loads((FIXTURES / "sample_job.json").read_text(encoding="utf-8")))
    same = _device("space_temperature", "Space Temperature Sensor", 5, id="UH.d1")
    other = _device("space_temperature", "Space Temperature Sensor", 33, id="UH.d2")
    unknown = _device("control_valve", "Control Valve", None, id="UH.d3")
    system = replace(
        job.system("CUH"),
        name="Unit Heaters",
        families=(Family("Cabinet Unit Heater", ("UH-1", "UH-2", "UH-3", "UH-4", "UH-5"), Count(value=5)),),
        devices=(same, other, unknown),
        clarifications=(),
        sequence_refs=(),
    )
    return replace(job, systems=(system,))


def _lines(draft):
    return [line.text for line in draft.lines if line.system_key == "CUH"]


def test_a_different_count_prints_in_front_and_the_block_count_does_not():
    playbook = load_playbook(FIXTURES / "sample_playbook")
    lines = _lines(render(_job(), playbook))
    assert lines[0].startswith("Qty (5) ") and "UH-1" in lines[0]
    assert "Furnish and Install Space Temperature Sensor" in lines
    assert "Qty (33) Furnish and Install Space Temperature Sensor" in lines
    assert "Qty (##) " in "\n".join(lines) and "Control Valve" in "\n".join(lines)
    assert not any("(Qty." in line or line.endswith("Total)") for line in lines)


def test_an_unknown_device_count_is_flagged():
    playbook = load_playbook(FIXTURES / "sample_playbook")
    job = _job()
    hits = [f for f in validate(job, playbook) if f.device_id == "UH.d3"]
    assert hits and all(f.severity == "check" for f in hits)
    assert any("unknown" in f.what for f in hits)


def test_the_price_sheet_prices_both_forms_and_holds_an_unknown_line(workup_template, tmp_path):
    playbook = load_playbook(FIXTURES / "sample_playbook")
    book = PriceBook([BookItem("Space Temperature Sensor", 142.0, "Std Parts Pricing", 2, False)], {})
    draft = render(_job(), playbook)
    old = Block(header="Qty (1) Spare Sensor", items=(
        Item(kind="line", quantity=6, runs=(Run(text="Furnish and Install Space Temperature Sensor (Qty. 6 Total)"),)),
        Item(kind="line", quantity=33, runs=(Run(text="Qty (33) Install and Wire Space Temperature Sensor"),)),
    ))
    scope = draft.scope.model_copy(update={"blocks": draft.scope.blocks + (old,)})
    path = write_price_sheet(scope, workup_template, book, tmp_path / "s.xlsx")
    from openpyxl import load_workbook
    ws = load_workbook(path)["Price Sheet 1"]
    rows = {str(ws.cell(r, 4).value): r for r in range(6, 80) if ws.cell(r, 4).value}
    same = next(r for text, r in rows.items() if text.endswith("Furnish and Install Space Temperature Sensor")
                and "Qty (33)" not in text and "(Qty." not in text)
    prefixed = next(r for text, r in rows.items() if "Qty (33) Furnish and Install Space Temperature Sensor" in text)
    unknown = next(r for text, r in rows.items() if "Qty (##)" in text and "Control Valve" in text)
    assert ws.cell(same, 3).value == 5 and ws.cell(same, 5).value == 142.0
    assert ws.cell(prefixed, 3).value == 33 and ws.cell(prefixed, 5).value == 142.0
    assert ws.cell(unknown, 8).value == QUANTITY_PENDING
    assert ws.cell(unknown, 3).value is None and ws.cell(unknown, 5).value is None
    color = ws.cell(unknown, 4).font.color
    assert color is not None and color.type != "theme" and str(color.rgb).endswith("FF0000")
    old_row = next(r for text, r in rows.items() if "(Qty. 6 Total)" in text)
    assert ws.cell(old_row, 5).value == 142.0
    install = next(r for text, r in rows.items() if "Qty (33) Install and Wire" in text)
    assert ws.cell(install, 8).value == PRICE_PENDING and ws.cell(install, 5).value is None


def test_price_key_takes_the_prefix_off_and_keeps_the_old_end_form():
    assert price_key("Qty (33) Furnish and Install Space Temperature Sensor") == "space temperature sensor"
    assert price_key("Qty (##) Furnish and Install Space Temperature Sensor") == "space temperature sensor"
    assert price_key("Furnish and Install Space Temperature Sensor (Qty. 6 Total)") == "space temperature sensor"
    assert price_key("Qty (33) Install and Wire Space Temperature Sensor") == (
        "install and wire space temperature sensor")
    assert price_key("o  Qty (33) Furnish and Install Space Temperature Sensor") == "space temperature sensor"
