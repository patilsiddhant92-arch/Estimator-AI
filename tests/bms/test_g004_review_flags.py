"""G-004: one review question per family, a red "Qty (##)" on an unknown device count, and a note for a system
with no devices. Invented data only."""
from __future__ import annotations

import json
import re
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from bms import review
from bms.model import (Count, Decision, DemolitionLine, Device, Family, Finding, FrontEnd, Job, OtherCount, Ref,
                       ScopeException, System)
from bms.playbook import load_playbook
from bms.render import _holds_phrase, render
from bms.review import _names_device
from estimator.price_sheet.price_book import BookItem, PriceBook
from estimator.price_sheet.write import QUANTITY_PENDING, write_price_sheet

REF = Ref("S1", "AHU row", "M-501, Invented.pdf page 6")
OTHER = Ref("S2", "AHU rows 3", "M-601, Invented.pdf page 7")
FLAG_OWNERSHIP = "Who furnishes, installs or wires this line is not settled."


def _settled(n: int = 1) -> Count:
    return Count(n, (REF,))


def _disputed() -> Count:
    return Count(5, (REF,), (OtherCount(3, OTHER),))


def _device(i: int, flag: str = "ownership", count: Count | None = None, system: str = "AHU") -> Device:
    return Device(id=f"{system}.d{i}", slot="other_device", name=f"Point {i}", verb="Furnish and Install",
                  furnished_by="AME", installed_by="AME", wired_by="AME", quantity=count or _settled(),
                  refs=(REF,), flag=flag)


def _job(*devices: Device, families=None, name: str = "Air Handling Unit", key: str = "AHU",
         front_end: FrontEnd | None = None) -> Job:
    families = families if families is not None else (Family("Air Handling Unit", ("AHU-1",), _settled()),)
    system = System(key=key, name=name, shape="Air Handling Unit", control="field", families=families,
                    devices=devices)
    return Job(name="Invented", systems=(system,), front_end=front_end)


def _page(job: Job, findings=(), tmp_path=None) -> str:
    return review.review_page(job, findings, tmp_path).read_text(encoding="utf-8")


def _headings(page: str) -> list[str]:
    return re.findall(r"<h3>(.*?)</h3>", page)


def _groups(page: str) -> list[str]:
    return re.findall(r'<section class="group(?: extra)?">(.*?)</section>', page, re.S)


def _qdata(fragment: str) -> list[dict]:
    return [json.loads(raw) for raw in re.findall(r'class="qdata">(.*?)</script>', fragment, re.S)]


def test_one_family_with_seven_flagged_devices_is_one_group(tmp_path):
    devices = tuple(_device(i) for i in range(1, 8))
    page = _page(_job(*devices), tmp_path=tmp_path)
    assert page.count('<section class="card"') == 1
    assert _headings(page) == ["Air Handling Unit: 7 points to check"]
    assert "q1" not in _headings(page)[0]
    (group,) = _groups(page)
    data = _qdata(group)
    assert [row["device_id"] for row in data] == [f"AHU.d{i}" for i in range(1, 8)]
    assert group.count('data-act="accept"') == 7
    assert group.count('data-act="change"') == 7
    assert "1 of 1 question(s) shown" in page


def test_accept_on_one_sub_point_changes_only_that_device(tmp_path):
    job = _job(_device(1), _device(2), _device(3))
    page = _page(job, tmp_path=tmp_path)
    assert _headings(page) == ["Air Handling Unit: 3 points to check"]
    assert len(_qdata(page)) == 3
    after = review.apply_decisions(job, [Decision("AHU", "accept", device_id="AHU.d2")])
    flags = {device.id: device.flag for device in after.system("AHU").devices}
    assert flags == {"AHU.d1": "ownership", "AHU.d2": "none", "AHU.d3": "ownership"}


def test_a_finding_that_names_a_flagged_device_keeps_one_sub_point(tmp_path):
    named = Finding("interlock", "AHU", "AHU.d1", "check", "The interlock for AHU.d2 is open.", "Cite AHU.d2.")
    kept = Finding("quote", "AHU", "AHU.d23", "check", "The note names AHU.d23 only.", "Keep the note.")
    job = _job(_device(1, flag="none"), _device(2), _device(20), _device(23, flag="none"))
    page = _page(job, (named, kept), tmp_path)
    data = _qdata(page)
    assert [row["device_id"] for row in data] == ["AHU.d20", "AHU.d1", "AHU.d23"]
    a = next(block for block in _groups(page) for row in _qdata(block) if row["device_id"] == "AHU.d1")
    assert FLAG_OWNERSHIP in a
    assert "AHU.d2" not in [row["device_id"] for row in data]
    own = Finding("interlock", "AHU", "AHU.d2", "must_fix", "AHU.d2 has no citation.", "Cite it.")
    other = Finding("interlock", "AHU", "AHU.d1", "check", "See also AHU.d2.", "Check AHU.d2.")
    both = _page(_job(_device(1, flag="none"), _device(2)), (own, other), tmp_path)
    assert {row["device_id"] for row in _qdata(both)} == {"AHU.d1", "AHU.d2"}


def test_two_families_put_other_questions_in_the_system_group(tmp_path):
    families = (Family("Air Handling Unit", ("AHU-1",), _disputed()),
                Family("Energy Recovery Unit", ("ERV-1",), _disputed()))
    finding = Finding("shape", "AHU", None, "check", "The block has no tech hours.", "Set the hours.")
    page = _page(_job(_device(1), families=families, name="Air Side"), (finding,), tmp_path)
    groups = {heading: block for heading, block in zip(_headings(page), _groups(page))}
    assert "Air Handling Unit: 1 point to check" in groups
    assert "Energy Recovery Unit: 1 point to check" in groups
    assert "Air Side: 2 points to check" in groups
    assert "How many Air Handling Unit units?" in groups["Air Handling Unit: 1 point to check"]
    assert "How many Energy Recovery Unit units?" in groups["Energy Recovery Unit: 1 point to check"]
    system = groups["Air Side: 2 points to check"]
    assert "Point 1 (Furnish and Install)" in system
    assert "The block has no tech hours." in system
    assert "How many" not in system


def test_a_group_counts_as_one_question_toward_the_cap(tmp_path):
    page = _page(_job(*(_device(i) for i in range(1, 17))), tmp_path=tmp_path)
    assert "1 of 1 question(s) shown" in page
    assert page.count('class="q extra"') == 0
    assert len(_qdata(page)) == 16


def test_show_all_hides_whole_groups(tmp_path):
    systems = []
    for n in range(16):
        key = f"S{n}"
        device = _device(1, system=key)
        systems.append(System(key=key, name=f"System {n}", shape="Air Handling Unit", control="field",
                              families=(Family(f"Family {n}", (f"T-{n}",), _settled()),), devices=(device,)))
    job = Job(name="Invented", systems=tuple(systems))
    page = _page(job, tmp_path=tmp_path)
    assert "15 of 16 question(s) shown" in page
    assert page.count('<section class="group extra">') == 1
    assert page.count('<section class="group">') == review.MAX_QUESTIONS
    assert 'class="q extra"' not in page
    hidden = re.search(r'<section class="group extra">(.*?)</section>', page, re.S).group(1)
    assert 'class="q"' in hidden and 'class="q extra"' not in hidden
    assert ".group.extra{display:none}" in page
    assert "body.all .group.extra{display:block}" in page
    assert '1 more question(s) under "Show all questions".' in page


def test_a_system_with_no_family_uses_the_system_name(tmp_path):
    device = _device(1, system="CB")
    system = System(key="CB", name="Coil Bank", shape="Coil Bank", control="none", families=(), devices=(device,))
    page = _page(Job(name="Invented", systems=(system,)), tmp_path=tmp_path)
    assert _headings(page) == ["Coil Bank: 1 point to check"]


def test_front_end_questions_use_one_group_named_front_end(tmp_path):
    device = _device(1, system="FE")
    job = Job(name="Invented", front_end=FrontEnd(mode="new-system", devices=(device,)))
    page = _page(job, tmp_path=tmp_path)
    assert _headings(page) == ["Front end: 1 point to check"]
    assert _qdata(page)[0]["device_id"] == "FE.d1"


# ---- Task 2: an unknown device count prints "Qty (##)" in red ---------------------------------------------------


def _scope_device(slot: str, name: str, qty: int | None, id: str, mark: str = "") -> Device:
    return Device(id=id, slot=slot, name=name, verb="Furnish and Install", furnished_by="AME", installed_by="AME",
                  wired_by="AME", quantity=Count(qty, (REF,)), refs=(REF,), mark=mark)


def _unit_job() -> Job:
    known = _scope_device("space_temperature", "Space Temperature Sensor", 4, "UH.d1")
    future = _scope_device("discharge_air", "Discharge Air Sensor", 2, "UH.d2", mark="future")
    unknown = _scope_device("control_valve", "Control Valve", None, "UH.d3")
    system = System(key="UH", name="Unit Heaters", shape="Unit Heater", control="field",
                    families=(Family("Cabinet Unit Heater", ("UH-1",), Count(None, (REF,))),),
                    devices=(known, future, unknown))
    base = Job.from_dict(json.loads((Path(__file__).resolve().parent / "fixtures" / "sample_job.json")
                                    .read_text(encoding="utf-8")))
    return replace(base, systems=(system,), demolition=(), front_end=FrontEnd(mode="new-system"))


def _playbook():
    return load_playbook(Path(__file__).resolve().parent / "fixtures" / "sample_playbook")


def _line(block, start: str):
    return next(item for item in block.items if item.kind == "line" and item.text.startswith(start))


def test_an_unknown_device_count_prints_red_qty_hashes():
    block = render(_unit_job(), _playbook()).scope.blocks[0]
    unknown = _line(block, "Qty (##) ")
    known = _line(block, "Qty (4) ")
    future = _line(block, "Qty (2) ")
    assert unknown.runs[0].red and unknown.quantity is None and unknown.depth == 1
    assert unknown.text.startswith("Qty (##) ") and "Control Valve" in unknown.text
    assert not known.runs[0].red and known.text.startswith("Qty (4) ")
    assert future.runs[0].red and future.text.startswith("Qty (2) ")
    assert block.header.startswith("Qty (##) ")
    assert all(item.text != block.header for item in block.items)


def test_a_front_end_unknown_count_prints_red_qty_hashes():
    unknown = _scope_device("other_device", "Operator Workstation", None, "FE.d1")
    known = _scope_device("other_device", "Server Graphics", 1, "FE.d2")
    job = replace(_unit_job(), front_end=FrontEnd(mode="new-system", devices=(unknown, known)))
    items = render(job, _playbook()).scope.front_end.items
    red = next(item for item in items if item.kind == "line" and item.text.startswith("Qty (##) "))
    black = next(item for item in items if item.kind == "line" and item.text.endswith("Server Graphics"))
    assert red.runs[0].red and red.quantity is None and "Operator Workstation" in red.text
    assert not black.runs[0].red and not black.text.startswith("Qty (##)")


def test_the_price_sheet_row_for_an_unknown_count_is_red(workup_template, tmp_path):
    book = PriceBook([BookItem("Control Valve", 10.0, "Std Parts Pricing", 1, False)], {})
    scope = render(_unit_job(), _playbook()).scope
    path = write_price_sheet(scope, workup_template, book, tmp_path / "sheet.xlsx")
    ws = load_workbook(path)["Price Sheet 1"]
    rows = {str(ws.cell(r, 4).value): r for r in range(6, 80) if ws.cell(r, 4).value}

    def rgb(row: int) -> str:
        color = ws.cell(row, 4).font.color
        if color is None or color.type == "theme" or color.rgb is None:
            return ""
        return str(color.rgb)

    unknown = next(r for text, r in rows.items() if "Qty (##)" in text and "Control Valve" in text)
    known = next(r for text, r in rows.items() if "Qty (4)" in text and "Space Temperature" in text)
    assert ws.cell(unknown, 8).value == QUANTITY_PENDING
    assert ws.cell(unknown, 3).value is None
    assert rgb(unknown).endswith("FF0000")
    assert not rgb(known).endswith("FF0000") and not rgb(known).endswith("C00000")


# ---- Task 3: a system with no devices prints one note ----------------------------------------------------------

_COIL_REASON = "The coil valves are counted with ERV-1."
_COIL_STAR = "*Coil Bank: The coil valves are counted with ERV-1."


def _carrier(key: str, name: str, tag: str) -> System:
    device = _scope_device("space_temperature", "Space Temperature Sensor", 1, f"{key}.d1")
    return System(key=key, name=name, shape="Unit Heater", control="field",
                  families=(Family(name, (tag,), Count(1, (REF,))),), devices=(device,))


def _silent(key: str, name: str, reason: str | None, alternate: str = "base") -> System:
    # A known family count keeps this system in the equipment block order, so a test can put it before the carrier.
    return System(key=key, name=name, shape="", control="none",
                  families=(Family(name, (), Count(1, (REF,))),), devices=(), alternate=alternate,
                  no_scope_reason=reason)


def _job_of(*systems: System, demolition: tuple[DemolitionLine, ...] = ()) -> Job:
    base = Job.from_dict(json.loads((Path(__file__).resolve().parent / "fixtures" / "sample_job.json")
                                    .read_text(encoding="utf-8")))
    return replace(base, systems=systems, demolition=demolition, exceptions=(), set_aside=(),
                   front_end=FrontEnd(mode="new-system"))


def _demo() -> tuple[DemolitionLine, ...]:
    return (DemolitionLine("UH", "demo", 1, (REF,)),)


def test_a_reason_that_names_one_block_prints_a_star_line_on_that_block():
    erv = _carrier("ERV", "Energy Recovery Ventilator", "ERV-1")
    coil = _silent("CB", "Coil Bank", _COIL_REASON)
    for systems in ((coil, erv), (erv, coil)):
        draft = render(_job_of(*systems), _playbook())
        assert len(draft.scope.blocks) == 1
        assert draft.scope.blocks[0].items[-1].kind == "note"
        assert draft.scope.blocks[0].items[-1].text == _COIL_STAR
        assert [line.text for line in draft.lines if line.system_key == "ERV" and line.kind == "note"
                and line.text == _COIL_STAR] == [_COIL_STAR]
        coil_lines = [line.text for line in draft.lines if line.system_key == "CB"]
        assert len(coil_lines) == 1 and "left out, no lines written" in coil_lines[0]
        assert "The coil valves are counted with ERV-1" in coil_lines[0]


def test_a_reason_that_holds_the_system_name_does_not_repeat_it():
    reason = "Coil Bank valves are counted with ERV-1."
    draft = render(_job_of(_silent("CB", "Coil Bank", reason),
                           _carrier("ERV", "Energy Recovery Ventilator", "ERV-1")), _playbook())
    assert draft.scope.blocks[0].items[-1].text == f"*{reason}"


def test_a_reason_that_names_no_block_prints_no_bms_scope_before_demolition():
    reason = "The documents give this system to others."
    draft = render(_job_of(_carrier("ERV", "Energy Recovery Ventilator", "ERV-1"),
                           _silent("CB", "Coil Bank", reason), demolition=_demo()), _playbook())
    blocks = draft.scope.blocks
    note = f"No BMS scope: {reason}"
    assert blocks[-1].header == "Demolition Scope"
    assert blocks[-2].items[-1].kind == "note" and blocks[-2].items[-1].text == note
    assert all("No BMS scope" not in item.text for item in blocks[-1].items)
    texts = [line.text for line in draft.lines]
    assert texts.index(note) < texts.index("Demolition Scope")
    assert any(line.system_key == "CB" and "left out, no lines written" in line.text for line in draft.lines)


def test_a_reason_that_names_two_blocks_prints_no_bms_scope_on_the_last_block_only():
    reason = "The valves are counted with ERV-1 and AHU-1."
    draft = render(_job_of(_carrier("ERV", "Energy Recovery Ventilator", "ERV-1"),
                           _carrier("AHU", "Air Handling Unit", "AHU-1"),
                           _silent("CB", "Coil Bank", reason), demolition=_demo()), _playbook())
    note = f"No BMS scope: {reason}"
    equipment = [block for block in draft.scope.blocks if block.header != "Demolition Scope"]
    assert equipment[-1].items[-1].text == note
    assert all(item.text != note for item in equipment[0].items)
    assert not any(item.text.startswith("*Coil Bank") for block in draft.scope.blocks for item in block.items)


def test_a_key_inside_a_longer_tag_does_not_name_that_block():
    reason = "The coil valves are counted with ERV-1."
    draft = render(_job_of(_silent("CB", "Coil Bank", reason),
                           _carrier("ERV", "Energy Recovery Ventilator", "ERV-2")), _playbook())
    assert draft.scope.blocks[0].items[-1].text == f"No BMS scope: {reason}"
    assert not any(item.text.startswith("*Coil Bank") for item in draft.scope.blocks[0].items)


def test_a_named_exception_prints_no_scope_note():
    line = "EXCEPTION: Qty (1) CB-1 - not in this contract"
    job = replace(_job_of(_carrier("ERV", "Energy Recovery Ventilator", "ERV-1"),
                          _silent("CB", "Coil Bank", _COIL_REASON)),
                  exceptions=(ScopeException("CB", line, (REF,)),))
    draft = render(job, _playbook())
    assert [item.text for item in draft.lines if item.system_key == "CB"] == [line]
    assert draft.scope.named_exceptions == (line,)
    assert not any("No BMS scope" in item.text or item.text.startswith("*Coil Bank")
                   for block in draft.scope.blocks for item in block.items)


def test_a_system_with_no_reason_prints_no_scope_note():
    draft = render(_job_of(_carrier("ERV", "Energy Recovery Ventilator", "ERV-1"),
                           _silent("CB", "Coil Bank", None)), _playbook())
    assert not any("No BMS scope" in item.text or "Coil Bank" in item.text
                   for block in draft.scope.blocks for item in block.items)
    assert [line.text for line in draft.lines if line.system_key == "CB"] == [
        "Coil Bank: left out, no lines written"]


def test_an_alternate_with_no_devices_prints_no_scope_note_on_the_base_bid():
    draft = render(_job_of(_carrier("ERV", "Energy Recovery Ventilator", "ERV-1"),
                           _silent("CB", "Coil Bank", _COIL_REASON, alternate="add")), _playbook())
    assert not any("No BMS scope" in item.text or item.text.startswith("*Coil Bank")
                   for block in draft.scope.blocks for item in block.items)
    assert any(line.system_key == "CB" and "left out, no lines written" in line.text for line in draft.lines)


def test_a_long_reason_prints_as_short_points():
    reason = ("The coil valves are counted with ERV-1 because the documents show one set of valves for both units. "
              "The second sentence keeps this invented reason past the short point limit of one hundred and fifty "
              "characters.")
    assert len(reason) > 150
    draft = render(_job_of(_silent("CB", "Coil Bank", reason),
                           _carrier("ERV", "Energy Recovery Ventilator", "ERV-1")), _playbook())
    note = draft.scope.blocks[0].items[-1].text
    assert note.startswith("*Coil Bank: ") and "\n- " in note and "ERV-1" in note
    assert any(line.system_key == "ERV" and line.kind == "note" and line.text == note for line in draft.lines)


def test_a_device_id_matches_only_as_a_whole_id():
    must = ("The interlock for AHU.d2 is open.", "Cite AHU.d2.", "AHU.d2 stands alone")
    must_not = ("The note names AHU.d23 only.", "See AHU.d2.1 for the point.", "AHU.d20 is a different point.")
    assert all(_names_device(text, "AHU.d2") for text in must)
    assert not any(_names_device(text, "AHU.d2") for text in must_not)


def test_a_block_name_matches_only_as_a_whole_phrase():
    must = ("The coil valves are counted with ERV-1.", "Counted with erv-1 today.", "ERV-1 leads the list.")
    must_not = ("The coil valves are counted with ERV-1.", "Tag ERV-10 is another unit.", "The ERVX unit is different.")
    assert all(_holds_phrase(text, "ERV-1") for text in must)
    assert not any(_holds_phrase(text, "ERV") for text in must_not)
    assert not _holds_phrase("The Coil Banks stay in the room.", "Coil Bank")


def test_a_composed_line_over_160_characters_is_short_points():
    reason = ("The coil valves are counted with ERV-1 because the documents show one set word word word word word. "
              "The second sentence finishes this invented reason.")
    assert len(reason) == 150 and len("*Coil Bank: " + reason) > 160
    draft = render(_job_of(_silent("CB", "Coil Bank", reason),
                           _carrier("ERV", "Energy Recovery Ventilator", "ERV-1")), _playbook())
    note = draft.scope.blocks[0].items[-1].text
    assert "\n- " in note and note.startswith("*Coil Bank: ")
    loose = ("The documents give this system to others and the note stays one invented sentence word word word. "
             "The second sentence finishes this invented reason.")
    assert len(loose) <= 150 and len("No BMS scope: " + loose) > 160
    other = render(_job_of(_carrier("ERV", "Energy Recovery Ventilator", "ERV-1"),
                           _silent("CB", "Coil Bank", loose)), _playbook())
    assert "\n- " in other.scope.blocks[0].items[-1].text


def test_a_job_with_no_block_keeps_todays_left_out_note():
    draft = render(_job_of(_silent("CB", "Coil Bank", "The documents give this system to others.")), _playbook())
    assert draft.scope.blocks == ()
    assert [line.text for line in draft.lines if line.system_key == "CB"] == [
        "Coil Bank: left out, no lines written (The documents give this system to others)"]
    assert not any("No BMS scope" in line.text for line in draft.lines)
