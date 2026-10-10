"""The Evidence Package (bms.evidence): its rows, on the invented sample job, and the workbook's layout."""
from __future__ import annotations

import dataclasses as dc
import json
from pathlib import Path

import openpyxl
import pytest

from bms.evidence import (CITATION_HEADINGS, COUNT_HEADINGS, FINDING_HEADINGS, NEW_TABS, NOTE_HEADINGS,
                          SCOPE_HEADINGS, UNKNOWN_QTY, citation_rows, count_rows, finding_rows, note_rows,
                          quantity_note, scope_rows, split_where, system_summaries)
from bms.model import Count, Finding, Job, OtherCount, Ref

FIX = Path(__file__).resolve().parent / "fixtures"
AT = {h: i for i, h in enumerate(SCOPE_HEADINGS)}


def _job() -> Job:
    return Job.from_dict(json.loads((FIX / "sample_job.json").read_text(encoding="utf-8")))


def _devices(job: Job):
    out = [d for s in job.systems for d in s.devices]
    return out + (list(job.front_end.devices) if job.front_end else [])


# ---- Scope lines: device lines only (EXCEL-01) ---------------------------------------------------------------

def test_scope_lines_has_one_row_per_device_line_and_nothing_else():
    job = _job()
    rows = scope_rows(job)
    assert [r[AT["Device id"]] for r in rows] == [d.id for d in _devices(job)]
    assert all(len(r) == len(SCOPE_HEADINGS) for r in rows)
    assert [r[AT["#"]] for r in rows] == list(range(1, len(rows) + 1))


def test_a_system_with_no_lines_is_on_system_notes_with_its_reason_not_on_scope_lines():
    """RND-8: a system the writer gave no lines prints nothing in the proposal; the package keeps it, with the
    writer's reason, on System notes, so a filter on Scope lines never mixes it with device rows."""
    job = _job()
    job = job.with_system(dc.replace(job.system("DH"), devices=(),
                                     no_scope_reason="the invented heaters are controlled by their own panel"))
    assert not any("(DH)" in r[AT["System"]] for r in scope_rows(job))
    left = [r for r in note_rows(job) if r[1] == "Left out"]
    assert len(left) == 1 and len(left[0]) == len(NOTE_HEADINGS)
    assert "(DH)" in left[0][0] and "the invented heaters are controlled by their own panel" in left[0][2]


def test_a_system_with_lines_has_no_left_out_note():
    assert [r for r in note_rows(_job()) if r[1] == "Left out"] == []


def test_quantity_is_a_number_and_an_unknown_one_is_none_with_its_note():
    job = _job()
    rows = {r[AT["Device id"]]: r for r in scope_rows(job)}
    for d in _devices(job):
        q = rows[d.id][AT["Qty"]]
        assert q == d.quantity.highest() and (q is None or isinstance(q, int))
        if q is None:
            assert rows[d.id][AT["Qty note"]] == UNKNOWN_QTY


def test_a_count_known_only_from_another_count_prints_that_count_and_its_source():
    """RND-7 (estimator ruling D9): the package prints the number the proposal prints (`Count.highest()`), and the
    note keeps the other count with where it is; Citations keeps its source id and quote."""
    job = _job()
    dh = job.system("DH")
    count = Count(value=None, others=(OtherCount(2, Ref("S2", "invented quote", "Q-9.00, Invented.pdf page 1")),))
    job = job.with_system(dc.replace(dh, devices=(dc.replace(dh.devices[0], quantity=count),) + dh.devices[1:]))
    row = next(r for r in scope_rows(job) if r[AT["Device id"]] == dh.devices[0].id)
    assert row[AT["Qty"]] == 2
    assert "Q-9.00" in row[AT["Qty note"]] and "highest" in row[AT["Qty note"]]
    cites = [r for r in citation_rows(job) if r[1] == dh.devices[0].id and r[4] == "S2"]
    assert cites and cites[0][-1] == "invented quote" and "other count (2)" in cites[0][3]


def test_the_same_citation_twice_on_a_line_is_one_source_and_one_citation_row():
    """EVID-01: a ref that is both a device ref and a quantity source is listed once, its uses joined."""
    job = _job()
    rtu = job.system("RTU")
    d = rtu.devices[3]                               # its ref and its quantity source are the same S5 quote
    assert d.refs[0] == d.quantity.sources[0]
    row = next(r for r in scope_rows(job) if r[AT["Device id"]] == d.id)
    assert row[AT["Source ids"]] == "S5" and "more" not in row[AT["Cited at"]]
    cites = [r for r in citation_rows(job) if r[1] == d.id]
    assert len(cites) == 1 and cites[0][3] == "calls for the line, backs the quantity"


def test_open_findings_show_on_their_line():
    job = _job()
    f = Finding("Shipped-loose devices", "RTU", "RTU.d4", "must_fix", "invented what", "invented fix")
    row = next(r for r in scope_rows(job, findings=[f]) if r[AT["Device id"]] == "RTU.d4")
    assert row[AT["Open findings"]].startswith("[MUST FIX] Shipped-loose devices: invented what")
    other = next(r for r in scope_rows(job, findings=[f]) if r[AT["Device id"]] == "RTU.d1")
    assert other[AT["Open findings"]] is None


def test_printed_line_falls_back_to_the_device_name():
    job = _job()
    rows = scope_rows(job, printed={"RTU.d1": "Furnish and Install invented printed line (Qty. 2 Total)"})
    assert rows[0][AT["Printed line"]].startswith("Furnish and Install invented")
    assert rows[1][AT["Printed line"]] == job.system("RTU").devices[1].name


# ---- the other tabs ---------------------------------------------------------------------------------------------

def test_findings_put_must_fix_first():
    fs = [Finding("A", "RTU", None, "check", "c", "f"), Finding("B", "", None, "must_fix", "m", "f")]
    rows = finding_rows(fs)
    assert [r[1] for r in rows] == ["Must fix", "Check"] and rows[0][3] == "Whole job"
    assert all(len(r) == len(FINDING_HEADINGS) for r in rows)


def test_counts_name_the_spread_between_sources():
    rows = count_rows(_job())
    assert all(len(r) == len(COUNT_HEADINGS) for r in rows)
    rtu = next(r for r in rows if r[1] == "Rooftop Unit")
    assert rtu[3] == 2 and rtu[6] == 1 and "higher prints" in rtu[7]


def test_internal_takeoff_questions_are_labelled_internal():
    job = dc.replace(_job(), questions=("invented question?",))
    q = [r for r in note_rows(job) if r[2] == "invented question?"]
    assert q and q[0][1] == "Takeoff question (internal)"


def test_split_where():
    assert split_where("M-201, Drawings.pdf page 3") == ("M-201", "Drawings.pdf page 3")
    assert split_where("M-201") == ("M-201", "")
    assert split_where("") == ("", "")


def test_quantity_note():
    assert quantity_note(Count(value=None)) == UNKNOWN_QTY
    assert quantity_note(Count(value=3)) is None
    note = quantity_note(Count(value=3, others=(OtherCount(1, Ref("S1", "q", "M-1, a.pdf page 1")),)))
    assert note == "3 printed; also 1 per M-1, a.pdf page 1"


def test_system_summaries_carry_tech_hours_and_status():
    """LABOR-01: each system's tech hours are on the Summary."""
    job = _job()
    fs = [Finding("X", "RTU", None, "must_fix", "w", "f")]
    rows = {s.key: s for s in system_summaries(job, fs)}
    assert rows["RTU"].status == "Must fix" and rows["RTU"].must_fix == 1
    assert rows["EF"].tech_hours == job.system("EF").tech_hours
    assert rows[""].tech_hours == job.front_end.tech_hours


# ---- the workbook -----------------------------------------------------------------------------------------------

class _Catalogue:
    """The least the old workbook writes from: no documents, no sheets."""
    class job:
        documents, left_out, nested_jobs, other_files = (), (), (), ()
        name = "invented"
    flags, pages, sheets, sections, tags = (), (), (), (), ()


def test_workbook_layout(tmp_path, monkeypatch):
    import bms.evidence as ev

    def old(path, catalogue, reading, ledger=None):
        book = openpyxl.Workbook()
        book.active.title = "Bid check"
        book.active.append(["Bid check — invented"])
        book.create_sheet("Pages").append(["Document", "Page"])
        book.save(str(path))
        return path

    monkeypatch.setattr(ev, "write_old_evidence", old)
    job = dc.replace(_job(), questions=("=SUM(1,2)",))
    fs = [Finding("Shipped-loose devices", "RTU", "RTU.d4", "must_fix", "w", "f")]
    path = ev.write_evidence(tmp_path / "E.xlsx", _Catalogue, None, None, job, fs, {})
    book = openpyxl.load_workbook(path)
    assert book.sheetnames[:len(NEW_TABS)] == list(NEW_TABS)
    assert book.sheetnames[len(NEW_TABS):] == ["Bid check", "Pages"]
    scope = book["Scope lines"]
    assert [c.value for c in scope[1]] == SCOPE_HEADINGS
    assert scope.auto_filter.ref.startswith("A1:") and scope.freeze_panes == "A2"
    kinds = {type(scope.cell(r, AT["Qty"] + 1).value) for r in range(2, scope.max_row + 1)}
    assert kinds <= {int, str}                         # a number, or "##" when not known
    finding = book["Findings"]
    assert finding.cell(2, FINDING_HEADINGS.index("Device id") + 1).hyperlink is not None
    notes = book["System notes"]
    formula = [c for row in notes.iter_rows() for c in row if c.value == "=SUM(1,2)"]
    assert formula and formula[0].data_type == "s"     # job text never becomes a formula
    assert book.active.title == "Summary"


def test_an_exception_on_a_planned_system_shows_that_system_label():
    """RUN-10: the System cell reads like every other row's, not the bare key."""
    from bms.model import ScopeException
    job = dc.replace(_job(), exceptions=(ScopeException("EF", "invented exception", ()),
                                         ScopeException("XX", "not planned", ())))
    labels = [r[0] for r in note_rows(job) if r[1] == "Exception"]
    assert labels == ["Exhaust Fans (EF)", "XX"]
