"""The estimate command: from a job folder to the proposal, the price sheet and the Evidence Package."""
import json
import re

import openpyxl
import pytest
from docx import Document

from estimator.__main__ import _parser, ai_mode, main
from estimator.evidence.workbook import ESTIMATE_TABS, LEDGER_TAB, PLAN_TAB, QC_TAB, READ_TAB, TABS
from estimator.proposal.sections import scope_texts
from pdfs import FILLER, Page, make_pdf
from estimator.writing.whole_job import Contradiction, QcCheck, WrittenLine as GemLine, WrittenScope as GemScope
from test_assemble import Dossier, FakeJudge, JobPlan, PlannedSystem
from test_schedules import grid

W, H = 3024, 2160
VAV_ROWS = [["VAV BOX SCHEDULE", "", "", "", ""],
            ["TAG", "SERVES", "MAX CFM", "HW GPM", "MANUFACTURER"],
            ["VAV-1", "OFFICE 101", "600", "1.5", "TITUS"],
            ["VAV-2", "OFFICE 102", "700", "1.5", "TITUS"],
            ["VAV-3", "OFFICE 103", "500", "1.0", "TITUS"]]
INTERNAL = ("confidence", "ASSUMPTION", "CONFIRMED", "rule:", "AI:", "tech hour", "PENDING", "ADDED", "EXTRA",
            # what the estimator's GEM writes for the estimator: flag glyphs, source ids, its word for the controls
            # contractor, its reasoning words and labels written twice
            "🔴", "🟡", "🟠", "🟢", "◊", "(S170)", "S480", "[Q7]", "I14", "CC",
            "factory-controls rule", "field DDC default", "Mandatory minimum", "term inferred", "per GEM", "Quick Scan",
            "C1", "C2", "C3", "C4", "C5", "C6", "QC-", "*Assumed: ASSUMED:", "Confirm ownership: CONFIRM QUANTITY")


def _leaks(word: str, text: str) -> bool:
    """The word is in the text: a word starting with a letter, as a word; anything else, anywhere."""
    return bool(re.search(rf"\b{re.escape(word)}" if word[0].isalpha() else re.escape(word), text))


def _schedule_sheet() -> Page:
    lines, texts = grid(300, 300, [90, 120, 70, 70, 100], VAV_ROWS)
    return Page(W, H, ((200, 200, FILLER, 12), *texts,
                       (2640, 1860, "HVAC SCHEDULES", 20), (2800, 2000, "M-601", 40)), lines=lines)


PLAN = Page(W, H, ((200, 300, FILLER, 12), (900, 900, "VAV-1", 12), (900, 1000, "VAV-2", 12),
                   (900, 1100, "VAV-3", 12), (2640, 1860, "SECOND FLOOR PLAN", 20), (2800, 2000, "M-201", 40)))
SEQUENCE = Page(texts=((72, 72, "SECTION 23 09 93 - SEQUENCE OF OPERATIONS\n\n"
                                "A. VAV boxes: each VAV box shall maintain the space temperature and CO2 "
                                "setpoints.\n\n23 09 93 - 1", 10),))


@pytest.fixture
def job(tmp_path):
    root = tmp_path / "26-9999, 1 Test Street"
    base = root / "2. Takeoffs" / "1. Base Takeoff"
    make_pdf(base / "Mech.pdf", _schedule_sheet(), PLAN)
    make_pdf(base / "230993.pdf", SEQUENCE)
    return root


@pytest.fixture
def estimated(job, tmp_path, capsys):
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--no-ai", "--workers", "1", "--method", "rules"]) == 0
    return out, capsys.readouterr().out


def _proposal_text(path) -> str:
    document = Document(str(path))
    texts = [p.text for p in document.paragraphs]
    texts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(texts)


def test_the_estimate_writes_the_proposal_the_price_sheet_and_the_evidence_package(estimated):
    out, printed = estimated
    proposal, sheet = out / "26-9999 Proposal.docx", out / "26-9999 Price Sheet.xlsx"
    package = out / "26-9999 Evidence Package.xlsx"
    assert proposal.exists() and sheet.exists() and package.exists()

    text = _proposal_text(proposal)
    assert "Qty (3)" in text
    assert "Discharge Air Temperature Sensor" in text
    assert "Space Temperature / CO2 Sensor" in text

    names = openpyxl.load_workbook(package).sheetnames
    assert names[:3] == [TABS[0], *ESTIMATE_TABS]
    assert ESTIMATE_TABS == ("Review list", "Scope lines")

    assert "Systems written: 1 in the proposal" in printed
    assert "PENDING" in printed and "review list" in printed.lower()
    assert "AI: not used" in printed


def test_the_api_without_a_key_is_not_used_and_its_lines_say_why(job, tmp_path, capsys, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "api", "--workers", "1", "--method", "rules"]) == 0
    assert "AI: not used — no AI key set" in capsys.readouterr().out
    assert (out / "26-9999 Proposal.docx").exists()


@pytest.mark.parametrize("flags, key, mode", [
    ([], None, "chat"),                        # no key: the chat answers
    ([], "not-a-real-key", "api"),             # a key: the API answers
    (["--ai", "chat"], "not-a-real-key", "chat"),
    (["--ai", "api"], None, "api"),
    (["--ai", "off"], "not-a-real-key", "off"),
    (["--no-ai"], "not-a-real-key", "off"),     # the old flag still works
])
def test_the_ai_flag_picks_who_answers(flags, key, mode, monkeypatch):
    if key:
        monkeypatch.setenv("ANTHROPIC_API_KEY", key)
    else:
        monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    args = _parser().parse_args(["estimate", "JOB", "--out", "OUT", *flags])
    assert ai_mode(args) == mode


def test_no_ai_and_ai_cannot_both_be_given(capsys):
    with pytest.raises(SystemExit):
        _parser().parse_args(["estimate", "JOB", "--out", "OUT", "--no-ai", "--ai", "chat"])


def test_in_chat_mode_the_questions_wait_in_the_requests_folder(job, tmp_path, capsys, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setattr("anthropic.Anthropic", lambda *a, **k: pytest.fail("the chat mode never calls the API"))
    out, cache = tmp_path / "out", tmp_path / "cache"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "chat", "--cache", str(cache),
                 "--workers", "1", "--method", "rules"]) == 0
    printed = capsys.readouterr().out
    requests = cache / "judge-chat" / "requests"
    written = list(requests.glob("*.json"))
    assert written
    assert (f"AI: answered in the chat for 0 systems; {len(written)} waiting — requests in {requests}. "
            "Answer them, then run estimate again.") in printed
    assert (out / "26-9999 Proposal.docx").exists()


def test_nothing_internal_reaches_the_proposal(estimated):
    out, _ = estimated
    scope = "\n".join(scope_texts(out / "26-9999 Proposal.docx"))
    for word in INTERNAL:
        assert not _leaks(word, scope), word
    whole = _proposal_text(out / "26-9999 Proposal.docx")
    for word in ("confidence", "ASSUMPTION", "CONFIRMED", "rule:", "AI:", "tech hour", "PENDING"):
        assert word not in whole, word


def test_the_evidence_package_lists_the_review_and_every_scope_line(estimated):
    out, _ = estimated
    workbook = openpyxl.load_workbook(out / "26-9999 Evidence Package.xlsx")
    review = [[c.value for c in row] for row in workbook["Review list"].iter_rows()]
    assert review[0][:6] == ["Rank", "System", "Category", "What to decide", "Evidence", "Suggested action"]
    assert review[1][0] == 1
    assert any(row[2] == "ASSUMPTION" for row in review[1:])
    assert any(row[1] == "Cover page" for row in review[1:])            # the MEP engineer could not be read

    lines = [[c.value for c in row] for row in workbook["Scope lines"].iter_rows()]
    headings = lines[0]
    assert headings == ["System", "Line", "Quantity", "Tech hours", "Category", "Confidence", "Source", "Evidence",
                        "In the proposal?", "Reference line", "Flag", "Estimator decision", "Change to / why"]
    valve = next(row for row in lines if row[1] == "Furnish and Wire Hot Water Control Valve and Actuator")
    assert (valve[2], valve[4], valve[8]) == (3, "CONFIRMED", "yes")
    assert "HW GPM" in valve[7]
    header = next(row for row in lines if row[1] and str(row[1]).startswith("Qty (3)"))
    assert header[3] == 12


# --- the whole-job method (the default) ----------------------------------------------------------------------

DOSSIER_TEXT = "\n".join([
    "1. Job", "[S1] 26-9999, 1 Test Street", "",
    "2. Sequence of operation",
    "[S2] A. VAV boxes: each VAV box shall maintain the space temperature and CO2 setpoints.", "",
    "4. Schedules", "[S3] M-601 VAV BOX SCHEDULE",
    "[S4] VAV-1 | OFFICE 101 | 600 | 1.5 | TITUS", "[S5] VAV-2 | OFFICE 102 | 700 | 1.5 | TITUS",
    "[S6] VAV-3 | OFFICE 103 | 500 | 1.0 | TITUS", "",
    "[S7] Exhaust fans EF-1 and EF-2 run on the line-voltage thermostat.", ""])


def _fake_dossier(catalogue, reading, cache_dir):
    return Dossier("26-9999", DOSSIER_TEXT, (), {"S2": "23 09 93 SEQUENCE OF OPERATIONS, 230993.pdf page 1",
                                                  "S3": "M-601 VAV BOX SCHEDULE, Mech.pdf page 1", "S7": "M-201"},
                   tuple(reading.totals), {}, {"VAV": {"schedule_rows": 3, "quantity_column": None,
                                                       "plan_placements": 3, "riser": None, "distinct_tags": 3}}, {})


def _fake_ledger(save_to=None):
    from estimator.ledger.facts import SETTLED, Ledger
    ledger = Ledger({"S1": "the job"})
    ledger.add("family", "VAV", "Variable Air Volume Box", "[S1]", "VAV", "seed", SETTLED)
    ledger.understanding = {"schedules": "A school with VAV boxes."}
    ledger.complete = True
    ledger.job_name = "26-9999"
    if save_to is not None:
        ledger.save(save_to)
    return ledger


JOB_PLAN = JobPlan((
    PlannedSystem("VAV", "Variable Air Volume Box VAV-1 thru 3", "Variable Air Volume Box", ("VAV-1", "VAV-2", "VAV-3"),
                  3, "S3", "M-601 VAV BOX SCHEDULE", ("S2",), ("S3", "S4", "S5", "S6"), why="three scheduled"),
    PlannedSystem("EF", "Exhaust Fan EF-1, 2", "Exhaust Fan", ("EF-1", "EF-2"), 2, "S7", "Exhaust fans EF-1 and EF-2",
                  ("S7",), exception="EXCEPTION TO ANY WORK WITH Qty (2) EF-1, 2 - LINE-VOLTAGE THERMOSTAT"),
))
# the GEM's reasoning, written into every field of its answer: none of it may reach the proposal
GEM_WORDS = ("per the factory-controls rule, field DDC default, term inferred, Mandatory minimum, per GEM "
             "(Quick Scan; QC-04; S480; I14) (S170)")
SCOPES = {"VAV": GemScope((
    GemLine("Furnish DDC Controller to the Factory for Installation", 1, None, "each VAV box shall maintain", "S2",
            "a VAV box controller", 0.9),
    GemLine("Furnish and Install Space Temperature / CO2 Sensor", 1, None, "space temperature and CO2 setpoints", "S2",
            "the sequence names CO2", 0.9),
    GemLine("Furnish and Wire Hot Water Control Valve and Actuator (confidence 0.9)", 1, None, "1.5", "S4", "", 0.9),
    GemLine("Furnish and Install VAV Transformer", 1, None, "each VAV box", "S2", "", 0.9),
    GemLine("Furnish and Install Discharge Air Temperature Sensor (S170)", 1, None, "each VAV box shall maintain", "S2",
            "", 0.9, "Mandatory minimum point — per GEM", "orange",
            f"ASSUMED: CC furnishes it {GEM_WORDS} 🟠"),
    GemLine("Furnish Fan Enable [Q7]", 1, None, "each VAV box shall maintain", "S2", "", 0.9, "", "red",
            f"CONFIRM QUANTITY: {GEM_WORDS} 🔴"),
    GemLine("Wire Fire Alarm Relay (Relay Furnished by FA)", 1, None, "each VAV box shall maintain", "S2", "", 0.9, "",
            "yellow", f"CONFIRM QUANTITY: {GEM_WORDS} 🟡"),
    GemLine("Controls Contractor requires mfr. rep to assist in integration, cost to be incurred by mech contractor", 2,
            None, "each VAV box shall maintain", "S2", "", 0.9),
), ("*Assumed: ASSUMED: the sensor by CC (S170)", "*Confirm ownership: CONFIRM QUANTITY: ◊ per GEM",
    "**Wiring of the VAV power is by the Electrical Contractor (S214) 🔴"), None, (), "m",
    contradictions=tuple(Contradiction(f"C{n}", f"Discharge air sensor furnished by CC — {GEM_WORDS}", ("q",),
                                       "orange") for n in range(1, 7)),
    misc_responsibilities=(f"Program the VAV schedule {GEM_WORDS} 🟢",
                           f"Provide additional controllers {GEM_WORDS}. 🟠",
                           "Wiring of the VAV power is by the Electrical Contractor (S214) 🔴"),
    exclusions=(f"Airflow stations by CC — {GEM_WORDS} 🟡",
                "Supply airflow measuring stations (furnished with the unit, S170)."),
    qc=(QcCheck("QC-01", f"CC term: BMS contractor (C1) {GEM_WORDS}", False, "per GEM"),))}


@pytest.fixture
def whole_job(monkeypatch):
    monkeypatch.setattr("estimator.__main__.build_dossier", _fake_dossier)
    monkeypatch.setattr("estimator.__main__.read_ledger",
                        lambda judge, dossier, catalogue, save_to=None, pictures=True: _fake_ledger(save_to))
    monkeypatch.setattr("estimator.__main__.ledger_plan", lambda ledger: JOB_PLAN)
    judges = []

    def make(cache, mode, job=None):
        judges.append(FakeJudge(JOB_PLAN, SCOPES, mode=mode))
        return judges[-1]

    monkeypatch.setattr("estimator.__main__.Judge", make)
    return judges


def test_no_pictures_is_an_estimate_option():
    assert _parser().parse_args(["estimate", "JOB", "--out", "OUT", "--no-pictures"]).no_pictures is True
    assert _parser().parse_args(["estimate", "JOB", "--out", "OUT"]).no_pictures is False


def test_whole_job_is_the_default_method():
    assert _parser().parse_args(["estimate", "JOB", "--out", "OUT"]).method == "whole-job"
    assert _parser().parse_args(["estimate", "JOB", "--out", "OUT", "--method", "rules"]).method == "rules"


def test_with_the_ai_off_the_whole_job_method_writes_the_dossier_and_the_evidence_package_only(
        job, tmp_path, capsys, whole_job):
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "off", "--workers", "1"]) == 0
    printed = capsys.readouterr().out
    assert (out / "26-9999 Dossier.txt").read_text(encoding="utf-8") == DOSSIER_TEXT
    assert (out / "26-9999 Evidence Package.xlsx").exists()
    assert not (out / "26-9999 Proposal.docx").exists() and not (out / "26-9999 Price Sheet.xlsx").exists()
    assert "the AI writes the scope; run with --ai chat or api" in printed
    assert whole_job == []                                     # no judge was made


def test_the_whole_job_method_writes_the_planned_blocks_and_nothing_internal(job, tmp_path, capsys, whole_job):
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "api", "--workers", "1"]) == 0
    printed = capsys.readouterr().out
    assert whole_job[0].asked == ["VAV"]                      # the exception needs no scope
    proposal = out / "26-9999 Proposal.docx"
    assert proposal.exists() and (out / "26-9999 Price Sheet.xlsx").exists()
    text = _proposal_text(proposal)
    assert "Qty (3) Variable Air Volume Box VAV-1 thru 3" in text
    assert "Space Temperature / CO2 Sensor" in text
    assert "EXCEPTION TO ANY WORK WITH Qty (2) EF-1, 2 - LINE-VOLTAGE THERMOSTAT" in text
    assert "Hot Water Control Valve" not in text               # it named its confidence: dropped
    scope = "\n".join(scope_texts(proposal))
    for word in INTERNAL + ("[S", "whole job", "DISCREPANCY"):
        assert not _leaks(word, scope), word
    for word in INTERNAL:                                      # the front end, the exceptions and the notes too
        assert not _leaks(word, text), word
    for word in ("confidence", "ASSUMPTION", "CONFIRMED", "rule:", "AI:", "tech hour", "PENDING", "AI (whole job)"):
        assert word not in text, word

    workbook = openpyxl.load_workbook(out / "26-9999 Evidence Package.xlsx")
    assert workbook.sheetnames[:7] == [TABS[0], READ_TAB, *ESTIMATE_TABS, QC_TAB, PLAN_TAB, LEDGER_TAB]
    plan = [[c.value for c in row] for row in workbook[PLAN_TAB].iter_rows()]
    vav = next(row for row in plan if row[0] == "VAV")
    assert (vav[4], vav[5], vav[11]) == (3, "CONFIRMED", 3)
    lines = [[c.value for c in row] for row in workbook["Scope lines"].iter_rows()]
    sensor = next(row for row in lines if row[1] == "Furnish and Install Space Temperature / CO2 Sensor")
    assert sensor[6] == "AI (whole job) [S2]" and sensor[8] == "yes"
    review = [[c.value for c in row] for row in workbook["Review list"].iter_rows()]
    assert any("dropped" in f"{row[3]} {row[4]}" for row in review)

    assert "Systems planned: 2; 1 written in the proposal, 1 put in the exceptions." in printed
    assert "AI: 1 call to fake-model, 0 answered from the cache at no cost." in printed
    assert f"Wrote {out / '26-9999 Dossier.txt'}" in printed


def test_in_chat_mode_the_whole_job_method_waits_for_the_answers(job, tmp_path, capsys, monkeypatch, whole_job):
    monkeypatch.setattr("estimator.__main__.Judge", lambda cache, mode, job=None: whole_job.append(_Waiting(mode)) or
                        whole_job[-1])
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "chat", "--workers", "1"]) == 0
    printed = capsys.readouterr().out
    assert "1 request waiting — requests in requests. Answer them, then run estimate again." in printed
    assert (out / "26-9999 Evidence Package.xlsx").exists() and (out / "26-9999 Dossier.txt").exists()
    assert not (out / "26-9999 Proposal.docx").exists()


class _Waiting(FakeJudge):
    """The chat has not answered the system's scope yet."""

    def __init__(self, mode):
        super().__init__(None, {}, mode=mode)

    def write_system(self, system, *args, **kwargs):
        self.waiting = 1
        return None


def test_the_whole_job_method_writes_the_ledger_and_its_tabs(job, tmp_path, capsys, whole_job):
    main(["estimate", str(job), "--out", str(tmp_path / "out"), "--ai", "api"])
    out = tmp_path / "out"
    ledger = next(out.glob("* Ledger.json"))
    assert json.loads(ledger.read_text(encoding="utf-8"))["complete"] is True
    from openpyxl import load_workbook
    book = load_workbook(next(out.glob("* Evidence Package.xlsx")))
    assert "How the job was read" in book.sheetnames and "Job ledger" in book.sheetnames
    printed = capsys.readouterr().out
    assert "Reading the job: 1 fact settled in 1 pass; 0 contradictions." in printed


def test_when_a_pass_is_waiting_the_run_stops_after_the_evidence_package(job, tmp_path, capsys, monkeypatch, whole_job):
    def waiting(judge, dossier, catalogue, save_to=None, pictures=True):
        ledger = _fake_ledger(save_to)
        ledger.complete, ledger.stopped_at = False, "plans"
        judge.last_error = "judgement pending — waiting for an answer in the chat"
        return ledger

    monkeypatch.setattr("estimator.__main__.read_ledger", waiting)
    main(["estimate", str(job), "--out", str(tmp_path / "out"), "--ai", "api"])
    out = tmp_path / "out"
    assert next(out.glob("* Evidence Package.xlsx"), None) is not None
    assert next(out.glob("*.docx"), None) is None
    assert "Reading the job: stopped at the Plans pass — judgement pending" in capsys.readouterr().out


def test_the_run_asks_the_qc_pass_and_writes_its_tab(job, tmp_path, capsys, whole_job):
    main(["estimate", str(job), "--out", str(tmp_path / "out"), "--ai", "api"])
    from openpyxl import load_workbook
    book = load_workbook(next((tmp_path / "out").glob("* Evidence Package.xlsx")))
    assert "Proposal QC" in book.sheetnames
    rows = list(book["Proposal QC"].iter_rows(values_only=True))
    assert rows[0][:7] == ("System", "Line", "Kind", "Severity", "What is wrong", "Change to", "Rests on")
    assert any(r[2] == "ownership" for r in rows[1:])
    read = list(book["How the job was read"].iter_rows(values_only=True))
    assert read[1][0] == "Job brief"
    assert "QC: 1 findings (1 must-fix)" in capsys.readouterr().out.replace("1 finding (", "1 findings (")


def test_no_qc_skips_the_pass(job, tmp_path, capsys, whole_job):
    main(["estimate", str(job), "--out", str(tmp_path / "out"), "--ai", "api", "--no-qc"])
    assert "QC: skipped" in capsys.readouterr().out
    assert whole_job[-1].qc_calls == 0


def test_a_failed_qc_call_in_api_mode_does_not_hold_the_proposal(job, tmp_path, capsys, monkeypatch, whole_job):
    def no_answer(self, *args, **kwargs):
        self.qc_calls += 1
        self.last_error = "the API timed out"
        return None

    monkeypatch.setattr(FakeJudge, "qc_job", no_answer)
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "api"]) == 0
    assert "QC: not answered — the API timed out." in capsys.readouterr().out
    assert (out / "26-9999 Proposal.docx").exists() and (out / "26-9999 Price Sheet.xlsx").exists()
    book = openpyxl.load_workbook(out / "26-9999 Evidence Package.xlsx")
    rows = list(book["Proposal QC"].iter_rows(values_only=True))
    assert rows[1][1] == "The QC pass gave no answer: the API timed out" and len(rows) == 2
    read = [r[0] for r in book["How the job was read"].iter_rows(values_only=True)]
    assert "The QC pass gave no answer: the API timed out" in read


def test_in_chat_mode_a_waiting_qc_request_holds_the_proposal(job, tmp_path, capsys, monkeypatch, whole_job):
    def waiting(self, *args, **kwargs):
        self.qc_calls += 1
        self.waiting += 1
        self.last_error = "judgement pending — waiting for an answer in the chat"
        return None

    monkeypatch.setattr(FakeJudge, "qc_job", waiting)
    out = tmp_path / "out"
    assert main(["estimate", str(job), "--out", str(out), "--ai", "chat"]) == 0
    printed = capsys.readouterr().out
    assert "QC: waiting for the chat's answer." in printed
    assert "1 request waiting" in printed
    assert (out / "26-9999 Evidence Package.xlsx").exists()
    assert not (out / "26-9999 Proposal.docx").exists()
    assert whole_job[-1].qc_calls == 1


def test_each_write_request_gets_the_class_of_the_unit_a_duct_heater_serves(job, tmp_path, monkeypatch, whole_job):
    seen = []
    original = FakeJudge.write_system

    def recording(self, system, *args, **kwargs):
        seen.append(kwargs.get("heater_class"))
        return original(self, system, *args, **kwargs)

    monkeypatch.setattr(FakeJudge, "write_system", recording)
    monkeypatch.setattr("estimator.__main__.served_class", lambda ledger, plan, system: "packaged")
    main(["estimate", str(job), "--out", str(tmp_path / "out"), "--ai", "api"])
    assert seen and set(seen) == {"packaged"}


def test_each_write_request_gets_the_heater_block_that_writes_its_duct_heaters(job, tmp_path, monkeypatch, whole_job):
    seen = []
    original = FakeJudge.write_system

    def recording(self, system, *args, **kwargs):
        seen.append(kwargs.get("heater_home"))
        return original(self, system, *args, **kwargs)

    monkeypatch.setattr(FakeJudge, "write_system", recording)
    monkeypatch.setattr("estimator.__main__.heater_home", lambda ledger, plan, system: ("EDH-1", "Electric Duct Heater EDH-1"))
    main(["estimate", str(job), "--out", str(tmp_path / "out"), "--ai", "api"])
    assert seen and set(seen) == {("EDH-1", "Electric Duct Heater EDH-1")}
