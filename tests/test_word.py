import pytest
from docx.oxml.ns import qn
from docx.table import _Cell

from estimator.proposal.word import WordLists, field_result_text, merged_runs, visible_text
from estimator.text import squash


def _paragraph(document, text):
    return next(p for p in document.paragraphs if squash(p.text) == text)


def _date_paragraph(document):
    table = document.tables[0]
    return next(p for tc in table._tbl.iter(qn("w:tc")) for p in _Cell(tc, table).paragraphs
                if p.text.strip().startswith("Date:"))


def test_scope_line_list_steps_in_by_360(template_document):
    lists = WordLists(template_document)
    assert [lists.level_indent(12, level) for level in range(4, 9)] == [1800, 2160, 2520, 2880, 3240]


def test_front_end_lists(template_document):
    lists = WordLists(template_document)
    assert [lists.level_indent(11, 1), lists.level_indent(11, 2)] == [1440, 2160]
    assert [lists.level_indent(10, 1), lists.level_indent(10, 2)] == [1440, 2160]
    assert lists.has_level(12, 8) and not lists.has_level(12, 9)


def test_pr2_numbers_its_paragraphs(template_document):
    assert WordLists(template_document).style_numbering("PR2") == (4, 5)


def test_a_block_header(sent_document):
    lists = WordLists(sent_document)
    header = _paragraph(sent_document, "Qty (4) Packaged Air Handling Unit AHU-1, 2, 3, 4")
    assert header.style.name == "PR2"
    assert WordLists.written_numbering(header) is None
    assert lists.numbering(header) == (4, 5)
    assert lists.left_indent(header) == 1566


def test_the_scope_heading_switches_numbering_off(sent_document):
    heading = next(p for p in sent_document.paragraphs if squash(p.text).startswith("SCOPE OF WORK"))
    assert WordLists.written_numbering(heading) == (0, 0)
    assert WordLists(sent_document).numbering(heading) is None


@pytest.mark.parametrize("text, indent", [
    ("Map Available Points Via Mfr. Provided BMS Interface or Connect to Nearby DDC Controller", 1800),
    ("Main Electrical Meter (Furnished and Installed by Others) – (Qty. 1 Total)", 2160),
    ("Booster Pump Set General Fault Alarm", 2520),
    ("DB-ELEV – (Qty. 1 Total)", 2880),
    ("Intrusion Alarm", 2160),
])
def test_nested_lines_resolve_to_the_indent_word_draws(sent_document, text, indent):
    lists = WordLists(sent_document)
    paragraph = _paragraph(sent_document, text)
    assert lists.numbering(paragraph) is not None
    assert lists.left_indent(paragraph) == indent


def test_emphasis_inside_a_line_is_kept(sent_document):
    paragraph = _paragraph(sent_document, "Furnish and Install (Rack Mounted) BACnet Server")
    stretches = merged_runs(paragraph)
    assert [text.strip() for text, bold in stretches if bold] == ["(Rack Mounted)"]
    assert squash("".join(text for text, _ in stretches)) == "Furnish and Install (Rack Mounted) BACnet Server"


def test_visible_text_matches_the_runs_for_an_ordinary_line(sent_document):
    paragraph = _paragraph(sent_document, "Furnish and Install (Rack Mounted) BACnet Server")
    assert squash(visible_text(paragraph)) == squash(paragraph.text)


def test_the_date_field_gives_its_stored_value(template_document, sent_document):
    assert field_result_text(_date_paragraph(template_document), "DATE") == "3/6/2026"
    assert field_result_text(_date_paragraph(sent_document), "DATE") == "6/8/2026"


def test_a_paragraph_without_the_field(sent_document):
    assert field_result_text(_paragraph(sent_document, "Intrusion Alarm"), "DATE") is None
