import shutil

import docx
import pytest
from docx.oxml.ns import qn
from docx.table import _Cell

from estimator.proposal import layout
from estimator.proposal.write import TemplateError, _delete_from_until, check_template, write_proposal
from estimator.scope import FrontEnd
from estimator.text import squash


def _cover_runs(path):
    document = docx.Document(str(path))
    table = document.tables[0]
    return [p for tc in table._tbl.iter(qn("w:tc")) for p in _Cell(tc, table).paragraphs if p.text.strip()]


def _texts(path):
    return [squash(p.text) for p in docx.Document(str(path)).paragraphs if p.text.strip()]


def test_the_template_is_recognised(template_document):
    check_template(template_document)


def test_a_changed_template_is_refused(proposal_template, tmp_path):
    document = docx.Document(str(proposal_template))
    heading = next(p for p in document.paragraphs if squash(p.text) == layout.ADD_ALTERNATES)
    heading.runs[0].text = "OPTIONAL EXTRAS:"
    with pytest.raises(TemplateError, match="ADD ALTERNATES"):
        check_template(document)


def test_the_cover_page_is_filled(bare_job, proposal_template, tmp_path):
    out = write_proposal(bare_job, proposal_template, tmp_path / "p.docx")
    assert [squash(p.text) for p in _cover_runs(out)] == [
        "To: All Bidders",
        "Date: 6/8/2026",
        "Quote #: 26-9999",
        "Reference: Test Building, 1 Main St",
        "MEP Engineer: Test MEP",
        "Based on Plans and Specs:",
        "M-001 thru M-100 Dated: 01/02/2026",
        "Specs 23 09 00",
    ]


def test_the_date_is_written_down_not_left_as_a_field(bare_job, proposal_template, tmp_path):
    out = write_proposal(bare_job, proposal_template, tmp_path / "p.docx")
    date = next(p for p in _cover_runs(out) if p.text.startswith("Date:"))
    assert date._p.find(f".//{qn('w:fldChar')}") is None


def test_labels_keep_their_red_and_values_stay_dark(bare_job, proposal_template, tmp_path):
    out = write_proposal(bare_job, proposal_template, tmp_path / "p.docx")
    quote = next(p for p in _cover_runs(out) if p.text.startswith("Quote"))
    colours = {r.text: r.font.color.rgb for r in quote.runs}
    assert str(colours["Quote #: "]) == "EC1E27"
    assert str(colours["26-9999"]) == "231F20"


def test_a_new_system_keeps_only_its_own_opening(bare_job, proposal_template, tmp_path):
    texts = _texts(write_proposal(bare_job, proposal_template, tmp_path / "p.docx"))
    assert sum(t.startswith(layout.OPENING) for t in texts) == 1
    assert sum(t.startswith(layout.SCOPE_HEADING) for t in texts) == 1
    assert layout.NEW_SYSTEM_HEADING in texts
    assert layout.OR_LINE not in texts
    assert not [t for t in texts if t.startswith(layout.EXISTING_SYSTEM_HEADING) or "IT Infrastructure" in t]


def test_extending_an_existing_system(bare_job, proposal_template, tmp_path):
    job = bare_job.model_copy(update={"front_end": FrontEnd(mode="extend-existing", existing_system_name="Honeywell WEBs")})
    out = write_proposal(job, proposal_template, tmp_path / "p.docx")
    texts = _texts(out)
    openings = [t for t in texts if t.startswith(layout.OPENING)]
    assert len(openings) == 1 and "extension of the existing" in openings[0]
    assert "Seamlessly Tie in to the Existing Honeywell WEBs BMS" in texts
    assert layout.NEW_SYSTEM_HEADING not in texts and layout.OR_LINE not in texts
    heading = next(p for p in docx.Document(str(out)).paragraphs if squash(p.text).startswith(layout.EXISTING_SYSTEM_HEADING))
    assert not [r for r in heading.runs if r.font.color.rgb is not None and str(r.font.color.rgb) == "FF0000"]


def test_scope_heading_suffix(bare_job, proposal_template, tmp_path):
    assert "SCOPE OF WORK – PREVAILING WAGE" in _texts(write_proposal(bare_job, proposal_template, tmp_path / "p.docx"))


def test_the_rest_of_the_template_is_untouched(bare_job, proposal_template, tmp_path):
    out_texts = _texts(write_proposal(bare_job, proposal_template, tmp_path / "p.docx"))
    template_texts = _texts(proposal_template)
    start = next(i for i, t in enumerate(template_texts) if t == layout.ADD_ALTERNATES)
    assert out_texts[out_texts.index(layout.ADD_ALTERNATES):] == template_texts[start:]


def test_delete_from_until_refuses_a_stop_before_first(proposal_template):
    document = docx.Document(str(proposal_template))
    openings = [p for p in document.paragraphs if squash(p.text).startswith(layout.OPENING)]
    with pytest.raises(TemplateError, match="runs to the end"):
        _delete_from_until(openings[1], openings[0])
