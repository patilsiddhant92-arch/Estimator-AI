from dataclasses import replace
from pathlib import Path

import openpyxl

from estimator.catalogue.documents import Document, JobFolder
from estimator.catalogue.job import Catalogue
from estimator.catalogue.overview import build_overview
from estimator.catalogue.qualify import BidCheck, Mention
from estimator.catalogue.register import PageRecord
from estimator.catalogue.sheet_types import Sheet
from estimator.catalogue.spec_sections import SpecSection
from estimator.catalogue.tags import TagEntry
from estimator.evidence.workbook import TABS, write_evidence


def _catalogue() -> Catalogue:
    documents = (Document(Path("m.pdf"), "2. Takeoffs/1. Base Takeoff/Mech.pdf", "drawings", "base", 2),
                 Document(Path("s.pdf"), "2. Takeoffs/1. Base Takeoff/15970.pdf", "specs", "base", 1),
                 Document(Path("c.pdf"), "2. Takeoffs/1. Base Takeoff/copy.pdf", "drawings", "base", 2,
                          duplicate_of="2. Takeoffs/1. Base Takeoff/Mech.pdf"))
    job = JobFolder("26-9999", Path("26-9999"), documents, (), ("3. Proposal",),
                    ("2. Takeoffs/1. Base Takeoff/26-9998",), ("2. Takeoffs/1. Base Takeoff/notes.xlsx",))
    pages = (PageRecord(documents[0].relative, 1, "text", "AHU-1", 90), PageRecord(documents[0].relative, 2, "unreadable", "", 0),
             PageRecord(documents[1].relative, 1, "ocr", "SECTION 15970", 0))
    sheets = (Sheet(documents[0].relative, 1, "M-601", "HVAC SCHEDULES", "Mechanical", ("Schedules",)),
              Sheet(documents[0].relative, 2, None, None, "Unknown", ()))
    mention = Mention(documents[1].relative, 1, "The BMS shall be built on the Niagara 4 Framework.", ("Niagara N4",), (), (), False)
    bid = BidCheck("Qualifies", "The spec names Niagara N4 (15970.pdf, page 1).", ("Siemens",), (mention,), ("15970 TEMPERATURE CONTROL",))
    sections = (SpecSection("15970", "TEMPERATURE CONTROL", documents[1].relative, 1, 1),)
    tags = (TagEntry("AHU-1", "Air Handling Unit", ("M-601",), (), 1),)
    return Catalogue(job, pages, sheets, sections, bid, tags, ("Mech.pdf: 1 page could not be read (page 2)",))


def _values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_the_tabs_in_order(tmp_path):
    out = write_evidence(tmp_path / "26-9999 Evidence Package.xlsx", _catalogue())
    assert openpyxl.load_workbook(out).sheetnames == list(TABS)


def test_the_bid_check_comes_first_and_quotes_the_spec(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue()))["Bid check"]
    rows = _values(ws)
    assert rows[0][0] == "Bid check — 26-9999"
    assert rows[1][0] == ("This is the tool's first reading of the spec. Every passage it found is "
                         "quoted below; the US team reads them and makes the call.")
    assert sum("makes the" in str(r[0]) for r in rows) == 1
    assert ["Finding", "Qualifies"] == rows[3][:2]
    assert ["Other controls manufacturers named (for the US team)", "Siemens"] == rows[5][:2]
    assert ["15970.pdf, page 1", "Niagara N4 (a line AME bids)",
            "The BMS shall be built on the Niagara 4 Framework."] in rows


def test_flags_documents_pages_sheets_sections_and_tags(tmp_path):
    wb = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue()))
    assert _values(wb["Flags"])[1] == ["Mech.pdf: 1 page could not be read (page 2)"]
    assert _values(wb["Documents"])[3] == ["2. Takeoffs/1. Base Takeoff/copy.pdf", "Drawings", "base", 2,
                                           "2. Takeoffs/1. Base Takeoff/Mech.pdf"]
    assert _values(wb["Documents"])[4] == ["3. Proposal", "Left out — AME's own work, not a bid document", None, None, None]
    assert _values(wb["Documents"])[5] == ["2. Takeoffs/1. Base Takeoff/26-9998",
                                           "Left out — another job, catalogue it on its own", None, None, None]
    assert _values(wb["Documents"])[6] == ["2. Takeoffs/1. Base Takeoff/notes.xlsx", "Not read — not a PDF (.xlsx)",
                                           None, None, None]
    assert _values(wb["Pages"])[2] == ["2. Takeoffs/1. Base Takeoff/Mech.pdf", 2, "could not be read", None]
    assert _values(wb["Pages"])[1][3] == "M-601"
    assert _values(wb["Sheets"])[1] == ["M-601", "HVAC SCHEDULES", "Mechanical", "Schedules",
                                        "2. Takeoffs/1. Base Takeoff/Mech.pdf", 1]
    assert _values(wb["Spec sections"])[1] == ["15970", "TEMPERATURE CONTROL", "yes", "2. Takeoffs/1. Base Takeoff/15970.pdf", 1, 1]
    assert _values(wb["Equipment tags"])[1] == ["AHU-1", "Air Handling Unit", "M-601", None, 1]


def test_nothing_to_flag_says_so(tmp_path):
    from dataclasses import replace
    wb = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", replace(_catalogue(), flags=())))
    assert _values(wb["Flags"])[1] == ["Nothing to flag."]


def test_a_flag_starting_with_equals_is_kept_as_plain_text(tmp_path):
    # M1: openpyxl treats a string starting with "=" as a formula unless told otherwise.
    from dataclasses import replace
    catalogue = replace(_catalogue(), flags=("=SUM(A1) looks like a formula but is only a quote",))
    out = write_evidence(tmp_path / "e.xlsx", catalogue)
    ws = openpyxl.load_workbook(out)["Flags"]
    cell = ws["A2"]
    assert cell.data_type == "s"
    assert cell.value == "=SUM(A1) looks like a formula but is only a quote"


def test_headings_are_bold_and_frozen(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue()))["Sheets"]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


def _catalogue_with_overview() -> Catalogue:
    doc = "2. Takeoffs/1. Base Takeoff/Specs.pdf"
    page = PageRecord(doc, 1, "text", "Actuators shall be Belimo or approved equal. Provide BACnet MS/TP.", 0)
    sections = (SpecSection("23 09 23", "INSTRUMENTATION AND CONTROL FOR HVAC", doc, 1, 1),)
    overview = build_overview([page], [], sections)
    return replace(_catalogue(), overview=overview)


def test_the_spec_overview_tab_is_the_second_tab(tmp_path):
    out = write_evidence(tmp_path / "e.xlsx", _catalogue_with_overview())
    assert openpyxl.load_workbook(out).sheetnames == list(TABS)
    assert TABS[1] == "Spec overview"


def test_the_spec_overview_shows_found_and_not_found(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue_with_overview()))["Spec overview"]
    rows = _values(ws)
    assert rows[0][0] == "Spec overview — 26-9999"
    assert any(row[0] == "BACnet MS/TP" and row[1] == "Found" for row in rows)
    assert any(row[0] == "Modbus" and row[1] == "Not found" for row in rows)
    assert any(row[0] == "Actuators and dampers" and row[1] == "Belimo" and row[2] == "Yes" for row in rows)
    assert any("Modbus" in str(row[0]) and "not found" in str(row[0]) for row in rows)


def test_a_found_cell_is_green_and_a_not_found_cell_is_not(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue_with_overview()))["Spec overview"]
    for row in ws.iter_rows():
        values = [c.value for c in row]
        if values[0] == "BACnet MS/TP" and values[1] == "Found":
            assert row[1].fill.fgColor.rgb == "00C6EFCE"
        if values[0] == "Modbus" and values[1] == "Not found":
            assert row[1].fill.fgColor.rgb in (None, "00000000")


def test_ame_makers_get_vendors_and_an_email_draft(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue_with_overview()))["Spec overview"]
    rows = _values(ws)
    belimo = next(row for row in rows if row[0] == "Actuators and dampers" and row[1] == "Belimo")
    assert "Broudy Precision" in str(belimo[7])
    assert str(belimo[8]).startswith("Could you please provide the pricing for the Belimo - Actuators and dampers?")


def _catalogue_with_scope() -> Catalogue:
    doc = "2. Takeoffs/1. Base Takeoff/Specs.pdf"
    page = PageRecord(doc, 1, "text", "BACnet IP is not allowed anywhere in the building. The BMS contractor shall "
                                      "furnish and install all water flow meters.", 0)
    sections = (SpecSection("25 05 00", "INTEGRATED AUTOMATION", doc, 1, 1),)
    return replace(_catalogue(), overview=build_overview([page], [], sections))


def test_a_ruled_out_cell_is_amber(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue_with_scope()))["Spec overview"]
    row = next(r for r in ws.iter_rows() if r[0].value == "BACnet/IP")
    assert row[1].value == "Ruled out"
    assert row[1].fill.fgColor.rgb == "00FFEB9C"
    assert row[4].value == "BACnet IP is not allowed anywhere in the building."


def test_the_standard_exclusions_the_job_asks_for_are_on_the_tab(tmp_path):
    ws = openpyxl.load_workbook(write_evidence(tmp_path / "e.xlsx", _catalogue_with_scope()))["Spec overview"]
    rows = _values(ws)
    headings = [r[0] for r in rows if isinstance(r[0], str) and r[0][:2] in ("1.", "2.", "3.")]
    assert headings[:2] == ["1. Controls platform", "2. AME normally excludes this — the job asks for it"]
    flow = next(r for r in ws.iter_rows() if r[0].value == "All water flow meters.")
    assert flow[1].value == "Found" and flow[1].fill.fgColor.rgb == "00C6EFCE"
    assert flow[3].value == "Controls spec"
    assert flow[4].value.startswith("The BMS contractor shall furnish and install all water flow meters")
