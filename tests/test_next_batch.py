"""The next batch (.superpowers/sdd/2026-09-26-whole-job-writing/next-batch-brief.md): clarifications in brackets at
the end of their line, the AME practice rules in the GEM prompt, the estimator's corrections read back, the
--documents input folder, the scoreboard and the requests folder cleaned. Synthetic data only."""
from __future__ import annotations

from estimator.writing.assemble import maker_line, with_clarifications
from test_gem_round2 import STANDARD_NOTE, Contradiction, _assemble, _gl, _items, _planned, _scope, _scope_lines

EDH = _planned("EDH", "Electric Duct Heater EDH-1 thru 3", ("EDH-1", "EDH-2", "EDH-3"), 3, kind="Electric Duct Heater")
EDH_HEADER = "Qty (3) Electric Duct Heater EDH-1 thru 3"


# --- 1. clarifications in brackets at the end of their line -----------------------------------------------------------

def test_a_clarification_goes_in_brackets_at_the_end_of_its_line():
    assert with_clarifications("Furnish and Install EDH DDC Controller (Qty. 3 Total)",
                               ["*Assumed: One electric duct heater DDC controller per heater assumed"]) == (
        "Furnish and Install EDH DDC Controller (Qty. 3 Total) "
        "(*Assumed: One electric duct heater DDC controller per heater assumed)")
    assert with_clarifications("Furnish Heater Enable", []) == "Furnish Heater Enable"
    twice = with_clarifications("Furnish X", ["*Confirm quantity: X", "*Confirm quantity: X"])
    assert twice == "Furnish X (*Confirm quantity: X)"


def test_the_assumed_and_confirm_quantity_notes_print_once_for_the_block_after_its_lines():
    scope = _scope(_gl("Furnish and Install EDH DDC Controller (Qty. 3 Total)", quantity=3, quote="Job 26-9999",
                       flag="orange", flag_text="One electric duct heater DDC controller per heater assumed — because the schedule lists one heater per controller"),
                   _gl("Furnish and Install Airflow Switch", flag="red", flag_text="No count on the drawings"))
    job, systems, _, _ = _assemble([EDH], {"EDH": scope})
    items = _items(job, EDH_HEADER)
    texts = [t for _, t in items]
    # one clarification per block (2026-10-03): the lines carry no bracket; one note each, after the lines
    assert "Furnish and Install EDH DDC Controller (Qty. 3 Total)" in texts
    assert "Furnish and Install Airflow Switch (Qty. per Mechanical Schedule)" in texts
    stars = [t for kind, t in items if kind == "note" and t.startswith("*") and not t.startswith("**")]
    assert stars[0].startswith("*Confirm quantity: ") and "Airflow Switch" in stars[0]
    assert stars[1] == "*Assumed: One electric duct heater DDC controller per heater assumed, because the schedule lists one heater per controller"
    assert max(i for i, (kind, _) in enumerate(items) if kind == "line") < items.index(("note", stars[0]))


def test_a_c5_assumption_goes_in_the_block_note():
    scope = _scope(_gl("Furnish and Install EDH DDC Controller"),
                   contradictions=(Contradiction("C5", "One EDH DDC controller per heater assumed, because the schedule lists one per heater", (), "orange"),))
    job, _, _, _ = _assemble([EDH], {"EDH": scope})
    texts = [t for _, t in _items(job, EDH_HEADER)]
    assert "Furnish and Install EDH DDC Controller" in texts and any(t.startswith("*Assumed: ") for t in texts), texts


def test_a_manufacturer_device_ame_does_nothing_with_is_its_own_bracketed_line():
    assert maker_line("*REFRIGERANT LEAK DETECTION BY UNIT MFR.") == "(*Refrigerant Leak Detection by Unit Mfr.)"
    assert maker_line("*VAV CONTROLLER AND CO2 SENSOR BY UNIT MFR.") == "(*VAV Controller and CO2 Sensor by Unit Mfr.)"
    assert maker_line("**Owner to provide the network drop") == "**Owner to provide the network drop"


def test_the_furnished_by_unit_mfr_line_stays_in_the_house_form():
    scope = _scope(_gl("Install and Wire Space Thermostat (Furnished by Unit Mfr.)"),
                   quick_scan={"unit_type": "Factory/OEM Controls"})
    job, _, _, _ = _assemble([EDH], {"EDH": scope})
    (text,) = [t for kind, t in _items(job, EDH_HEADER) if kind == "line"]
    assert text.lower() == "install and wire space thermostat (furnished by unit mfr.)"


def test_the_standard_note_is_a_bracket_too():
    assert with_clarifications("Furnish and Install VAV Transformer", [STANDARD_NOTE]) == (
        "Furnish and Install VAV Transformer (*Assumed per AME standard — confirm)")


# --- 2. the AME practice rules in the GEM prompt -----------------------------------------------------------------------

def test_the_gem_prompt_carries_the_ame_practice_rules_after_connecting_the_dots():
    from estimator.writing.whole_job import PRACTICE_FILE, WRITE_PROMPT
    assert PRACTICE_FILE.name == "ame_practice.txt" and PRACTICE_FILE.exists()
    for words in ("FACTORY CONTROLS — AME PRACTICE: the controller and every end device on a factory/OEM-controls unit "
                  "belong to the unit manufacturer",
                  "'Install and Wire <device> (Furnished by Unit Mfr.)' or 'Wire <device> (Furnished by Unit Mfr.)'",
                  "explicit assignment always wins",
                  "SPLIT SYSTEMS AND VRF — AME PRACTICE: write 'Low Voltage Interlock Wiring between Indoor and "
                  "Outdoor Unit' when the units are field-connected",
                  "write the wiring lines for them unless another party is named"):
        assert words in WRITE_PROMPT, words
    assert WRITE_PROMPT.index("CONNECTING THE DOTS") < WRITE_PROMPT.index("FACTORY CONTROLS — AME PRACTICE")
    assert WRITE_PROMPT.index("SPLIT SYSTEMS AND VRF") < WRITE_PROMPT.index("OUTPUT\n")


def test_a_rule_the_user_adds_to_the_practice_file_reaches_the_prompt(tmp_path):
    from estimator.writing.whole_job import GEM_FILE, gem_prompt
    rules = tmp_path / "ame_practice.txt"
    rules.write_text("AME PRACTICE\nPUMPS — AME PRACTICE: a synthetic rule for the test.\n", encoding="utf-8")
    assert "PUMPS — AME PRACTICE: a synthetic rule for the test." in gem_prompt(GEM_FILE, rules)
    assert "FACTORY CONTROLS — AME PRACTICE" not in gem_prompt(GEM_FILE, tmp_path / "none.txt")


# --- 6. the requests folder cleaned ------------------------------------------------------------------------------------

def test_a_chat_run_moves_requests_with_no_job_recorded_and_this_jobs_older_ones_into_old(tmp_path):
    import json
    from estimator.writing.judge import Judge
    judge = Judge(tmp_path / "cache", mode="chat", job="99-0001, Synthetic School")
    requests = judge.requests_dir
    requests.mkdir(parents=True)
    files = {"mine": {"key": "a", "kind": "write_system", "job": "99-0001, Synthetic School", "prompt": "x"},
             "other": {"key": "b", "kind": "write_system", "job": "99-0002, Other Synthetic", "prompt": "x"},
             "nojob": {"key": "c", "kind": "write_system", "prompt": "no job recorded"},
             "nojob-plan": {"key": "d", "kind": "plan_job", "prompt": "The dossier:\n[S1] Job: 99-0002"}}
    for name, data in files.items():
        (requests / f"{name}.json").write_text(json.dumps(data), encoding="utf-8")
    (requests / "broken.json").write_text("{not json", encoding="utf-8")
    assert judge.retire_requests() == 4
    assert sorted(p.name for p in requests.glob("*.json")) == ["other.json"]
    assert sorted(p.name for p in (requests / "old").glob("*.json")) == ["broken.json", "mine.json", "nojob-plan.json",
                                                                        "nojob.json"]
    assert judge.stale_requests() == 1 and judge.unrecorded == 3


# --- 4. the --documents input folder -----------------------------------------------------------------------------------

def _job_folder(tmp_path):
    from pdfs import Page, make_pdf
    root = tmp_path / "99-0001, Synthetic School"
    drawing = Page(width=3024, height=2160, texts=((200, 200, "FLOOR PLAN", 20),))
    make_pdf(root / "1. Project Documents" / "1. Base" / "Mech.pdf", drawing)
    make_pdf(root / "1. Project Documents" / "2. Addenda" / "Addendum 1.pdf", Page(texts=((72, 72, "Words", 12),)))
    make_pdf(root / "2. Takeoffs" / "1. Base Takeoff" / "Marked_Mech.pdf", drawing)
    return root


def test_only_the_documents_subfolder_is_read_and_the_name_comes_from_the_job_folder(tmp_path):
    from estimator.catalogue.documents import find_documents
    root = _job_folder(tmp_path)
    whole = find_documents(root)
    assert len(whole.documents) == 3
    folder = find_documents(root, "1. Project Documents")
    assert folder.name == "99-0001" and folder.root == root
    assert [d.relative for d in folder.documents] == ["1. Project Documents/1. Base/Mech.pdf",
                                                      "1. Project Documents/2. Addenda/Addendum 1.pdf"]
    import pytest
    with pytest.raises(FileNotFoundError):
        find_documents(root, "9. Nowhere")


def test_catalogue_and_estimate_take_documents(tmp_path, monkeypatch):
    import estimator.__main__ as cli
    from estimator.catalogue.job import catalogue_job
    root = _job_folder(tmp_path)
    result = catalogue_job(root, tmp_path / "cache", documents="1. Project Documents")
    assert result.job.name == "99-0001" and len(result.job.documents) == 2
    for command in ("catalogue", "estimate"):
        assert cli._parser().parse_args([command, "JOB", "--out", "OUT"]).documents is None
        args = cli._parser().parse_args([command, "JOB", "--out", "OUT", "--documents", "1. Project Documents"])
        assert args.documents == "1. Project Documents"
    seen = []

    def _fake(job, cache, workers, **kw):
        seen.append(kw)
        raise SystemExit

    monkeypatch.setattr(cli, "catalogue_job", _fake)
    import pytest
    for command in ("catalogue", "estimate"):
        with pytest.raises(SystemExit):
            cli.main([command, str(root), "--out", str(tmp_path / "out"), "--documents", "1. Project Documents"])
    assert seen == [{"documents": "1. Project Documents"}] * 2


# --- 3. the estimator's corrections are read back ----------------------------------------------------------------------

def _corrected_package(tmp_path):
    """A synthetic job written, its Evidence Package filled in the way the estimator fills it."""
    import openpyxl
    from estimator.evidence.workbook import write_evidence
    from test_evidence import _catalogue as evidence_catalogue
    scope = _scope(_gl("Furnish and Install EDH DDC Controller (Qty. 3 Total)", quantity=3),
                   _gl("Furnish Heater Enable"),
                   _gl("Furnish Heater Status"),
                   _gl("Furnish and Install Airflow Switch"))
    job, systems, general, ranked = _assemble([EDH], {"EDH": scope})
    path = write_evidence(tmp_path / "99-0001 Evidence Package.xlsx", evidence_catalogue(), systems=systems,
                          review_flags=ranked)
    book = openpyxl.load_workbook(path)
    ws = book["Scope lines"]
    heads = [c.value for c in ws[1]]
    at, change = heads.index("Estimator decision") + 1, heads.index("Change to / why") + 1
    fill = {"Furnish Heater Status": ("Delete", "Status comes from the heater's own current switch, never a separate "
                                               "point on EDH-1 thru 3 (Job 99-0001, M-601)."),
            "Furnish Heater Enable": ("change", "Furnish Heater Enable / Disable why: AME writes enable and disable "
                                                "as one line."),
            "Furnish and Install Airflow Switch": ("keep", "")}
    for row in range(2, ws.max_row + 1):
        text = ws.cell(row, 2).value
        if text in fill:
            ws.cell(row, at).value, ws.cell(row, change).value = fill[text]
    review = book["Review list"]
    heads = [c.value for c in review[1]]
    decide_at = heads.index("Estimator decision") + 1
    first = review.cell(2, heads.index("What to decide") + 1).value
    review.cell(2, decide_at).value = "delete"
    book.save(path)
    return path, job, systems, ranked, first


def test_the_evidence_package_has_the_estimators_two_columns(tmp_path):
    import openpyxl
    path, *_ = _corrected_package(tmp_path)
    book = openpyxl.load_workbook(path)
    for tab in ("Scope lines", "Review list"):
        heads = [c.value for c in book[tab][1]]
        assert heads[-2:] == ["Estimator decision", "Change to / why"], tab
        assert book[tab].data_validations.dataValidation, tab            # keep / delete / change to choose from


def test_the_filled_columns_are_read_saved_under_out_and_merged(tmp_path):
    from estimator.corrections import load_corrections, read_corrections, save_corrections
    path, _, _, _, first = _corrected_package(tmp_path)
    corrections, problems = read_corrections(path)
    assert problems == [] and corrections["job"] == "26-9999"
    by_text = {e["text"]: e for e in corrections["lines"]}
    assert set(by_text) == {"Furnish Heater Status", "Furnish Heater Enable", "Furnish and Install Airflow Switch"}
    assert by_text["Furnish Heater Status"]["decision"] == "delete"
    assert by_text["Furnish Heater Enable"]["change_to"] == "Furnish Heater Enable / Disable"
    assert by_text["Furnish Heater Enable"]["why"] == "AME writes enable and disable as one line."
    assert by_text["Furnish and Install Airflow Switch"]["decision"] == "keep"
    assert [(e["text"], e["decision"]) for e in corrections["review"]] == [(first, "delete")]
    folder = tmp_path / "out" / "corrections"
    saved = save_corrections(corrections, folder)
    assert saved == folder / "26-9999.json"
    later = {"job": "26-9999", "lines": [dict(by_text["Furnish and Install Airflow Switch"], decision="delete")],
             "review": []}
    data = load_corrections(save_corrections(later, folder))
    decisions = {e["text"]: e["decision"] for e in data["lines"]}
    assert decisions == {"Furnish Heater Status": "delete", "Furnish Heater Enable": "change",
                         "Furnish and Install Airflow Switch": "delete"}          # the new decision replaces the old


def test_each_reason_is_proposed_as_a_general_rule_with_no_job_content(tmp_path):
    from estimator.corrections import general_rule, propose_rules, read_corrections
    path, *_ = _corrected_package(tmp_path)
    corrections, _ = read_corrections(path)
    rules = tmp_path / "ame_practice_proposed.txt"
    added = propose_rules(corrections, rules)
    assert added == ["- Status comes from the heater's own current switch, never a separate point.",
                     "- AME writes enable and disable as one line."]
    text = rules.read_text(encoding="utf-8")
    for leak in ("EDH-1", "99-0001", "M-601", "26-9999"):
        assert leak not in text, leak
    assert propose_rules(corrections, rules) == []                                # once each
    assert general_rule('Per note "EF-2 shall be interlocked" [S12] on sheet M-101.', "") is None
    assert general_rule("Write the chime interlock for every “leak” alarm the plans show.", "") == (
        "Write the chime interlock for every alarm the plans show.")


def test_the_next_run_applies_the_corrections(tmp_path):
    from estimator.corrections import apply_corrections, read_corrections
    path, job, systems, ranked, first = _corrected_package(tmp_path)
    corrections, _ = read_corrections(path)
    job, systems, ranked = apply_corrections(job, systems, ranked, corrections)
    texts = [t for _, t in _items(job, EDH_HEADER)]
    assert "Furnish Heater Status" not in texts
    assert "Furnish Heater Enable / Disable" in texts and "Furnish Heater Enable" not in texts
    assert "Furnish and Install Airflow Switch" in texts
    assert "Furnish and Install EDH DDC Controller (Qty. 3 Total)" in texts            # no decision: as it was
    lines = {line.text: line for s in systems for line in s.lines}
    assert lines["Furnish Heater Status"].client is False
    assert "was “Furnish Heater Enable”" in lines["Furnish Heater Enable / Disable"].evidence
    rows = [r for s in systems for r in s.review]
    assert any(r.decide == "Deleted by the estimator: “Furnish Heater Status”" for r in rows)
    assert first not in [r.decide for r in ranked]


def test_the_corrections_command_and_the_estimate_that_applies_them(tmp_path, capsys, monkeypatch):
    import estimator.__main__ as cli
    import estimator.writing.whole_job as wj
    path, *_ = _corrected_package(tmp_path)
    assert cli.main(["corrections", str(path)]) == 0
    printed = capsys.readouterr().out
    assert "Read 3 line decisions and 1 review decision for 26-9999." in printed
    assert (cli.CORRECTIONS / "26-9999.json").exists()
    assert "Added 2 proposed rules" in printed and wj.PROPOSED_PRACTICE_FILE.exists()
    empty = tmp_path / "empty"
    empty.mkdir()
    from estimator.evidence.workbook import write_evidence
    from test_evidence import _catalogue as evidence_catalogue
    blank = write_evidence(empty / "e.xlsx", evidence_catalogue(), systems=[], review_flags=[])
    assert cli.main(["corrections", str(blank), "--folder", str(empty)]) == 0
    assert "nothing to read" in capsys.readouterr().out
    assert cli.main(["corrections", str(tmp_path / "missing.xlsx")]) == 1


from test_estimate import job, whole_job, _proposal_text     # noqa: E402,F401  (the estimate's synthetic job)


def test_the_next_estimate_of_the_job_applies_its_corrections_file(job, tmp_path, capsys, whole_job):  # noqa: F811
    import json
    import openpyxl
    import estimator.__main__ as cli
    sensor = "Furnish and Install Space Temperature / CO2 Sensor"
    cli.CORRECTIONS.mkdir(parents=True)
    (cli.CORRECTIONS / "26-9999.json").write_text(json.dumps({"job": "26-9999", "review": [], "lines": [
        {"kind": "lines", "system": "", "text": sensor, "decision": "delete", "change_to": "",
         "why": "Synthetic reason."}]}), encoding="utf-8")
    out = tmp_path / "out"
    assert cli.main(["estimate", str(job), "--out", str(out), "--ai", "api", "--workers", "1"]) == 0
    assert "Applied the estimator's corrections" in capsys.readouterr().out
    assert sensor not in _proposal_text(out / "26-9999 Proposal.docx")
    book = openpyxl.load_workbook(out / "26-9999 Evidence Package.xlsx")
    lines = [[c.value for c in row] for row in book["Scope lines"].iter_rows()]
    row = next(r for r in lines if r[1] == sensor)
    assert row[8] == "no" and row[6] == "Deleted by the estimator"
    review = [[c.value for c in row] for row in book["Review list"].iter_rows()]
    assert any(r[3] == f"Deleted by the estimator: “{sensor}”" for r in review)


# --- 5. the scoreboard -------------------------------------------------------------------------------------------------

def _line(text, depth=1):
    from estimator.scope import Item, Run
    return Item(kind="line", depth=depth, runs=(Run(text=text),))


def _proposals(tmp_path, bare_job, proposal_template):
    """A synthetic draft and a synthetic sent proposal, both written on the AME template."""
    from estimator.proposal.write import write_proposal
    from estimator.scope import Block
    sent = bare_job.model_copy(update={"blocks": (
        Block(header="Qty (2) Exhaust Fan EF-1, 2", items=(
            _line("Furnish Fan Start / Stop and Status"), _line("Furnish and Install Damper Actuator (Qty. 2 Total)"),
            _line("Wire Fire Alarm Shutdown Relay"))),
        Block(header="Qty (3) Electric Duct Heater EDH-1 thru 3", items=(
            _line("Furnish and Install EDH DDC Controller (Qty. 3 Total)"), _line("Furnish Heater Enable"))),
        Block(header="Qty (1) Fire Pump FP-1", items=(_line("Furnish and Wire Fire Pump Running Alarm"),)))})
    draft = bare_job.model_copy(update={"blocks": (
        Block(header="Qty (2) Exhaust Fan EF-1, 2", items=(
            _line("Furnish Exhaust Fan Start/Stop and Status"),
            _line("Furnish and Install Damper Actuator (Qty. 4 Total) (*Assumed: One damper per fan)"),
            _line("Furnish and Install Space Temperature Sensor"))),
        Block(header="Qty (2) Electric Duct Heater EDH-1 thru 3", items=(
            _line("Furnish and Install EDH DDC Controller (Qty. 3 Total)"), _line("Furnish Heater Enable"))),
        Block(header="Qty (1) Generator GEN-1", items=(_line("Furnish Generator Run Status"),)))})
    return (write_proposal(draft, proposal_template, tmp_path / "99-0001 Proposal.docx"),
            write_proposal(sent, proposal_template, tmp_path / "sent.docx"))


def test_lines_are_compared_by_verb_and_device_words():
    from estimator.score import line_overlap, line_quantity, line_tokens
    assert line_tokens("Furnish and Install Damper Actuator (Qty. 4 Total) (*Assumed: One per fan)") == (
        "furnish and install", frozenset({"damper", "actuator"}))
    assert line_overlap("Furnish Fan Start / Stop and Status", "Furnish Exhaust Fan Start/Stop and Status") >= 0.6
    assert line_overlap("Furnish Fan Start / Stop and Status", "Wire Fire Alarm Shutdown Relay") < 0.6
    assert (line_quantity("Furnish X (Qty. 4 Total)"), line_quantity("Furnish X - (x2)"), line_quantity("Furnish X")) \
        == (4, 2, None)


def test_tags_are_normalised_and_stemmed_by_type():
    from estimator.score import normalize_tag, tag_stem
    assert normalize_tag("RTU-CR-1") == normalize_tag("RTU_CR 1") == "RTUCR1"
    assert tag_stem("HEL05_01_CRAH01") == "CRAH"
    assert tag_stem("RTU-CR-1") == tag_stem("RTU-CR1") == "RTUCR"
    assert tag_stem("CRAH-#") == "CRAH"
    assert tag_stem("DOAU##") == "DOAU"
    assert tag_stem("FCU-#") == "FCU"


def test_kind_words_expand_abbreviations_both_ways():
    from estimator.score import _block
    vav = _block(["Qty (57) VAV with Electric Reheat"], [])
    vav_box = _block(["Qty (57) VAV Box VAV-100 thru VAV-521"], [])
    assert {"variable", "air", "volume"} <= vav.kind
    assert {"variable", "air", "volume"} <= vav_box.kind
    chillers = _block(["Qty (112) Chillers"], [])
    air_cooled = _block(["Qty (112) Air-Cooled Chiller HEL05_01_CH01"], [])
    assert chillers.kind <= air_cooled.kind and "chiller" in chillers.kind


def test_the_scoreboard_compares_systems_lines_and_counts(tmp_path, bare_job, proposal_template):
    from estimator.score import read_blocks, score
    draft, sent = _proposals(tmp_path, bare_job, proposal_template)
    result = score(read_blocks(draft), read_blocks(sent), "99-0001")
    systems = {(row[0], row[1]): row[2] for row in result.systems}
    assert systems[("Qty (2) Exhaust Fan EF-1, 2", "Qty (2) Exhaust Fan EF-1, 2")] == "matched"
    assert systems[("Qty (3) Electric Duct Heater EDH-1 thru 3", "Qty (2) Electric Duct Heater EDH-1 thru 3")] \
        == "quantity differs"
    assert systems[("Qty (1) Fire Pump FP-1", None)] == "missing"
    assert systems[(None, "Qty (1) Generator GEN-1")] == "extra"
    lines = {(row[1], row[2]): row[3] for row in result.lines}
    assert lines[("Furnish Fan Start / Stop and Status", "Furnish Exhaust Fan Start/Stop and Status")] == \
        "worded differently"
    assert lines[("Furnish Heater Enable", "Furnish Heater Enable")] == "matched"
    assert lines[("Wire Fire Alarm Shutdown Relay", None)] == "missing"
    assert lines[(None, "Furnish and Install Space Temperature Sensor")] == "extra"
    damper = next(row for row in result.lines if row[1] == "Furnish and Install Damper Actuator (Qty. 2 Total)")
    assert (damper[3], damper[5], damper[6], damper[7]) == ("matched", 2, 4, "differ")
    assert result.summary == "systems 2/3, lines 4/6, counts 2/4"


def test_the_score_command_writes_the_workbook_and_the_summary_line(tmp_path, bare_job, proposal_template, capsys):
    import docx
    import openpyxl
    import estimator.__main__ as cli
    draft, sent = _proposals(tmp_path, bare_job, proposal_template)
    out = tmp_path / "score"
    assert cli.main(["score", str(draft), str(sent), "--out", str(out)]) == 0
    printed = capsys.readouterr().out
    assert "Score for 99-0001: systems 2/3, lines 4/6, counts 2/4." in printed
    book = openpyxl.load_workbook(out / "99-0001 Score.xlsx")
    assert book.sheetnames == ["Summary", "Systems", "Lines"]
    assert book["Summary"]["A2"].value == "systems 2/3, lines 4/6, counts 2/4"
    assert "not a judge" in book["Summary"]["A3"].value and "miscount" in book["Summary"]["A3"].value
    summary_rows = [row[0].value for row in book["Summary"].iter_rows()]
    assert "Systems matched by kind or tag type only (check these)" in summary_rows
    assert [p.name for p in out.iterdir()] == ["99-0001 Score.xlsx"]              # nothing else is written
    code = cli.main(["score", str(draft), str(sent), "--out", str(cli.ROOT / "tests" / "no-score-here")])
    assert code == 1 and "under" in capsys.readouterr().err
    assert not (cli.ROOT / "tests" / "no-score-here").exists()
    # the older "(xN)" style is read by the phrasing bank's fallback
    old = docx.Document()
    for text in ("SCOPE OF WORK – ", "(x2) Exhaust Fan EF-1, 2", "Furnish Fan Start / Stop and Status - (x2)",
                 "CLARIFICATIONS"):
        old.add_paragraph(text)
    old.save(tmp_path / "old.docx")
    from estimator.score import read_blocks
    (block,) = read_blocks(tmp_path / "old.docx")
    assert (block.header, block.quantity, block.tags) == ("(x2) Exhaust Fan EF-1, 2", 2, frozenset({"EF1", "EF2"}))


def test_house_short_forms_are_matched_against_long_names_and_joined_tags():
    # Sent proposals write systems in the house short form; the draft writes long names and joined tags.
    from estimator.score import _block, score
    sent = [
        _block(["Qty (8) DOAU##"], []),
        _block(["Qty (80) ITE (Data Hall) CRAH-#"], []),
        _block(["Qty (112) Chillers"], []),
        _block(["Qty (3) RTU-CR-1, RTU-CR-2, RTU-CR-3"], []),
    ]
    draft = [
        _block(["Qty (8) Dedicated Outdoor Air Unit HEL05_01_DOAU01, 02 thru 08"], []),
        _block(["Qty (80) Computer Room Air Handler HEL05_01_CRAH01 thru 80"], []),
        _block(["Qty (112) Air-Cooled Chiller HEL05_01_CH01 thru 28"], []),
        _block(["Qty (3) Rooftop Unit RTU-CR1 thru CR3"], []),
    ]
    result = score(draft, sent, "99-0002")
    systems = {row[0]: row[2] for row in result.systems}
    assert systems["Qty (8) DOAU##"] == "matched"
    assert systems["Qty (80) ITE (Data Hall) CRAH-#"] == "matched"
    assert systems["Qty (112) Chillers"] == "matched"
    assert systems["Qty (3) RTU-CR-1, RTU-CR-2, RTU-CR-3"] == "matched"


def test_two_unrelated_blocks_do_not_pair():
    from estimator.score import _block, score
    result = score([_block(["Qty (1) Generator GEN-1"], [])], [_block(["Qty (1) Boiler B-1"], [])], "99-0004")
    assert result.systems == [
        ["Qty (1) Boiler B-1", None, "missing", 1, None, None],
        [None, "Qty (1) Generator GEN-1", "extra", None, 1, None],
    ]


def test_a_family_split_lets_each_part_match_a_different_sent_block():
    from estimator.score import _block, score
    sent = [
        _block(["Qty (57) VAV with Electric Reheat"], [("Furnish VAV-100 DDC Controller", 1)]),
        _block(["Qty (9) Constant Air Volume Box"], [("Furnish CAV-1 DDC Controller", 1)]),
    ]
    draft = [_block(
        ["Qty (57) VAV Box VAV-100 thru VAV-521", "Qty (9) Constant Air Volume Box CAV-1 thru 9"],
        [("Furnish VAV-100 DDC Controller", 1), ("Furnish CAV-1 DDC Controller", 1)],
    )]
    result = score(draft, sent, "99-0003")
    systems = {row[0]: (row[1], row[2]) for row in result.systems}
    assert systems["Qty (57) VAV with Electric Reheat"] == ("Qty (57) VAV Box VAV-100 thru VAV-521", "matched")
    assert systems["Qty (9) Constant Air Volume Box"] == \
        ("Qty (9) Constant Air Volume Box CAV-1 thru 9", "matched")
    lines = {(row[1], row[2]): row[3] for row in result.lines}
    assert lines[("Furnish VAV-100 DDC Controller", "Furnish VAV-100 DDC Controller")] == "matched"
    assert lines[("Furnish CAV-1 DDC Controller", "Furnish CAV-1 DDC Controller")] == "matched"


def test_a_qc_line_with_its_quantity_matches_the_scope_line_without_it():
    from estimator.corrections import apply_corrections
    scope = _scope(_gl("Furnish Heater Enable"), _gl("Furnish Heater Status"))
    job, systems, general, ranked = _assemble([EDH], {"EDH": scope})
    corrections = {"job": "26-9999", "review": [],
                   "lines": [{"kind": "lines", "system": "EDH", "text": "Furnish Heater Status (Qty. 3 Total)",
                              "decision": "change", "change_to": "Furnish Heater Status Alarm", "why": ""}]}
    job, systems, ranked = apply_corrections(job, systems, ranked, corrections)
    texts = [t for _, t in _items(job, EDH_HEADER)]
    assert "Furnish Heater Status Alarm" in texts and "Furnish Heater Status" not in texts


def test_a_kept_qc_finding_is_taken_off_the_review_list_next_run(tmp_path):
    import openpyxl
    from estimator.corrections import apply_corrections, read_corrections
    from estimator.evidence.workbook import write_evidence
    from estimator.writing.assemble import review_list
    from estimator.writing.qc import QcFinding, QcReport, qc_review_rows
    from test_evidence import _catalogue as evidence_catalogue
    finding = QcFinding("EDH", "", "missing", "No airflow switch is written for the duct heater.",
                        "Furnish and Install Airflow Switch", "check", "")
    report = QcReport((finding,), {}, "", "m", ())
    rows = review_list([], qc_review_rows(report))
    assert rows[0].decide == finding.say                     # the review row is worded as the finding
    path = write_evidence(tmp_path / "e.xlsx", evidence_catalogue(), systems=[], qc=report)
    book = openpyxl.load_workbook(path)
    book["Proposal QC"].cell(2, 8).value = "keep"
    book.save(path)
    corrections, problems = read_corrections(path)
    assert problems == []
    scope = _scope(_gl("Furnish Heater Enable"))
    job, systems, _, _ = _assemble([EDH], {"EDH": scope})
    _, _, left = apply_corrections(job, systems, rows, corrections)
    assert finding.say not in [r.decide for r in left]
