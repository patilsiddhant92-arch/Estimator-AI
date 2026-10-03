"""The GEM round (.superpowers/sdd/2026-09-26-whole-job-writing/gem-round-brief.md): systems are the sequence's
sections, the write step runs the estimator's GEM, and the assembly follows the GEM instead of the trimming rules."""
from __future__ import annotations

import json

from estimator.reading.sequences import POINTS, SEQUENCE
from estimator.writing.dossier import build_dossier, is_section_heading, sections_in
from estimator.writing.whole_job import PLAN_PROMPT, Family, parse_plan, plan_schema
from estimator.writing.style import load_house_style
from test_dossier import SPEC, _backbone, _catalogue, _doc, _job, _page, _reading

STYLE = load_house_style()


# --- Part 1: the dossier marks the sequence's headings --------------------------------------------------------------

def test_a_lettered_or_numbered_title_is_a_section_heading_and_a_sentence_is_not():
    for heading in ("C. DRYER EXHAUST AND MAKE-UP AIR FANS", "1. RTU-1 AND RTU-2", "3.1 AIR HANDLING UNITS",
                    "D. Outdoor Air Fan:", "B. VRF Split System Units.", "A. Energy Recovery Ventilation Unit",
                    "I. Parking Garage Fans and Associated CO Control System",
                    "C. Rooftop Units with Heat Recovery (Heat Recovery where indicated on schedules):"):
        assert is_section_heading(heading), heading
    for sentence in ("A. AHU-1 shall run continuously in the occupied mode.", "4. Warm-up:", "1. DDC points:",
                     "a. Outdoor air temperature", "H. E/P: Electric to pneumatic.", "12. Bill of materials.",
                     "F. Provide surge protection for wiring and all system components.",
                     "A. The BMS contractor to provide the following global points to be available to all ddc "
                     "controllers on the BMS.", "SECTION 23 09 93 - SEQUENCE OF OPERATIONS"):
        assert not is_section_heading(sentence), sentence


DRYER_PAGE = """C. DRYER EXHAUST AND MAKE-UP AIR FANS

1. The dryer exhaust fans DXF-1 and DXF-2 shall run with the dryers.

2. The make-up air fans MUAF-1 and MUAF-2 shall be interlocked with the dryer exhaust fans.

D. EMERGENCY GENERATOR

1. The BMS shall monitor the generator run status and common alarm.
"""


def _sectioned_job():
    catalogue, reading = _job()
    catalogue.job.documents += (_doc("Specs/Seq2.pdf", "specs", 1),)
    catalogue.pages += (_page("Specs/Seq2.pdf", 1, DRYER_PAGE),)
    reading.backbone += (_backbone(SEQUENCE, "23 09 93 SEQUENCE PART 2 (Seq2.pdf, page 1)", "Specs/Seq2.pdf", 1, 1),)
    return catalogue, reading


def test_each_heading_of_the_sequence_starts_a_section_block_with_a_stable_id(tmp_path):
    d = build_dossier(*_sectioned_job(), tmp_path)
    sequence = d.text.split("== 2. SEQUENCE OF OPERATION ==")[1].split("== 3.")[0]
    assert "--- [Q1] 3.1 AIR HANDLING UNITS ---" in sequence
    assert "--- [Q2] 3.2 EXHAUST FANS ---" in sequence
    assert "--- [Q3] C. DRYER EXHAUST AND MAKE-UP AIR FANS ---" in sequence
    assert "--- [Q4] D. EMERGENCY GENERATOR ---" in sequence
    heading, ids = d.sections["Q3"]
    assert heading == "C. DRYER EXHAUST AND MAKE-UP AIR FANS"
    assert len(ids) == 3                                   # the heading's paragraph and the two under it
    texts = [sequence.split(f"[{i}]", 1)[1].split("\n", 1)[0] for i in ids]
    assert "DRYER EXHAUST" in texts[0] and "DXF-1" in texts[1] and "MUAF-1" in texts[2]
    assert all(i.startswith("S") for i in ids)            # every paragraph keeps its own [S..] id
    assert d.sections["Q4"][0] == "D. EMERGENCY GENERATOR" and len(d.sections["Q4"][1]) == 2
    # the section of the first heading holds its paragraphs up to the next heading, and no further
    first = d.sections["Q1"][1]
    assert any("AHU-1 shall run" in sequence.split(f"[{i}]", 1)[1].split("\n", 1)[0] for i in first)
    assert not any("EF-1 and EF-2" in sequence.split(f"[{i}]", 1)[1].split("\n", 1)[0] for i in first)


def test_the_sections_are_read_back_from_the_text(tmp_path):
    d = build_dossier(*_sectioned_job(), tmp_path)
    assert sections_in(d.text) == d.sections
    assert "Q3" in d.sources and "DRYER EXHAUST" in d.sources["Q3"]


def test_a_points_list_with_per_system_headings_gets_sections_too(tmp_path):
    catalogue, reading = _job()
    points = "A. AHU-1 POINTS\n\n1. Supply fan status.\n\n2. Discharge air temperature.\n\nB. EF-1 POINTS\n\n1. Fan status.\n"
    catalogue.pages = tuple(p if (p.document, p.page) != (SPEC, 3) else _page(SPEC, 3, points) for p in catalogue.pages)
    d = build_dossier(catalogue, reading, tmp_path)
    listed = d.text.split("== 3. POINTS LISTS ==")[1].split("== 4.")[0]
    assert "] A. AHU-1 POINTS" in listed
    q = next(k for k, (heading, _) in d.sections.items() if heading == "A. AHU-1 POINTS")
    assert f"--- [{q}] A. AHU-1 POINTS ---" in listed and len(d.sections[q][1]) == 3


def test_a_job_with_no_headings_has_no_sections(tmp_path):
    documents = [_doc(SPEC, "specs", 1)]
    catalogue = _catalogue(documents, [_page(SPEC, 1, "The fans shall run.\n\nThe pumps shall run.")])
    reading = _reading(backbone=[_backbone(SEQUENCE, "SEQ (Specs.pdf, page 1)", SPEC, 1, 1)])
    d = build_dossier(catalogue, reading, tmp_path)
    assert d.sections == {} and "[Q" not in d.text


# --- Part 1: the plan names each system's sequence section and its families ----------------------------------------

def test_the_plan_prompt_makes_one_system_per_sequence_section_with_one_header_per_family():
    for words in ("One system per sequence section.",
                  "A section that describes several equipment families together (dryer exhaust fans and their "
                  "make-up air fans; a fan and its damper; a chiller plant) is ONE system with one header per family.",
                  "Give each system its `Q` id(s), and every equipment family it covers with the family's count.",
                  "A family that no sequence section covers becomes its own system, citing its schedule row or plan "
                  "count."):
        assert words in PLAN_PROMPT, words
    system = plan_schema()["properties"]["systems"]["items"]
    assert {"sequence_section", "families"} <= set(system["required"])
    family = system["properties"]["families"]["items"]
    assert set(family["required"]) == {"header", "tags", "quantity", "quantity_source", "quantity_quote"}


def _dryer_answer(**families):
    fams = families.get("families", [
        {"header": "Dryer Exhaust Fan DXF-1-1, 2", "tags": ["DXF-1-1", "DXF-1-2"], "quantity": 2,
         "quantity_source": "[S5]", "quantity_quote": "DXF-1-1 | DXF-1-2"},
        {"header": "Make-up Air Fan MUAF-1-1, 2", "tags": ["MUAF-1-1", "MUAF-1-2"], "quantity": 2,
         "quantity_source": "[S6]", "quantity_quote": "MUAF-1-1 | MUAF-1-2"}])
    return {"systems": [{
        "key": "DXF-MUAF", "header": "Dryer Exhaust and Make-up Air Fans", "kind": "Exhaust Fan",
        "tags": ["DXF-1-1", "DXF-1-2", "MUAF-1-1", "MUAF-1-2"], "quantity": 2, "quantity_source": "[S5]",
        "quantity_quote": "DXF-1-1 | DXF-1-2", "sequence_sources": ["[S2]"], "schedule_sources": ["[S5]", "[S6]"],
        "group": "", "alternate": "", "exception": "", "why": "One sequence section runs them together.",
        "sequence_section": "[Q3]", "families": fams}],
        "front_end": {"mode": "new-system", "existing_name": "", "lines": [], "notes": [], "sources": []},
        "unresolved": []}


KNOWN = frozenset({"S2", "S5", "S6", "Q3"})


def test_a_system_of_several_families_keeps_its_section_and_each_family_with_its_count():
    plan = parse_plan(_dryer_answer(), "m", STYLE, KNOWN)
    (system,) = plan.systems
    assert system.sequence_section == "Q3"
    assert system.families == (
        Family("Dryer Exhaust Fan DXF-1-1, 2", ("DXF-1-1", "DXF-1-2"), 2, "[S5]", "DXF-1-1 | DXF-1-2"),
        Family("Make-up Air Fan MUAF-1-1, 2", ("MUAF-1-1", "MUAF-1-2"), 2, "[S6]", "MUAF-1-1 | MUAF-1-2"))


def test_a_family_count_with_no_words_stating_it_is_cleared_and_asked_but_the_system_stays():
    answer = _dryer_answer()
    answer["systems"][0]["families"][1]["quantity_quote"] = ""
    plan = parse_plan(answer, "m", STYLE, KNOWN)
    (system,) = plan.systems
    assert system.families[1].quantity is None
    assert any("Make-up Air Fan" in q for q in plan.unresolved)


def test_one_family_is_the_system_itself():
    answer = _dryer_answer(families=[{"header": "Dryer Exhaust Fan DXF-1-1, 2", "tags": ["DXF-1-1", "DXF-1-2"],
                                      "quantity": 2, "quantity_source": "[S5]", "quantity_quote": "DXF-1-1 | DXF-1-2"}])
    (system,) = parse_plan(answer, "m", STYLE, KNOWN).systems
    assert system.families == ()
    assert system.header == "Dryer Exhaust Fan DXF-1-1, 2" and system.tags == ("DXF-1-1", "DXF-1-2")


def test_an_old_plan_answer_without_sections_or_families_still_reads():
    answer = _dryer_answer()
    for key in ("sequence_section", "families"):
        del answer["systems"][0][key]
    (system,) = parse_plan(json.loads(json.dumps(answer)), "m", STYLE, KNOWN).systems
    assert system.sequence_section == "" and system.families == ()


# --- Part 2: the write step runs the GEM ---------------------------------------------------------------------------

from estimator.writing.judge import Judge                                            # noqa: E402
from estimator.writing.whole_job import (GEM_FILE, WRITE_PROMPT, PlannedSystem, WrittenScope, gem_prompt,  # noqa: E402
                                         parse_scope, write_prompt, write_schema)
from estimator.writing.assemble import system_evidence                                # noqa: E402
from test_assemble import Dossier                                                     # noqa: E402
from test_whole_job import FakeClient, _response                                      # noqa: E402


def test_the_gem_is_copied_into_the_assets():
    assert GEM_FILE.name == "BMS_DDC_GEM_v3.2.txt" and GEM_FILE.parent.name == "gem"
    assert "DEVICE OWNERSHIP RESOLUTION" in GEM_FILE.read_text(encoding="utf-8")


def test_the_system_prompt_is_the_gem_without_its_chat_mechanics():
    prompt = WRITE_PROMPT
    assert prompt == gem_prompt()
    assert prompt.startswith("WHO YOU ARE")
    assert prompt.rstrip().endswith("END OF INSTRUCTIONS — v3.2")
    for kept in ("DEVICE OWNERSHIP RESOLUTION", "SELF-QC", "CONTRADICTION AND ASSUMPTION SCAN",
                 "DAMPER ACTUATOR RULE — DISTINCT TYPES, DISTINCT LINES", "CONSOLIDATING IDENTICAL SCOPE LINES",
                 "No named party + packaged/factory controls unit",
                 "POSITION 19 — INTEGRATION LINES (always last within scope)",
                 "If any check shows FAIL: state the issue and the fix applied.", "STEP 0 — PROJECT QUICK SCAN",
                 "Unknown quantities: (Qty. per Mechanical Schedule) / (Qty. per Drawing)"):
        assert kept in prompt, kept
    for gone in ("RESET", "AUTO-START", "Reply 'highlight'", "PHASE 2 — CLEAN SOW", "DISPLAY FORMAT",
                 "SEQUENCE HARDWARE HIGHLIGHT", "HARDWARE POINTS IDENTIFIED", "[?]", "CHANGES IN v3.2"):
        assert gone not in prompt, gone
    assert ("OUTPUT\nAnswer as JSON in the schema given. Put the Quick Scan, the System Explanation, the scope lines "
            "(each with its reference line and flag), the Contradiction and Assumption scan, the Miscellaneous "
            "Responsibilities, the Exclusions and the Self-QC results in their fields.") in prompt
    assert prompt.index("OUTPUT\nAnswer as JSON") < prompt.index("END OF INSTRUCTIONS")


def test_the_gem_connects_the_dots_with_the_other_documents_after_its_input_handling():
    prompt = WRITE_PROMPT
    heading = "CONNECTING THE DOTS — THE OTHER DOCUMENTS (addendum for use inside the estimating tool)"
    assert heading in prompt
    assert (prompt.index("INPUT HANDLING — SEQUENCE AND POINT LIST") < prompt.index("PRIORITY RULE:")
            < prompt.index(heading) < prompt.index("POINT LIST INTEGRATION RULE (MANDATORY)"))
    section = prompt[prompt.index(heading):prompt.index("POINT LIST INTEGRATION RULE (MANDATORY)")]
    for words in ("Priority 3 — Equipment schedules and their notes", "Priority 6 — The specification",
                  "flag C1 / C2 with both snippets quoted", "(Quantity rules, Rule 1 and 2)",
                  "Counts from the drawings", "(Qty. N Total)", "(Qty. per Mechanical Schedule) or (Qty. per Drawing)",
                  "Factory/OEM Controls", "Field-Installed DDC", "(As per Plan — not called in sequence)",
                  "BY UNIT MFR.", "Category B", "(Schedule note 4, H601)", "(Plan count, M-101 × 10 floors)",
                  "(Spec 23 09 23 §3.4)"):
        assert words in section, words


GEM_TEXT = "\n".join([
    "== 1. JOB ==", "[S1] Job: 26-9999", "",
    "== 2. SEQUENCE OF OPERATION ==",
    "--- 23 09 93 SEQUENCE, Specs.pdf page 3 ---",
    "--- [Q1] I. PARKING GARAGE FANS AND ASSOCIATED CO CONTROL SYSTEM ---",
    "[S10] I. PARKING GARAGE FANS AND ASSOCIATED CO CONTROL SYSTEM",
    "[S11] 1. The garage exhaust system consists of the CO detector system and exhaust fans GXF-C-1.",
    "[S12] 2. The CO/NO2 sensors shall start the fan on a high reading.",
    "--- [Q2] J. DRYER EXHAUST AND MAKE-UP AIR FANS ---",
    "[S13] J. DRYER EXHAUST AND MAKE-UP AIR FANS",
    "[S14] 1. The dryer exhaust fans shall run with the dryers.", "",
    "== 3. POINTS LISTS ==",
    "[S20] GXF-C-1 status, GXF-C-1 start/stop, CO level high alarm.",
    "[S21] DXF-1-1 status.", "",
    "== 4. SCHEDULES ==",
    "[S30] GARAGE FAN SCHEDULE — M-601, Mech.pdf page 5 (2 rows)",
    "    Columns: TAG | CFM",
    "[S31] GXF-C-1 | 12000",
    "[S32] DXF-1-1 | 400",
    "[S33] Note: 1. FAN FURNISHED WITH VFD.", "",
    "== 5. EQUIPMENT COUNTS ==",
    "[S40] GXF (Fan): schedule rows 1; QUANTITY column total —; distinct tags 1; plan placements 1 (M-100.00 1); "
    "riser —. Tags: GXF-C-1.",
    "[S41] DXF (Fan): schedule rows 1; QUANTITY column total —; distinct tags 1; plan placements 0; riser —.", "",
    "== 6. DEVICES ON PLANS ==",
    "[S50] CO (CO sensor): 9 in all — M-100.00 9",
    "[S51] SP (Static pressure sensor): 2 in all — M-111.00 2", "",
    "== 7. NOTES ==",
    "--- M-100.00, Mech.pdf page 2 ---",
    "[S60] KEY NOTES 4: PROVIDE CO/NO2 SENSORS FOR GXF-C-1, SEE SEQUENCE.",
    "[S61] KEY NOTES 5: DXF-1-1 ON THE ROOF.", "",
    "== 8. CONTROLS SPEC HIGHLIGHTS ==", "(No controls spec section found.)", ""])
GEM_SOURCES = {f"S{n}": f"place {n}" for n in (1, 10, 11, 12, 13, 14, 20, 21, 30, 31, 32, 33, 40, 41, 50, 51, 60, 61)}
GEM_COUNTS = {"GXF": {"schedule_rows": 1, "quantity_column": None, "plan_placements": 1, "riser": None,
                      "distinct_tags": 1}}


def _gem_dossier():
    return Dossier("26-9999", GEM_TEXT, (), GEM_SOURCES, (), {"CO": 9, "SP": 2}, GEM_COUNTS, {})


def _garage(**kw):
    values = dict(key="GARAGE", header="Garage Exhaust Fan GXF-C-1", kind="Exhaust Fan", tags=("GXF-C-1",),
                  quantity=1, quantity_source="[S31]", quantity_quote="GXF-C-1 | 12000", sequence_sources=("[S12]",),
                  schedule_sources=("[S31]",), group="", alternate="", exception="", why="the garage section",
                  sequence_section="Q1")
    values.update(kw)
    return PlannedSystem(**values)


def test_the_evidence_gives_the_whole_section_then_points_schedules_notes_and_counts_in_that_order():
    evidence = system_evidence(_garage(), _gem_dossier())
    order = ["[S10] I. PARKING GARAGE", "[S11] 1. The garage exhaust system",           # not cited by the plan
             "[S12] 2. The CO/NO2 sensors", "[S20] GXF-C-1 status", "[S30] GARAGE FAN SCHEDULE", "[S31] GXF-C-1 | 12000",
             "[S33] Note: 1. FAN FURNISHED WITH VFD.", "[S60] KEY NOTES 4",
             "Counts from the drawings (a count on the drawings is a confirmed quantity — Quantity rules, Rule 1 and "
             "2)", "[S40] GXF (Fan)", "[S50] CO (CO sensor): 9 in all"]
    at = [evidence.index(part) for part in order]
    assert at == sorted(at), [p for p, a in zip(order, at)]
    for other in ("[S14]", "[S21]", "[S32]", "[S41]", "[S61]", "[S51]"):     # another system's, or another sheet's
        assert other not in evidence, other


def test_the_write_request_carries_the_system_its_section_and_the_template_in_the_gems_words():
    system = _garage()
    prompt = write_prompt(system, system_evidence(system, _gem_dossier()), [], ["Furnish and Install DDC Controller"],
                          [], "STYLE", "Exhaust Fan: 4 h per unit")
    assert "Garage Exhaust Fan GXF-C-1" in prompt and "Sequence section: Q1" in prompt
    assert ("The Standard Template defines the ORDER OF CATEGORIES in the output. It does not define device names."
            in prompt)
    assert "- Furnish and Install DDC Controller" in prompt and "Exhaust Fan: 4 h per unit" in prompt
    assert prompt.index("[S10]") < prompt.index("ORDER OF CATEGORIES")


def test_the_write_request_lists_each_family_with_its_count():
    system = _garage(families=(Family("Dryer Exhaust Fan DXF-1-1, 2", ("DXF-1-1", "DXF-1-2"), 2, "[S32]", "DXF"),
                               Family("Make-up Air Fan MUAF-1-1, 2", ("MUAF-1-1", "MUAF-1-2"), None, "", "")))
    prompt = write_prompt(system, "", [], [], [], "STYLE")
    assert "Dryer Exhaust Fan DXF-1-1, 2: 2" in prompt and "Make-up Air Fan MUAF-1-1, 2: not stated" in prompt


def test_the_answer_schema_has_the_gems_fields():
    schema = write_schema()
    assert {"quick_scan", "explanation", "lines", "contradictions", "misc_responsibilities", "exclusions", "qc",
            "client_notes", "tech_hours"} <= set(schema["required"])
    line = schema["properties"]["lines"]["items"]
    assert set(line["required"]) == {"text", "reference", "flag", "flag_text", "quantity", "quote", "source_id",
                                     "depth", "fact_ids"}
    assert line["properties"]["flag"]["enum"] == ["", "red", "yellow", "orange", "green"]
    scan = schema["properties"]["quick_scan"]
    assert set(scan["required"]) == {"systems_identified", "unit_type", "controller", "submittal_status",
                                     "blanket_wiring_clause", "fa_interface_present", "point_list_provided",
                                     "contractor_name_used"}
    contradiction = schema["properties"]["contradictions"]["items"]
    assert contradiction["properties"]["type"]["enum"] == ["C1", "C2", "C3", "C4", "C5", "C6"]
    qc = schema["properties"]["qc"]["items"]
    assert qc["properties"]["check"]["enum"] == [f"QC-{n:02d}" for n in range(1, 11)]


def _gem_line(text, reference="The CO/NO2 sensors shall start the fan", flag="", flag_text="", quantity=None,
              quote="The CO/NO2 sensors shall start the fan", source_id="[S12]", depth=1):
    return {"text": text, "reference": reference, "flag": flag, "flag_text": flag_text, "quantity": quantity,
            "quote": quote, "source_id": source_id, "depth": depth}


DIAMOND = "◊ Controls Contractor requires mfr. rep to assist in integration, cost to be incurred by mech contractor"


def gem_answer(lines=None, **overrides):
    answer = {
        "quick_scan": {"systems_identified": "Garage exhaust fan, CO/NO2 detection", "unit_type": "Field-Installed DDC",
                       "controller": "CC Furnishes and Installs DDC Controller", "submittal_status": "Not Available",
                       "blanket_wiring_clause": "No", "fa_interface_present": "No", "point_list_provided": "Yes",
                       "contractor_name_used": "BMS contractor"},
        "explanation": "The garage exhaust fan runs on the CO/NO2 sensors.",
        "lines": lines if lines is not None else [
            _gem_line("Furnish and Install DDC Controller"),
            _gem_line("Furnish and Install CO/NO2 Sensor (Qty. 9 Total)", quantity=9, quote="CO (CO sensor): 9 in all",
                      source_id="[S50]"),
            _gem_line("Furnish Garage Exhaust Fan Start/Stop and Status",
                      reference="Mandatory minimum point — fan under DDC control. Not explicitly called in sequence."),
            _gem_line("Map Available Points Via Mfr. Provided BMS Interface"),
            _gem_line(DIAMOND, depth=1)],
        "contradictions": [{"type": "C5", "item": "Garage Exhaust Fan VFD Control",
                            "quotes": ["FAN FURNISHED WITH VFD."], "flag": "orange"}],
        "misc_responsibilities": ["Coordinate the CO/NO2 sensor mounting heights with the architect"],
        "exclusions": ["Line voltage power wiring to the fan by the Electrical Contractor"],
        "qc": [{"check": "QC-01", "evidence": "CC term: BMS contractor", "pass": True, "fix": ""},
               {"check": "QC-04", "evidence": "Garage Exhaust Fan: Start/Stop and Status yes, VFD Control no",
                "pass": False, "fix": "VFD Control not written: the schedule gives no VFD"}],
        "client_notes": [], "tech_hours": None, "review_notes": ["NOTE: the fan is on the garage level"],
    }
    answer.update(overrides)
    return answer


GEM_KNOWN = frozenset({"S12", "S50", "S31", "Q1"})


def test_the_gems_answer_is_read_whole():
    scope = parse_scope(gem_answer(), "m", STYLE, GEM_KNOWN)
    assert isinstance(scope, WrittenScope)
    texts = [(line.text, line.depth) for line in scope.lines]
    assert texts[1] == ("Furnish and Install CO/NO2 Sensor (Qty. 9 Total)", 1)
    assert scope.lines[1].quantity == 9 and scope.lines[1].source_id == "[S50]"
    assert scope.lines[2].reference.startswith("Mandatory minimum point")
    # the GEM's diamond line is the Master Template's own line, one level deeper
    assert texts[-1] == ("AME requires mfr. Rep to assist in integration, cost to be incurred by mech contractor", 2)
    assert scope.quick_scan["contractor_name_used"] == "BMS contractor"
    assert scope.explanation.startswith("The garage exhaust fan")
    (c5,) = scope.contradictions
    assert (c5.type, c5.item, c5.quotes, c5.flag) == ("C5", "Garage Exhaust Fan VFD Control",
                                                        ("FAN FURNISHED WITH VFD.",), "orange")
    assert scope.misc_responsibilities == ("Coordinate the CO/NO2 sensor mounting heights with the architect",)
    assert scope.exclusions == ("Line voltage power wiring to the fan by the Electrical Contractor",)
    assert [(q.check, q.passed) for q in scope.qc] == [("QC-01", True), ("QC-04", False)]
    assert scope.qc[1].fix.startswith("VFD Control not written")
    assert scope.refused == ()


def test_a_line_with_no_reference_is_kept_flagged_orange_never_dropped():
    lines = [_gem_line("Furnish and Install Space Temperature Sensor", reference=""),
             _gem_line("Some wiring maybe"),                                  # no house verb: the GEM's words stand
             _gem_line("Furnish Fan Status (4 tech hours)")]                  # the client must never see hours
    scope = parse_scope(gem_answer(lines), "m", STYLE, GEM_KNOWN)
    assert [line.text for line in scope.lines] == ["Furnish and Install Space Temperature Sensor", "Some wiring maybe"]
    sensor = scope.lines[0]
    assert sensor.flag == "orange" and "no reference" in sensor.flag_text.lower()
    assert scope.lines[1].flag == ""
    assert len(scope.refused) == 1 and "tech hours" in scope.refused[0]


def test_a_flag_written_into_the_line_is_taken_out_of_its_text():
    lines = [_gem_line("Furnish and Install Spill Air Damper Actuator (Qty. per Plenum) \U0001F534", flag=""),
             _gem_line("Furnish and Install Generator Remote Annunciator", flag="\U0001F7E1",
                       flag_text="Factory unit, no named party")]
    scope = parse_scope(gem_answer(lines), "m", STYLE, GEM_KNOWN)
    assert scope.lines[0].text == "Furnish and Install Spill Air Damper Actuator (Qty. per Plenum)"
    assert scope.lines[0].flag == "red"
    assert scope.lines[1].flag == "yellow"


def test_the_judge_asks_the_gem_and_reads_its_answer(tmp_path):
    fake = FakeClient(_response(gem_answer()))
    judge = Judge(tmp_path / "cache", model="test-model", client=fake)
    evidence = system_evidence(_garage(), _gem_dossier())
    scope = judge.write_system(_garage(), evidence, (), ["Furnish and Install DDC Controller"], [], "STYLE")
    (call,) = fake.calls
    assert call["system"] == WRITE_PROMPT
    assert set(call["output_config"]["format"]["schema"]["required"]) >= {"quick_scan", "qc", "contradictions"}
    assert scope.lines[1].source_id == "[S50]" and scope.quick_scan and scope.qc


# --- Part 3: assembly follows the GEM ------------------------------------------------------------------------------

import openpyxl                                                                       # noqa: E402

from estimator.evidence.workbook import GEM_TAB, write_evidence                       # noqa: E402
from estimator.price_sheet.write import check_fit                                    # noqa: E402
from estimator.writing.assemble import (TIER_QUANTITY, assemble_job, plan_checks, review_list,  # noqa: E402
                                        shape_header)
from estimator.writing.whole_job import Contradiction, JobPlan, QcCheck, WrittenLine  # noqa: E402
from test_assemble import (JOB_READING, READING, TEMPLATES, TODAY, _dossier, _headers,  # noqa: E402
                           _job_catalogue)


def _gl(text, reference="Job 26-9999", flag="", flag_text="", quote="Job 26-9999", sid="S9", depth=1, quantity=None):
    return WrittenLine(text, depth, quantity, quote, sid, reference, 0.85, reference, flag, flag_text)


def _scope(*lines, **kw):
    return WrittenScope(tuple(lines), (), None, (), "m", **kw)


def _planned(key, header, tags=(), quantity=1, **kw):
    values = dict(key=key, header=header, kind=kw.pop("kind", header.split(" ")[0]), tags=tuple(tags),
                  quantity=quantity, quantity_source="S9", quantity_quote="Job 26-9999", sequence_sources=("S9",),
                  schedule_sources=(), group="", alternate="", exception="", why="")
    values.update(kw)
    return PlannedSystem(**values)


def _assemble(systems, scopes, dossier=None):
    plan = JobPlan(tuple(systems), {}, (), "m")
    return assemble_job(_job_catalogue(), JOB_READING, dossier or _dossier(), plan, scopes, TODAY, TEMPLATES)


def _items(job, header):
    at = next(n for n, b in enumerate(job.blocks) if b.header == header)
    for block in job.blocks[at:]:
        if block.items:
            return [(i.kind, i.text) for i in block.items]
    return []


FANS = _planned("EF", "Exhaust Fan EF-1, 2", ("EF-1", "EF-2"), 2, kind="Exhaust Fan")


def test_the_round_four_and_five_trimming_is_gone():
    lines = (_gl("Furnish EDH Status"), _gl("Furnish Electric Duct Heater Status"),
             _gl("Monitor Fan Status"), _gl("Monitor Fan Fail Alarm"),
             _gl("Interlock with Make-Up Air System (as per sequence)"),
             _gl("Interlock Dryer Exhaust Fan with Make-Up Air Fan"))
    job, systems, general = _assemble([FANS], {"EF": _scope(*lines)})
    texts = [t for _, t in _items(job, "Qty (2) Exhaust Fan EF-1, 2")]
    for text in ("Furnish EDH Status", "Furnish Electric Duct Heater Status", "Monitor Fan Status",
                 "Monitor Fan Fail Alarm", "Interlock with Make-Up Air System",
                 "Interlock Dryer Exhaust Fan with Make-Up Air Fan"):
        assert text in texts, text
    said = " ".join(r.decide for s in systems for r in s.review)
    assert "Repeated lines printed once" not in said and "long block" not in said


def test_an_exact_repeat_is_printed_once():
    job, _, _ = _assemble([FANS], {"EF": _scope(_gl("Furnish Fan Start/Stop and Status"),
                                                _gl("Furnish Fan Start/Stop and Status"))})
    assert [t for _, t in _items(job, "Qty (2) Exhaust Fan EF-1, 2")].count("Furnish Fan Start/Stop and Status") == 1


def test_identical_lines_across_the_units_are_one_line_with_the_total():
    chillers = _planned("CH", "Chiller CH-1, 2, 3", ("CH-1", "CH-2", "CH-3"), 3, kind="Chiller")
    lines = (_gl("Furnish CH-1 Enable"), _gl("Furnish CH-2 Enable"), _gl("Furnish CH-3 Enable"),
             _gl("Furnish and Install Minimum Outdoor Air Damper Actuator"),
             _gl("Furnish and Install Return Air Damper Actuator"),
             _gl("Wire Fire Alarm Stop Relay to Supply Fan VFD (Relay Furnished by FA)"),
             _gl("Wire Fire Alarm Start Relay to DDC Controller (Relay Furnished by FA)"))
    job, _, _ = _assemble([chillers], {"CH": _scope(*lines)})
    texts = [t for _, t in _items(job, "Qty (3) Chiller CH-1, 2, 3")]
    assert "Furnish Chiller Enable (Qty. 3 Total)" in texts
    assert not any("CH-1 Enable" in t or "CH-2 Enable" in t for t in texts)
    for distinct in ("Minimum Outdoor Air Damper Actuator", "Return Air Damper Actuator", "Stop Relay", "Start Relay"):
        assert sum(distinct in t for t in texts) == 1, distinct


def test_the_lines_follow_the_gems_sow_positions_with_integration_last():
    shuffled = (_gl("Map Available Points Via Mfr. Provided BMS Interface"),
                _gl("AME requires mfr. Rep to assist in integration, cost to be incurred by mech contractor", depth=2),
                _gl("Wire Fire Alarm Stop Relay to Supply Fan VFD (Relay Furnished by FA)"),
                _gl("Furnish Supply Fan Start/Stop and Status"),
                _gl("Furnish and Wire Chilled Water Coil Control Valve and Actuator"),
                _gl("Install and Wire Space Thermostat (Furnished by Unit Mfr.)"),
                _gl("Furnish and Install Supply Air Temperature Sensor"),
                _gl("Furnish BMS Communication Wiring to Factory Controls"),
                _gl("Furnish and Wire Filter Differential Pressure Switch"),
                _gl("Furnish and Install Current Switch"),
                _gl("Furnish Outdoor Air Damper Position Feedback"),
                _gl("Furnish and Install Return Air Damper Actuator"),
                _gl("Furnish and Install Freezestat"),
                _gl("Furnish and Install DDC Controller"))
    job, _, _ = _assemble([FANS], {"EF": _scope(*shuffled)})
    texts = [t for _, t in _items(job, "Qty (2) Exhaust Fan EF-1, 2")]
    expected = ["Furnish and Install DDC Controller", "Furnish and Install Supply Air Temperature Sensor",
                "Furnish and Install Freezestat", "Furnish and Install Current Switch",
                "Furnish and Install Return Air Damper Actuator", "Furnish Supply Fan Start/Stop and Status",
                "Furnish Outdoor Air Damper Position Feedback",
                "Wire Fire Alarm Stop Relay to Supply Fan VFD (Relay Furnished by FA)",
                "Furnish and Wire Filter Differential Pressure Switch",
                "Furnish and Wire Chilled Water Coil Control Valve and Actuator",
                "Install and wire Space Thermostat (Furnished by Unit Mfr.)",
                "Furnish BMS Communication Wiring to Factory Controls",
                "Map Available Points Via Mfr. Provided BMS Interface",
                "AME requires mfr. Rep to assist in integration, cost to be incurred by mech contractor"]
    assert texts == expected
    block = next(b for b in job.blocks if b.items)
    assert block.items[-1].depth == 2                    # the rep line, one level deeper


def test_the_position_table_is_in_the_house_style():
    positions = {entry["position"] for entry in json.loads(
        (STYLE_FILE_PATH).read_text(encoding="utf-8"))["sow_positions"]["positions"]}
    assert set(range(1, 20)) <= positions
    assert STYLE.sow_position("Furnish BMS Communication Wiring to VFD Controls") == 19
    assert STYLE.sow_position("Furnish and Install DDC Controller") == 1


from estimator.writing.style import STYLE_FILE as STYLE_FILE_PATH                     # noqa: E402


def test_flags_print_per_the_doctrine():
    lines = (_gl("Furnish and Install Spill Air Damper Actuator", flag="red", flag_text="No count for the plenums"),
             _gl("Furnish Generator Run Status", flag="yellow", flag_text="Generator is a factory unit; no named party"),
             _gl("Furnish and Install Space Temperature Sensor", flag="orange",
                 flag_text="Sensor assumed from the space temperature alarm"),
             _gl("Furnish Fan Start/Stop and Status", flag="green", flag_text="Fan runs continuously"))
    job, systems, general = _assemble([FANS], {"EF": _scope(*lines)})
    items = _items(job, "Qty (2) Exhaust Fan EF-1, 2")
    texts = [t for _, t in items]
    red = next(t for t in texts if "Spill Air Damper Actuator" in t)
    assert red == ("Furnish and Install Spill Air Damper Actuator (Qty. per Mechanical Schedule) "
                   "(*Confirm quantity: Spill Air Damper Actuator)")
    assert "Furnish Generator Run Status" in texts          # yellow: an ownership question, review only, no bracket
    assert ("Furnish and Install Space Temperature Sensor (*Assumed: Sensor assumed from the space temperature alarm)"
            in texts)
    assert not any(kind == "note" and t.startswith("*") for kind, t in items)
    assert not any("Fan runs continuously" in t for t in texts)                      # green: review only
    ranked = review_list(systems, general, _dossier())
    red_row = next(r for r in ranked if "Spill Air Damper Actuator" in r.decide)
    assert red_row.tier == TIER_QUANTITY and red_row.category == "PENDING"
    for words in ("Generator Run Status", "Space Temperature Sensor", "Fan runs continuously"):
        row = next(r for r in ranked if words in f"{r.decide} {r.evidence}")
        assert ranked.index(red_row) < ranked.index(row), words


def test_the_contradiction_scan_and_the_self_qc_reach_the_review_list_and_c5_prints():
    scope = _scope(_gl("Furnish Fan Start/Stop and Status"), _gl("Furnish Fan VFD Control"),
                   contradictions=(Contradiction("C1", "Duct smoke detector", ("by FA", "by BMS"), "yellow"),
                                   Contradiction("C5", "Fan VFD Control", ("fan with VFD",), "orange")),
                   qc=(QcCheck("QC-01", "CC term: BMS contractor", True, ""),
                       QcCheck("QC-08", "FA relays 2 vs 1 Wire line", False, "Added the start relay")))
    red = _gl("Furnish and Install Current Switch", flag="red", flag_text="No count")
    job, systems, general = _assemble([FANS], {"EF": dataclasses.replace(scope, lines=scope.lines + (red,))})
    items = _items(job, "Qty (2) Exhaust Fan EF-1, 2")
    assert ("line", "Furnish Fan VFD Control (*Assumed: Fan VFD Control)") in items
    ranked = review_list(systems, general, _dossier())
    c1 = next(r for r in ranked if "Duct smoke detector" in r.decide)
    qty = next(r for r in ranked if "Current Switch" in r.decide)
    assert ranked.index(qty) < ranked.index(c1)
    assert "C1" in c1.decide and "by FA" in c1.evidence
    assert not any("QC-" in r.decide for r in ranked)            # the Self-QC stays in the GEM QC tab


def test_misc_responsibilities_are_notes_at_the_end_and_exclusions_are_review_rows():
    scope = _scope(_gl("Furnish Fan Start/Stop and Status"),
                   misc_responsibilities=("Demolition of the existing pneumatic tubing by others",
                                          "Program the fan schedule"),
                   exclusions=("Fire alarm relays by the Fire Alarm Contractor",))
    job, systems, general = _assemble([FANS], {"EF": scope})
    items = _items(job, "Qty (2) Exhaust Fan EF-1, 2")
    assert items[-1] == ("note", "**Demolition of the existing pneumatic tubing by others")
    assert not any("Program the fan schedule" in t for _, t in items)
    assert not any("Fire alarm relays" in t for _, t in items) and not job.named_exceptions
    rows = [r for s in systems for r in s.review]
    assert any("Fire alarm relays by the Fire Alarm Contractor" in r.decide and "exception" in r.action.lower()
               for r in rows)


def test_a_unit_the_quick_scan_says_has_no_cc_scope_is_a_named_exception():
    generator = _planned("EDG", "Emergency Generator", (), 1, kind="Generator")
    scope = _scope(quick_scan={"unit_type": "Factory/OEM Controls", "controller": "Factory Controls — No CC DDC "
                               "Controller", "systems_identified": "Emergency generator — not CC scope"},
                   exclusions=("Generator controls by the generator manufacturer",))
    job, systems, _ = _assemble([generator, FANS], {"EDG": scope, "EF": _scope(_gl("Furnish Fan Start/Stop"))})
    assert "Qty (1) Emergency Generator" not in _headers(job)
    assert job.named_exceptions == ("EXCEPTION TO ANY WORK WITH Qty (1) EMERGENCY GENERATOR - GENERATOR CONTROLS BY "
                                    "THE GENERATOR MANUFACTURER",)


def test_a_line_total_from_the_drawings_is_kept():
    garage = _planned("GAR", "Garage Exhaust Fan GXF-C-1", ("GXF-C-1",), 1, kind="Exhaust Fan")
    job, _, _ = _assemble([garage], {"GAR": _scope(_gl("Furnish and Install CO/NO2 Sensor (Qty. 9 Total)",
                                                       quantity=9, quote="Job 26-9999"))})
    assert ("line", "Furnish and Install CO/NO2 Sensor (Qty. 9 Total)") in _items(job, "Qty (1) Garage Exhaust Fan "
                                                                                    "GXF-C-1")
    block = next(b for b in job.blocks if b.items)
    assert block.items[0].quantity == 9


DRYERS = _planned("DXF", "Dryer Exhaust and Make-up Air Fans", ("DXF-1-1", "DXF-1-2", "MUAF-1-1", "MUAF-1-2"), 2,
                  kind="Exhaust Fan", sequence_section="Q2",
                  families=(Family("Dryer Exhaust Fan DXF-1-1, 2", ("DXF-1-1", "DXF-1-2"), 2, "S9", "Job 26-9999"),
                            Family("Make-up Air Fan MUAF-1-1, 2", ("MUAF-1-1", "MUAF-1-2"), None, "", "")))


def test_one_block_per_planned_system_with_one_header_per_family_stacked_as_the_plan_gives_them():
    assert shape_header(DRYERS, 2).headers == ("Qty (2) Dryer Exhaust Fan DXF-1-1, 2",
                                               "Qty (##) Make-up Air Fan MUAF-1-1, 2")
    job, systems, general = _assemble([DRYERS], {"DXF": _scope(_gl("Furnish Dryer Exhaust Fan Status"),
                                                               _gl("Interlock Make-up Air Fan with Dryer Exhaust Fan"))})
    check_fit(job)
    assert _headers(job) == ["Qty (2) Dryer Exhaust Fan DXF-1-1, 2", "Qty (##) Make-up Air Fan MUAF-1-1, 2"]
    assert job.blocks[1].joined and job.blocks[0].items == ()
    assert any("Make-up Air Fan MUAF-1-1, 2" in r.decide and r.category == "PENDING"
               for s in systems for r in s.review)
    [check] = plan_checks(JobPlan((DRYERS,), {}, (), "m"), _dossier(), READING)[0]
    assert len(check.families) == 2


def test_the_evidence_package_has_each_lines_reference_and_flag_and_a_gem_qc_tab(tmp_path):
    scope = _scope(_gl("Furnish Fan Start/Stop and Status", reference="The fans shall run on the BMS schedule"),
                   _gl("Furnish and Install Current Switch", flag="red", flag_text="No count"),
                   quick_scan={"unit_type": "Field-Installed DDC", "contractor_name_used": "BMS contractor"},
                   explanation="The fans run on the BMS schedule.",
                   contradictions=(Contradiction("C6", "Current switch count", ("no count",), "red"),),
                   qc=(QcCheck("QC-04", "Fan: Start/Stop and Status yes", True, ""),))
    _, systems, general = _assemble([FANS], {"EF": scope})
    from test_evidence import _catalogue as evidence_catalogue
    path = write_evidence(tmp_path / "e.xlsx", evidence_catalogue(), systems=systems, review_flags=general)
    book = openpyxl.load_workbook(path)
    lines = [[c.value for c in row] for row in book["Scope lines"].iter_rows()]
    assert lines[0][9:11] == ["Reference line", "Flag"]
    start = next(row for row in lines if row[1] == "Furnish Fan Start/Stop and Status")
    assert start[9:11] == ["The fans shall run on the BMS schedule", None]
    switch = next(row for row in lines if row[1] and row[1].startswith("Furnish and Install Current Switch"))
    assert switch[10] == "red"
    gem = [[c.value for c in row] for row in book[GEM_TAB].iter_rows()]
    flat = " ".join(str(v) for row in gem for v in row if v is not None)
    for words in ("Exhaust Fan EF-1, 2", "Field-Installed DDC", "BMS contractor", "C6", "Current switch count",
                  "QC-04", "The fans run on the BMS schedule."):
        assert words in flat, words


import dataclasses                                                                   # noqa: E402
