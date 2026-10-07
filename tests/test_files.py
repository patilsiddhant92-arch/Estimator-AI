"""Opening a Word or Excel file: a file that isn't one is refused in plain words, never a library's own error."""
import docx
import openpyxl
import pytest

from estimator.files import UnreadableFileError, open_docx, open_workbook


def _docx(path):
    document = docx.Document()
    document.add_paragraph("Scope of work")
    document.save(str(path))
    return path


def _xlsx(path):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Price"
    workbook.save(str(path))
    return path


def test_each_opener_given_the_other_ones_file_refuses_it_in_plain_words(tmp_path):
    # LDG-13: a workbook named .docx and a Word file named .xlsx (both real zip files)
    with pytest.raises(UnreadableFileError, match="real Word file"):
        open_docx(_xlsx(tmp_path / "price.docx"))
    with pytest.raises(UnreadableFileError, match="real Excel file"):
        open_workbook(_docx(tmp_path / "proposal.xlsx"))
    # and given the other's file under its own name
    with pytest.raises(UnreadableFileError):
        open_docx(_xlsx(tmp_path / "price.xlsx"))
    with pytest.raises(UnreadableFileError):
        open_workbook(_docx(tmp_path / "proposal.docx"))


def test_each_opener_still_opens_its_own_file(tmp_path):
    assert open_docx(_docx(tmp_path / "a.docx")).paragraphs[0].text == "Scope of work"
    assert open_workbook(_xlsx(tmp_path / "a.xlsx")).active["A1"].value == "Price"
