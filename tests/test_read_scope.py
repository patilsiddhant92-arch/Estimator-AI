import copy
from collections import Counter

import docx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from conftest import PROPOSAL_TEMPLATE, SENT_PROPOSAL
from estimator.proposal.read import ProposalFormatError, read_proposal
from estimator.text import squash

MISC_MONITORING_DEPTHS = [1, 2, 3, 2, 3, 2, 3, 2, 4, 4, 4, 4, 4, 4, 4, 3, 2, 2, 2, 3, 3, 2, 3, 3, 2, 3, 3, 3, 3, 4, 2]


def test_scope_heading_and_front_end(sent_scope):
    suffix, front_end, *_ = sent_scope
    assert suffix == "PREWAILING WAGE"
    assert (front_end.mode, front_end.existing_system_name) == ("new-system", None)
    assert [item.kind for item in front_end.items] == ["line"] * 7 + ["note"] * 5
    assert front_end.items[0].text == "Furnish Network Level Control Panel"
    assert front_end.items[1].text == "Furnish PC Workstation (School Operator Console )"
    assert front_end.items[5].text == "Furnish 3-D Graphics"
    assert all(run.bold for item in front_end.items for run in item.runs)


def test_blocks(sent_scope):
    blocks = sent_scope[2]
    assert len(blocks) == 45
    assert blocks[0].header == "*Qty (1) BACnet Data Cabinet"
    assert blocks[-1].header == "Provide 2 Year Guarantee for Workmanship and Material – 15970-49"
    items = [item for block in blocks for item in block.items]
    assert len(items) == 203
    assert Counter((i.kind, i.depth) for i in items) == {
        ("line", 1): 137, ("line", 2): 31, ("line", 3): 15, ("line", 4): 8, ("note", 1): 12}


def test_stacked_headers_are_joined_and_run_on_from_each_other(sent_scope):
    """EF-2 / EF-3,8 / EF-5 / EF-6,10 stack with no blank line between them in the sent copy —
    only the first of the group starts a fresh (non-joined) block."""
    blocks = {b.header: b for b in sent_scope[2]}
    assert not blocks["Qty (1) Shower Room Exhaust Fan EF-2"].joined
    assert blocks["Qty (2) Storage Room Exhaust Fan EF-3, 8"].joined
    assert blocks["Qty (1) Water Service Room Exhaust Fan EF-5"].joined
    assert blocks["Qty (2) Locker Room Exhaust Fan EF-6, 10"].joined
    assert blocks["Qty (1) Energy Recovery Ventilator Unit ERV-502"].joined
    assert blocks["Qty (2) Gravity Ventilator GV-502-IN, EX"].joined
    assert blocks["Qty (1) Electric Duct Heater EDH-502"].joined
    assert blocks["Qty (2) Kitchen Hood"].joined


def test_kitchen_exhaust_system_is_an_unnumbered_group_heading(sent_scope):
    """A bold PR2 heading with numbering switched off, carrying 10 hours in the workup — not a
    note of the block before it."""
    blocks = {b.header: b for b in sent_scope[2]}
    heading = blocks["Kitchen Exhaust System"]
    assert not heading.numbered and not heading.joined and heading.items == ()
    assert blocks["Qty (1) Kitchen Exhaust Fan KEF-2"].joined


def test_a_non_bold_unnumbered_pr2_paragraph_is_not_a_group_heading(tmp_path):
    """_is_group_heading's own docstring says the paragraph mark must be bold, like a block
    header's — a same-sized but non-bold mark must be read the way the reader treated an
    unnumbered PR2 paragraph before the group-heading feature existed: as a plain note under
    whatever block is open, not as a heading that starts a new (unnumbered) block of its own."""
    document = docx.Document(str(SENT_PROPOSAL))
    heading = next(p for p in document.paragraphs if squash(p.text) == "Kitchen Exhaust System")
    rPr = heading._p.pPr.find(qn("w:rPr"))
    rPr.remove(rPr.find(qn("w:b")))
    edited = tmp_path / "kitchen-not-bold.docx"
    document.save(str(edited))
    job = read_proposal(edited, PROPOSAL_TEMPLATE)
    assert "Kitchen Exhaust System" not in [b.header for b in job.blocks]
    note = next(item for block in job.blocks for item in block.items if item.text == "Kitchen Exhaust System")
    assert note.kind == "note"


def test_nesting_is_read_from_the_indent(sent_scope):
    misc = next(b for b in sent_scope[2] if b.header == "Misc Monitoring Systems")
    assert [item.depth for item in misc.items] == MISC_MONITORING_DEPTHS
    assert all(item.kind == "line" for item in misc.items)


def test_emphasis_survives(sent_scope):
    cabinet = sent_scope[2][0]
    server = next(i for i in cabinet.items if squash(i.text) == "Furnish and Install (Rack Mounted) BACnet Server")
    assert [run.text.strip() for run in server.runs if run.bold] == ["(Rack Mounted)"]


def test_alternates(sent_scope):
    add, deduct = sent_scope[3], sent_scope[4]
    assert [a.label for a in add] == ["ALT 1: Controls As per Sequence of operation 15985 – 18"]
    assert [b.header for b in add[0].blocks] == ["*Qty (1) Can Wash Kitchen Exhaust Fan EF-9"]
    assert [i.text for i in add[0].blocks[0].items] == [
        "Furnish and Install DDC Controller", "Furnish Fan Start / Stop and Status"]
    assert [a.label for a in deduct] == ["ALT #:"]
    assert [b.header for b in deduct[0].blocks] == ["Qty (1) Shower Room Exhaust Fan EF-2"]
    assert [i.kind for i in deduct[0].blocks[0].items] == ["line", "line", "note"]


def test_the_whole_proposal(sent_job):
    assert sent_job.header.quote_number == "26-0246"
    assert sent_job.scope_heading_suffix == "PREWAILING WAGE"
    assert len(sent_job.blocks) == 45
    assert sent_job.named_exceptions == ("EXCEPTION TO ANY WORK WITH EXHAUST FAN EF-4 (24/7 OPERATION)",)
    assert sent_job.front_end.tech_hours is None and sent_job.blocks[0].items[0].quantity is None


def test_a_table_inside_the_scope_is_refused(tmp_path):
    document = docx.Document(str(SENT_PROPOSAL))
    table = document.add_table(rows=1, cols=1)
    heading = next(p for p in document.paragraphs if squash(p.text).startswith("SCOPE OF WORK"))
    heading._p.addnext(table._tbl)
    edited = tmp_path / "with-table.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="table"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def _hide_runs_in_tracked_insertion(paragraph, runs) -> None:
    """Move the given <w:r> children of the paragraph into a <w:ins> (an unaccepted tracked
    insertion) — text python-docx's ordinary run access can't see, but still on the page."""
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "901")
    ins.set(qn("w:author"), "test")
    for run in runs:
        paragraph._p.remove(run)
        ins.append(run)
    paragraph._p.append(ins)


def test_a_wholly_hidden_item_is_refused(tmp_path):
    """A paragraph whose entire text sits inside a tracked insertion has no visible run at all —
    it must still be caught, not silently dropped for having no ordinary text."""
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "Furnish and Install UPS")
    _hide_runs_in_tracked_insertion(paragraph, paragraph._p.findall(qn("w:r")))
    edited = tmp_path / "wholly-hidden-item.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def test_a_partly_hidden_block_header_is_refused(tmp_path):
    """The hidden-text check must run on block headers (and every other kind of scope paragraph),
    not only on items — a block header read from truncated text would misname the block."""
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "*Qty (1) BACnet Data Cabinet")
    runs = paragraph._p.findall(qn("w:r"))
    _hide_runs_in_tracked_insertion(paragraph, runs[-1:])
    edited = tmp_path / "partly-hidden-header.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError):
        read_proposal(edited, PROPOSAL_TEMPLATE)


@pytest.mark.parametrize("target_text, indent", [
    ("Furnish and Install UPS", Twips(1800 + 360 * 6)),                 # depth 7 under a block; limit 5
    ("Furnish Network Level Control Panel", Twips(1440 + 720 * 2)),     # depth 3 in the front end; limit 2
])
def test_a_line_indented_past_the_depth_limit_is_refused(tmp_path, target_text, indent):
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == target_text)
    paragraph.paragraph_format.left_indent = indent
    edited = tmp_path / "too-deep.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="depth"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


# --- I4: anything between the headings that isn't a plain paragraph --------------------------

@pytest.mark.parametrize("tag, match", [
    ("w:sdt", "content control"),
    ("w:customXml", "custom XML"),
    ("w:altChunk", "embedded content"),
])
def test_a_foreign_body_element_inside_the_scope_is_refused(tmp_path, tag, match):
    document = docx.Document(str(SENT_PROPOSAL))
    heading = next(p for p in document.paragraphs if squash(p.text).startswith("SCOPE OF WORK"))
    heading._p.addnext(OxmlElement(tag))
    edited = tmp_path / f"with-{tag.split(':')[1]}.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match=match):
        read_proposal(edited, PROPOSAL_TEMPLATE)


@pytest.mark.parametrize("tag", ["w:bookmarkStart", "w:bookmarkEnd", "w:commentRangeStart", "w:commentRangeEnd"])
def test_a_bookmark_or_comment_marker_inside_the_scope_is_not_refused(tmp_path, tag):
    document = docx.Document(str(SENT_PROPOSAL))
    heading = next(p for p in document.paragraphs if squash(p.text).startswith("SCOPE OF WORK"))
    marker = OxmlElement(tag)
    marker.set(qn("w:id"), "1")
    heading._p.addnext(marker)
    edited = tmp_path / f"with-{tag.split(':')[1]}.docx"
    document.save(str(edited))
    read_proposal(edited, PROPOSAL_TEMPLATE)   # does not raise


def test_a_foreign_body_element_inside_the_exclusions_is_refused(tmp_path):
    """A content control among the exclusions is named for that region, not misreported as
    sitting inside the scope of work."""
    document = docx.Document(str(SENT_PROPOSAL))
    heading = next(p for p in document.paragraphs if squash(p.text).startswith("Exceptions and Exclusions"))
    heading._p.addnext(OxmlElement("w:sdt"))
    edited = tmp_path / "with-sdt-in-exclusions.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match=r"a content control sits among the exclusions"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


# --- I5: hidden text, and tracked deletions and moves --------------------------------------

def test_hidden_vanish_text_is_refused(tmp_path):
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "Furnish and Install UPS")
    for run in paragraph.runs:
        run.font.hidden = True
    edited = tmp_path / "vanish.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="hidden"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def test_web_hidden_text_is_refused(tmp_path):
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "Furnish and Install UPS")
    for run in paragraph._p.findall(qn("w:r")):
        run.get_or_add_rPr().append(OxmlElement("w:webHidden"))
    edited = tmp_path / "web-hidden.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="hidden"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def test_a_paragraph_with_an_unaccepted_deletion_is_refused(tmp_path):
    """A paragraph that keeps some ordinary, visible text (so it's not simply dropped as blank)
    but also carries an unaccepted tracked deletion — resolve the tracked change first."""
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "Furnish and Install UPS")
    extra = copy.deepcopy(paragraph._p.findall(qn("w:r"))[-1])
    for t in extra.findall(qn("w:t")):
        t.tag = qn("w:delText")
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "903")
    deletion.set(qn("w:author"), "test")
    deletion.append(extra)
    paragraph._p.append(deletion)
    edited = tmp_path / "partial-deletion.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="deletion"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def test_a_paragraph_with_an_unaccepted_move_is_refused(tmp_path):
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "Furnish and Install UPS")
    extra = copy.deepcopy(paragraph._p.findall(qn("w:r"))[-1])
    for t in extra.findall(qn("w:t")):
        t.tag = qn("w:delText")
    moved = OxmlElement("w:moveFrom")
    moved.set(qn("w:id"), "904")
    moved.set(qn("w:author"), "test")
    moved.append(extra)
    paragraph._p.append(moved)
    edited = tmp_path / "moved-from.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="deletion"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def _wrap_in_deletion(paragraph, runs) -> None:
    """Turn the given <w:r> children into a <w:del> (an unaccepted tracked deletion): each run's
    <w:t> becomes <w:delText> — the words a real Word "Reject" would put back — and the runs move
    inside <w:del>, leaving no ordinary run text on the paragraph at all."""
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "905")
    deletion.set(qn("w:author"), "test")
    for run in runs:
        for t in run.findall(qn("w:t")):
            t.tag = qn("w:delText")
        paragraph._p.remove(run)
        deletion.append(run)
    paragraph._p.append(deletion)


def test_a_wholly_deleted_item_is_refused(tmp_path):
    """A paragraph whose entire text sits inside a tracked deletion has no visible text at all —
    it has to be caught by name, not silently dropped as though it were just a blank line."""
    document = docx.Document(str(SENT_PROPOSAL))
    paragraph = next(p for p in document.paragraphs if squash(p.text) == "Furnish and Install UPS")
    _wrap_in_deletion(paragraph, paragraph._p.findall(qn("w:r")))
    edited = tmp_path / "wholly-deleted-item.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match=r"Furnish and Install UPS.*tracked deletion"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def test_a_wholly_deleted_paragraph_in_the_exclusions_is_refused(tmp_path):
    """The same check has to run in the exclusions region too, not only the scope of work."""
    document = docx.Document(str(SENT_PROPOSAL))
    heading = next(p for p in document.paragraphs if squash(p.text).startswith("Exceptions and Exclusions"))
    new_paragraph = copy.deepcopy(heading._p)
    for run in new_paragraph.findall(qn("w:r")):
        run.remove(run.find(qn("w:t")))
        t = OxmlElement("w:t")
        t.text = "A WHOLLY DELETED EXCEPTION"
        run.append(t)
    heading._p.addnext(new_paragraph)
    wrapped = docx.text.paragraph.Paragraph(new_paragraph, heading._parent)
    _wrap_in_deletion(wrapped, new_paragraph.findall(qn("w:r")))
    edited = tmp_path / "wholly-deleted-exclusion.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match=r"A WHOLLY DELETED EXCEPTION.*tracked deletion"):
        read_proposal(edited, PROPOSAL_TEMPLATE)


def test_a_wholly_inserted_exclusion_is_refused(tmp_path):
    """A named exception whose paragraph sits wholly inside a tracked insertion has no text
    python-docx's ordinary run access can see at all — it must still be refused, not silently
    read as no exception."""
    document = docx.Document(str(SENT_PROPOSAL))
    heading = next(p for p in document.paragraphs if squash(p.text).startswith("Exceptions and Exclusions"))
    new_paragraph = copy.deepcopy(heading._p)
    for run in new_paragraph.findall(qn("w:r")):
        run.remove(run.find(qn("w:t")))
        t = OxmlElement("w:t")
        t.text = "A WHOLLY NEW EXCEPTION"
        run.append(t)
    heading._p.addnext(new_paragraph)
    wrapped = docx.text.paragraph.Paragraph(new_paragraph, heading._parent)
    _hide_runs_in_tracked_insertion(wrapped, new_paragraph.findall(qn("w:r")))
    edited = tmp_path / "wholly-inserted-exclusion.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError):
        read_proposal(edited, PROPOSAL_TEMPLATE)
