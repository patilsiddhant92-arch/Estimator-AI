from estimator.price_sheet.rows import align, read_sent_rows, sheet_plan
from estimator.scope import Alternate, Block
from estimator.text import squash, strip_list_marker


def test_the_base_sheet_follows_the_proposal(small_job):
    base = sheet_plan(small_job)[0]
    assert base.title == "Price Sheet 1"
    assert [row.text for row in base.rows if row.text] == [
        "SCOPE OF WORK – PREVAILING WAGE",
        "•  Front End Web Enabled BMS System",
        "o  Furnish Network Level Control Panel",
        "o  Furnish BMS Server",
        "*CONNECTION TO OWNER WAN BY OTHERS",
        "1.  Qty (2) Exhaust Fan EF-1, 2",
        "o  Furnish and Install DDC Controller",
        "➢  Mfr. to Provide Relay",
        "*DAMPERS BY OTHERS",
        "2.  Qty (1) Air Handling Unit AHU-1",
        "o  Furnish and Install Space Temperature Sensor",
        "o  Furnish and Wire Filter DP Switch",
        "o  Furnish and Install (Duct Mounted) Static Pressure Sensor",
    ]


def test_hours_sit_on_headings_and_quantities_on_lines(small_job):
    rows = {row.text: row for row in sheet_plan(small_job)[0].rows if row.text}
    assert rows["SCOPE OF WORK – PREVAILING WAGE"].hours == 40.0
    assert rows["1.  Qty (2) Exhaust Fan EF-1, 2"].hours == 8.0
    assert rows["o  Furnish and Install DDC Controller"].quantity == 2
    assert rows["o  Furnish and Install Space Temperature Sensor"].ref == ("blocks", 1, "items", 0)


def test_alternates_get_their_own_sheet_and_numbering_runs_on(small_job):
    alternate = sheet_plan(small_job)[1]
    assert alternate.title == "ADD ALT"
    assert [row.text for row in alternate.rows if row.text] == [
        "ADD ALTERNATES:", "1.  ALT 1: Controls for EF-3", "3.  Qty (1) Exhaust Fan EF-3",
        "o  Furnish and Install DDC Controller"]


def test_a_joined_block_has_no_blank_row_and_keeps_its_own_number(small_job):
    extra = Block(header="Qty (1) Exhaust Fan EF-9", joined=True)
    job = small_job.model_copy(update={"blocks": small_job.blocks + (extra,)})
    all_rows = sheet_plan(job)[0].rows
    index = next(i for i, row in enumerate(all_rows) if row.text == "3.  Qty (1) Exhaust Fan EF-9")
    assert all_rows[index - 1].text != ""   # no blank row was inserted before it


def test_an_unnumbered_group_heading_has_no_number_and_does_not_use_one_up(small_job):
    heading = Block(header="Kitchen Exhaust System", numbered=False, tech_hours=10.0)
    after = Block(header="Qty (1) Exhaust Fan EF-9")
    job = small_job.model_copy(update={"blocks": small_job.blocks + (heading, after)})
    plan = sheet_plan(job)[0]
    rows = {row.text: row for row in plan.rows if row.text}
    assert rows["Kitchen Exhaust System"].hours == 10.0
    assert "3.  Kitchen Exhaust System" not in rows
    assert "3.  Qty (1) Exhaust Fan EF-9" in rows   # the unnumbered heading didn't use up a number


def test_several_alternates_are_numbered_sheets(small_job):
    second = Alternate(label="ALT 2: Controls for EF-4", blocks=small_job.add_alternates[0].blocks)
    job = small_job.model_copy(update={"add_alternates": small_job.add_alternates + (second,)})
    assert [s.title for s in sheet_plan(job)] == ["Price Sheet 1", "ADD ALT 1", "ADD ALT 2"]


def test_reading_what_was_sent(sent_workup):
    sheets = read_sent_rows(sent_workup)
    assert {"Price Sheet 1", "ADD ALT", "DEDUCT ALT"} <= set(sheets)
    first = sheets["Price Sheet 1"][0]
    assert (first.row, squash(first.text), first.hours) == (6, "SCOPE OF WORK – PREWAILING WAGE", 140.0)
    ups = next(r for r in sheets["Price Sheet 1"] if strip_list_marker(r.text) == "Furnish and Install UPS")
    assert (ups.row, ups.quantity, ups.unit_price) == (28, 1, 350.0)
    router = next(r for r in sheets["Price Sheet 1"] if "BACnet IP Router" in r.text)
    assert (router.quantity, router.unit_price) == (7, None)
    assert [strip_list_marker(r.text) for r in sheets["ADD ALT"]][:3] == [
        "ADD ALTERNATES:", "ALT 1: Controls As per Sequence of operation 15985 – 18",
        "*Qty (1) Can Wash Kitchen Exhaust Fan EF-9"]


def test_align_pairs_rows_in_order():
    assert align(["a", "b", "c"], ["x", "a", "c"]) == [(0, 1), (2, 2)]


def test_align_ignores_markers_spacing_and_case():
    assert align(["o  Furnish UPS", "44.  *Qty (1) EF-9"], ["o\xa0\xa0 furnish ups", "45.\xa0\xa0 *qty (1) ef-9 "]) == [(0, 0), (1, 1)]


def test_a_line_whose_quantity_is_still_to_fill_in_a_counted_block_is_flagged(small_job, workup_template,
                                                                              price_book, tmp_path):
    """RND-2: a "(Qty. ##)" line under a block with a known count is not priced, and says so."""
    import openpyxl
    from estimator.price_sheet.write import QUANTITY_PENDING, write_price_sheet
    from estimator.scope import Item, Run
    block = Block(header="Qty (2) Exhaust Fan EF-9", items=(
        Item(kind="line", runs=(Run(text="Furnish and Install Space Temperature Sensor (Qty. ##)"),)),))
    job = small_job.model_copy(update={"blocks": small_job.blocks + (block,)})
    wb = openpyxl.load_workbook(str(write_price_sheet(job, workup_template, price_book, tmp_path / "line.xlsx")))
    ws = wb["Price Sheet 1"]
    row = next(r for r in range(6, 60) if "(Qty. ##)" in (ws.cell(r, 4).value or ""))
    assert ws.cell(row, 8).value == QUANTITY_PENDING
    assert (ws.cell(row, 3).value, ws.cell(row, 5).value) == (None, None)
    assert QUANTITY_PENDING in wb["Summary"]["B34"].value
    assert "1 line has no quantity yet" in wb["Summary"]["B34"].value
