import datetime as dt

import docx
import pytest

from conftest import PROPOSAL_TEMPLATE, SENT_PROPOSAL
from estimator.proposal.read import ProposalFormatError, read_cover, read_named_exceptions
from estimator.scope import Header
from estimator.text import squash


def test_cover_page_of_what_was_sent(sent_document):
    assert read_cover(sent_document) == Header(
        to="All Bidders",
        date=dt.date(2026, 6, 8),
        quote_number="26-0246",
        reference="PS Q517, 1 Beach 39th St, Queens NY - New Bldg",
        mep_engineer="SCA",
        based_on=("M001.00 Thru M505.00 Dated: 03/31/2026",),
    )


def test_the_template_cover_is_placeholders(template_document):
    header = read_cover(template_document)
    assert (header.quote_number, header.reference, header.mep_engineer) == ("26-#####", "########", "#######")
    assert header.date == dt.date(2026, 3, 6)
    assert header.based_on == ("########",)


def test_named_exceptions_of_what_was_sent(sent_document):
    assert read_named_exceptions(sent_document, PROPOSAL_TEMPLATE) == (
        "EXCEPTION TO ANY WORK WITH EXHAUST FAN EF-4 (24/7 OPERATION)",
    )


def test_the_template_itself_has_none(template_document):
    assert read_named_exceptions(template_document, PROPOSAL_TEMPLATE) == ()


def test_an_edited_standard_exclusion_is_refused(tmp_path):
    document = docx.Document(str(SENT_PROPOSAL))
    target = next(p for p in document.paragraphs if squash(p.text) == "All gas meters.")
    target.runs[0].text = "All gas meters, except the one on the roof."
    edited = tmp_path / "edited.docx"
    document.save(str(edited))
    with pytest.raises(ProposalFormatError, match="standard exclusions"):
        read_named_exceptions(docx.Document(str(edited)), PROPOSAL_TEMPLATE)
