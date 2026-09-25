"""Write each system from the Master Template: lines, quantities, tech hours, confidence and review."""
import datetime as dt
import re
from pathlib import Path
from types import SimpleNamespace

from estimator.catalogue.overview import SpecOverview
from estimator.price_sheet.write import check_fit
from estimator.reading.devices import DeviceCount
from estimator.reading.inventory import InventoryItem
from estimator.reading.keynotes import Note
from estimator.reading.schedules import Schedule, ScheduleRow
from estimator.scope import Job
import pytest

from estimator.writing.judge import Choice, Composition, Judgement, TechHoursHint, WrittenLine
from estimator.writing.systems import BundleRow, Family, Passage, SystemBundle
from estimator.writing.templates import integration_trio, load_templates
from estimator.writing.write import build_job, tag_list, write_systems

TEMPLATES = load_templates()
DWG = "Drawings/Mech.pdf"
CATEGORY_WORDS = ("CONFIRMED", "DISCREPANCY", "ADDED", "EXTRA", "MISSING", "ORPHANED", "CONFLICT", "EXISTING",
                  "ASSUMPTION", "PENDING")


def _catalogue(overview=None):
    return SimpleNamespace(overview=overview or SpecOverview(), pages=(), sheets=(), sections=(), tags=())


def _rows(tags, headers, cells, title="VAV BOX SCHEDULE", sheet="M-601"):
    return tuple(BundleRow(t, title, sheet, DWG, 5, tuple(headers), {"TAG": t, **cells}, None, None) for t in tags)


def _bundle(prefix, kind, tags, rows=(), passages=(), controls_by=(), notes=(), devices=(), flags=(),
            quantity="count", schedule_notes=()):
    n = len(tags) if quantity == "count" else quantity
    family = Family(prefix, kind, tuple(tags), n, {"CONFIRMED": len(tags)})
    return SystemBundle(key=prefix, name=kind, families=(family,), schedule_rows=tuple(rows),
                        schedule_notes=tuple(schedule_notes), controls_by=tuple(controls_by), notes=tuple(notes),
                        passages=tuple(passages), devices=tuple(devices), sheets=("M-601", "M-101"),
                        flags=tuple(flags))


SEQUENCE = Passage("23 09 93 SEQUENCE OF OPERATIONS, page 44", "Specs.pdf", 44,
                   "VAV boxes: each VAV box shall maintain the space temperature and CO2 setpoints.")
VAV_TAGS = ["VAV-1", "VAV-2", "VAV-3"]


def _vav(**kw):
    rows = kw.pop("rows", _rows(VAV_TAGS, ("TAG", "MAX CFM", "HW GPM"), {"MAX CFM": "600", "HW GPM": "1.5"}))
    return _bundle("VAV", "Variable Air Volume Box", VAV_TAGS, rows=rows, passages=kw.pop("passages", (SEQUENCE,)),
                   **kw)


def _write(bundle, judge=None, catalogue=None, reading=None):
    [system] = write_systems(catalogue or _catalogue(), reading or SimpleNamespace(), [bundle], TEMPLATES, judge)
    return system


def _line(system, text):
    return next(line for line in system.lines if line.text == text)


def _printed(system):
    return [item.text for block in system.blocks for item in block.items]


# --- the VAV template ---------------------------------------------------------------------------

def test_a_hot_water_reheat_vav_with_a_co2_sequence():
    system = _write(_vav())
    assert system.printed
    [block] = system.blocks
    assert block.header == "Qty (3) Variable Air Volume Box VAV-1 thru 3"
    assert _printed(system)[:5] == [
        "Furnish DDC Controller to the Factory for Installation",
        "Furnish and Install Space Temperature / CO2 Sensor",
        "Furnish and Install Discharge Air Temperature Sensor",
        "Furnish and Wire Hot Water Control Valve and Actuator",
        "Furnish and Install VAV Transformer",
    ]
    assert all(item.quantity == 3 for item in block.items)
    assert block.tech_hours == 12          # 4 h per hot-water reheat VAV

    sensor = _line(system, "Furnish and Install Space Temperature / CO2 Sensor")
    assert sensor.category == "CONFIRMED" and sensor.client
    assert "23 09 93 SEQUENCE OF OPERATIONS, page 44" in sensor.evidence
    assert "CO2" in sensor.evidence
    assert sensor.source.startswith("rule: space_co2")

    valve = _line(system, "Furnish and Wire Hot Water Control Valve and Actuator")
    assert valve.category == "CONFIRMED" and valve.confidence == 0.8
    assert "M-601" in valve.evidence and "HW GPM" in valve.evidence

    ddc = _line(system, "Furnish DDC Controller to the Factory for Installation")
    assert (ddc.category, ddc.confidence, ddc.source) == ("ASSUMPTION", 0.5, "default")
    assert system.confidence == 0.5
    assert "Furnish Fan Enable" not in _printed(system)          # not fan powered: the schedule says nothing of it


def test_a_cooling_only_vav_is_three_hours_and_has_no_valve():
    rows = _rows(VAV_TAGS, ("TAG", "MAX CFM"), {"MAX CFM": "600"})
    system = _write(_vav(rows=rows))
    assert "Furnish and Wire Hot Water Control Valve and Actuator" not in _printed(system)
    assert "Furnish and Install Discharge Air Temperature Sensor" not in _printed(system)
    assert system.blocks[0].tech_hours == 9


def test_a_gpm_column_left_empty_does_not_make_a_reheat_vav():
    rows = _rows(VAV_TAGS, ("TAG", "HW GPM"), {"HW GPM": "-"})
    assert "Furnish and Wire Hot Water Control Valve and Actuator" not in _printed(_write(_vav(rows=rows)))


def test_with_no_sequence_the_sensor_is_the_default_and_the_sequence_only_lines_are_pending():
    system = _write(_vav(passages=()))
    sensor = _line(system, "Furnish and Install Space Temperature Sensor")
    assert sensor.category == "ASSUMPTION" and sensor.client
    # 24 V by the BMS can't be checked with no sequence and no notes: the choice falls to the default.
    assert _line(system, "Furnish and Install VAV Transformer").category == "ASSUMPTION"
    assert not any(line.category == "PENDING" and line.client for line in system.lines)


def test_a_line_that_cannot_be_checked_is_pending_and_not_printed():
    # The AHU's supply static line looks only in the sequence; with no sequence it can't be decided.
    rows = _rows(["AHU-1"], ("TAG", "CFM"), {"CFM": "5000"}, title="AIR HANDLING UNIT SCHEDULE")
    system = _write(_bundle("AHU", "Air Handling Unit", ["AHU-1"], rows=rows))
    line = _line(system, "Furnish and Install Supply Static Pressure Sensor")
    assert (line.category, line.client, line.quantity, line.confidence) == ("PENDING", False, 1, 0.0)
    assert "judgement pending" in line.evidence
    assert "Furnish and Install Supply Static Pressure Sensor" not in _printed(system)
    assert any(item.category == "PENDING" and "Supply Static Pressure Sensor" in item.decide for item in system.review)


def test_the_air_handler_is_two_hours_a_hardware_point():
    rows = _rows(["AHU-1", "AHU-2"], ("TAG", "CFM"), {"CFM": "5000"}, title="AIR HANDLING UNIT SCHEDULE")
    passage = Passage("23 09 93, page 50", "Specs.pdf", 50,
                      "AHU-1 and AHU-2: the supply fan shall run to hold the duct static pressure. "
                      "The mixed air temperature is monitored.")
    system = _write(_bundle("AHU", "Air Handling Unit", ["AHU-1", "AHU-2"], rows=rows, passages=(passage,)))
    printed = _printed(system)
    assert "Furnish and Install Supply Static Pressure Sensor" in printed
    assert "Furnish and Install Mixed Air Temperature Sensor" in printed
    points = [line for line in system.lines if line.client and line.text in printed and "DDC Controller" not in line.text
              and line.text.startswith(("Furnish and Install", "Furnish and Wire", "Furnish "))]
    assert system.blocks[0].tech_hours == 2 * sum(line.quantity for line in points)


# --- factory controls -----------------------------------------------------------------------------

def test_a_packaged_rtu_with_a_factory_control_package_gets_the_trio_and_the_install_and_wire_lines():
    rows = _rows(["RTU-1", "RTU-2"], ("TAG", "CFM"), {"CFM": "4000"}, title="ROOFTOP UNIT SCHEDULE")
    system = _write(_bundle("RTU", "Rooftop Unit", ["RTU-1", "RTU-2"], rows=rows,
                            controls_by=("UNIT SHALL BE PROVIDED WITH A FACTORY CONTROL PACKAGE.",)))
    printed = _printed(system)
    first, second, _ = integration_trio("")
    assert first in printed and second in printed
    assert "Install and wire mfr. Provided sensors shipped loose w/ unit" in printed
    trio = _line(system, first)
    assert trio.category == "CONFIRMED" and "FACTORY CONTROL PACKAGE" in trio.evidence
    assert system.blocks[0].tech_hours == 32           # 16 h a unit, from the estimator
    assert not any("No tech hours" in str(item) for item in system.review)


def test_factory_controls_replace_the_ddc_controller_lines():
    system = _write(_vav(controls_by=("VAV CONTROLS BY UNIT MANUFACTURER.",)))
    printed = _printed(system)
    assert "Furnish DDC Controller to the Factory for Installation" not in printed
    first, second, third = integration_trio("Variable Air Volume Box")
    assert printed[:3] == [first, second, third]
    depths = [item.depth for item in system.blocks[0].items[:3]]
    assert depths == [1, 1, 2]


# --- quantities ------------------------------------------------------------------------------------

def test_a_counted_line_takes_the_device_count_on_the_systems_sheets():
    rows = _rows(["SPF-1"], ("TAG", "CFM"), {"CFM": "9000"}, title="FAN SCHEDULE")
    devices = (DeviceCount("M-101", DWG, 3, "Static pressure sensor", "SP", 4, 1, 4, True),)
    system = _write(_bundle("SPF", "Stair Pressurization Fan", ["SPF-1"], rows=rows, devices=devices))
    sensor = _line(system, "Furnish and Install Stairwell Pressure Sensor")
    assert (sensor.quantity, sensor.category, sensor.confidence) == (4, "CONFIRMED", 0.9)


def test_a_counted_line_with_no_count_falls_back_to_one_per_unit_as_an_assumption():
    rows = _rows(["SPF-1", "SPF-2"], ("TAG", "CFM"), {"CFM": "9000"}, title="FAN SCHEDULE")
    system = _write(_bundle("SPF", "Stair Pressurization Fan", ["SPF-1", "SPF-2"], rows=rows))
    sensor = _line(system, "Furnish and Install Stairwell Pressure Sensor")
    assert (sensor.quantity, sensor.category) == (2, "ASSUMPTION")


def test_a_system_whose_units_could_not_be_counted_is_not_printed():
    system = _write(_vav(quantity=None))
    assert not system.printed
    assert all(line.category == "PENDING" and not line.client for line in system.lines)
    assert any("could not be counted" in str(item) for item in system.review)


# --- the AI judge ----------------------------------------------------------------------------------

class FakeJudge:
    available = True
    last_error = None

    def __init__(self, choices, notes=(), dropped=()):
        self.choices, self.notes, self.dropped = choices, notes, dropped
        self.asked = []

    def judge(self, system, evidence, options, rules):
        self.asked.append((system, evidence, options, rules))
        return Judgement(tuple(c for c in self.choices if c.line in options), tuple(self.notes), "fake",
                         tuple(self.dropped), extra_lines=tuple(getattr(self, "extra", ())))


def test_the_ai_overrides_an_assumption_and_a_disagreement_makes_a_review_line():
    judge = FakeJudge([
        Choice("Furnish DDC Controller to the Factory for Installation", True, None,
               "The sequence names the controller.", "each VAV box shall maintain", 0.9),
        Choice("Furnish and Wire Hot Water Control Valve and Actuator", False, None,
               "The boxes are cooling only.", "maintain the space temperature and CO2 setpoints", 0.9),
        Choice("Furnish and Install VAV Transformer", False, None, "no evidence", "", 0.3),
    ], notes=["**Space sensors mounted at 48 inches", "**VAV boxes maintain the space temperature"],
        dropped=["Furnish a coffee machine"])
    system = _write(_vav(), judge=judge)
    [(name, evidence, options, rules)] = judge.asked
    assert name == "VAV / Fan VAV"
    assert "23 09 93 SEQUENCE OF OPERATIONS, page 44" in evidence and "HW GPM" in evidence
    assert "Furnish Fan Enable" in options and len(options) == len(set(options))
    assert any("GPM" in rule for rule in rules)

    ddc = _line(system, "Furnish DDC Controller to the Factory for Installation")
    assert (ddc.category, ddc.confidence) == ("CONFIRMED", 0.9)
    assert ddc.source.startswith("AI:")

    valve = _line(system, "Furnish and Wire Hot Water Control Valve and Actuator")
    assert valve.category == "CONFIRMED" and valve.client          # the rule's answer is kept
    assert any("The rule and the AI disagree on Furnish and Wire Hot Water Control Valve and Actuator" in str(item)
               for item in system.review)

    transformer = _line(system, "Furnish and Install VAV Transformer")   # a weak answer changes nothing
    assert transformer.category == "ASSUMPTION"

    notes = [item.text for item in system.blocks[0].items if item.kind == "note"]
    assert notes == ["**VAV boxes maintain the space temperature"]   # only the note the evidence bears out
    assert any("Space sensors mounted at 48 inches" in str(item) for item in system.review)
    assert any("Furnish a coffee machine" in item.evidence for item in system.review)


def test_the_ai_can_decide_a_pending_line_from_a_drawing_note():
    rows = _rows(["AHU-1"], ("TAG", "CFM"), {"CFM": "5000"}, title="AIR HANDLING UNIT SCHEDULE")
    note = Note("M-101", DWG, 3, "KEY NOTES", "4", "PROVIDE DUCT PRESSURE SENSOR FOR AHU-1.", True, "")
    judge = FakeJudge([Choice("Furnish and Install Supply Static Pressure Sensor", True, None,
                              "A key note asks for it.", "PROVIDE DUCT PRESSURE SENSOR FOR AHU-1", 0.8)])
    system = _write(_bundle("AHU", "Air Handling Unit", ["AHU-1"], rows=rows, notes=(note,)), judge=judge)
    line = _line(system, "Furnish and Install Supply Static Pressure Sensor")
    assert (line.category, line.client, line.confidence) == ("ADDED", True, 0.7)
    assert "M-101" in line.evidence


def test_an_ai_that_fails_leaves_the_lines_pending_with_its_reason():
    class Down(FakeJudge):
        def judge(self, *args):
            self.last_error = "the AI service is busy (too many requests); try again later"
            return None
    rows = _rows(["AHU-1"], ("TAG", "CFM"), {"CFM": "5000"}, title="AIR HANDLING UNIT SCHEDULE")
    system = _write(_bundle("AHU", "Air Handling Unit", ["AHU-1"], rows=rows), judge=Down([]))
    line = _line(system, "Furnish and Install Supply Static Pressure Sensor")
    assert line.category == "PENDING" and "too many requests" in line.evidence


def test_with_no_ai_key_the_pending_lines_say_so_once(tmp_path, monkeypatch):
    from estimator.writing.judge import Judge
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    rows = _rows(["AHU-1"], ("TAG", "CFM"), {"CFM": "5000"}, title="AIR HANDLING UNIT SCHEDULE")
    system = _write(_bundle("AHU", "Air Handling Unit", ["AHU-1"], rows=rows), judge=Judge(tmp_path))
    line = _line(system, "Furnish and Install Supply Static Pressure Sensor")
    assert line.evidence.endswith("judgement pending — no AI key set")
    assert "judgement pending — judgement pending" not in line.evidence


# --- no template, existing, flags -----------------------------------------------------------------

WAITING = "waiting for the AI to write this system (no Master Template entry)"


def test_with_no_ai_a_missing_template_gives_a_pending_block_that_is_not_printed():
    system = _write(_bundle("GV", "Gravity Ventilator", ["GV-1"]))
    assert not system.printed
    assert system.blocks[0].header == f"Gravity Ventilator — {WAITING}"
    assert any(item.category == "PENDING" and WAITING in item.decide and "Gravity Ventilator" in item.decide
               for item in system.review)
    [line] = system.lines
    assert line.category == "PENDING" and not line.client


# --- writing a system the Master Template has no entry for -----------------------------------------

HEATER_TAGS = ["EDH-1", "EDH-2"]
HEATER_ROWS = _rows(HEATER_TAGS, ("TAG", "KW", "CONTROL"), {"KW": "5", "CONTROL": "SCR"},
                    title="ELECTRIC DUCT HEATER SCHEDULE")
HEATER_SEQUENCE = Passage("23 09 93, page 12", "Specs.pdf", 12,
                          "Electric duct heaters: the heater shall modulate by SCR to hold the discharge air "
                          "temperature. Provide two high limit switches at each heater.")


def _heater(**kw):
    return _bundle("EDH", "Electric Duct Heater", HEATER_TAGS, rows=HEATER_ROWS,
                   passages=kw.pop("passages", (HEATER_SEQUENCE,)), **kw)


def _written(text, quote, confidence=0.9, quantity=None, depth=1):
    return WrittenLine(text, depth, quantity, "The evidence calls for it.", quote, confidence)


def _composition(lines=None, notes=(), hours=TechHoursHint("unit", 2.0, "EUH / CUH / UH"), refused=()):
    lines = lines if lines is not None else [
        _written("Furnish Electric Heat SCR Control (AO)", "modulate by SCR"),
        _written("Furnish and Install Discharge Air Temperature Sensor", "hold the discharge air temperature", 0.6),
        _written("Furnish and Install High Limit Switch", "two high limit switches at each heater", quantity=4),
    ]
    return Composition("Electric Duct Heater", tuple(lines), tuple(notes), hours, "fake", tuple(refused))


class FakeWriter(FakeJudge):
    """A fake judge that also writes systems with no template."""

    def __init__(self, composition, choices=()):
        super().__init__(list(choices))
        self.composition = composition
        self.composed = []

    def compose(self, system, evidence, examples, style, analogs):
        self.composed.append((system, evidence, examples, style, analogs))
        return self.composition


def test_a_system_with_no_template_is_written_from_the_evidence_as_added():
    judge = FakeWriter(_composition(notes=["**Heaters modulate by SCR to hold the discharge air temperature",
                                           "**Heaters mounted in the ceiling"]))
    system = _write(_heater(), judge=judge)
    [(name, evidence, examples, style, analogs)] = judge.composed
    assert name == "Electric Duct Heater"
    assert "ELECTRIC DUCT HEATER SCHEDULE" in evidence and "modulate by SCR" in evidence and "EDH-1" in evidence
    assert [a.name for a in analogs][0] == "EUH / CUH / UH" and len(analogs) <= 3
    assert examples and len(examples) <= 4 and all("EDH" in e["prefixes"] or e["kind_words"] for e in examples)
    assert "Furnish and Install" in style.verbs

    assert system.printed
    [block] = system.blocks
    assert block.header == "Qty (2) Electric Duct Heater EDH-1, 2"
    assert [(i.text, i.quantity) for i in block.items if i.kind == "line"] == [
        ("Furnish Electric Heat SCR Control (AO)", 2),
        ("Furnish and Install Discharge Air Temperature Sensor", 2),
        ("Furnish and Install High Limit Switch", 4),                 # the count the evidence states
    ]
    for line in (l for l in system.lines if l.text.startswith("Furnish")):
        assert line.category == "ADDED" and line.client and line.confidence <= 0.7
        assert line.source == "written from the evidence (like EUH / CUH / UH)"
        assert "23 09 93, page 12" in line.evidence
    assert _line(system, "Furnish and Install Discharge Air Temperature Sensor").confidence == 0.6

    assert block.tech_hours == 4 and system.tech_hours == 4               # 2 h a unit, like the unit heaters
    review = [str(item) for item in system.review]
    assert review.count("Electric Duct Heater: New system written from the evidence — check the wording") == 1
    assert "Electric Duct Heater: Tech hours taken from EUH / CUH / UH — confirm" in review
    assert any(i.category == "ASSUMPTION" and i.decide == "Tech hours taken from EUH / CUH / UH"
               for i in system.review)

    notes = [i.text for i in block.items if i.kind == "note"]
    assert notes == ["**Heaters modulate by SCR to hold the discharge air temperature"]   # the evidence bears it out
    assert any("Heaters mounted in the ceiling" in str(i) for i in system.review)


def test_tech_hours_like_no_real_template_are_left_for_the_estimator():
    system = _write(_heater(), judge=FakeWriter(_composition(hours=TechHoursHint("unit", 2.0, "Space Heaters"))))
    assert system.printed and system.tech_hours is None and system.blocks[0].tech_hours is None
    assert any("No tech hours" in item.decide and "Electric Duct Heater" in item.decide for item in system.review)
    assert all(line.source == "written from the evidence (like EUH / CUH / UH)"
               for line in system.lines if line.text.startswith("Furnish"))     # the nearest template


def test_a_stated_quantity_without_its_quote_in_the_job_is_not_trusted():
    lines = [_written("Furnish and Install High Limit Switch", "six limit switches per heater", quantity=12)]
    system = _write(_heater(), judge=FakeWriter(_composition(lines=lines)))
    line = _line(system, "Furnish and Install High Limit Switch")
    assert line.quantity == 2                         # one a unit: the quote is not in the documents
    assert line.confidence <= 0.5 and "not found" in line.evidence


def test_the_ais_refused_lines_reach_the_review_list():
    refused = ["“Occupancy Sensor” — it does not start with a house verb"]
    system = _write(_heater(), judge=FakeWriter(_composition(refused=refused)))
    assert any("break the house style" in item.decide and "Occupancy Sensor" in item.evidence
               for item in system.review)


def test_an_ai_that_cannot_write_it_leaves_the_pending_block_with_its_reason():
    class Waiting(FakeWriter):
        def compose(self, *args):
            self.last_error = "judgement pending — waiting for an answer in the chat (requests folder: X)"
            return None
    system = _write(_heater(), judge=Waiting(None))
    assert not system.printed
    assert system.blocks[0].header == f"Electric Duct Heater — {WAITING}"
    assert any("waiting for an answer in the chat" in item.evidence for item in system.review)


def test_a_composition_with_no_usable_line_is_not_printed():
    system = _write(_heater(), judge=FakeWriter(_composition(lines=[])))
    assert not system.printed
    assert any("no line" in item.evidence for item in system.review if item.category == "PENDING")


def test_a_judge_that_cannot_compose_still_leaves_the_pending_block():
    system = _write(_heater(), judge=FakeJudge([]))           # an older judge: choose only
    assert not system.printed and WAITING in system.blocks[0].header


def test_the_chat_writes_a_compose_request_and_picks_up_its_answer(tmp_path, monkeypatch):
    import json
    from estimator.writing.judge import Judge
    monkeypatch.setattr("anthropic.Anthropic", lambda *a, **k: pytest.fail("the chat never calls the API"))
    system = _write(_heater(), judge=Judge(tmp_path, mode="chat"))
    assert not system.printed
    [request_file] = (tmp_path / "judge-chat" / "requests").glob("*.json")
    request = json.loads(request_file.read_text(encoding="utf-8"))
    assert request["kind"] == "compose" and request["system"] == "Electric Duct Heater"
    assert request["analogs"][0] == "EUH / CUH / UH"
    assert set(request["schema"]["required"]) == {"header", "lines", "client_notes", "tech_hours"}
    answer = {"header": "Electric Duct Heater",
              "lines": [{"text": "Furnish Electric Heat SCR Control (AO)", "depth": 1, "quantity": None,
                         "why": "The sequence modulates by SCR.", "quote": "modulate by SCR", "confidence": 0.9}],
              "client_notes": [], "tech_hours": {"per": "unit", "hours": 2, "like": "EUH / CUH / UH"}}
    answers = tmp_path / "judge-chat" / "answers"
    answers.mkdir(parents=True)
    (answers / request_file.name).write_text(json.dumps(answer), encoding="utf-8")

    system = _write(_heater(), judge=Judge(tmp_path, mode="chat"))
    assert system.printed and system.blocks[0].header == "Qty (2) Electric Duct Heater EDH-1, 2"
    line = _line(system, "Furnish Electric Heat SCR Control (AO)")
    assert (line.category, line.quantity, line.confidence) == ("ADDED", 2, 0.7)
    assert system.tech_hours == 4


# --- lines beyond the Master Template on a templated system ---------------------------------------------

OCC_DEVICES = (DeviceCount("M-101", DWG, 3, "Occupancy sensor", "OCC", 5, 1, 5, True),)


def test_an_extra_line_on_a_templated_system_is_added_and_reviewed():
    judge = FakeJudge([])
    judge.extra = (WrittenLine("Furnish and Install Occupancy Sensor", 1, None,
                               "The plans show occupancy sensors in the offices.", "OCC", 0.9),)
    system = _write(_vav(devices=OCC_DEVICES), judge=judge)
    [(_, evidence, _, _)] = judge.asked
    assert "OCC" in evidence and "Occupancy sensor" in evidence          # the device counts reach the AI
    line = _line(system, "Furnish and Install Occupancy Sensor")
    assert (line.category, line.client, line.quantity) == ("ADDED", True, 5)   # counted on the plans
    assert line.confidence <= 0.7 and line.source.startswith("AI:")
    assert "M-101" in line.evidence
    assert "Furnish and Install Occupancy Sensor" in _printed(system)
    assert any(str(item) == "Variable Air Volume Box: Line added beyond the Master Template — check"
               and "Occupancy Sensor" in item.evidence for item in system.review)


def test_an_extra_line_with_no_device_count_is_one_a_unit():
    passage = Passage("23 09 93, page 45", "Specs.pdf", 45, "Each VAV box zone shall have an occupancy sensor.")
    judge = FakeJudge([])
    judge.extra = (WrittenLine("Furnish and Install Occupancy Sensor", 1, None, "The sequence asks for one.",
                               "shall have an occupancy sensor", 0.9),)
    system = _write(_vav(passages=(SEQUENCE, passage)), judge=judge)
    line = _line(system, "Furnish and Install Occupancy Sensor")
    assert (line.quantity, line.category, line.confidence) == (3, "ADDED", 0.7)


def test_an_extra_line_whose_quote_is_not_in_the_job_is_only_reviewed():
    judge = FakeJudge([])
    judge.extra = (WrittenLine("Furnish and Install Occupancy Sensor", 1, None, "made up", "OCC-7 on M-999", 0.9),)
    system = _write(_vav(), judge=judge)
    assert "Furnish and Install Occupancy Sensor" not in _printed(system)
    assert any("Occupancy Sensor" in item.decide and "not in the job's documents" in item.decide
               for item in system.review)


def test_existing_units_are_not_written_as_new_work():
    family = Family("VAV", "Variable Air Volume Box", ("VAV-9",), 1, {"EXISTING": 1})
    bundle = SystemBundle("Existing VAV", "Existing Variable Air Volume Box", (family,), (), (), (), (), (), (),
                          (), ("1 VAV tag marked existing on the schedule (VAV-9) — not new work; confirm any "
                               "removal or reuse scope",), existing=True)
    system = _write(bundle)
    assert not system.printed
    assert any("VAV-9" in str(item) for item in system.review)


def test_bundle_flags_reach_the_review_list_with_their_category():
    flag = "2 more VAV tags on the drawings but not on the schedule (EXTRA) — not priced until confirmed"
    system = _write(_vav(flags=(flag,)))
    [item] = [i for i in system.review if "on the drawings but not on the schedule" in i.decide]
    assert item.category == "EXTRA" and item.action == "not priced until confirmed"


def test_nothing_internal_reaches_the_client_blocks():
    judge = FakeJudge([Choice("Furnish DDC Controller to the Factory for Installation", True, None,
                              "The sequence names it.", "each VAV box shall maintain", 0.9)])
    judge.extra = (WrittenLine("Furnish and Install Occupancy Sensor", 1, None, "The plans show it.", "OCC", 0.9),)
    composed = FakeWriter(_composition(notes=["**Heaters modulate by SCR to hold the discharge air temperature"]))
    systems = [_write(_vav(devices=OCC_DEVICES), judge=judge), _write(_vav(passages=())),
               _write(_bundle("RTU", "Rooftop Unit", ["RTU-1"], controls_by=("FACTORY CONTROL PACKAGE",))),
               _write(_heater(), judge=composed)]
    assert all(system.printed for system in systems)
    assert "Furnish and Install Occupancy Sensor" in _printed(systems[0])
    hints = [hint for t in TEMPLATES for hint in t.hints]
    for system in systems:
        for block in system.blocks:
            texts = [block.header] + [item.text for item in block.items]
            for text in texts:
                assert not any(re.search(rf"\b{word}\b", text) for word in CATEGORY_WORDS), text
                assert "confidence" not in text.lower() and "rule:" not in text and "AI:" not in text
                assert text not in hints and "tech hour" not in text.lower()


def test_tag_list_is_compact():
    assert tag_list(["VAV-21-01", "VAV-21-02", "VAV-21-03"]) == "VAV-21-01 thru 03"
    assert tag_list(["AHU-1", "AHU-2"]) == "AHU-1, 2"
    assert tag_list(["EF-3", "EF-8", "TEF-1"]) == "EF-3, 8, TEF-1"


# --- the job ---------------------------------------------------------------------------------------

def test_build_job_gives_a_job_that_fits_the_price_sheet():
    schedule = Schedule(DWG, 5, "M-601", "VAV BOX SCHEDULE", ("TAG", "HW GPM"),
                        tuple(ScheduleRow(t, {"TAG": t, "HW GPM": "1.5"}, None, None) for t in VAV_TAGS),
                        False, (), ())
    items = tuple(InventoryItem(t, "Variable Air Volume Box", "CONFIRMED", 1, "M-601 VAV BOX SCHEDULE", ("M-101",),
                                "", "") for t in VAV_TAGS)
    items += (InventoryItem("ZZ-1", "Mystery Box", "EXTRA", None, None, ("M-101",), "", ""),)
    catalogue = SimpleNamespace(job=SimpleNamespace(name="26-9999", root=Path("26-9999, 1 Test Street")),
                                overview=SpecOverview(), pages=(), sheets=(), sections=(), tags=())
    reading = SimpleNamespace(schedules=(schedule,), inventory=items, notes=(), backbone=(), devices=(), versions=(),
                              flags=())
    job, systems, flags = build_job(catalogue, reading, None, dt.date(2026, 9, 25))
    assert isinstance(job, Job)
    check_fit(job)
    assert job.header.reference == "1 Test Street"
    assert [b.header for b in job.blocks] == ["Qty (3) Variable Air Volume Box VAV-1 thru 3"]
    assert systems[0].bundle_key == "Front end"
    assert any("ZZ-1" in flag for flag in flags)            # a family on no schedule reaches the review list


# --- schedule columns and notes beat the template's defaults ------------------------------------------

def _heaters(cells, notes=(), passages=()):
    rows = _rows(["UH-1", "UH-2"], ("TAG", *cells), cells, title="UNIT HEATER SCHEDULE")
    return _bundle("UH", "Unit Heater", ["UH-1", "UH-2"], rows=rows, schedule_notes=notes, passages=passages)


def test_an_scr_control_column_gives_the_scr_line():
    system = _write(_heaters({"KW": "5", "CONTROL": "SCR"}))
    line = _line(system, "Furnish Electric Heat SCR Control (AO)")
    assert (line.category, line.client, line.quantity) == ("CONFIRMED", True, 2)
    assert line.source.startswith("schedule:") and "CONTROL" in line.evidence


def test_scr_in_another_column_is_not_the_control_column():
    system = _write(_heaters({"KW": "5", "REMARKS": "SEE SCR NOTE ON E-101"}))
    assert "Furnish Electric Heat SCR Control (AO)" not in _printed(system)


def test_a_transformer_by_the_unit_manufacturer_replaces_the_vav_transformer():
    system = _write(_vav(schedule_notes=("1. PROVIDE MANUFACTURER'S CONTROL POWER TRANSFORMER.",)))
    printed = _printed(system)
    assert "24V Transformer by Unit Mfr." in printed
    assert "Furnish and Install VAV Transformer" not in printed and "Run 24V Power Loop to VAVs" not in printed
    assert _line(system, "24V Transformer by Unit Mfr.").category == "CONFIRMED"


def test_a_thermostat_by_the_manufacturer_takes_the_install_and_wire_form():
    ddc = Passage("23 09 93, page 50", "Specs.pdf", 50, "Unit heaters UH-1 and UH-2: the DDC controller cycles the fan.")
    system = _write(_heaters({"KW": "5", "REMARKS": "7-DAY PROGRAMMABLE THERMOSTAT FURNISHED BY MANUFACTURER"},
                             passages=(ddc,)))
    printed = _printed(system)
    assert "Install and wire Space Thermostat (Furnished by Unit Mfr.)" in printed
    assert "Furnish and Install DDC Controller or Communicating Thermostat" not in printed
    assert _line(system, "Install and wire Space Thermostat (Furnished by Unit Mfr.)").category == "CONFIRMED"


def test_a_bacnet_gateway_kit_gives_the_integration_trio():
    system = _write(_vav(schedule_notes=("2. PROVIDE BACNET GATEWAY KIT FOR EACH UNIT.",)))
    first, second, third = integration_trio("Variable Air Volume Box")
    assert [first, second, third] == _printed(system)[:3]
    assert "Furnish DDC Controller to the Factory for Installation" not in _printed(system)


def test_leak_detection_by_the_unit_manufacturer_is_a_client_note():
    system = _write(_heaters({"KW": "5"}, notes=("3. REFRIGERANT LEAK DETECTION SENSOR PROVIDED BY UNIT MFR.",)))
    notes = [i.text for i in system.blocks[0].items if i.kind == "note"]
    assert "**LEAK DETECTION BY UNIT MFR." in notes


def test_the_schedule_note_beats_the_sequence_and_the_review_quotes_both():
    rows = _rows(["ERV-1", "ERV-2"], ("TAG", "CFM"), {"CFM": "2000"}, title="ENERGY RECOVERY VENTILATOR SCHEDULE")
    seq = Passage("23 09 93, page 2", "Specs.pdf", 2,
                  "1. The ERV shall be furnished with factory DDC controls for complete stand-alone operation.")
    bundle = _bundle("ERV", "Energy Recovery Ventilator", ["ERV-1", "ERV-2"], rows=rows, passages=(seq,),
                     controls_by=("FURNISH AND INSTALL CONTROLLER FOR EACH ERV.",))
    system = _write(bundle)
    printed = _printed(system)
    assert "Furnish and Install DDC Controller" in printed
    assert not any("Factory Controls" in t or "Map Available Points" in t for t in printed)
    [item] = [i for i in system.review if "disagree" in i.decide]
    assert "FURNISH AND INSTALL CONTROLLER FOR EACH ERV" in item.evidence and "factory DDC controls" in item.evidence


def test_an_ai_answer_does_not_undo_a_schedule_override():
    judge = FakeJudge([Choice("24V Transformer by Unit Mfr.", False, None, "The sequence says nothing.",
                              "each VAV box shall maintain", 0.9)])
    system = _write(_vav(schedule_notes=("1. TRANSFORMER BY UNIT MFR.",)), judge=judge)
    assert "24V Transformer by Unit Mfr." in _printed(system)


# --- a rule fires only on a passage that names the system ---------------------------------------------

def _fans(passages, tags=("EF-1", "EF-2")):
    rows = _rows(list(tags), ("TAG", "CFM"), {"CFM": "800"}, title="FAN SCHEDULE")
    return _bundle("EF", "Exhaust Fan", list(tags), rows=rows, passages=tuple(passages))


def test_a_rule_does_not_fire_on_the_paragraph_after_the_one_naming_the_system():
    passage = Passage("23 09 93, page 5", "Specs.pdf", 5,
                      "EF-1 and EF-2 run continuously.\n\nThe duct static pressure sensor sets the supply fan speed.")
    printed = _printed(_write(_fans([passage])))
    assert "Furnish and Install Duct Static Pressure Sensor" not in printed
    assert "Furnish and Install DDC Controller" not in printed


def test_a_rule_fires_on_its_own_section_even_where_a_paragraph_does_not_name_it():
    passage = Passage("23 09 93, page 5", "Specs.pdf", 5,
                      "C. Fans:\n\n1. The duct static pressure sensor sets the fan speed.", section="Fans")
    assert "Furnish and Install Duct Static Pressure Sensor" in _printed(_write(_fans([passage])))


def test_the_ais_quote_naming_the_units_beats_a_rule_whose_passage_does_not():
    general = Passage("23 09 93, page 5", "Specs.pdf", 5, "Exhaust fans shall hold the space temperature.")
    own = Passage("23 09 93, page 6", "Specs.pdf", 6,
                  "EF-1 and EF-2 shall be started by their line-voltage thermostat; no BMS points.")
    judge = FakeJudge([Choice("Furnish and Install DDC Controller", False, None,
                              "The fans run on their own thermostat.", "EF-1 and EF-2 shall be started by their "
                              "line-voltage thermostat", 0.9)])
    system = _write(_fans([general, own]), judge=judge)
    printed = _printed(system)
    assert "Furnish and Install DDC Controller" not in printed
    assert "Interlock with Associated AHU / RTU /MAU Unit Operation" in printed      # the choice falls back
    [item] = [i for i in system.review if "The AI's answer is kept" in i.decide]
    assert "Furnish and Install DDC Controller" in item.decide and "EF-1 and EF-2 shall be started" in item.evidence


def test_the_rule_is_kept_when_its_own_passage_names_the_units():
    own = Passage("23 09 93, page 5", "Specs.pdf", 5, "EF-1 and EF-2 shall hold the space temperature.")
    judge = FakeJudge([Choice("Furnish and Install DDC Controller", False, None, "No controller.",
                              "EF-1 and EF-2 shall hold the space temperature", 0.9)])
    system = _write(_fans([own]), judge=judge)
    assert "Furnish and Install DDC Controller" in _printed(system)
    assert any(i.category == "CONFLICT" for i in system.review)


# --- a system made from a sequence heading ------------------------------------------------------------

def _heading_job(text):
    from estimator.catalogue.register import PageRecord
    from estimator.reading.sequences import SEQUENCE as SEQ, Backbone
    schedule = Schedule(DWG, 5, "M-601", "VAV BOX SCHEDULE", ("TAG", "HW GPM"),
                        tuple(ScheduleRow(t, {"TAG": t, "HW GPM": "1.5"}, None, None) for t in VAV_TAGS),
                        False, (), ())
    items = tuple(InventoryItem(t, "Variable Air Volume Box", "CONFIRMED", 1, "M-601 VAV BOX SCHEDULE", ("M-101",),
                                "", "") for t in VAV_TAGS)
    catalogue = SimpleNamespace(job=SimpleNamespace(name="26-9999", root=Path("26-9999, 1 Test Street")),
                                overview=SpecOverview(), pages=(PageRecord("Specs.pdf", 44, "text", text, 0),),
                                sheets=(), sections=(), tags=())
    backbone = (Backbone(SEQ, "23 09 93 SEQUENCE OF OPERATIONS (Specs.pdf, page 44)", "Specs.pdf", 44, 44,
                         "spec section"),)
    reading = SimpleNamespace(schedules=(schedule,), inventory=items, notes=(), backbone=backbone, devices=(),
                              versions=(), flags=())
    return catalogue, reading


HEADINGS = ("A. VAV Boxes:\n1. Each VAV box maintains the space temperature.\n"
            "B. Gas Booster:\n1. The BMS shall monitor the gas booster pump status.\n"
            "C. Air Curtains:\n1. The air curtain runs when the door opens.\n")


def test_a_heading_with_no_template_goes_to_the_ai_to_write_and_waits_for_a_count():
    catalogue, reading = _heading_job(HEADINGS)
    judge = FakeWriter(_composition(lines=[_written("Furnish Gas Booster Pump Status (BI)",
                                                    "monitor the gas booster pump status")]))
    job, systems, _ = build_job(catalogue, reading, judge, dt.date(2026, 9, 25))
    by_key = {s.bundle_key: s for s in systems}
    booster = by_key["Gas Booster"]
    assert [c[0] for c in judge.composed] == ["Gas Booster"]
    assert "monitor the gas booster pump status" in judge.composed[0][1]
    assert not booster.printed                                        # no tags named: PENDING, kept out
    assert any(item.category == "PENDING" and "Gas Booster" in item.decide for item in booster.review)
    assert by_key["Air Curtain"].template == "Air Curtain"               # a template matches its name: choose
    assert "Air Curtain" in [asked[0] for asked in judge.asked]
    assert [b.header for b in job.blocks] == ["Qty (3) Variable Air Volume Box VAV-1 thru 3"]


# --- a Sensaphone becomes an exception when the network control panel sends the alarms ---------------

ALARM = "HVAC Alarm Panel w/ Auto Dialer"
ALARM_PASSAGE = Passage("23 09 93 SEQUENCE OF OPERATIONS, page 45", "Specs.pdf", 45,
                        "Alarms: provide a Sensaphone auto dialer to call the owner on a high room temperature alarm.")
AUTO_DIALER_EXCEPTION = "EXCEPTION TO AUTO DIALER (BMS CAN SEND OUT ALARMS VIA EMAIL AND SMS TEXT)"


def _alarm_bundle(passages=(ALARM_PASSAGE,)):
    return _bundle("AP", "HVAC Alarm Panel", ["AP-1"], passages=passages)


def _with_panel():
    from estimator.writing.front_end import front_end
    fe, _, _ = front_end(_catalogue())
    assert "Furnish Network Level Control Panel" in [item.text for item in fe.items]
    return fe


def test_the_exception_is_the_templates_own_line():
    [alarm] = [t for t in TEMPLATES if t.name == ALARM]
    assert alarm.becomes_exception_when is not None
    assert alarm.becomes_exception_when.front_end_line == "Furnish Network Level Control Panel"
    assert alarm.becomes_exception_when.exception == AUTO_DIALER_EXCEPTION
    assert AUTO_DIALER_EXCEPTION in [line.text for line in alarm.lines]


def test_with_the_panel_a_sensaphone_passage_becomes_an_exception_not_a_block():
    from estimator.writing.write import apply_exceptions
    bundle = _alarm_bundle()
    assert _write(bundle).printed                               # a block, were it not for the rule
    to_write, excepted, exceptions = apply_exceptions(_with_panel(), [bundle], _catalogue(), TEMPLATES)

    assert to_write == [] and exceptions == (AUTO_DIALER_EXCEPTION,)
    [alarm] = excepted
    assert alarm.blocks == () and alarm.template == ALARM
    [item] = [i for i in alarm.review if "Sensaphone / auto dialer called for" in i.decide]
    assert item.system == ALARM
    assert item.decide == ("Sensaphone / auto dialer called for (23 09 93 SEQUENCE OF OPERATIONS, page 45) — put in "
                           "the exceptions, because the network control panel sends the alarm emails. Check.")
    assert "Sensaphone auto dialer" in item.evidence
    [line] = alarm.lines
    assert line.text == AUTO_DIALER_EXCEPTION and line.client and "23 09 93" in line.evidence


def test_with_no_network_control_panel_the_alarm_panel_block_is_written():
    from estimator.scope import FrontEnd, Item, Run
    from estimator.writing.write import apply_exceptions
    standalone = FrontEnd(mode="new-system", items=(Item(kind="line", runs=(Run(text="Furnish 3-D Graphics"),)),))
    bundle = _alarm_bundle()
    to_write, excepted, exceptions = apply_exceptions(standalone, [bundle], _catalogue(), TEMPLATES)

    assert (to_write, excepted, exceptions) == ([bundle], [], ())
    alarm = _write(bundle)
    assert alarm.printed and alarm.blocks[0].header == "Qty (1) HVAC Alarm Panel w/ Auto Dialer"
    assert "Furnish and Install Sensaphone" in _printed(alarm)


def test_with_the_panel_but_no_sensaphone_called_for_the_block_stays():
    from estimator.writing.write import apply_exceptions
    bundle = _alarm_bundle(passages=())
    assert apply_exceptions(_with_panel(), [bundle], _catalogue(), TEMPLATES) == ([bundle], [], ())


def _sensaphone_in_the_spec_job():
    from estimator.catalogue.overview import FOUND, Found, Place
    found = Found("Network and front end", "Auto dialer / Sensaphone",
                  (Place("Specs.pdf", 12, None, "23 09 23", "Controls spec"),),
                  "The BAS shall include a Sensaphone for remote alarms.", FOUND)
    schedule = Schedule(DWG, 5, "M-601", "VAV BOX SCHEDULE", ("TAG", "HW GPM"),
                        tuple(ScheduleRow(t, {"TAG": t, "HW GPM": "1.5"}, None, None) for t in VAV_TAGS),
                        False, (), ())
    items = tuple(InventoryItem(t, "Variable Air Volume Box", "CONFIRMED", 1, "M-601 VAV BOX SCHEDULE", ("M-101",),
                                "", "") for t in VAV_TAGS)
    catalogue = SimpleNamespace(job=SimpleNamespace(name="26-9999", root=Path("26-9999, 1 Test Street")),
                                overview=SpecOverview(topics=(found,)), pages=(), sheets=(), sections=(), tags=())
    reading = SimpleNamespace(schedules=(schedule,), inventory=items, notes=(), backbone=(), devices=(), versions=(),
                              flags=())
    return build_job(catalogue, reading, None, dt.date(2026, 9, 25))


def test_the_spec_overview_finding_a_sensaphone_makes_the_exception_in_the_job():
    job, systems, _ = _sensaphone_in_the_spec_job()
    assert job.named_exceptions == (AUTO_DIALER_EXCEPTION,)
    assert not any("Alarm Panel" in block.header for block in job.blocks)
    [item] = [i for s in systems for i in s.review if "Sensaphone / auto dialer called for" in i.decide]
    assert "Specs.pdf · 23 09 23 · page 12" in item.decide
    assert item.evidence == "The BAS shall include a Sensaphone for remote alarms."


def test_the_proposal_prints_the_named_exception(proposal_template, tmp_path):
    from estimator.proposal.write import write_proposal
    from docx import Document
    job, _, _ = _sensaphone_in_the_spec_job()
    out = write_proposal(job, proposal_template, tmp_path / "proposal.docx")
    texts = [p.text.strip() for p in Document(str(out)).paragraphs]
    assert AUTO_DIALER_EXCEPTION in texts


def test_the_proposal_holds_nothing_internal_from_a_system_the_ai_wrote(proposal_template, tmp_path):
    from docx import Document
    from estimator.proposal.write import write_proposal
    vav = Schedule(DWG, 5, "M-601", "VAV BOX SCHEDULE", ("TAG", "HW GPM"),
                   tuple(ScheduleRow(t, {"TAG": t, "HW GPM": "1.5"}, None, None) for t in VAV_TAGS), False, (), ())
    heaters = Schedule(DWG, 5, "M-601", "ELECTRIC DUCT HEATER SCHEDULE", ("TAG", "KW"),
                       tuple(ScheduleRow(t, {"TAG": t, "KW": "5"}, None, None) for t in HEATER_TAGS), False, (), ())
    items = tuple(InventoryItem(t, "Variable Air Volume Box", "CONFIRMED", 1, "M-601 VAV BOX SCHEDULE", ("M-101",),
                                "", "") for t in VAV_TAGS)
    items += tuple(InventoryItem(t, "Electric Duct Heater", "CONFIRMED", 1, "M-601 ELECTRIC DUCT HEATER SCHEDULE",
                                 ("M-101",), "", "") for t in HEATER_TAGS)
    catalogue = SimpleNamespace(job=SimpleNamespace(name="26-9999", root=Path("26-9999, 1 Test Street")),
                                overview=SpecOverview(), pages=(), sheets=(), sections=(), tags=())
    reading = SimpleNamespace(schedules=(vav, heaters), inventory=items, notes=(), backbone=(), devices=(),
                              versions=(), flags=())
    lines = [_written("Furnish Electric Heat Enable (DO)", "KW: 5")]
    judge = FakeWriter(_composition(lines=lines, notes=["**KW 5 heaters"]))
    judge.extra = (WrittenLine("Furnish and Install Occupancy Sensor", 1, None, "The schedule shows it.", "HW GPM",
                               0.9),)
    job, systems, _ = build_job(catalogue, reading, judge, dt.date(2026, 9, 25))
    check_fit(job)
    headers = [b.header for b in job.blocks]
    assert "Qty (2) Electric Duct Heater EDH-1, 2" in headers
    out = write_proposal(job, proposal_template, tmp_path / "proposal.docx")
    whole = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Furnish Electric Heat Enable (DO)" in whole and "Furnish and Install Occupancy Sensor" in whole
    for word in ("confidence", "ASSUMPTION", "CONFIRMED", "ADDED", "PENDING", "rule:", "AI:", "tech hour",
                 "written from the evidence", "Master Template", "The schedule shows it", "KW: 5"):
        assert word not in whole, word


def test_the_summary_counts_a_system_put_in_the_exceptions():
    from estimator.__main__ import estimate_lines
    _, systems, flags = _sensaphone_in_the_spec_job()
    [first, *_] = estimate_lines(systems, flags, None, True)
    assert first == "Systems written: 1 in the proposal, 1 put in the exceptions."
