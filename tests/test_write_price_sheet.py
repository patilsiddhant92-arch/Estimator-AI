import re
import zipfile

import openpyxl
import pytest

from estimator.price_sheet.rows import read_sent_rows
from estimator.price_sheet.write import (PRICE_NEEDS_REVIEW, PRICE_PENDING, QUANTITY_PENDING, PriceSheetError,
                                        write_price_sheet)
from estimator.proposal import layout
from estimator.proposal.sections import scope_texts
from estimator.proposal.write import write_proposal
from estimator.text import strip_list_marker


@pytest.fixture
def written(small_job, workup_template, price_book, tmp_path):
    out = write_price_sheet(small_job, workup_template, price_book, tmp_path / "sheet.xlsx")
    return out, openpyxl.load_workbook(str(out))


def _formulas(workbook, rename=None):
    rename = rename or {}
    return {(rename.get(ws.title, ws.title), c.coordinate): c.value
            for ws in workbook.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")}


def test_rows_land_in_order(written):
    ws = written[1]["Price Sheet 1"]
    assert (ws["D6"].value, ws["B6"].value) == ("SCOPE OF WORK – PREVAILING WAGE", 40)
    assert ws["D7"].value == "•  Front End Web Enabled BMS System"
    assert (ws["D12"].value, ws["B12"].value) == ("1.  Qty (2) Exhaust Fan EF-1, 2", 8)
    assert (ws["D13"].value, ws["C13"].value) == ("o  Furnish and Install DDC Controller", 2)
    assert ws["D11"].value is None                     # the blank line before a block


def test_prices_come_from_the_price_book(written):
    ws = written[1]["Price Sheet 1"]
    assert (ws["E9"].value, ws["H9"].value) == (2500, None)      # Furnish BMS Server
    assert (ws["E18"].value, ws["H18"].value) == (142, None)     # Space Temperature Sensor
    assert (ws["E20"].value, ws["H20"].value) == (None, None)    # no quantity, so not priced


def test_a_line_the_book_cannot_price_is_price_pending(written):
    ws = written[1]["Price Sheet 1"]
    assert (ws["E8"].value, ws["H8"].value) == (None, PRICE_PENDING)     # Network Level Control Panel
    assert (ws["E13"].value, ws["H13"].value) == (None, PRICE_PENDING)   # DDC Controller


def test_a_price_the_book_itself_questions_is_marked(small_job, workup_template, price_book, tmp_path):
    """The revised workbook questions no part the small job uses, so the book here marks Filter DP Switch itself."""
    from dataclasses import replace
    from estimator.price_sheet.price_book import PriceBook
    book = PriceBook([replace(i, needs_review=i.name == "Filter DP Switch") for i in price_book.items], {})
    ws = openpyxl.load_workbook(str(write_price_sheet(small_job, workup_template, book, tmp_path / "s.xlsx")))["Price Sheet 1"]
    assert (ws["E19"].value, ws["H19"].value) == (50, PRICE_NEEDS_REVIEW)  # Filter DP Switch


def test_the_summary_says_how_many_lines_have_no_price(written):
    assert written[1]["Summary"]["B34"].value.startswith("3 lines with a quantity have no price yet")


def test_the_summary_is_filled_and_the_leftover_job_name_is_gone(written, small_job):
    wb = written[1]
    summary = wb["Summary"]
    assert summary["B7"].value == small_job.header.reference
    assert summary["B9"].value == "26-9999"
    assert summary["B15"].value.date() == small_job.header.date
    assert (summary["B21"].value, summary["B22"].value) == ("Base Bid", "ADD ALT")
    leftovers = [c.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and "MCCC LA" in c.value]
    assert leftovers == []


def test_an_alternate_sheet_is_renamed_and_the_summary_follows(written):
    wb = written[1]
    assert "ADD ALT" in wb.sheetnames and "Price Sheet 2" not in wb.sheetnames
    assert wb["Summary"]["C22"].value == "='ADD ALT'!B401"


def test_every_template_formula_survives(written, workup_template):
    before = _formulas(openpyxl.load_workbook(str(workup_template)), {"Price Sheet 2": "ADD ALT"})
    expected = {key: value.replace("'Price Sheet 2'!", "'ADD ALT'!") for key, value in before.items()}
    assert len(expected) == 4199
    assert _formulas(written[1]) == expected


def test_no_formula_points_at_a_sheet_that_does_not_exist(written):
    wb = written[1]
    named = {name for formula in _formulas(wb).values() for name in re.findall(r"'([^']+)'!", formula)}
    assert named <= set(wb.sheetnames)


def test_the_logo_is_kept(written):
    assert "xl/media/image1.png" in zipfile.ZipFile(written[0]).namelist()


def test_excel_recalculates_when_the_file_opens(written):
    assert written[1].calculation.fullCalcOnLoad is True


def test_a_job_too_long_for_one_sheet_is_refused(small_job, workup_template, price_book, tmp_path):
    long_job = small_job.model_copy(update={"blocks": tuple(small_job.blocks[0] for _ in range(100))})
    with pytest.raises(PriceSheetError, match="rows"):
        write_price_sheet(long_job, workup_template, price_book, tmp_path / "long.xlsx")


def test_the_price_sheet_is_the_proposal_row_for_row(small_job, proposal_template, workup_template, price_book, tmp_path):
    proposal = write_proposal(small_job, proposal_template, tmp_path / "p.docx")
    sheet = write_price_sheet(small_job, workup_template, price_book, tmp_path / "s.xlsx")
    rows = read_sent_rows(sheet)
    in_sheet = [strip_list_marker(r.text) for title in ("Price Sheet 1", "ADD ALT") for r in rows[title]]
    in_proposal = scope_texts(proposal)
    in_proposal = in_proposal[: in_proposal.index(layout.DEDUCT_ALTERNATES)]   # small_job has no deduct alternates
    assert in_sheet == in_proposal


def test_a_block_whose_count_is_still_to_fill_is_flagged_and_its_lines_are_not_priced(small_job, workup_template,
                                                                                       price_book, tmp_path):
    from estimator.scope import Block, Item, Run
    block = Block(header="Qty (##) Exhaust Fan EF-9", items=(
        Item(kind="line", runs=(Run(text="Furnish and Install Space Temperature Sensor"),)),))
    job = small_job.model_copy(update={"blocks": small_job.blocks + (block,)})
    wb = openpyxl.load_workbook(str(write_price_sheet(job, workup_template, price_book, tmp_path / "hash.xlsx")))
    ws = wb["Price Sheet 1"]
    header = next(r for r in range(6, 60) if (ws.cell(r, 4).value or "").endswith("Qty (##) Exhaust Fan EF-9"))
    assert ws.cell(header, 8).value == QUANTITY_PENDING
    assert (ws.cell(header + 1, 3).value, ws.cell(header + 1, 5).value) == (None, None)
    assert "1 block has no count yet" in wb["Summary"]["B34"].value
