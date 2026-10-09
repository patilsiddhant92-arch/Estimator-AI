"""Bug 2 (plan.py `_front_end`, ~1398-1401): with no front_end fact the planner set mode "new-system" and the
proposal printed "Front End Web Enabled BMS System" with no source. It must stay unresolved: no front end chosen
(Job.front_end None), the question kept, and the rendered proposal must not print either template front end
heading (owner ruling 2026-10-09: with no front-end source no front-end line is printed). PROVISIONAL ruling R2:
with no BMS or front end shown, the scope prints standalone with one question, word for word.
Invented ledger only."""
import datetime as dt
from pathlib import Path

from estimator.ledger.facts import SETTLED, Ledger
from estimator.proposal.layout import EXISTING_SYSTEM_HEADING, NEW_SYSTEM_HEADING
from estimator.writing.dossier import Dossier

from bms.model import Header
from bms.plan import plan_job
from bms.playbook import load_playbook
from bms.render import render

PLAYBOOK = load_playbook()
SOURCES = {"S3": "M-501, Invented.pdf page 6 (schedule)", "S4": "23 09 00, Invented Specs.pdf page 4"}


def _dossier():
    return Dossier(job_name="Invented Job", text="", images=(), sources={}, families=(), device_totals={},
                   counts={}, stats={})


def _ledger(mode=None):
    ledger = Ledger(dict(SOURCES))
    ledger.job_name = "Invented Job"
    if mode:
        ledger.add("front_end", "mode", mode, "S4", "the BMS", "spec", SETTLED)
    ledger.add("family", "EF", "Exhaust Fan", "S3", "EF Exhaust Fan", "schedules", SETTLED)
    ledger.add("unit", "EF", "EF-1", "S3", "Tags: EF-1", "schedules", SETTLED)
    ledger.add("count", "EF", 1, "S3", "EF rows 1", "schedules", SETTLED)
    return ledger


def _header():
    return Header(quote_number="Q-0000", reference="Invented Job", mep_engineer="Invented Engineer",
                  based_on=("Invented drawings",), date=dt.date(2026, 10, 9).isoformat())


def test_no_front_end_fact_leaves_the_front_end_unresolved():
    job = plan_job(_ledger(), None, _dossier(), PLAYBOOK)
    assert job.front_end is None


STANDALONE = "No BMS or front end is shown in the documents; confirm standalone controls."


def test_r2_no_bms_shown_asks_one_standalone_question():
    """PROVISIONAL ruling R2 (GrokTeam/discussion.md): no BMS or front end shown -> ONE question, word for word,
    and no other front-end question beside it."""
    job = plan_job(_ledger(), None, _dossier(), PLAYBOOK)
    assert [q for q in job.questions if q == STANDALONE] == [STANDALONE], job.questions
    assert not [q for q in job.questions if "new or extends an existing" in q or "front end" in q.lower()
                and q != STANDALONE], job.questions


def test_r2_a_bms_shown_without_its_mode_asks_new_or_existing_and_assumes_nothing():
    """PROVISIONAL ruling R2, negative partner: a front_end fact shows a BMS, so it is not standalone; only its mode
    is open, asked without assuming a new system."""
    ledger = _ledger()
    ledger.add("front_end", "note", "BMS front end per spec", "S4", "the BMS front end", "spec", SETTLED)
    job = plan_job(ledger, None, _dossier(), PLAYBOOK)
    asks = [q for q in job.questions if "new or extends an existing" in q]
    assert len(asks) == 1 and "assumes a new system" not in asks[0], job.questions
    assert STANDALONE not in job.questions


def test_r2_the_standalone_scope_still_prints_its_blocks():
    """REGRESSION GUARD (passes on current code) for PROVISIONAL ruling R2: the systems print standalone."""
    import dataclasses

    from bms.model import Count, Device, Ref

    job = plan_job(_ledger(), None, _dossier(), PLAYBOOK)
    ref = Ref("S3", "EF Exhaust Fan", "M-501, Invented.pdf page 6")
    line = Device("EF.d1", "start_stop_status", "Start / Stop and Status", "Furnish and Install", "AME", "AME", "AME",
                  Count(1, (ref,)), (ref,))
    ef = dataclasses.replace(job.systems[0], devices=(line,))
    job = dataclasses.replace(job, systems=(ef,), header=_header())
    blocks = render(job, PLAYBOOK).scope.blocks
    assert len(blocks) == 1 and blocks[0].items


def test_a_sourced_mode_is_still_used():
    job = plan_job(_ledger("new-system"), None, _dossier(), PLAYBOOK)
    assert job.front_end is not None and job.front_end.mode == "new-system"


def test_the_rendered_scope_has_no_front_end_line_when_unresolved():
    import dataclasses

    job = dataclasses.replace(plan_job(_ledger(), None, _dossier(), PLAYBOOK), header=_header())
    scope = render(job, PLAYBOOK).scope
    assert scope.front_end.mode not in ("new-system", "extend-existing")
    assert scope.front_end.items == ()


def test_the_price_sheet_prints_no_front_end_heading_when_unresolved():
    import dataclasses
    from estimator.price_sheet.rows import sheet_plan

    job = dataclasses.replace(plan_job(_ledger(), None, _dossier(), PLAYBOOK), header=_header())
    texts = [r.text for sheet in sheet_plan(render(job, PLAYBOOK).scope) for r in sheet.rows if r.text]
    assert not [t for t in texts if NEW_SYSTEM_HEADING in t or EXISTING_SYSTEM_HEADING in t], texts


def test_the_written_proposal_has_neither_template_front_end_heading(tmp_path):
    import dataclasses
    import docx
    from estimator.proposal.write import write_proposal

    root = Path(__file__).resolve().parents[2]
    template = root / "assets" / "templates" / "AME_nordomatic Proposal Template.docx"
    job = dataclasses.replace(plan_job(_ledger(), None, _dossier(), PLAYBOOK), header=_header())
    out = write_proposal(render(job, PLAYBOOK).scope, template, tmp_path / "Proposal.docx")
    texts = [p.text.strip() for p in docx.Document(str(out)).paragraphs]
    assert not any(t.startswith(NEW_SYSTEM_HEADING) or t.startswith(EXISTING_SYSTEM_HEADING) for t in texts)
    # owner ruling: no front-end line at all, not even a placeholder; the template's default lines are gone too
    assert not any(t.startswith("Furnish Network Level Control Panel") or "front end" in t.lower() for t in texts)


def test_the_written_proposal_reads_back_with_no_front_end(tmp_path):
    import dataclasses
    import docx
    from estimator.proposal.read import read_scope
    from estimator.proposal.write import write_proposal

    root = Path(__file__).resolve().parents[2]
    template = root / "assets" / "templates" / "AME_nordomatic Proposal Template.docx"
    job = dataclasses.replace(plan_job(_ledger(), None, _dossier(), PLAYBOOK), header=_header())
    scope = render(job, PLAYBOOK).scope
    out = write_proposal(scope, template, tmp_path / "Proposal.docx")
    _suffix, front_end, blocks, _add, _deduct = read_scope(docx.Document(str(out)))
    assert front_end.items == () and front_end.mode == scope.front_end.mode
    assert len(blocks) == len(scope.blocks)
